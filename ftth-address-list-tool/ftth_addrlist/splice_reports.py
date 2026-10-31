"""Splice-report readers.

Two layouts are recognised:

* "Splice Connections Report" - the spatial-system export (primary source).
  Each connection row gives both fibre ends (cable, segment, tube.fibre(abs)),
  colours, the circuit usage text and the circuit start/end enclosures.
* "Fiber Optic Splice Guide" - a hand-made guide found in some packs. It has
  no fibre numbers on the outgoing side and no circuit data, so fibre numbers
  are derived from colours and every connection is treated as low-trust.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from .colors import ColorCode
from .model import (CableFiber, Connection, DeviceInfo, DevicePort, Issue, Location,
                    SegmentInfo, Side, SourceRef, SpliceReport)

FMT_REPORT = "connections_report"
FMT_GUIDE = "splice_guide"

SEG_ROW_RE = re.compile(r"^Segment\s+(\d+)\s+of\s+(.+?)\s+-\s+([\w /]+)$")
SEG_REF_RE = re.compile(r"^(.*\S)\s*:\s*Seg\s*(\d+)$")
CABLE_POS_RE = re.compile(r"^(\d+)\.(\d+)\((\d+)\)$")
DEVICE_POS_RE = re.compile(r"^(\d+)\((\d+)\)$")
GUIDE_SEG_RE = re.compile(r"^(.*\S)\s+SEG-?\s*(\d+)$", re.I)
GUIDE_INOUT_RE = re.compile(r"^(.*\S)\s*:\s*Seg\s*(\d+)\s*-\s*(\d+)\s*ct$", re.I)

# Column header sequence of the connection table (occurrence-aware).
CONN_COLUMNS = [
    ("Segment", 1, "a_seg"), ("Group 2", 1, "a_tube"), ("Fiber", 1, "a_fiber"), ("#", 1, "a_pos"),
    ("Dir", 1, "dir"), ("#", 2, "b_pos"), ("Fiber", 2, "b_fiber"), ("Group 2", 2, "b_tube"),
    ("Segment", 2, "b_seg"), ("Usage", 1, "usage"), ("State", 1, "state"), ("Owner", 1, "owner"),
    ("Start", 1, "start"), ("End", 1, "end"),
]
REQUIRED_CONN = {"a_seg", "a_pos", "b_pos", "b_seg", "usage", "state", "start", "end"}

META_LABELS = {
    "Prepared By:": "prepared_by", "On:": "prepared_on", "Splice name:": "splice_name",
    "Location:": "location", "Address:": "address", "Coordinates:": "coordinates",
    "Latitude/Longitude:": "lat_lon", "Enclosure Type:": "enclosure_type",
}


class ReportError(Exception):
    pass


def _clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = " ".join(v.split())
        return v or None
    return v


def _int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


class NotAReport(ReportError):
    """The workbook is an address list or another output of this tool, not a splice report."""


OUTPUT_SHEETS = {"Row Status", "Field Evidence", "Run Info", "Difference classes", "Field differences"}


def read_report(path: str | Path, cfg: dict) -> SpliceReport:
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # corrupt / not xlsx
        raise ReportError(f"{path.name}: cannot be opened as an Excel workbook ({exc})") from exc
    try:
        sheets = set(wb.sheetnames)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    first = [str(_clean(v) or "") for v in (rows[0] if rows else [])]
    head = {str(_clean(v) or "").upper() for r in rows[:3] for v in r}
    if sheets & OUTPUT_SHEETS or {"PON NAME", "FIBER #"} <= head:
        raise NotAReport(f"{path.name}: an address-list workbook (or another output of this tool), not a splice "
                         f"report - skipped")
    if any(c == "Splice Connections Report" for c in first):
        return _read_connections_report(path, rows, cfg)
    if any(c == "Fiber Optic Splice Guide" for c in first):
        return _read_splice_guide(path, rows, cfg)
    raise ReportError(f"{path.name}: unrecognised layout (expected 'Splice Connections Report' "
                      f"or 'Fiber Optic Splice Guide' in row 1)")


# --------------------------------------------------------------------------- report

def _parse_endpoint(seg_text, pos_text, src: str, issues: list[Issue]):
    seg_text = _clean(seg_text)
    pos_text = _clean(str(pos_text)) if pos_text is not None else None
    if not seg_text or not pos_text:
        return None
    m = SEG_REF_RE.match(seg_text)
    if m:
        p = CABLE_POS_RE.match(pos_text)
        if not p:
            issues.append(Issue("WARNING", "BAD_FIBER_POSITION",
                                f"Fibre position {pos_text!r} not understood", source=src))
            return None
        return CableFiber(m.group(1), int(m.group(2)), int(p.group(3)), int(p.group(1)), int(p.group(2)))
    p = DEVICE_POS_RE.match(pos_text)
    if p:
        return DevicePort(seg_text, int(p.group(2)))
    issues.append(Issue("WARNING", "BAD_PORT_POSITION", f"Port position {pos_text!r} not understood",
                        source=src))
    return None


def _read_connections_report(path: Path, rows: list[list], cfg: dict) -> SpliceReport:
    issues: list[Issue] = []
    header: dict = {}
    seg_hdr = conn_hdr = None
    for i, r in enumerate(rows):
        a = _clean(r[0]) if r else None
        if isinstance(a, str):
            key = META_LABELS.get(a)
            if key:
                vals = [v for v in r[1:] if v is not None]
                header[key] = vals[0] if len(vals) == 1 else (tuple(vals) if vals else None)
            if a == "Segment" and len(r) > 1 and _clean(r[1]) == "Count":
                seg_hdr = i
            if a == "Segment" and any(_clean(v) == "Usage" for v in r):
                conn_hdr = i
    closure = _clean(header.get("splice_name"))
    if not closure:
        closure = path.stem
        issues.append(Issue("WARNING", "NO_SPLICE_NAME",
                            f"'Splice name' missing; closure name taken from file name {closure!r}",
                            source=path.name))
    elif closure != path.stem:
        issues.append(Issue("INFO", "FILE_NAME_DIFFERS",
                            f"File name {path.stem!r} differs from splice name {closure!r}", source=path.name))
    if conn_hdr is None:
        raise ReportError(f"{path.name}: connection table header (Segment ... Usage) not found")

    rep = SpliceReport(path, FMT_REPORT, closure, header, issues=issues)

    # segment / device table
    if seg_hdr is not None:
        split_re = re.compile(cfg["network"]["splitter_regex"])
        for i in range(seg_hdr + 1, conn_hdr):
            r = rows[i]
            name = _clean(r[0]) if r else None
            if not name or name in ("A Cable",):
                continue
            src = SourceRef(path.name, i + 1)
            m = SEG_ROW_RE.match(name)
            if m:
                rep.segments.append(SegmentInfo(
                    cable=m.group(2), seg=int(m.group(1)), count=_int(r[1]),
                    start=Location.parse(r[2]), end=Location.parse(r[3]),
                    placement=m.group(3), source=src))
                continue
            sm = split_re.match(name)
            if sm:
                label = sm.group("label")
                rep.devices.append(DeviceInfo(
                    device_id=label.replace(" ", ""), label=label, count=_int(r[1]),
                    ratio=_int(sm.groupdict().get("ratio")), index=_int(sm.groupdict().get("index")),
                    source=src))
            else:
                rep.devices.append(DeviceInfo(name.replace(" ", ""), name, _int(r[1]), None, None, src))
                issues.append(Issue("INFO", "UNKNOWN_DEVICE", f"Equipment row not recognised as a "
                                    f"splitter: {name!r}", source=str(src)))
    else:
        issues.append(Issue("WARNING", "NO_SEGMENT_TABLE", "Segment table not found", source=path.name))

    # map connection columns
    hdr = [_clean(v) for v in rows[conn_hdr]]
    colmap: dict[str, int] = {}
    for name, occ, key in CONN_COLUMNS:
        seen = 0
        for j, h in enumerate(hdr):
            if h == name:
                seen += 1
                if seen == occ:
                    colmap[key] = j
                    break
    missing = REQUIRED_CONN - colmap.keys()
    if missing:
        raise ReportError(f"{path.name}: connection table is missing column(s) "
                          f"{sorted(missing)} (header row {conn_hdr + 1})")

    def g(r, k):
        j = colmap.get(k)
        return r[j] if j is not None and j < len(r) else None

    for i in range(conn_hdr + 1, len(rows)):
        r = rows[i]
        if not r or all(v is None for v in r):
            continue
        src = SourceRef(path.name, i + 1)
        a = _parse_endpoint(g(r, "a_seg"), g(r, "a_pos"), str(src), issues)
        b = _parse_endpoint(g(r, "b_seg"), g(r, "b_pos"), str(src), issues)
        if a is None or b is None:
            issues.append(Issue("WARNING", "INCOMPLETE_CONNECTION",
                                "Connection row has an unreadable or missing side; ignored", source=str(src)))
            continue
        rep.connections.append(Connection(
            closure=closure,
            a=Side(a, _lower(g(r, "a_tube")), _lower(g(r, "a_fiber"))),
            b=Side(b, _lower(g(r, "b_tube")), _lower(g(r, "b_fiber"))),
            usage=_clean(g(r, "usage")), state=_clean(g(r, "state")), owner=_clean(g(r, "owner")),
            circuit_start=Location.parse(g(r, "start")), circuit_end=Location.parse(g(r, "end")),
            direction=_clean(g(r, "dir")), source=src, fmt=FMT_REPORT))
    return rep


def _lower(v):
    v = _clean(v)
    return v.lower() if isinstance(v, str) else None


# --------------------------------------------------------------------------- guide

def _read_splice_guide(path: Path, rows: list[list], cfg: dict) -> SpliceReport:
    colors = ColorCode(cfg)
    issues: list[Issue] = [Issue("WARNING", "SPLICE_GUIDE_FORMAT",
                                 "Hand-made 'Fiber Optic Splice Guide' layout: no outgoing fibre numbers "
                                 "or circuit data; fibre numbers derived from colours (low trust)",
                                 source=path.name)]
    header: dict = {}
    table_hdr = None
    in_col = out_col = None
    for i, r in enumerate(rows[:20]):
        for j, v in enumerate(r):
            v = _clean(v)
            if v == "In":
                in_col = j
            elif v == "Out":
                out_col = j
            elif v in ("Splice Name:", "Date Created:", "Enclosure Make/Model:", "Location:",
                       "Latitude:", "Longitude:", "Pole/Ped Number:"):
                rest = [x for x in r[j + 1:] if _clean(x) is not None]
                header[v.rstrip(":").lower().replace(" ", "_").replace("/", "_")] = rest[0] if rest else None
        if r and _clean(r[1] if len(r) > 1 else None) == "ID":
            table_hdr = i
    closure = _clean(header.get("splice_name")) or path.stem
    rep = SpliceReport(path, FMT_GUIDE, closure, header, issues=issues)

    # In / Out cable segments listed in the header block
    for i, r in enumerate(rows[:12]):
        for col, role in ((in_col, "in"), (out_col, "out")):
            if col is None or col >= len(r):
                continue
            v = _clean(r[col])
            m = GUIDE_INOUT_RE.match(v) if isinstance(v, str) else None
            if not m:
                continue
            loc = Location("Splice Case", closure)
            rep.segments.append(SegmentInfo(
                cable=m.group(1), seg=int(m.group(2)), count=int(m.group(3)),
                start=None if role == "in" else loc, end=loc if role == "in" else None,
                placement=None, source=SourceRef(path.name, i + 1)))
    if table_hdr is None:
        raise ReportError(f"{path.name}: splice-guide table header (Cable ID / Buffer / Fiber) not found")

    used_b: dict[tuple, int] = {}
    for i in range(table_hdr + 1, len(rows)):
        r = rows[i] + [None] * 16
        num, a_cab, a_buf, a_fib = _int(r[0]), _clean(r[1]), r[5], r[6]
        b_cab, b_buf, b_fib, hot, addr = _clean(r[7]), r[11], r[12], _clean(r[13]), _clean(r[14])
        if not a_cab or not b_cab:
            continue
        src = SourceRef(path.name, i + 1)
        ma, mb = GUIDE_SEG_RE.match(a_cab), GUIDE_SEG_RE.match(b_cab)
        if not ma or not mb:
            issues.append(Issue("WARNING", "BAD_GUIDE_CABLE", f"Cable id not understood: {a_cab!r} / {b_cab!r}",
                                source=str(src)))
            continue
        fa = colors.fiber_number(a_buf, a_fib)
        fb = colors.fiber_number(b_buf, b_fib)
        if fa is None or fb is None:
            issues.append(Issue("WARNING", "BAD_GUIDE_COLOR", "Colour code not understood", source=str(src)))
            continue
        notes = []
        if num is not None and num != fa:
            notes.append(f"row number {num} disagrees with colour-derived fibre {fa}")
        bkey = (" ".join(mb.group(1).split()), int(mb.group(2)), fb)
        if bkey in used_b:
            notes.append(f"outgoing fibre {fb} also used on row {used_b[bkey]}")
            issues.append(Issue("WARNING", "GUIDE_DUPLICATE_FIBER",
                                f"Outgoing {bkey[0]} Seg{bkey[1]} fibre {fb} assigned twice "
                                f"(rows {used_b[bkey]} and {i + 1})", source=str(src)))
        used_b.setdefault(bkey, i + 1)
        a = CableFiber(" ".join(ma.group(1).split()), int(ma.group(2)), fa)
        b = CableFiber(bkey[0], bkey[1], fb)
        rep.connections.append(Connection(
            closure=closure,
            a=Side(a, colors.normalise(a_buf), colors.normalise(a_fib)),
            b=Side(b, colors.normalise(b_buf), colors.normalise(b_fib)),
            usage=addr, state="Hot" if hot else None, owner=None, circuit_start=None, circuit_end=None,
            direction=None, source=src, fmt=FMT_GUIDE, notes=notes))
    return rep


def load_reports(paths, cfg: dict):
    """Read many report files. Returns (reports, issues)."""
    reports: list[SpliceReport] = []
    issues: list[Issue] = []
    for p in paths:
        try:
            rep = read_report(p, cfg)
        except NotAReport as exc:   # completed lists / earlier outputs kept with the reports: ignored, not an error
            issues.append(Issue("INFO", "NOT_A_SPLICE_REPORT", str(exc), source=Path(p).name))
            continue
        except ReportError as exc:
            issues.append(Issue("ERROR", "REPORT_UNREADABLE", str(exc), source=Path(p).name))
            continue
        reports.append(rep)
        issues.extend(rep.issues)
    return reports, issues
