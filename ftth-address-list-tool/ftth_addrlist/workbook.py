"""Write the address-list workbook (template layout) plus review, evidence and run sheets."""

from __future__ import annotations

import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import __version__
from .builder import FIELD_LABELS, FIELDS, OlmResult
from .config import resolve_tool_path
from .model import APPROVED, CONFIRMED, PROVISIONAL, UNRESOLVED

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
FONT = Font(name="Aptos Narrow", size=11)
HDR_FONT = Font(name="Aptos Narrow", size=11, bold=True)
HDR_FILL = PatternFill("solid", fgColor="FFCCCCCC")
STATUS_STYLE = {   # Excel built-in "Neutral", "Bad" and a blue for approved values
    PROVISIONAL: (PatternFill("solid", fgColor="FFFFEB9C"), Font(name="Aptos Narrow", size=11, color="FF9C5700")),
    UNRESOLVED: (PatternFill("solid", fgColor="FFFFC7CE"), Font(name="Aptos Narrow", size=11, color="FF9C0006")),
    APPROVED: (PatternFill("solid", fgColor="FFDDEBF7"), Font(name="Aptos Narrow", size=11, color="FF1F4E78")),
}
TEXT_FIELDS = {"address", "tube_color", "fiber_color", "cable_name", "tap_name", "fiber_splice"}

ACTIONS = {
    "SPLICE_CLOSURE_UNKNOWN": "Obtain splice reports for the candidate closures, or approve the closure in overrides",
    "SEGMENT_NUMBER_MISMATCH": "Re-export both closure reports from the same spatial-system revision",
    "TAP_CONFLICT": "Confirm the tap in the spatial system; correct the usage text or approve",
    "USAGE_TEXT_VARIANT": "Address/tap renamed between report revisions; re-export reports together",
    "DOWNSTREAM_CIRCUIT_DIFFERS": "Downstream reports show another circuit on this fibre; check revisions",
    "CIRCUIT_CONFLICT": "Downstream records contradict the hub; resolve in the spatial system",
    "TRACE_BROKEN": "Live fibre missing from a closure report; check the report export",
    "TRACE_STOPPED": "Connectivity incomplete; check segment tables of the reports",
    "DUPLICATE_ADDRESS": "Same address on several fibres; confirm units/services",
    "ADDRESS_SUSPECT": "Usage text has a tap id where the address should be; correct in spatial system",
    "USAGE_NONSTANDARD": "PON-port usage text is not '<tap>:<port> <address>'; correct or approve",
    "PON_WITHOUT_ADDRESS": "PON port in use with no address; supply the address",
    "COLOR_MISMATCH": "Report colour differs from colour-code position; verify cable data",
    "SPLICE_GUIDE_FORMAT": "Replace the hand-made splice guide with a spatial-system Splice Connections Report",
    "GUIDE_DUPLICATE_FIBER": "Hand-made guide assigns one fibre twice; obtain the system report",
    "DUPLICATE_REPORT": "Remove the outdated duplicate report",
    "REPORT_UNREADABLE": "Re-export the report as .xlsx",
    "NOT_A_SPLICE_REPORT": "Information only: an address list or earlier output in the report folder was ignored",
    "NO_HUB": "Supply the hub (splitter) closure report for this OLM",
    "PON_PORTS_FREE": "Information only",
    "FIBER_SPLICE_FROM_DRAWING": "Closure located from the design drawing; confirm with that closure's splice "
                                 "report, or accept FS-DRAWING in the approvals policy",
    "DRAWING_DISAGREES": "Report-derived closure and drawing differ; check the closure report or the drawing",
    "DRAWN_DROP_DISAGREES": "Drawn drop and splice reports name different closures; check the drop or the reports",
    "DRAWN_DROP_UNREACHABLE": "Drawn drop points to a closure the fibre cannot reach; check the drop or the reports",
    "DRAWING_NO_CANDIDATE": "No unclaimed tap enclosure reachable in the drawing; supply the splice reports",
    "DOWNSTREAM_PREMISE_DIFFERS": "Hub and downstream reports name different premises on one circuit; correct "
                                  "the spatial system and re-export",
    "ADDRESS_NO_NUMBER": "Address has no house number; correct it in the spatial system",
    "ADDRESS_REPEATED_WORD": "Repeated word in the address (e.g. 'AVE AVE'); correct it in the spatial system",
}


def _hdr_cell(c):
    c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _template_styles(cfg: dict):
    """Header and body styles from the template workbook (first sheet, rows 1-2), if present."""
    p = cfg["template"].get("path")
    if not p:
        return None
    path = resolve_tool_path(p)
    if not path.exists():
        return None
    ws = openpyxl.load_workbook(path).worksheets[0]
    cols = cfg["template"]["columns"]
    hdr, body, widths = [], [], []
    for j in range(1, len(cols) + 1):
        h, b = ws.cell(1, j), ws.cell(2, j)
        hdr.append((h.value, copy.copy(h.font), copy.copy(h.fill), copy.copy(h.border), copy.copy(h.alignment)))
        body.append((copy.copy(b.font), copy.copy(b.fill), copy.copy(b.border), copy.copy(b.alignment)))
        widths.append(ws.column_dimensions[get_column_letter(j)].width)
    return {"header": hdr, "body": body, "widths": widths, "height": ws.row_dimensions[1].height, "path": str(path)}


def write_address_sheet(ws, result: OlmResult, cfg: dict):
    cols = cfg["template"]["columns"]
    tpl = _template_styles(cfg)
    mark = cfg["template"].get("mark_review_cells", True)
    for j, col in enumerate(cols, 1):
        c = ws.cell(1, j, col["header"])
        if tpl and tpl["header"][j - 1][0] is not None:
            _, c.font, c.fill, c.border, c.alignment = tpl["header"][j - 1]
        else:
            _hdr_cell(c)
        width = (tpl["widths"][j - 1] if tpl and tpl["widths"][j - 1] else None) or col.get("width", 12)
        ws.column_dimensions[get_column_letter(j)].width = width
    if tpl and tpl["height"]:
        ws.row_dimensions[1].height = tpl["height"]
    if cfg["template"].get("autofit_columns", True):
        for j, col in enumerate(cols, 1):
            longest = max([len(str(r.fields[col["field"]].value)) for r in result.rows
                           if r.fields[col["field"]].value is not None] + [0])
            letter = get_column_letter(j)
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                     min(60, longest * 1.1 + 2))
    if cfg["template"].get("print_fit_width", True):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
    for i, r in enumerate(result.rows, 2):
        for j, col in enumerate(cols, 1):
            fv = r.fields[col["field"]]
            v = fv.value
            if col["field"] == "pon_name" and isinstance(v, str) and v.isdigit():
                v = int(v)           # numeric like the reference lists
            c = ws.cell(i, j, v)
            if tpl:
                c.font, c.fill, c.border, c.alignment = (copy.copy(x) for x in tpl["body"][j - 1])
            else:
                c.font, c.border, c.alignment = FONT, BORDER, CENTER
            if col["field"] in TEXT_FIELDS and v is not None:
                c.number_format = "@"      # keep identifiers as text (leading zeros preserved)
                c.value = str(v)
            if mark and fv.status in STATUS_STYLE and (v is not None or fv.status == UNRESOLVED):
                c.fill, c.font = STATUS_STYLE[fv.status]


def _table(ws, headers, rows, widths=None):
    for j, h in enumerate(headers, 1):
        _hdr_cell(ws.cell(1, j, h))
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def write_workbook(result: OlmResult, cfg: dict, out_path: Path, run_info: dict,
                   crosscheck: list[dict] | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg["template"]["sheet_name"].format(olm=result.olm)[:31]
    write_address_sheet(ws, result, cfg)

    # Row status
    rs = wb.create_sheet("Row Status")
    hdr = ["Sheet row", "Section", "Cable", "Fiber #", "Row kind", "Row status", "Fields to review", "Reasons",
           "Fibre trace"]
    rows = []
    for i, r in enumerate(result.rows, 2):
        review = [FIELD_LABELS[k] for k in FIELDS if r.fields[k].status != CONFIRMED]
        reasons = [f"{FIELD_LABELS[k]}: {r.fields[k].note or r.fields[k].rule}" for k in FIELDS
                   if r.fields[k].status != CONFIRMED]
        rows.append([i, r.section, r.cable, r.fiber, r.kind, r.status, ", ".join(review), " | ".join(reasons),
                     " ; ".join(r.trace)])
    _table(rs, hdr, rows, [9, 8, 42, 8, 14, 13, 30, 90, 120])
    for i, row in enumerate(rows, 2):
        st = row[5]
        if st in STATUS_STYLE:
            rs.cell(i, 6).fill, rs.cell(i, 6).font = STATUS_STYLE[st]

    # Field evidence (long format)
    fe = wb.create_sheet("Field Evidence")
    hdr = ["Sheet row", "Cable", "Fiber #", "Field", "Value", "Status", "Rule code", "Derivation", "Evidence (file row)",
           "Note"]
    rows = []
    for i, r in enumerate(result.rows, 2):
        for k in FIELDS:
            fv = r.fields[k]
            if fv.code == "N/A" and fv.value is None:
                continue
            rows.append([i, r.cable, r.fiber, FIELD_LABELS[k], fv.value, fv.status, fv.code, fv.rule,
                         "; ".join(dict.fromkeys(fv.evidence)), fv.note])
    _table(fe, hdr, rows, [9, 42, 8, 14, 30, 13, 18, 70, 60, 70])

    # Exceptions
    ex = wb.create_sheet("Exceptions")
    hdr = ["Severity", "Code", "OLM", "Cable", "Fiber #", "Message", "Source", "Action needed"]
    order = {"ERROR": 0, "WARNING": 1, "INFO": 2}
    issues = sorted(result.issues, key=lambda x: (order.get(x.severity, 3), x.code))
    _table(ex, hdr, [[x.severity, x.code, x.olm, x.cable, x.fiber, x.message, x.source, ACTIONS.get(x.code, "")]
                     for x in issues], [10, 26, 10, 42, 8, 110, 30, 70])

    need = {}
    for r in result.rows:
        for c in r.needs_reports:
            need.setdefault(c, []).append(r)
    if need:
        rq = wb.create_sheet("Reports to Request")
        _table(rq, ["Closure", "Rows it would confirm", "Example addresses"],
               [[c, len(rs), "; ".join(str(x.value("address")) for x in rs[:3])]
                for c, rs in sorted(need.items(), key=lambda kv: (-len(kv[1]), kv[0]))], [34, 22, 90])

    if crosscheck is not None:
        cx = wb.create_sheet("Drawing Cross-check")
        _table(cx, ["Check", "Identifier", "Found in drawing", "Detail"],
               [[c["check"], c["identifier"], c["found"], c.get("detail", "")] for c in crosscheck], [26, 46, 16, 80])

    # Run info
    ri = wb.create_sheet("Run Info")
    c = result.counts
    verdict = ("FINAL - every row confirmed or approved, no errors" if result.is_final else
               "DRAFT - NOT FOR ISSUE: rows need review (see Row Status / Exceptions)")
    info = [
        ("Status", verdict), ("OLM", result.olm), ("Hub closure", result.hub),
        ("Rows", len(result.rows)), ("Confirmed rows", c[CONFIRMED]), ("Approved rows", c[APPROVED]),
        ("Provisional rows", c[PROVISIONAL]), ("Unresolved rows", c[UNRESOLVED]),
        ("Errors", len(result.errors)), ("Sections", "; ".join(f"{s[0]} ({s[2]}F)" for s in result.sections)),
        ("Tool version", __version__), ("Run at", run_info.get("run_at")), ("Configuration", cfg.get("_source")),
        ("Template", run_info.get("template")), ("Overrides", run_info.get("overrides") or "none"),
        ("Accepted rule policy", ", ".join(cfg.get("approvals", {}).get("accepted_rules", [])) or "none"),
        ("Legend", "Cell colours on the address-list sheet: amber = PROVISIONAL (derived, needs review), "
                   "red = UNRESOLVED (missing/conflicting data), blue = APPROVED (override or accepted rule), "
                   "no colour = CONFIRMED by explicit splice-report records."),
    ]
    for i, (k, v) in enumerate(info, 1):
        ri.cell(i, 1, k).font = HDR_FONT
        ri.cell(i, 2, v)
    if not result.is_final:
        ri.cell(1, 2).fill, ri.cell(1, 2).font = STATUS_STYLE[UNRESOLVED]
    start = len(info) + 2
    ri.cell(start, 1, "Input files").font = HDR_FONT
    for j, h in enumerate(["File", "Format", "Closure", "Prepared", "SHA-256", "Bytes"], 1):
        _hdr_cell(ri.cell(start + 1, j, h))
    for i, f in enumerate(run_info.get("inputs", []), start + 2):
        for j, k in enumerate(["file", "format", "closure", "prepared_on", "sha256", "bytes"], 1):
            ri.cell(i, j, f.get(k))
    ri.column_dimensions["A"].width = 34
    ri.column_dimensions["B"].width = 90
    for col in "CDEF":
        ri.column_dimensions[col].width = 22

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
