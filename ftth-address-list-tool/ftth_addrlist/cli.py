"""Command-line interface.

  python -m ftth_addrlist run      --reports <folder> --out <folder> [--olm ABC-R01] [--config ...]
                                   [--overrides overrides.csv] [--prefix "..."] [--reference <folder>]
                                   [--drawing design.dxf|.dwg]
  python -m ftth_addrlist batch    --jobs jobs.csv
  python -m ftth_addrlist inspect  --reports <folder> [--drawing ...]
  python -m ftth_addrlist compare  --reports <folder> --reference <file or folder> --out <folder> [--olm ...]
  python -m ftth_addrlist make-template --out templates/FTTH_Address_List_Template.xlsx
"""

from __future__ import annotations

import argparse
import collections
import sys
from pathlib import Path

from .config import ConfigError, load_config
from .pipeline import InputError, find_report_files, run, run_batch


def _inspect(args) -> int:
    from .network import Network
    from .splice_reports import load_reports
    cfg = load_config(args.config)
    files = find_report_files(Path(args.reports))
    reps, issues = load_reports(files, cfg)
    net = Network(reps, cfg)
    fmt = collections.Counter(r.fmt for r in reps)
    print(f"Files: {len(files)}  readable reports: {len(reps)}  formats: {dict(fmt)}")
    for olm in net.olms():
        hubs = net.find_hubs(olm)
        splitters = [d.label for (h, _), d in net.devices.items() if hubs and h == hubs[0]]
        print(f"  OLM {olm}: hub {hubs or 'NOT FOUND'}; splitters {splitters}")
    for i in issues + net.issues:
        if i.severity != "INFO":
            print(f"  [{i.severity}] {i.code}: {i.message}")
    if args.drawing:
        from .drawing import read_drawing_inventory
        inv, d_issues = read_drawing_inventory(Path(args.drawing), cfg)
        for i in d_issues:
            print(f"  [{i.severity}] {i.code}: {i.message}")
        if inv:
            print(f"Drawing {inv['source']}: {len(inv['closures'])} closure ids, {len(inv['cables'])} cable names, "
                  f"{len(inv['taps'])} tap names ({inv['method']})")
    return 0


def _make_template(args) -> int:
    import openpyxl
    from .workbook import BORDER, CENTER, FONT, _hdr_cell
    from openpyxl.utils import get_column_letter
    cfg = load_config(args.config)
    out = Path(args.out)
    if out.exists() and not args.force:
        print(f"{out} exists; use --force to replace it")
        return 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OLM"
    for j, col in enumerate(cfg["template"]["columns"], 1):
        _hdr_cell(ws.cell(1, j, col["header"]))
        body = ws.cell(2, j)
        body.font, body.border, body.alignment = FONT, BORDER, CENTER
        ws.column_dimensions[get_column_letter(j)].width = col.get("width", 12)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Template written: {out}")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="ftth_addrlist", description="FTTH address-list generator")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p, out=True):
        p.add_argument("--reports", required=True, help="folder (searched recursively) or single .xlsx splice report")
        if out:
            p.add_argument("--out", required=True, help="output folder (existing files are never overwritten)")
        p.add_argument("--olm", action="append", help="OLM name, e.g. ABC-R01 (repeatable; default: all found)")
        p.add_argument("--config", help="regional configuration .toml")
        p.add_argument("--drawing", help="design drawing (.dwg AutoCAD 2010-2018 or .dxf): cross-checks and Fiber "
                       "Splice where the reports stop (FS-DRAWING)")

    p = sub.add_parser("run", help="generate address-list workbook(s)")
    common(p)
    p.add_argument("--overrides", help="approved overrides CSV")
    p.add_argument("--prefix", help="output file-name prefix, e.g. '123456_Example City ST'")
    p.add_argument("--reference", help="completed workbook(s) to compare against (optional)")
    p.add_argument("--template", help="template workbook (.xlsx) for header/styles")
    p = sub.add_parser("batch", help="run several regions from a jobs CSV")
    p.add_argument("--jobs", required=True)
    p = sub.add_parser("inspect", help="check inputs without generating")
    common(p, out=False)
    p = sub.add_parser("compare", help="generate and compare with reference workbook(s)")
    common(p)
    p.add_argument("--reference", required=True)
    p.add_argument("--overrides")
    p = sub.add_parser("make-template", help="write a blank template workbook")
    p.add_argument("--out", required=True)
    p.add_argument("--config")
    p.add_argument("--force", action="store_true")
    args = ap.parse_args(argv)
    try:
        if args.cmd == "inspect":
            return _inspect(args)
        if args.cmd == "make-template":
            return _make_template(args)
        if args.cmd == "batch":
            outs = run_batch(args.jobs)
        else:
            outs = run(args.reports, args.out, args.olm, args.config, getattr(args, "overrides", None),
                       getattr(args, "prefix", None), getattr(args, "reference", None), args.drawing,
                       getattr(args, "template", None))
    except (InputError, ConfigError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    # exit code: 0 all final, 1 drafts produced, 3 something failed
    if any(o.workbook is None for o in outs):
        return 3
    return 0 if all(o.final for o in outs) else 1


if __name__ == "__main__":
    sys.exit(main())
