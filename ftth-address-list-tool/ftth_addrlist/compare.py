"""Compare a generated address list with a reference (manually prepared) workbook."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .builder import FIELDS, OlmResult


def _norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split()).upper()


def _exact(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


@dataclass
class FieldDiff:
    key: tuple
    field: str
    generated: object
    reference: object
    ref_row: int | None
    gen_status: str
    kind: str          # cosmetic (whitespace/case only) | substantive


@dataclass
class Comparison:
    olm: str
    reference: str
    ref_rows: int = 0
    gen_rows: int = 0
    matched: int = 0
    only_generated: list = field(default_factory=list)
    only_reference: list = field(default_factory=list)
    duplicate_ref_keys: list = field(default_factory=list)
    diffs: list[FieldDiff] = field(default_factory=list)
    field_stats: dict = field(default_factory=dict)    # field -> {compared, exact, normalised}
    order_identical: bool = False
    section_order_ref: list = field(default_factory=list)
    section_order_gen: list = field(default_factory=list)
    header_ref: list = field(default_factory=list)
    status_breakdown: dict = field(default_factory=dict)   # status -> {rows, rows_fully_matching}


def read_reference(path: str | Path, cfg: dict):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    by_header = {_norm(c["header"]): c["field"] for c in cfg["template"]["columns"]}
    colmap = {}
    for j, h in enumerate(hdr):
        f = by_header.get(_norm(h))
        if f:
            colmap[f] = j
    rows = []
    for r in ws.iter_rows(min_row=2):
        vals = [c.value for c in r]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
            continue
        rows.append(({f: (vals[j] if j < len(vals) else None) for f, j in colmap.items()}, r[0].row))
    return ws.title, hdr, colmap, rows


def compare(result: OlmResult, ref_path: str | Path, cfg: dict) -> Comparison:
    title, hdr, colmap, ref_rows = read_reference(ref_path, cfg)
    cmp = Comparison(result.olm, str(Path(ref_path).name), header_ref=hdr)
    cmp.ref_rows = len(ref_rows)
    cmp.gen_rows = len(result.rows)
    ref_by_key = {}
    for vals, rn in ref_rows:
        key = (_norm(vals.get("cable_name")), _norm(vals.get("fiber_no")))
        if key in ref_by_key:
            cmp.duplicate_ref_keys.append((key, rn))
            continue
        ref_by_key[key] = (vals, rn)
    gen_by_key = {(_norm(r.cable), _norm(r.fiber)): r for r in result.rows}
    for f in FIELDS:
        cmp.field_stats[f] = {"compared": 0, "exact": 0, "normalised": 0}
    status = {}
    for key, r in gen_by_key.items():
        st = status.setdefault(r.status, {"rows": 0, "rows_fully_matching": 0, "matched": 0})
        st["rows"] += 1
        if key not in ref_by_key:
            cmp.only_generated.append((key, r.status))
            continue
        cmp.matched += 1
        st["matched"] += 1
        vals, rn = ref_by_key[key]
        all_ok = True
        for f in FIELDS:
            if f not in colmap:
                continue
            g, x = r.value(f), vals.get(f)
            s = cmp.field_stats[f]
            s["compared"] += 1
            if _exact(g) == _exact(x):
                s["exact"] += 1
                s["normalised"] += 1
                continue
            all_ok = False
            if _norm(g) == _norm(x):
                s["normalised"] += 1
                kind = "cosmetic"
            else:
                kind = "substantive"
            cmp.diffs.append(FieldDiff(key, f, g, x, rn, r.fields[f].status, kind))
        if all_ok:
            st["rows_fully_matching"] += 1
    for key, (vals, rn) in ref_by_key.items():
        if key not in gen_by_key:
            cmp.only_reference.append((key, rn))
    ref_order = [(_norm(v.get("cable_name")), _norm(v.get("fiber_no"))) for v, _ in ref_rows]
    gen_order = [(_norm(r.cable), _norm(r.fiber)) for r in result.rows]
    cmp.order_identical = ref_order == gen_order
    cmp.section_order_ref = list(dict.fromkeys(k[0] for k in ref_order))
    cmp.section_order_gen = list(dict.fromkeys(k[0] for k in gen_order))
    cmp.status_breakdown = status
    return cmp


def classify(d: FieldDiff) -> str:
    """Generic difference classes (root causes are investigated separately)."""
    if d.kind == "cosmetic":
        return "cosmetic (whitespace/case only)"
    if d.generated in (None, "") and d.gen_status == "UNRESOLVED":
        return "tool left unresolved (data missing)"
    if d.reference in (None, ""):
        return "reference blank, tool has value"
    if d.gen_status in ("PROVISIONAL", "UNRESOLVED"):
        return "differs - tool value provisional"
    return "differs - tool value confirmed by reports"


def write_comparison_workbook(comparisons: list[Comparison], path: Path):
    from .workbook import _table
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    hdr = ["OLM", "Reference", "Ref rows", "Generated rows", "Matched rows", "Only generated", "Only reference",
           "Row order identical"] + [f"{f} exact" for f in FIELDS] + [f"{f} normalised" for f in FIELDS]
    rows = []
    for c in comparisons:
        rows.append([c.olm, c.reference, c.ref_rows, c.gen_rows, c.matched, len(c.only_generated),
                     len(c.only_reference), c.order_identical] +
                    [f"{c.field_stats[f]['exact']}/{c.field_stats[f]['compared']}" for f in FIELDS] +
                    [f"{c.field_stats[f]['normalised']}/{c.field_stats[f]['compared']}" for f in FIELDS])
    _table(ws, hdr, rows)
    cl = wb.create_sheet("Difference classes")
    counts = {}
    for c in comparisons:
        for d in c.diffs:
            k = (c.olm, d.field, classify(d))
            counts[k] = counts.get(k, 0) + 1
    _table(cl, ["OLM", "Field", "Class", "Count"], [list(k) + [v] for k, v in sorted(counts.items())], [10, 14, 45, 8])
    df = wb.create_sheet("Field differences")
    _table(df, ["OLM", "Reference row", "Cable", "Fiber #", "Field", "Generated", "Reference", "Generated status",
                "Class"],
           [[c.olm, d.ref_row, d.key[0], d.key[1], d.field, d.generated, d.reference, d.gen_status, classify(d)]
            for c in comparisons for d in c.diffs], [10, 10, 42, 8, 14, 32, 32, 14, 40])
    un = wb.create_sheet("Unmatched rows")
    _table(un, ["OLM", "Side", "Cable", "Fiber #", "Row/status"],
           [[c.olm, "generated only", k[0], k[1], s] for c in comparisons for k, s in c.only_generated] +
           [[c.olm, "reference only", k[0], k[1], rn] for c in comparisons for k, rn in c.only_reference])
    wb.save(path)
    return path
