"""Synthetic splice reports in the spatial-system 'Splice Connections Report' layout.

The network (identifiers are fictitious):

  OLM TST-R01 --feeder FC024--> hub TST-00001G (splitter PLC COUPLER 64 1)
     power pigtail  TST-00001G TO TST-R01-POWER TAP FC004  (PON 1 port 1)
     cable A        TST-00001G TO TST-00010 FC012
        Seg1 TST-00001G-TST-00002 (report) Seg2 -TST-00003 (no report) Seg3 -TST-00004 (report)
        Seg4 -TST-00010 (no report, cable end)
     at TST-00002: pigtail TST-00002 TO TST-R-100 MAIN ST FC004 -> MAX10 enclosure

  cable A fibres: 1-2 transport, 3 -> tap at TST-00002 (confirmed), 4 -> continues to the cable end
  (bridged over TST-00003), 5 -> cut at TST-00003 (broken remnant), 6 -> lost after TST-00003
  (unresolved), 7-12 not spliced (spare)
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

COLORS = ["blue", "orange", "green", "brown", "slate", "white", "red", "black", "yellow", "violet", "rose", "aqua"]
CONN_HDR = ["Segment", "Group 3", "Group 2", "Fiber", "#", "Dir", "#", "Fiber", "Group 2", "Group 3", "Segment",
            "Usage", "State", "Owner", "Carrier type", "Start", "End", "Master Circuit", "Circuit ID"]

HUB, OLM = "TST-00001G", "TST-R01"
FEEDER = f"{OLM} TO {HUB} FC024"
POWER = f"{HUB} TO TST-R01-POWER TAP FC004"
CAB = f"{HUB} TO TST-00010 FC012"
PIG = "TST-00002 TO TST-R-100 MAIN ST FC004"
SPL = "PLCCOUPLER641"


def pos(fiber: int) -> str:
    return f"{(fiber - 1) // 12 + 1:02d}.{(fiber - 1) % 12 + 1:02d}({fiber})"


def cable_side(seg_ref: str, fiber: int):
    return {"seg": seg_ref, "tube": COLORS[(fiber - 1) // 12], "fiber": COLORS[(fiber - 1) % 12], "pos": pos(fiber)}


def dev_side(device: str, port: int):
    return {"seg": device, "pos": f"{port}({port})"}


def write_report(path: Path, closure: str, segments: list[tuple], connections: list[dict],
                 devices: list[tuple] = (), drop_column: str | None = None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Splice Connections Report"
    ws["A1"] = "Splice Connections Report"
    ws["A3"], ws["B3"] = "Prepared By:", "Test Author"
    ws["A4"], ws["B4"] = "On: ", "2026-05-07 10:00:00"
    ws["A5"], ws["B5"] = "Splice name:", closure
    ws["A10"], ws["B10"] = "Enclosure Type:", "SPLICE CASE"
    for j, h in enumerate(["Segment", "Count", "Start Splice", "End Splice", "Owner", "Cable Design Length",
                           "Cable Asbuilt Length"], 1):
        ws.cell(12, j, h)
    r = 13
    for seg_name, count, start, end, kind in segments:
        cable, n = seg_name.rsplit(": Seg", 1)
        ws.cell(r, 1, f"Segment {n} of {cable} - {kind}")
        ws.cell(r, 2, count)
        ws.cell(r, 3, start)
        ws.cell(r, 4, end)
        ws.cell(r, 5, "Owner")
        r += 1
    for label, count in devices:
        ws.cell(r, 1, label)
        ws.cell(r, 2, count)
        r += 1
    r += 1
    ws.cell(r, 1, "A Cable")
    r += 1
    hdr = [h for h in CONN_HDR if h != drop_column] if drop_column else CONN_HDR
    for j, h in enumerate(hdr, 1):
        ws.cell(r, j, h)
    for c in connections:
        r += 1
        a, b = c["a"], c["b"]
        vals = [a["seg"], None, a.get("tube"), a.get("fiber"), a["pos"], c.get("dir", "->"), b["pos"], b.get("fiber"),
                b.get("tube"), None, b["seg"], c.get("usage", "No Desc"), c.get("state", "Return"), "No Owner", None,
                c.get("start"), c.get("end"), None, None]
        if drop_column:
            vals = [v for h, v in zip(CONN_HDR, vals) if h != drop_column]
        for j, v in enumerate(vals, 1):
            ws.cell(r, j, v)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def usage(n: int) -> str:
    return f"TST-R-{n} MAIN ST:1 {n} MAIN ST"


def build_network(folder: Path, *, hub_extra: list[dict] | None = None, skip: set[str] = frozenset(),
                  tst4_extra: list[dict] | None = None, hub_segments_extra: list[tuple] = (),
                  tst2_extra: list[dict] | None = None, tst2_segments_extra: list[tuple] = ()) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    term = lambda n: f"MAX10 Enclosure(M Bulk): TST-R-{n} MAIN ST"
    hub_c = [
        {"a": cable_side(f"{FEEDER}: Seg1", 1), "b": cable_side(f"{CAB}: Seg1", 1),
         "usage": f"{OLM} OLM TRANSPORT RETURN", "state": "Return", "start": "Rack", "end": f"OLM Housing: {OLM}"},
        {"a": cable_side(f"{FEEDER}: Seg1", 2), "b": cable_side(f"{CAB}: Seg1", 2),
         "usage": f"{OLM} OLM TRANSPORT FORWARD", "state": "Forward", "start": f"OLM Housing: {OLM}", "end": "Rack"},
        {"a": cable_side(f"{FEEDER}: Seg1", 5), "b": dev_side(SPL, 1), "state": "Misc"},
        {"a": dev_side(SPL, 2), "b": cable_side(f"{POWER}: Seg1", 1), "usage": f"{OLM}-POWER PORT",
         "start": "MAX10 Enclosure(M Bulk): TST-R01-POWER TAP", "end": f"Splice Case: {HUB}"},
    ]
    for fib, port, n in ((3, 3, 100), (4, 4, 200), (5, 5, 300), (6, 6, 400)):
        hub_c.append({"a": dev_side(SPL, port), "b": cable_side(f"{CAB}: Seg1", fib), "usage": usage(n),
                      "start": term(n), "end": f"Splice Case: {HUB}"})
    hub_c += hub_extra or []
    if "hub" not in skip:
        write_report(folder / f"{HUB}.xlsx", HUB, [
            (f"{FEEDER}: Seg1", 24, f"OLM Housing: {OLM}", f"Splice Case: {HUB}", "Underground"),
            (f"{POWER}: Seg1", 4, f"Splice Case: {HUB}", "MAX10 Enclosure(M Bulk): TST-R01-POWER TAP", "Pigtail"),
            (f"{CAB}: Seg1", 12, f"Splice Case: {HUB}", "Splice Case: TST-00002", "Underground"),
            *hub_segments_extra,
        ], hub_c, devices=[("PLC COUPLER 64 1 FE0-F1PLC64", 65)])
    # TST-00002: fibre 3 to the tap pigtail, fibres 1,2,4,5,6 continue
    c2 = [{"a": cable_side(f"{CAB}: Seg1", 3), "b": cable_side(f"{PIG}: Seg1", 1), "usage": usage(100),
           "start": term(100), "end": f"Splice Case: {HUB}"}]
    for fib, u in ((1, f"{OLM} OLM TRANSPORT RETURN"), (2, f"{OLM} OLM TRANSPORT FORWARD"), (4, usage(200)),
                   (5, usage(300)), (6, usage(400))):
        c2.append({"a": cable_side(f"{CAB}: Seg1", fib), "b": cable_side(f"{CAB}: Seg2", fib), "usage": u})
    c2 += tst2_extra or []
    if "TST-00002" not in skip:
        write_report(folder / "TST-00002.xlsx", "TST-00002", [
            (f"{CAB}: Seg1", 12, f"Splice Case: {HUB}", "Splice Case: TST-00002", "Underground"),
            (f"{CAB}: Seg2", 12, "Splice Case: TST-00002", "Splice Case: TST-00003", "Underground"),
            (f"{PIG}: Seg1", 4, "Splice Case: TST-00002", "MAX10 Enclosure(M Bulk): TST-R-100 MAIN ST", "Pigtail"),
            *tst2_segments_extra,
        ], c2)
    # TST-00004 (TST-00003 has no report): fibre 4 continues with the same circuit, fibre 5 broken from TST-00003
    c4 = [{"a": cable_side(f"{CAB}: Seg3", 4), "b": cable_side(f"{CAB}: Seg4", 4), "usage": usage(200)},
          {"a": cable_side(f"{CAB}: Seg3", 5), "b": cable_side(f"{CAB}: Seg4", 5), "state": "Broken",
           "start": "Splice Case: TST-00003", "end": "Splice Case: TST-00010"}]
    c4 += tst4_extra or []
    if "TST-00004" not in skip:
        write_report(folder / "TST-00004.xlsx", "TST-00004", [
            (f"{CAB}: Seg3", 12, "Splice Case: TST-00003", "Splice Case: TST-00004", "Underground"),
            (f"{CAB}: Seg4", 12, "Splice Case: TST-00004", "Splice Case: TST-00010", "Underground"),
        ], c4)
    return folder
