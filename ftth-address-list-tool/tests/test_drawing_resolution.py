"""Drawing decoder primitives, design model and Fiber Splice from the drawing (synthetic data).

Synthetic drawing (feet), matching the fixture network in fixtures.py:

    TST-00001G (0,0) -- TST-00002 (300,0) -- TST-00003 (600,0) -- TST-00004 (900,0) -- TST-00010 (1200,0)
                                                  |
                                             TST-00020 (600,300)   branch cable, no report

  tap enclosures (joined by pigtails): TST-00002, TST-00003, TST-00010, TST-00020
  premises with drop symbols: 100 -> TST-00002, 300 -> TST-00003, 200 -> TST-00010, 400 -> TST-00020
"""

import csv
import itertools
import math
import random
import struct

import openpyxl
import pytest

from ftth_addrlist.builder import Builder
from ftth_addrlist.config import load_config
from ftth_addrlist.design import load_design, split_address, street_core
from ftth_addrlist.drawing_splice import assign, hungarian, learn_fiber_order, resolve_from_drawing
from ftth_addrlist.dwg_objects import Bits, handle_map
from ftth_addrlist.model import APPROVED, PROVISIONAL, FieldValue, OutputRow
from ftth_addrlist.network import Network
from ftth_addrlist.pipeline import run
from ftth_addrlist.splice_reports import load_reports

from fixtures import CAB, HUB, OLM, build_network

ezdxf = pytest.importorskip("ezdxf")

CLOSURES = {HUB: (0, 0), "TST-00002": (300, 0), "TST-00003": (600, 0), "TST-00004": (900, 0),
            "TST-00010": (1200, 0), "TST-00020": (600, 300)}
DEFAULT_TAPS = {"TST-00002": ((300, -20), (300, -150), 100), "TST-00003": ((600, -20), (600, -150), 300),
                "TST-00010": ((1200, -20), (1200, -150), 200), "TST-00020": ((620, 300), (750, 300), 400)}


def write_design(path, taps=None, extra_closures=None, branches=((600, 0), (600, 300)), more_branches=(),
                 drawn=(), markers=True, addresses=()):
    """taps: closure -> (enclosure xy, premise xy, house number); drawn: (ped xy, premise xy) drop lines
    as drawn by FX_DROPS (ped box edge -> PED 6.6 ft from the address point); addresses: (xy, number, street)."""
    taps = DEFAULT_TAPS if taps is None else taps
    closures = {**CLOSURES, **(extra_closures or {})}
    doc = ezdxf.new()
    msp = doc.modelspace()
    for name, attrs in (("MC_SPLICE.dwg", ["NAME"]), ("CX1_MDU_ENCLOSURE", ["NAME"]),
                        ("FMRESADD2", ["STREET_NUMBER", "STREET_NAME"]), ("CX1-DROP", []), ("CX1-MAX10", [])):
        blk = doc.blocks.new(name)
        for i, tag in enumerate(attrs):
            blk.add_attdef(tag, (0, i))
    for name, xy in closures.items():
        msp.add_blockref("MC_SPLICE.dwg", xy, dxfattribs={"layer": "Fiber Splice"}).add_auto_attribs({"NAME": name})
    main = [CLOSURES[c] for c in (HUB, "TST-00002", "TST-00003", "TST-00004", "TST-00010")]
    for pts in (main, list(branches), *more_branches):
        msp.add_lwpolyline(pts, dxfattribs={"layer": "Fiber Cable - Underground"})
    for closure, (enc, premise, num) in taps.items():
        msp.add_blockref("CX1_MDU_ENCLOSURE", enc, dxfattribs={"layer": "Fiber Node"}).add_auto_attribs({"NAME": ""})
        msp.add_lwpolyline([closures[closure], enc], dxfattribs={"layer": "Fiber Cable - Pigtail"})
        msp.add_blockref("FMRESADD2", premise).add_auto_attribs({"STREET_NUMBER": str(num), "STREET_NAME": "MAIN"})
        drop = (premise[0] + 0.42 * (enc[0] - premise[0]), premise[1] + 0.42 * (enc[1] - premise[1]))
        rot = math.degrees(math.atan2(premise[1] - enc[1], premise[0] - enc[0]))   # points to the premise
        if markers:
            msp.add_blockref("CX1-DROP", drop, dxfattribs={"layer": "RTE_UGDROP", "rotation": rot})
    for xy, num, street in addresses:
        msp.add_blockref("FMRESADD2", xy).add_auto_attribs({"STREET_NUMBER": num, "STREET_NAME": street})
    for ped, premise in drawn:
        msp.add_blockref("CX1-MAX10", ped, dxfattribs={"layer": "RTE_UGPOLES"})
        ux, uy = premise[0] - ped[0], premise[1] - ped[1]
        n = math.hypot(ux, uy)
        start = (ped[0] + 6 * ux / n, ped[1] + 6 * uy / n)
        end = (premise[0] - 6.6 * ux / n, premise[1] - 6.6 * uy / n)
        msp.add_lwpolyline([start, end], dxfattribs={"layer": "RTE_SpatialNet_UGDROP"})
    doc.saveas(path)
    return path


def built(folder, cfg):
    reps, _ = load_reports(sorted(folder.glob("*.xlsx")), cfg)
    net = Network(reps, cfg)
    return net, Builder(net, cfg).build(OLM)


def row(res, cable, fiber):
    return next(r for r in res.rows if r.cable == cable and r.fiber == fiber)


# ------------------------------------------------------------------ DWG decoder primitives
class BitWriter:
    def __init__(self):
        self.bits = []

    def put(self, value, n):
        self.bits += [(value >> (n - 1 - i)) & 1 for i in range(n)]
        return self

    def bytes_(self, data):
        for b in data:
            self.put(b, 8)
        return self

    def data(self):
        bits = self.bits + [0] * (-len(self.bits) % 8)
        return bytes(int("".join(map(str, bits[i:i + 8])), 2) for i in range(0, len(bits), 8))


def test_bit_codes():
    w = BitWriter()
    w.put(1, 2).bytes_([42])                    # BS code 1: one byte
    w.put(2, 2)                                 # BS code 2: zero
    w.put(0, 2).bytes_([0x34, 0x12])            # BS code 0: little-endian short
    w.put(1, 2)                                 # BD code 1: 1.0
    w.put(0, 2).bytes_(struct.pack("<d", 2.5))  # BD code 0: raw double
    w.bytes_([0x81, 0x01])                      # MC unsigned 129
    w.bytes_([0x45])                            # MC signed -5
    w.bytes_([0x05, 0x00])                      # MS 5
    w.bytes_([0x31, 0x7F])                      # handle: code 3, value 127
    b = Bits(w.data())
    assert (b.BS(), b.BS(), b.BS()) == (42, 0, 0x1234)
    assert (b.BD(), b.BD()) == (1.0, 2.5)
    assert b.MC(signed=False) == 129 and b.MC(signed=True) == -5 and b.MS() == 5
    assert b.H() == (3, 127)


def test_handle_map():
    payload = bytes([0x01, 0xE4, 0x00, 0x02, 0x32])       # (+1, +100), (+2, +50)
    section = (len(payload) + 2).to_bytes(2, "big") + payload + b"\0\0" + b"\x00\x02"
    assert handle_map(section) == {1: 100, 3: 150}


# ------------------------------------------------------------------ design model
def test_address_normalisation():
    assert split_address("C-2309 12TH AVE") == ("2309", "12TH")
    assert split_address("481 BLDG 2 ELM ST") == ("481", "ELM")
    assert split_address("12TH AVE AVE") == (None, None)
    assert street_core("HIGHWAY X9 X9") == "HIGHWAY X9" and street_core("S OAK ST") == "OAK"


def test_design_model_from_dxf(tmp_path):
    cfg = load_config()
    des, issues = load_design(write_design(tmp_path / "d.dxf"), cfg)
    assert des is not None and set(des.closures) == set(CLOSURES)
    assert set(des.enclosure_of) == set(DEFAULT_TAPS)
    assert [a["raw"] for a in des.address_points("400 MAIN ST")] == ["400 MAIN"]
    assert des.region({"TST-00003"}, {"TST-00002", "TST-00004"}) == {"TST-00003", "TST-00020"}
    assert des.drop_support((750, 300), des.enclosure_of["TST-00020"]["xy"]) is not None
    assert des.drop_support((750, 300), des.enclosure_of["TST-00003"]["xy"]) is None
    assert any(i.code == "DRAWING_MODEL" for i in issues)


# ------------------------------------------------------------------ resolver
def test_drawing_fills_and_corroborates(tmp_path):
    cfg = load_config()
    net, res = built(build_network(tmp_path / "r"), cfg)
    assert row(res, CAB, 6).fields["fiber_splice"].code == "FS-UNKNOWN"
    des, _ = load_design(write_design(tmp_path / "d.dxf"), cfg)
    summary = resolve_from_drawing([res], net, des, cfg)
    fs = row(res, CAB, 6).fields["fiber_splice"]
    assert fs.value == "TST-00020" and fs.code == "FS-DRAWING" and fs.status == PROVISIONAL
    assert any("drop symbol" in e for e in fs.evidence) and "TST-00020" in row(res, CAB, 6).needs_reports
    codes = {i.code for i in res.issues if i.fiber == 6}
    assert "FIBER_SPLICE_FROM_DRAWING" in codes and "SPLICE_CLOSURE_UNKNOWN" not in codes
    for fib, value in ((4, "TST-00010"), (5, "TST-00003")):     # report-derived values kept and corroborated
        f = row(res, CAB, fib).fields["fiber_splice"]
        assert f.value == value and f.code != "FS-DRAWING" and "corroborated" in f.note
    assert row(res, CAB, 3).fields["fiber_splice"].code == "FS-TRACE-TERMINAL"
    assert "located for 1 row" in summary[0].message


def test_branch_beyond_a_closure_without_tap_enclosure(tmp_path):
    # TST-00003 (the break point of fibre 5) has no tap enclosure; premise 300 hangs off a second
    # branch closure TST-00030: the fibre continues on that branch.
    cfg = load_config()
    taps = {k: v for k, v in DEFAULT_TAPS.items() if k != "TST-00003"}
    taps["TST-00030"] = ((620, -300), (750, -300), 300)
    path = write_design(tmp_path / "d.dxf", taps, {"TST-00030": (600, -300)},
                        more_branches=([(600, 0), (600, -300)],))
    net, res = built(build_network(tmp_path / "r"), cfg)
    des, _ = load_design(path, cfg)
    resolve_from_drawing([res], net, des, cfg)
    f5 = row(res, CAB, 5).fields["fiber_splice"]
    assert f5.value == "TST-00030" and f5.code == "FS-DRAWING" and "no tap enclosure" in f5.rule
    assert row(res, CAB, 6).value("fiber_splice") == "TST-00020"


def test_no_candidate_stays_unresolved(tmp_path):
    cfg = load_config()
    taps = {k: v for k, v in DEFAULT_TAPS.items() if k not in ("TST-00003", "TST-00020")}
    net, res = built(build_network(tmp_path / "r"), cfg)
    des, _ = load_design(write_design(tmp_path / "d.dxf", taps), cfg)
    resolve_from_drawing([res], net, des, cfg)
    r = row(res, CAB, 6)
    assert r.value("fiber_splice") is None and r.status == "UNRESOLVED"
    assert any(i.code == "DRAWING_NO_CANDIDATE" for i in r.issues)


def test_drawing_rule_policy(tmp_path):
    cfg = load_config()
    cfg["approvals"]["accepted_rules"] = ["FS-DRAWING"]
    net, res = built(build_network(tmp_path / "r"), cfg)
    des, _ = load_design(write_design(tmp_path / "d.dxf"), cfg)
    resolve_from_drawing([res], net, des, cfg)
    assert row(res, CAB, 6).fields["fiber_splice"].status == APPROVED


def test_pipeline_with_drawing(tmp_path):
    reports = build_network(tmp_path / "r")
    dxf = write_design(tmp_path / "design.dxf")
    out = tmp_path / "out"
    oc = run(reports, out, drawing=dxf, log=lambda *_: None)[0]
    ws = openpyxl.load_workbook(oc.workbook)[OLM]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    hit = next(r for r in rows if r["CABLE NAME"] == CAB and r["FIBER #"] == 6)
    assert hit["Fiber Splice"] == "TST-00020"
    with open(oc.exceptions_csv, encoding="utf-8-sig") as fh:
        codes = {r["code"] for r in csv.DictReader(fh)}
    assert "FIBER_SPLICE_FROM_DRAWING" in codes and "DRAWING_FIBER_SPLICE" in codes


# ------------------------------------------------------------------ assignment and fibre order
def test_hungarian_matches_brute_force():
    rnd = random.Random(7)
    for _ in range(150):
        n = rnd.randint(1, 5)
        m = rnd.randint(n, 7)
        cost = [[rnd.randint(0, 50) for _ in range(m)] for _ in range(n)]
        sol = hungarian(cost)
        best = min(sum(cost[i][p[i]] for i in range(n)) for p in itertools.permutations(range(m), n))
        assert sum(cost[i][sol[i]] for i in range(n)) == best and len(set(sol)) == n


def test_assignment_uses_order_penalty_to_break_ties():
    cost = {(d, c): 10.0 for d in "AB" for c in "XY"}
    # penalty: lower fibre (A) must not take the nearer closure X
    pen = lambda a: 1 if a["A"] == "X" and a["B"] == "Y" else 0
    assert assign(["A", "B"], ["X", "Y"], cost, 200.0, pen, 80.0) == {"A": "Y", "B": "X"}
    cost[("A", "X")] = 0.0                           # without the convention the cheaper pairing wins
    assert assign(["A", "B"], ["X", "Y"], cost, 200.0, pen, 0.0) == {"A": "X", "B": "Y"}


class _Res:
    def __init__(self, rows):
        self.rows = rows


def _traced(fiber, closure, route):
    r = OutputRow(OLM, 1, CAB, fiber, {"fiber_splice": FieldValue(closure, code="FS-TRACE-TERMINAL"),
                                       "tap_name": FieldValue(f"TAP-{closure}")}, kind="subscriber")
    r.route, r.trace_end = route, {"stop": "terminal"}
    return r


def test_learn_fiber_order():
    route = [HUB] + [f"C{i}" for i in range(1, 13)]
    rows = [_traced(f, route[13 - f], route[:14 - f]) for f in range(1, 13)]   # fibre 1 farthest
    direction, farther, nearer = learn_fiber_order([_Res(rows)])
    assert direction == "lower_farther" and farther == 66 and nearer == 0
    assert learn_fiber_order([_Res(rows[:3])])[0] == "none"                    # too few pairs to decide


# ------------------------------------------------------------------ drawn drop lines (FX_DROPS)
def test_drawn_drop_decides_without_drop_markers(tmp_path):
    # no drop markers at all; the drawn drop of 400 MAIN ST runs to a ped next to TST-00020
    cfg = load_config()
    path = write_design(tmp_path / "d.dxf", markers=False, drawn=[((640, 330), (750, 300))])
    net, res = built(build_network(tmp_path / "r"), cfg)
    des, issues = load_design(path, cfg)
    assert len(des.drawn_drops) == 1 and des.drawn_drops[0]["closure"] == "TST-00020"
    summary = resolve_from_drawing([res], net, des, cfg)
    fs = row(res, CAB, 6).fields["fiber_splice"]
    assert fs.value == "TST-00020" and fs.code == "FS-DROP" and fs.status == PROVISIONAL
    assert any("drawn drop line" in e for e in fs.evidence)
    assert "1 located value(s) taken from them" in summary[0].message


def test_drawn_drop_beats_cable_end_inference(tmp_path):
    # fibre 4 (200 MAIN ST) is inferred at the cable end TST-00010; the checked drop goes to a branch closure
    cfg = load_config()
    taps = dict(DEFAULT_TAPS)
    taps["TST-00040"] = ((1220, 300), (1350, 300), 201)
    path = write_design(tmp_path / "d.dxf", taps, {"TST-00040": (1200, 300)},
                        more_branches=([(1200, 0), (1200, 300)],), drawn=[((1240, 330), (1200, -150))])
    net, res = built(build_network(tmp_path / "r"), cfg)
    des, _ = load_design(path, cfg)
    resolve_from_drawing([res], net, des, cfg)
    fs = row(res, CAB, 4).fields["fiber_splice"]
    assert fs.value == "TST-00040" and fs.code == "FS-DROP" and "TST-00010" in fs.note
    cfg["drawing"]["drawn_drop_over_report_inference"] = False
    net, res = built(build_network(tmp_path / "r2"), cfg)
    resolve_from_drawing([res], net, des, cfg)
    assert row(res, CAB, 4).value("fiber_splice") == "TST-00010"
    assert any(i.code == "DRAWING_DISAGREES" for i in row(res, CAB, 4).issues)


def test_drawn_drop_contradicting_reports_is_flagged(tmp_path):
    # 100 MAIN ST is confirmed at TST-00002 by the reports; a drawn drop says TST-00003
    cfg = load_config()
    path = write_design(tmp_path / "d.dxf", drawn=[((630, -40), (300, -150))])
    net, res = built(build_network(tmp_path / "r"), cfg)
    des, _ = load_design(path, cfg)
    resolve_from_drawing([res], net, des, cfg)
    r = row(res, CAB, 3)
    assert r.value("fiber_splice") == "TST-00002" and r.fields["fiber_splice"].status == PROVISIONAL
    assert any(i.code == "DRAWN_DROP_DISAGREES" for i in r.issues)


def test_reported_closure_carrying_the_circuit_does_not_bound_the_search(tmp_path):
    # fibre 6 is lost at TST-00003; on the branch, TST-00020 has a report showing the same circuit on another
    # cable, and the tap is one closure further (TST-00021)
    cfg = load_config()
    br1, br2 = "TST-00003 TO TST-00021 FC012", "TST-00020 TO TST-00021 FC012"
    d = build_network(tmp_path / "r")
    from fixtures import cable_side, usage, write_report
    write_report(d / "TST-00020.xlsx", "TST-00020", [
        (f"{br1}: Seg1", 12, "Splice Case: TST-00003", "Splice Case: TST-00020", "Underground"),
        (f"{br2}: Seg1", 12, "Splice Case: TST-00020", "Splice Case: TST-00021", "Underground")],
        [{"a": cable_side(f"{br1}: Seg1", 1), "b": cable_side(f"{br2}: Seg1", 1), "usage": usage(400),
          "start": f"Splice Case: {HUB}", "end": "MAX10 Enclosure(M Bulk): TST-R-400 MAIN ST"}])
    taps = {k: v for k, v in DEFAULT_TAPS.items() if k != "TST-00020"}
    taps["TST-00021"] = ((620, 600), (750, 600), 400)
    path = write_design(tmp_path / "d.dxf", taps, {"TST-00021": (600, 600)}, branches=((600, 0), (600, 300), (600, 600)))
    net, res = built(d, cfg)
    des, _ = load_design(path, cfg)
    resolve_from_drawing([res], net, des, cfg)
    assert row(res, CAB, 6).value("fiber_splice") == "TST-00021"


def test_address_points_keep_directions_and_placeholders(tmp_path):
    cfg = load_config()
    path = write_design(tmp_path / "d.dxf", addresses=[((0, 900), "211", "OAK"), ((50, 900), "211", "S OAK"),
                                                       ((90, 900), "12TH", "AVE")])
    des, _ = load_design(path, cfg)
    assert [a["xy"] for a in des.address_points("211 OAK ST")] == [(0, 900)]
    assert [a["xy"] for a in des.address_points("211 S OAK ST")] == [(50, 900)]
    assert [a["xy"] for a in des.address_points("12TH AVE AVE")] == [(90, 900)]   # no house number
