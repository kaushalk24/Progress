#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FTTx Fiber Callout Studio
=========================

Generates the DES_FIBER_SPL fiber-number callouts of an FTTx OSP design from
its Splice Connections Reports and Address Lists, replicating the drafting
standard reverse-engineered from the Credit River (Mediacom NMK) combined DWG.

Reverse-engineered rules (see docs/PHASE1_REVERSE_ENGINEERING.md for proof):

  Fiber number   N = (PON NAME - 1) * PORTS_PER_PON + PON PORT
                 (numbered independently per OLM; PORTS_PER_PON = 64 for 1x64 PLC couplers)

  Terminal ("Fiber Number") callout
                 One MULTILEADER per terminal (Address List "Fiber Tap Name") and per
                 contiguous run of its fiber numbers. Text "N" or "Nmin-Nmax".
                 Arrowhead on the boundary (corner/edge) of the handhole block
                 (CX1-MAX5 / CX1-MAX10 ...) that hosts the drop splice closure
                 (Address List "Fiber Splice").

  Cable ("BFO") callout
                 One MULTILEADER per live distribution-cable segment leaving a closure
                 that has a splice report. Text "BFO <segment fiber count>\\P<Nmin>-<Nmax>"
                 where N runs over the live (non-broken, non-spare) fibers of the segment.
                 Arrowhead on the cable polyline close to the upstream closure.

  Entity spec    MULTILEADER (AcDbMLeader), layer DES_FIBER_SPL (ACI 6, Continuous),
                 entity color/linetype/lineweight ByLayer, MLEADERSTYLE "Standard",
                 MText content, text style "Standard" (txt.shx, width 1.0, oblique 0),
                 text height 4.0, rotation 0, straight leader, closed-filled arrow 4.0,
                 landing on, dogleg 0.36, landing gap 0.09, left/right attachment 6
                 (bottom of top line + underline), top/bottom attachment 9,
                 MText attachment top-left (dogleg +X) / top-right (dogleg -X),
                 text location = landing + dogleg*(0.36+0.09) + (0, 1.2*height).

Usage
-----
  GUI : python fiber_callout_studio.py
  CLI : python fiber_callout_studio.py --cli --drawing design.dxf --splice ./reports \
            --address ./address_lists --out ./output [--validate-only]

Requirements: Python 3.9+, ezdxf>=1.1, pandas, openpyxl, PyQt6 (GUI only).
DWG input/output additionally requires the free ODA File Converter.
"""
from __future__ import annotations

import argparse
import dataclasses
import datetime as _dt
import difflib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import threading
import traceback
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Set, Tuple

try:
    import ezdxf
    from ezdxf import bbox as ezbbox
    from ezdxf import colors as ezcolors
    from ezdxf import path as ezpath
    from ezdxf.math import Vec2, Vec3
    from ezdxf.render import mleader as mlr
except ImportError as _exc:  # pragma: no cover - dependency guard
    sys.stderr.write("Missing dependency 'ezdxf'. Install with: pip install -r requirements.txt\n")
    raise SystemExit(2) from _exc
try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as _exc:  # pragma: no cover - dependency guard
    sys.stderr.write("Missing dependency 'openpyxl'. Install with: pip install -r requirements.txt\n")
    raise SystemExit(2) from _exc
try:
    import pandas as pd
except ImportError as _exc:  # pragma: no cover - dependency guard
    sys.stderr.write("Missing dependency 'pandas'. Install with: pip install -r requirements.txt\n")
    raise SystemExit(2) from _exc

APP_NAME = "FTTx Fiber Callout Studio"
APP_VERSION = "1.0.0"


# =============================================================================
# Reverse-engineered drafting specification
# =============================================================================
@dataclass
class CalloutSpec:
    """Every value here was measured on the reference drawing (1,184 callouts)."""

    cable_layer: str = "DES_FIBER_SPL"           # BFO (cable) callouts
    terminal_layer: str = "DES_FIBER_SPL"        # fiber-number (terminal) callouts
    layer_color: int = 6                         # ACI 6 (magenta)
    layer_linetype: str = "Continuous"
    mleader_style: str = "Standard"
    text_style: str = "Standard"                 # font txt.shx, width factor 1.0, oblique 0
    text_style_font: str = "txt"
    text_height: float = 4.0
    arrow_size: float = 4.0
    landing_gap: float = 0.09
    dogleg_length: float = 0.36
    side_attachment: int = 6                     # bottom of top line + underline
    vertical_attachment: int = 9                 # center
    text_y_factor: float = 1.2                   # MText top = landing.y + 1.2 * height
    cable_prefix: str = "BFO"
    terminal_leader_length: float = 17.0         # median of reference terminal callouts
    cable_leader_length: float = 30.0            # median of reference BFO callouts
    cable_offset_along: float = 30.0             # arrow distance along cable from closure
    handhole_search_radius: float = 75.0
    zero_pad: int = 0                            # 0 = "3", 2 = "03"
    cable_range_mode: str = "span"               # "span" -> min-max (reference); "exact" -> runs
    split_terminal_runs: bool = True             # non-contiguous terminal -> one callout per run

    def to_dict(self) -> Dict[str, Any]:
        return dataclasses.asdict(self)


@dataclass
class ColumnMap:
    """Logical address-list fields -> actual header text (None = not mapped)."""

    pon: Optional[str] = None
    port: Optional[str] = None
    address: Optional[str] = None
    tap: Optional[str] = None
    splice: Optional[str] = None
    fiber: Optional[str] = None
    cable: Optional[str] = None
    olm: Optional[str] = None

    def to_dict(self) -> Dict[str, Optional[str]]:
        return dataclasses.asdict(self)


@dataclass
class CadMap:
    closure_layer: str = "Fiber Splice"
    closure_block: str = ""                      # "" = any block on the closure layer
    closure_name_tag: str = "NAME"
    handhole_layer: str = "RTE_UGPOLES"
    handhole_blocks: List[str] = field(default_factory=lambda: ["CX1-MAX5", "CX1-MAX10"])
    terminal_layer: str = "Fiber Node"
    cable_layers: List[str] = field(default_factory=lambda: [
        "Fiber Cable - Underground", "Fiber Cable - Aerial"])
    snap_tolerance: float = 2.0

    def to_dict(self) -> Dict[str, Any]:
        return dataclasses.asdict(self)


@dataclass
class RunConfig:
    drawing_path: str = ""
    splice_paths: List[str] = field(default_factory=list)
    address_paths: List[str] = field(default_factory=list)
    output_dir: str = ""
    columns: ColumnMap = field(default_factory=ColumnMap)
    cad: CadMap = field(default_factory=CadMap)
    spec: CalloutSpec = field(default_factory=CalloutSpec)
    ports_per_pon: int = 0                       # 0 = auto (next power of two >= max port)
    existing_mode: str = "replace"               # replace | skip | append
    export_dwg: bool = False
    dwg_version: str = "R2018"
    oda_path: str = ""

    def to_dict(self) -> Dict[str, Any]:
        d = dataclasses.asdict(self)
        return d


ADDRESS_FIELDS: Dict[str, Tuple[bool, List[str], str]] = {
    # key: (required, synonyms, description)
    "pon": (True, ["PON NAME", "PON", "PON #", "PON NO", "PON NUMBER", "SPLITTER", "SPLITTER #",
                   "SPLITTER NAME", "COUPLER", "PLC"], "PON / splitter index"),
    "port": (True, ["PON PORT", "PORT", "PORT #", "SPLITTER PORT", "OUTPUT PORT", "PON PORT #",
                    "COUPLER PORT"], "Port on that PON"),
    "address": (True, ["ADDRESS", "SERVICE ADDRESS", "SERVICE LOCATION", "LOCATION", "FULL ADDRESS",
                       "PREMISE", "PREMISE ADDRESS"], "Served address"),
    "tap": (True, ["FIBER TAP NAME", "TAP NAME", "TAP", "TERMINAL", "TERMINAL NAME", "MST", "NAP",
                   "DROP TERMINAL", "FIBER TAP"], "Terminal / tap name"),
    "splice": (True, ["FIBER SPLICE", "SPLICE", "SPLICE NAME", "CLOSURE", "SPLICE CLOSURE",
                      "SPLICE CASE", "DROP SPLICE", "ENCLOSURE"], "Drop splice closure"),
    "fiber": (False, ["FIBER #", "FIBER", "FIBER NO", "FIBER NUMBER", "STRAND", "FIBER NUM"],
              "Fiber position in first cable"),
    "cable": (False, ["CABLE NAME", "CABLE", "CABLE ID"], "Feeder/distribution cable"),
    "olm": (False, ["OLM", "OLT", "HUB", "CABINET", "FDH", "OLM NAME", "NODE"], "OLM / hub id"),
}


# =============================================================================
# Infrastructure: reporter, issues, helpers
# =============================================================================
class Cancelled(Exception):
    """Raised when the user cancels a running job."""


class Reporter:
    """Thread-safe log/progress sink. GUI passes Qt-signal emitters; CLI prints."""

    def __init__(self, log_fn: Optional[Callable[[str, str], None]] = None,
                 progress_fn: Optional[Callable[[int, str], None]] = None,
                 cancel_event: Optional[threading.Event] = None) -> None:
        self._log_fn = log_fn
        self._progress_fn = progress_fn
        self._cancel = cancel_event or threading.Event()
        self.lines: List[Tuple[str, str, str]] = []

    def log(self, msg: str, level: str = "INFO") -> None:
        stamp = _dt.datetime.now().strftime("%H:%M:%S")
        self.lines.append((stamp, level, msg))
        if self._log_fn:
            self._log_fn(msg, level)
        else:
            print(f"[{stamp}] {level:<7} {msg}", flush=True)

    def progress(self, pct: float, msg: str = "") -> None:
        if self._progress_fn:
            self._progress_fn(int(max(0, min(100, pct))), msg)

    def check_cancel(self) -> None:
        if self._cancel.is_set():
            raise Cancelled("Cancelled by user")


@dataclass
class Issue:
    severity: str          # ERROR | WARNING | INFO
    category: str
    obj: str
    message: str


class IssueLog:
    def __init__(self) -> None:
        self.items: List[Issue] = []

    def add(self, severity: str, category: str, obj: str, message: str) -> None:
        self.items.append(Issue(severity, category, str(obj), message))

    def count(self, severity: str) -> int:
        return sum(1 for i in self.items if i.severity == severity)

    def summary(self) -> str:
        return f"{self.count('ERROR')} error(s), {self.count('WARNING')} warning(s), {self.count('INFO')} info"


def norm_text(value: Any) -> str:
    """Upper-case, whitespace-collapsed comparison key."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    s = str(value).replace(" ", " ").strip().upper()
    s = re.sub(r"\s+", " ", s)
    return s.rstrip(".")


def norm_header(value: Any) -> str:
    return re.sub(r"[^A-Z0-9#]", "", norm_text(value))


def norm_cable(value: Any) -> str:
    s = norm_text(value)
    s = re.sub(r"\s*:\s*", ":", s)
    return s


def to_int(value: Any) -> Optional[int]:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int,)):
        return int(value)
    if isinstance(value, float):
        if math.isnan(value):
            return None
        return int(round(value)) if abs(value - round(value)) < 1e-6 else None
    text = str(value).strip().replace(",", "")
    m = re.match(r"^(\d+)$", text)
    if m:
        return int(m.group(1))
    try:
        f = float(text)
    except ValueError:
        return None
    if math.isnan(f) or abs(f - round(f)) > 1e-6:
        return None
    return int(round(f))


def number_runs(nums: Iterable[int]) -> List[Tuple[int, int]]:
    ordered = sorted(set(nums))
    runs: List[Tuple[int, int]] = []
    for n in ordered:
        if runs and n == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], n)
        else:
            runs.append((n, n))
    return runs


def fmt_run(run: Tuple[int, int], pad: int) -> str:
    a, b = run
    fa = str(a).zfill(pad) if pad else str(a)
    fb = str(b).zfill(pad) if pad else str(b)
    return fa if a == b else f"{fa}-{fb}"


def fmt_numbers(nums: Iterable[int], mode: str, pad: int) -> str:
    runs = number_runs(nums)
    if not runs:
        return ""
    if mode == "span":
        return fmt_run((runs[0][0], runs[-1][1]), pad)
    return ",".join(fmt_run(r, pad) for r in runs)


def canonical_callout_text(text: str) -> str:
    """Comparison form: paragraph breaks unified, numbers stripped of leading zeros."""
    t = text.replace("\\P", "\n").replace("\r", "")
    t = re.sub(r"\\[A-Za-z][^;\\]*;", "", t)
    t = re.sub(r"[{}]", "", t)
    t = re.sub(r"\d+", lambda m: str(int(m.group(0))), t)
    return "|".join(norm_text(x) for x in t.split("\n") if x.strip())


# =============================================================================
# Geometry helpers
# =============================================================================
def v2(p: Any) -> Vec2:
    return Vec2(float(p[0]), float(p[1]))


def seg_distance(p: Vec2, a: Vec2, b: Vec2) -> Tuple[float, float]:
    d = b - a
    L2 = d.x * d.x + d.y * d.y
    t = 0.0 if L2 == 0 else max(0.0, min(1.0, ((p.x - a.x) * d.x + (p.y - a.y) * d.y) / L2))
    q = a + d * t
    return (p - q).magnitude, t


def polyline_length(pts: Sequence[Vec2]) -> float:
    return sum((pts[i + 1] - pts[i]).magnitude for i in range(len(pts) - 1))


def point_at_length(pts: Sequence[Vec2], s: float) -> Tuple[Vec2, Vec2]:
    """Point and unit tangent at arc length s along the polyline."""
    acc = 0.0
    for i in range(len(pts) - 1):
        seg = pts[i + 1] - pts[i]
        L = seg.magnitude
        if L <= 1e-9:
            continue
        if acc + L >= s:
            t = (s - acc) / L
            return pts[i] + seg * t, seg / L
        acc += L
    last = pts[-1] - pts[-2] if len(pts) > 1 else Vec2(1, 0)
    return pts[-1], (last.normalize() if last.magnitude > 1e-9 else Vec2(1, 0))


Box = Tuple[float, float, float, float]  # xmin, ymin, xmax, ymax


def box_overlap(a: Box, b: Box) -> float:
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0.0


def box_of_points(pts: Iterable[Vec2], pad: float = 0.0) -> Box:
    xs = [p.x for p in pts]
    ys = [p.y for p in pts]
    return (min(xs) - pad, min(ys) - pad, max(xs) + pad, max(ys) + pad)


def segments_intersect(p1: Vec2, p2: Vec2, q1: Vec2, q2: Vec2) -> bool:
    def orient(a: Vec2, b: Vec2, c: Vec2) -> float:
        return (b.x - a.x) * (c.y - a.y) - (b.y - a.y) * (c.x - a.x)
    o1, o2 = orient(p1, p2, q1), orient(p1, p2, q2)
    o3, o4 = orient(q1, q2, p1), orient(q1, q2, p2)
    return (o1 * o2 < 0) and (o3 * o4 < 0)


def segment_hits_box(a: Vec2, b: Vec2, box: Box) -> bool:
    """True if segment a-b touches the axis-aligned box."""
    def inside(p: Vec2) -> bool:
        return box[0] <= p.x <= box[2] and box[1] <= p.y <= box[3]
    if inside(a) or inside(b):
        return True
    c = [Vec2(box[0], box[1]), Vec2(box[2], box[1]), Vec2(box[2], box[3]), Vec2(box[0], box[3])]
    return any(segments_intersect(a, b, c[i], c[(i + 1) % 4]) for i in range(4))


class GridIndex:
    """Uniform grid spatial index for boxes and points."""

    def __init__(self, cell: float = 50.0) -> None:
        self.cell = cell
        self.cells: Dict[Tuple[int, int], List[Tuple[Box, Any]]] = defaultdict(list)

    def _keys(self, b: Box) -> Iterable[Tuple[int, int]]:
        c = self.cell
        for ix in range(int(math.floor(b[0] / c)), int(math.floor(b[2] / c)) + 1):
            for iy in range(int(math.floor(b[1] / c)), int(math.floor(b[3] / c)) + 1):
                yield ix, iy

    def insert(self, b: Box, item: Any) -> None:
        for k in self._keys(b):
            self.cells[k].append((b, item))

    def query(self, b: Box) -> List[Tuple[Box, Any]]:
        seen: Set[int] = set()
        out: List[Tuple[Box, Any]] = []
        for k in self._keys(b):
            for entry in self.cells.get(k, ()):
                if id(entry) not in seen:
                    seen.add(id(entry))
                    out.append(entry)
        return out


# =============================================================================
# Input collection & table reading
# =============================================================================
class InputCollector:
    """Expands files / folders / ZIP archives into a flat list of data files."""

    def __init__(self) -> None:
        self._tmp: List[str] = []

    def expand(self, paths: Sequence[str], exts: Sequence[str], issues: IssueLog) -> List[Path]:
        out: List[Path] = []
        seen: Set[str] = set()
        exts_l = tuple(e.lower() for e in exts)

        def add(p: Path) -> None:
            key = str(p.resolve())
            if key in seen or p.name.startswith("~$") or p.name.startswith("."):
                return
            if p.suffix.lower() in exts_l:
                seen.add(key)
                out.append(p)

        for raw in paths:
            p = Path(raw)
            if not p.exists():
                issues.add("ERROR", "Input", raw, "Path does not exist")
                continue
            if p.is_dir():
                for f in sorted(p.rglob("*")):
                    if f.is_file():
                        if f.suffix.lower() == ".zip":
                            for z in self._extract(f, issues):
                                for g in sorted(Path(z).rglob("*")):
                                    if g.is_file():
                                        add(g)
                        else:
                            add(f)
            elif p.suffix.lower() == ".zip":
                for z in self._extract(p, issues):
                    for g in sorted(Path(z).rglob("*")):
                        if g.is_file():
                            add(g)
            else:
                add(p)
        return out

    def _extract(self, zpath: Path, issues: IssueLog) -> List[str]:
        try:
            tmp = tempfile.mkdtemp(prefix="fcs_zip_")
            self._tmp.append(tmp)
            root = Path(tmp).resolve()
            with zipfile.ZipFile(zpath) as zf:
                for member in zf.infolist():
                    target = (root / member.filename).resolve()
                    if root not in target.parents and target != root:
                        issues.add("WARNING", "Input", str(zpath), f"Skipped unsafe ZIP member {member.filename}")
                        continue
                    zf.extract(member, root)
            return [tmp]
        except (zipfile.BadZipFile, OSError) as exc:
            issues.add("ERROR", "Input", str(zpath), f"Cannot open ZIP: {exc}")
            return []

    def cleanup(self) -> None:
        for t in self._tmp:
            shutil.rmtree(t, ignore_errors=True)
        self._tmp.clear()


def read_tables(path: Path) -> List[Tuple[str, List[List[Any]]]]:
    """Return [(sheet_name, rows)] for xlsx/xlsm/xls/csv/txt files."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        finally:
            wb.close()
    if suffix == ".xls":
        sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        return [(name, df.where(pd.notna(df), None).values.tolist()) for name, df in sheets.items()]
    if suffix in (".csv", ".txt"):
        import csv
        raw = path.read_bytes()
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("Cannot decode text file")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [[(c if c != "" else None) for c in r] for r in csv.reader(text.splitlines(), dialect)]
        return [(path.stem, rows)]
    raise ValueError(f"Unsupported file type {suffix}")


# =============================================================================
# Address lists
# =============================================================================
@dataclass
class AddressRecord:
    olm: str
    pon: int
    port: int
    number: int
    address: str
    tap: str
    splice: str
    fiber: Optional[int]
    cable: str
    source: str
    row: int


@dataclass
class Terminal:
    olm: str
    tap: str
    splice: str
    numbers: List[int] = field(default_factory=list)
    records: List[AddressRecord] = field(default_factory=list)


def _match_header(headers: List[str], synonyms: List[str]) -> Optional[str]:
    normed = {norm_header(h): h for h in headers if h}
    syn_n = [norm_header(s) for s in synonyms]
    for s in syn_n:
        if s in normed:
            return normed[s]
    for s in syn_n:
        for hn, h in normed.items():
            if s and (hn.startswith(s) or s in hn) and len(s) >= 3:
                return h
    best, best_r = None, 0.0
    for s in syn_n:
        for hn, h in normed.items():
            r = difflib.SequenceMatcher(None, s, hn).ratio()
            if r > best_r:
                best, best_r = h, r
    return best if best_r >= 0.82 else None


def detect_header_row(rows: List[List[Any]]) -> Tuple[int, List[str]]:
    all_syn = [s for (_, syns, _) in ADDRESS_FIELDS.values() for s in syns]
    syn_n = {norm_header(s) for s in all_syn}
    best_i, best_score = 0, -1
    for i, r in enumerate(rows[:30]):
        score = sum(1 for c in r if c is not None and norm_header(c) in syn_n)
        if score > best_score:
            best_i, best_score = i, score
    headers = [str(c).strip() if c is not None else "" for c in rows[best_i]] if rows else []
    return best_i, headers


def auto_map_columns(headers: List[str]) -> ColumnMap:
    cm = ColumnMap()
    used: Set[str] = set()
    for key in ("pon", "port", "address", "tap", "splice", "fiber", "cable", "olm"):
        _, syns, _ = ADDRESS_FIELDS[key]
        h = _match_header([x for x in headers if x not in used], syns)
        if key == "olm" and h is not None and norm_header(h) in {norm_header(s) for s in ADDRESS_FIELDS["splice"][1]}:
            h = None
        if h:
            used.add(h)
        setattr(cm, key, h)
    return cm


OLM_TOKEN = re.compile(r"(?<![A-Z0-9])([A-Z]{2,}-[A-Z]{0,3}\d{1,4}[A-Z]?)(?![A-Z0-9])")
GENERIC_SHEETS = {"SHEET", "SHEET1", "SHEET2", "SHEET3", "DATA", "ADDRESS LIST", "ADDRESSES"}


def derive_olm(sheet: str, filename: str) -> str:
    """OLM id: a meaningful sheet name, else an OLM-like token (e.g. NMK-R02) in the file name."""
    s = norm_text(sheet)
    stem = norm_text(Path(filename).stem)
    if s and s not in GENERIC_SHEETS and s != stem:
        return s
    toks = OLM_TOKEN.findall(stem)
    return toks[-1] if toks else stem


class AddressBook:
    def __init__(self) -> None:
        self.records: List[AddressRecord] = []
        self.terminals: Dict[Tuple[str, str, str], Terminal] = {}
        self.by_tap: Dict[str, List[AddressRecord]] = defaultdict(list)
        self.by_tap_addr: Dict[Tuple[str, str], List[AddressRecord]] = defaultdict(list)
        self.by_addr: Dict[str, List[AddressRecord]] = defaultdict(list)
        self.by_number: Dict[Tuple[str, int], AddressRecord] = {}
        self.tap_names: Set[str] = set()
        self.headers: List[str] = []
        self.ports_per_pon = 64
        self.olms: List[str] = []

    @staticmethod
    def scan_headers(files: List[Path], issues: IssueLog) -> List[str]:
        headers: List[str] = []
        for f in files:
            try:
                for _sheet, rows in read_tables(f):
                    _, h = detect_header_row(rows)
                    for x in h:
                        if x and x not in headers:
                            headers.append(x)
            except Exception as exc:  # noqa: BLE001 - report every unreadable file
                issues.add("ERROR", "Address List", f.name, f"Unreadable: {exc}")
        return headers

    def load(self, files: List[Path], cmap: ColumnMap, ports_per_pon: int, issues: IssueLog,
             rep: Reporter) -> None:
        raw: List[Tuple[str, int, int, Dict[str, Any], str, int]] = []
        for f in files:
            rep.check_cancel()
            try:
                tables = read_tables(f)
            except Exception as exc:  # noqa: BLE001
                issues.add("ERROR", "Address List", f.name, f"Unreadable: {exc}")
                continue
            for sheet, rows in tables:
                if not rows:
                    continue
                hi, headers = detect_header_row(rows)
                idx = {h: i for i, h in enumerate(headers) if h}
                missing = [k for k, (req, _, _) in ADDRESS_FIELDS.items()
                           if req and (getattr(cmap, k) is None or getattr(cmap, k) not in idx)]
                if missing:
                    issues.add("ERROR", "Address List", f"{f.name}:{sheet}",
                               "Required column(s) not found/mapped: " + ", ".join(missing))
                    continue
                default_olm = derive_olm(sheet, f.name)
                for ri, r in enumerate(rows[hi + 1:], start=hi + 2):
                    def get(key: str) -> Any:
                        col = getattr(cmap, key)
                        if col is None or col not in idx or idx[col] >= len(r):
                            return None
                        return r[idx[col]]
                    pon, port = to_int(get("pon")), to_int(get("port"))
                    if pon is None and port is None:
                        continue
                    if pon is None or port is None or pon < 1 or port < 1:
                        issues.add("WARNING", "Address List", f"{f.name}:{sheet} row {ri}",
                                   f"Invalid PON/port ({get('pon')!r}, {get('port')!r}) - row skipped")
                        continue
                    olm = norm_text(get("olm")) or default_olm
                    raw.append((olm, pon, port, {k: get(k) for k in ADDRESS_FIELDS}, f"{f.name}:{sheet}", ri))
        if not raw:
            issues.add("ERROR", "Address List", "-", "No PON-assigned rows found in the address list(s)")
            return
        max_port = max(p for _, _, p, _, _, _ in raw)
        if ports_per_pon and ports_per_pon > 0:
            self.ports_per_pon = ports_per_pon
        else:
            self.ports_per_pon = next((s for s in (8, 16, 32, 64, 128, 256) if s >= max_port), max_port)
        if max_port > self.ports_per_pon:
            issues.add("ERROR", "Address List", "-",
                       f"PON PORT {max_port} exceeds ports-per-PON {self.ports_per_pon}")
        for olm, pon, port, vals, src, ri in raw:
            number = (pon - 1) * self.ports_per_pon + port
            rec = AddressRecord(olm=olm, pon=pon, port=port, number=number,
                                address=norm_text(vals["address"]), tap=norm_text(vals["tap"]),
                                splice=norm_text(vals["splice"]), fiber=to_int(vals["fiber"]),
                                cable=norm_text(vals["cable"]), source=src, row=ri)
            key = (olm, number)
            if key in self.by_number:
                issues.add("ERROR", "Address List", f"{src} row {ri}",
                           f"Duplicate fiber number {number} in {olm} (PON {pon} port {port})")
                continue
            self.by_number[key] = rec
            self.records.append(rec)
            if not rec.tap:
                issues.add("WARNING", "Address List", f"{src} row {ri}",
                           f"No terminal name for {rec.address}; address used as terminal key")
                rec.tap = rec.address
            if not rec.splice:
                issues.add("WARNING", "Address List", f"{src} row {ri}",
                           f"No drop splice closure for {rec.address} ({rec.tap})")
            other = [k for k in self.terminals if k[0] == olm and k[1] == rec.tap and k[2] != rec.splice]
            if other:
                issues.add("WARNING", "Splice Mismatch", rec.tap,
                           f"Terminal rows reference different drop closures ({other[0][2]} vs {rec.splice}); "
                           f"one callout is placed at each closure")
            t = self.terminals.setdefault((olm, rec.tap, rec.splice), Terminal(olm=olm, tap=rec.tap, splice=rec.splice))
            t.numbers.append(number)
            t.records.append(rec)
            self.by_tap_addr[(rec.tap, rec.address)].append(rec)
            self.by_addr[rec.address].append(rec)
            self.by_tap[rec.tap].append(rec)
            self.tap_names.add(rec.tap)
        self.olms = sorted({r.olm for r in self.records})
        for lst in self.by_tap.values():
            lst.sort(key=lambda r: (r.olm, r.number))
        rep.log(f"Address lists: {len(self.records)} PON-assigned rows, {len(self.terminals)} terminals, "
                f"OLMs {', '.join(self.olms)}, ports/PON {self.ports_per_pon}")

    @staticmethod
    def address_variants(address: str) -> List[str]:
        a = norm_text(address)
        out = [a]
        b = re.sub(r"^(?:[A-Z]{1,2}-)(?=\d)", "", a)                       # C-9030 ... -> 9030 ...
        c = re.sub(r"\s+(?:#\s*\d+[A-Z]?|BLDG\s*\d+[A-Z]?|UNIT\s*\w+|APT\s*\w+|STE\s*\w+)\b", "", b)
        for v in (b, c):
            v = re.sub(r"\s+", " ", v).strip()
            if v and v not in out:
                out.append(v)
        return out

    def resolve_candidates(self, tap: str, address: str, olm_hint: Optional[str],
                           port_idx: Optional[int] = None) -> List[AddressRecord]:
        """Resolve a splice-report fiber usage to address-list record(s).

        Order: exact (terminal, address) -> terminal port index (TAP:k = k-th fiber of the terminal)
        -> address only -> normalised address variants (C- prefix, unit/building tokens removed).
        Several candidates are returned only when the data is genuinely ambiguous."""
        tap_n, addr_n = norm_text(tap), norm_text(address)

        def scope(c: List[AddressRecord]) -> List[AddressRecord]:
            if olm_hint:
                in_olm = [x for x in c if x.olm == olm_hint]
                if in_olm:
                    return in_olm
            return c
        if tap_n:
            c = scope(self.by_tap_addr.get((tap_n, addr_n), []))
            if len(c) == 1:
                return c
            recs = scope(self.by_tap.get(tap_n, []))
            if recs and port_idx and 1 <= port_idx <= len(recs):
                olms = {r.olm for r in recs}
                if len(olms) == 1:
                    return [recs[port_idx - 1]]
            if len(c) > 1:
                return c
        for variant in self.address_variants(addr_n):
            c = scope(self.by_addr.get(variant, []))
            if len(c) == 1:
                return c
            if c:
                return c
        return []

    def resolve(self, tap: str, address: str, olm_hint: Optional[str], port_idx: Optional[int] = None) -> Optional[AddressRecord]:
        c = self.resolve_candidates(tap, address, olm_hint, port_idx)
        return c[0] if len(c) == 1 else None


# =============================================================================
# Splice reports (two vendor layouts supported)
# =============================================================================
@dataclass
class ReportSegment:
    cable: str
    seg: int
    count: Optional[int]
    kind: str
    start: Optional[str]
    end: Optional[str]
    direction: str              # out | in | thru


@dataclass
class ReportConn:
    a_seg: Optional[Tuple[str, int]]
    b_seg: Optional[Tuple[str, int]]
    usage: str
    state: str
    tap: str
    address: str
    port_idx: Optional[int] = None


@dataclass
class SpliceReport:
    name: str
    file: str
    fmt: str
    xy: Optional[Tuple[float, float]] = None
    segments: List[ReportSegment] = field(default_factory=list)
    conns: List[ReportConn] = field(default_factory=list)
    olm: Optional[str] = None


SEG_ROW_RE = re.compile(r"^\s*Segment\s+(\d+)\s+of\s+(.+?)\s+-\s+([A-Za-z][\w /]*)\s*$", re.I)
CONN_SEG_RE = re.compile(r"^(.+?)\s*:\s*Seg\s*(\d+)\s*$", re.I)
GUIDE_SEG_RE = re.compile(r"^(.+?)\s*:\s*Seg\s*(\d+)(?:\s*-\s*(\d+)\s*ct)?\s*$", re.I)
GUIDE_CABLE_RE = re.compile(r"^(.+?)\s+SEG\s*-?\s*(\d+)\s*$", re.I)
FC_RE = re.compile(r"\bFC\s*0*(\d+)\s*$", re.I)
USAGE_RE = re.compile(r"^(.*?):(\d+)\s+(.*)$")
DEAD_STATES = ("BROKEN", "SPARE", "DARK", "CUT", "DEAD")
NON_ADDRESS_RE = re.compile(r"TRANSPORT|SPARES?\b|POWER PORT|^NO DESC|COUPLER|\bOLM\b", re.I)


def _strip_kind(value: Any) -> Optional[str]:
    s = str(value or "").strip()
    if not s:
        return None
    return norm_text(s.split(":", 1)[1] if ":" in s else s)


def _fc_of(cable: str) -> Optional[int]:
    m = FC_RE.search(cable)
    return int(m.group(1)) if m else None


def _find_label_value(row: List[Any], label_idx: int) -> Any:
    for v in row[label_idx + 1:]:
        if v is not None and str(v).strip() != "":
            return v
    return None


def parse_connections_report(rows: List[List[Any]], fname: str) -> Optional[SpliceReport]:
    rep: Optional[SpliceReport] = None
    name = None
    xy = None
    seg_hdr = conn_hdr = None
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in r]
        low = [c.lower() for c in cells]
        if low and low[0].startswith("splice name"):
            name = norm_text(_find_label_value(r, 0))
        elif low and low[0].startswith("coordinates"):
            nums = [c for c in r[1:] if isinstance(c, (int, float))]
            if len(nums) >= 2:
                xy = (float(nums[0]), float(nums[1]))
        elif "segment" in low and "count" in low and "start splice" in low and seg_hdr is None:
            seg_hdr = (i, {k: low.index(k) for k in ("segment", "count", "start splice", "end splice")})
        elif "usage" in low and low.count("segment") >= 2 and conn_hdr is None:
            segs = [j for j, c in enumerate(low) if c == "segment"]
            hashes = [j for j, c in enumerate(low) if c == "#"]
            conn_hdr = (i, dict(a_seg=segs[0], b_seg=segs[1], usage=low.index("usage"),
                                state=low.index("state") if "state" in low else None,
                                a_pos=hashes[0] if hashes else None))
    if not name:
        return None
    rep = SpliceReport(name=name, file=fname, fmt="Splice Connections Report", xy=xy)
    if seg_hdr:
        i0, c = seg_hdr
        for r in rows[i0 + 1:]:
            if not r or r[c["segment"]] is None:
                break
            m = SEG_ROW_RE.match(str(r[c["segment"]]))
            if not m:
                continue
            cable = norm_cable(m.group(2))
            start = _strip_kind(r[c["start splice"]]) if c["start splice"] < len(r) else None
            end = _strip_kind(r[c["end splice"]]) if c["end splice"] < len(r) else None
            cnt = to_int(r[c["count"]]) if c["count"] < len(r) else None
            direction = "out" if start == name else ("in" if end == name else "thru")
            rep.segments.append(ReportSegment(cable=cable, seg=int(m.group(1)), count=cnt or _fc_of(cable),
                                              kind=m.group(3).strip().title(), start=start, end=end,
                                              direction=direction))
    if conn_hdr:
        i0, c = conn_hdr
        for r in rows[i0 + 1:]:
            if not any(v is not None and str(v).strip() for v in r):
                break

            def segkey(v: Any) -> Optional[Tuple[str, int]]:
                m = CONN_SEG_RE.match(str(v or "").strip())
                return (norm_cable(m.group(1)), int(m.group(2))) if m else None
            usage = str(r[c["usage"]] or "").strip() if c["usage"] < len(r) else ""
            state = str(r[c["state"]] or "").strip() if c["state"] is not None and c["state"] < len(r) else ""
            m = USAGE_RE.match(usage)
            tap, addr = (norm_text(m.group(1)), norm_text(m.group(3))) if m else ("", norm_text(usage))
            rep.conns.append(ReportConn(a_seg=segkey(r[c["a_seg"]]), b_seg=segkey(r[c["b_seg"]]) if c["b_seg"] < len(r) else None,
                                        usage=usage, state=state, tap=tap, address=addr,
                                        port_idx=int(m.group(2)) if m else None))
    return rep


def parse_guide_report(rows: List[List[Any]], fname: str) -> Optional[SpliceReport]:
    name = None
    title_olm = None
    in_col = out_col = None
    io_row = None
    hdr_row = None
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in r]
        low = [c.lower() for c in cells]
        for j, c in enumerate(low):
            if c.startswith("splice name") and name is None:
                name = norm_text(_find_label_value(r, j))
            if c == "in" and in_col is None:
                in_col, io_row = j, i
            if c == "out" and out_col is None:
                out_col = j
        if i == 1 and any(cells):
            toks = OLM_TOKEN.findall(norm_text(" ".join(cells)))
            title_olm = toks[-1] if toks else None
        if low.count("cable") >= 2 and "buffer" in low and hdr_row is None:
            hdr_row = i
    if not name:
        return None
    rep = SpliceReport(name=name, file=fname, fmt="Fiber Optic Splice Guide", olm=title_olm)
    if io_row is not None:
        stop = hdr_row if hdr_row is not None else len(rows)
        for col, direction in ((in_col, "in"), (out_col, "out")):
            if col is None:
                continue
            for r in rows[io_row + 1:stop]:
                if col < len(r) and r[col]:
                    m = GUIDE_SEG_RE.match(str(r[col]).strip())
                    if m:
                        cable = norm_cable(m.group(1))
                        cnt = int(m.group(3)) if m.group(3) else _fc_of(cable)
                        rep.segments.append(ReportSegment(
                            cable=cable, seg=int(m.group(2)), count=cnt,
                            kind="Pigtail" if (cnt is not None and cnt <= 4) else "Cable",
                            start=name if direction == "out" else None,
                            end=name if direction == "in" else None, direction=direction))
    if hdr_row is not None:
        hdr = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_row]]
        sub = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_row + 1]] if hdr_row + 1 < len(rows) else []
        cable_cols = [j for j, c in enumerate(hdr) if c == "cable"]
        hot_col = sub.index("hot") if "hot" in sub else None
        a_col, b_col = cable_cols[0], cable_cols[1]
        for r in rows[hdr_row + 2:]:
            if not any(v is not None and str(v).strip() for v in r):
                continue

            def ck(v: Any) -> Optional[Tuple[str, int]]:
                m = GUIDE_CABLE_RE.match(str(v or "").strip())
                return (norm_cable(m.group(1)), int(m.group(2))) if m else None
            hot = str(r[hot_col] or "").strip().upper() if hot_col is not None and hot_col < len(r) else ""
            addr = ""
            if hot_col is not None:
                for v in r[hot_col + 1:]:
                    if v is not None and str(v).strip():
                        addr = norm_text(v)
                        break
            state = "Live" if hot in ("H", "HOT", "Y", "YES", "1") else "Dark"
            rep.conns.append(ReportConn(a_seg=ck(r[a_col]) if a_col < len(r) else None,
                                        b_seg=ck(r[b_col]) if b_col < len(r) else None,
                                        usage=addr, state=state, tap="", address=addr))
    return rep


def load_splice_reports(files: List[Path], issues: IssueLog, rep: Reporter) -> List[SpliceReport]:
    reports: List[SpliceReport] = []
    for n, f in enumerate(files):
        rep.check_cancel()
        try:
            tables = read_tables(f)
        except Exception as exc:  # noqa: BLE001
            issues.add("ERROR", "Splice Report", f.name, f"Unreadable: {exc}")
            continue
        parsed_any = False
        for sheet, rows in tables:
            flat = " ".join(str(c) for r in rows[:6] for c in r if c is not None).lower()
            try:
                if "fiber optic splice guide" in flat:
                    r = parse_guide_report(rows, f.name)
                else:
                    r = parse_connections_report(rows, f.name)
                    if r is None:
                        r = parse_guide_report(rows, f.name)
            except Exception as exc:  # noqa: BLE001
                issues.add("ERROR", "Splice Report", f"{f.name}:{sheet}", f"Parse failure: {exc}")
                continue
            if r is not None:
                reports.append(r)
                parsed_any = True
        if not parsed_any:
            issues.add("ERROR", "Splice Report", f.name, "Unrecognized splice report layout")
        if n % 20 == 0:
            rep.progress(10 + 15 * n / max(1, len(files)), f"Reading splice reports ({n + 1}/{len(files)})")
    names = Counter(r.name for r in reports)
    for nm, c in names.items():
        if c > 1:
            issues.add("WARNING", "Splice Report", nm, f"{c} reports share this closure name; merged")
    fmts = Counter(r.fmt for r in reports)
    rep.log(f"Splice reports: {len(reports)} parsed ({', '.join(f'{k}: {v}' for k, v in fmts.items())})")
    return reports


# =============================================================================
# Drawing model
# =============================================================================
@dataclass
class Closure:
    name: str
    pos: Vec2
    handle: str
    box: Box


@dataclass
class Handhole:
    pos: Vec2
    handle: str
    block: str
    corners: List[Vec2]          # world-space box corners (counter-clockwise)
    box: Box


@dataclass
class CablePiece:
    handle: str
    layer: str
    pts: List[Vec2]
    a: Optional[str]
    b: Optional[str]


@dataclass
class ExistingCallout:
    handle: str
    dxftype: str
    layer: str
    text: str
    tip: Vec2
    box: Box


def open_drawing(path: str, oda_path: str, rep: Reporter) -> Any:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Drawing not found: {path}")
    if p.suffix.lower() == ".dwg":
        from ezdxf.addons import odafc
        if oda_path:
            key = "win_exec_path" if os.name == "nt" else "unix_exec_path"
            ezdxf.options.set("odafc-addon", key, oda_path)
        if not odafc.is_installed():
            raise RuntimeError(
                "DWG input needs the free ODA File Converter (https://www.opendesign.com/guestfiles/oda_file_converter). "
                "Install it (or set its path in the Inputs tab), or save the drawing as DXF in AutoCAD and load the DXF.")
        rep.log("Converting DWG -> DXF with ODA File Converter ...")
        return odafc.readfile(str(p))
    from ezdxf import recover
    doc, auditor = recover.readfile(str(p))
    if auditor.has_errors:
        rep.log(f"Drawing recovered with {len(auditor.errors)} structural error(s) fixed", "WARNING")
    return doc


def drawing_inventory(doc: Any) -> Dict[str, Any]:
    msp = doc.modelspace()
    inserts: Counter = Counter()
    tags: Dict[Tuple[str, str], Set[str]] = defaultdict(set)
    poly_layers: Counter = Counter()
    for e in msp:
        t = e.dxftype()
        if t == "INSERT":
            key = (e.dxf.layer, e.dxf.name)
            inserts[key] += 1
            if inserts[key] <= 25:
                for a in e.attribs:
                    tags[key].add(a.dxf.tag.upper())
        elif t in ("LWPOLYLINE", "POLYLINE", "LINE", "SPLINE", "ARC"):
            poly_layers[e.dxf.layer] += 1
    return dict(layers=sorted(l.dxf.name for l in doc.layers), inserts=inserts, tags=tags, poly_layers=poly_layers)


def _block_box_local(doc: Any, block_name: str) -> Optional[Tuple[Vec2, Vec2]]:
    try:
        blk = doc.blocks.get(block_name)
        if blk is None:
            return None
        ents = [e for e in blk if e.dxftype() in ("LINE", "LWPOLYLINE", "POLYLINE", "CIRCLE", "ARC", "SOLID",
                                                     "ELLIPSE", "SPLINE", "HATCH", "TRACE", "3DFACE")]
        if not ents:
            return None
        ext = ezbbox.extents(ents, fast=True)
        if not ext.has_data:
            return None
        return Vec2(ext.extmin.x, ext.extmin.y), Vec2(ext.extmax.x, ext.extmax.y)
    except Exception:  # noqa: BLE001 - malformed block -> treated as point symbol
        return None


def _insert_world_box(doc: Any, ins: Any, cache: Dict[str, Any], default_half: float) -> Tuple[List[Vec2], Box]:
    name = ins.dxf.name
    if name not in cache:
        cache[name] = _block_box_local(doc, name)
    loc = cache[name]
    m = ins.matrix44()
    if loc is None:
        lo, hi = Vec2(-default_half, -default_half), Vec2(default_half, default_half)
        base = Vec3(doc.blocks.get(name).base_point) if doc.blocks.get(name) is not None else Vec3()
        lo, hi = lo + Vec2(base.x, base.y), hi + Vec2(base.x, base.y)
    else:
        lo, hi = loc
    corners = [v2(m.transform(Vec3(x, y, 0))) for x, y in ((lo.x, lo.y), (hi.x, lo.y), (hi.x, hi.y), (lo.x, hi.y))]
    return corners, box_of_points(corners)


def _entity_points(e: Any) -> List[Vec2]:
    t = e.dxftype()
    if t == "LINE":
        return [v2(e.dxf.start), v2(e.dxf.end)]
    try:
        p = ezpath.make_path(e)
        pts = [v2(v) for v in p.flattening(distance=0.25)]
    except Exception:  # noqa: BLE001 - fall back to raw vertices
        if t == "LWPOLYLINE":
            pts = [Vec2(x, y) for x, y in e.get_points("xy")]
        elif t == "POLYLINE":
            pts = [v2(v.dxf.location) for v in e.vertices]
        else:
            pts = []
    out: List[Vec2] = []
    for p_ in pts:
        if not out or (p_ - out[-1]).magnitude > 1e-6:
            out.append(p_)
    return out


class DrawingModel:
    def __init__(self, doc: Any, cad: CadMap, spec: CalloutSpec, issues: IssueLog, rep: Reporter) -> None:
        self.doc = doc
        self.msp = doc.modelspace()
        self.cad = cad
        self.spec = spec
        self.issues = issues
        self.rep = rep
        self.closures: Dict[str, Closure] = {}
        self.handholes: List[Handhole] = []
        self.pieces: List[CablePiece] = []
        self.adj: Dict[str, List[int]] = defaultdict(list)
        self.existing: List[ExistingCallout] = []
        self.obstacles = GridIndex(50.0)
        self.cable_points = GridIndex(50.0)
        self.hh_index = GridIndex(100.0)
        self.closure_index = GridIndex(50.0)
        self._build()

    def _build(self) -> None:
        cad, spec = self.cad, self.spec
        cl_layer = cad.closure_layer.upper()
        hh_layer = cad.handhole_layer.upper()
        hh_blocks = {b.upper() for b in cad.handhole_blocks}
        term_layer = cad.terminal_layer.upper()
        cable_layers = {l.upper() for l in cad.cable_layers}
        out_layers = {spec.cable_layer.upper(), spec.terminal_layer.upper()}
        tag = cad.closure_name_tag.upper()
        cache: Dict[str, Any] = {}
        polylines: List[Any] = []
        n_unnamed = 0
        for e in self.msp:
            t = e.dxftype()
            layer = e.dxf.layer.upper()
            if t == "INSERT":
                if layer == cl_layer and (not cad.closure_block or e.dxf.name.upper() == cad.closure_block.upper()):
                    name = None
                    for a in e.attribs:
                        if a.dxf.tag.upper() == tag:
                            name = norm_text(a.dxf.text)
                    corners, box = _insert_world_box(self.doc, e, cache, 6.0)
                    if not name:
                        n_unnamed += 1
                        continue
                    if name in self.closures:
                        self.issues.add("WARNING", "Drawing", name, "Duplicate closure name in drawing; first kept")
                        continue
                    c = Closure(name=name, pos=v2(e.dxf.insert), handle=e.dxf.handle, box=box)
                    self.closures[name] = c
                    self.obstacles.insert(box, ("closure", name))
                    self.closure_index.insert((c.pos.x, c.pos.y, c.pos.x, c.pos.y), name)
                elif layer == hh_layer and e.dxf.name.upper() in hh_blocks:
                    corners, box = _insert_world_box(self.doc, e, cache, 6.0)
                    h = Handhole(pos=v2(e.dxf.insert), handle=e.dxf.handle, block=e.dxf.name, corners=corners, box=box)
                    self.handholes.append(h)
                    self.hh_index.insert(box, h)
                    self.obstacles.insert(box, ("handhole", e.dxf.handle))
                elif layer == term_layer:
                    corners, box = _insert_world_box(self.doc, e, cache, 3.5)
                    self.obstacles.insert(box, ("terminal", e.dxf.handle))
            elif t in ("LWPOLYLINE", "POLYLINE", "LINE", "SPLINE", "ARC") and layer in cable_layers:
                polylines.append(e)
            if t in ("MULTILEADER", "MTEXT", "TEXT") and layer in out_layers:
                ex = self._existing_callout(e)
                if ex:
                    self.existing.append(ex)
        if n_unnamed:
            self.issues.add("WARNING", "Drawing", cad.closure_layer,
                            f"{n_unnamed} closure block(s) without a '{cad.closure_name_tag}' attribute value ignored")
        self.rep.log(f"Drawing: {len(self.closures)} named closures, {len(self.handholes)} handholes, "
                     f"{len(polylines)} cable entities, {len(self.existing)} existing callouts on output layer(s)")
        if not self.closures:
            self.issues.add("ERROR", "Drawing", cad.closure_layer,
                            "No named closure blocks found - check the Layer Mapper (closure layer/block/tag)")
        self._split_cables(polylines)

    def _existing_callout(self, e: Any) -> Optional[ExistingCallout]:
        t = e.dxftype()
        try:
            if t == "MULTILEADER":
                ctx = e.context
                if ctx.mtext is None:
                    return None
                text = ctx.mtext.default_content
                tip = None
                for ld in ctx.leaders:
                    for ln in ld.lines:
                        if ln.vertices:
                            tip = v2(ln.vertices[0])
                            break
                    if tip is not None:
                        break
                loc = v2(ctx.mtext.insert)
                tip = tip or loc
                h = ctx.char_height or self.spec.text_height
            elif t == "MTEXT":
                text = e.text
                loc = v2(e.dxf.insert)
                tip = loc
                h = e.dxf.char_height
            else:
                text = e.dxf.text
                loc = v2(e.dxf.insert)
                tip = loc
                h = e.dxf.height
        except Exception:  # noqa: BLE001 - unreadable entity is ignored
            return None
        if not text or not re.search(r"\d", text):
            return None
        box = (loc.x - 8 * h, loc.y - 3 * h, loc.x + 8 * h, loc.y)
        return ExistingCallout(handle=e.dxf.handle, dxftype=t, layer=e.dxf.layer, text=text, tip=tip, box=box)

    def nearest_closure(self, p: Vec2, tol: float) -> Optional[str]:
        best, bd = None, tol
        for (_, name) in self.closure_index.query((p.x - tol, p.y - tol, p.x + tol, p.y + tol)):
            d = (self.closures[name].pos - p).magnitude
            if d <= bd:
                best, bd = name, d
        return best

    def _split_cables(self, entities: List[Any]) -> None:
        tol = self.cad.snap_tolerance
        for e in entities:
            pts = _entity_points(e)
            if len(pts) < 2:
                continue
            for i in range(0, len(pts), 2):
                p = pts[i]
                self.cable_points.insert((p.x, p.y, p.x, p.y), p)
            for i in range(len(pts) - 1):
                a, b = pts[i], pts[i + 1]
                L = (b - a).magnitude
                steps = int(L // 5.0)
                for k in range(1, steps + 1):
                    q = a + (b - a) * (k * 5.0 / L)
                    self.cable_points.insert((q.x, q.y, q.x, q.y), q)
            # insert split points where the cable passes through a closure
            enriched: List[Tuple[Vec2, Optional[str]]] = []
            for i in range(len(pts)):
                enriched.append((pts[i], self.nearest_closure(pts[i], tol)))
                if i + 1 < len(pts):
                    a, b = pts[i], pts[i + 1]
                    bx = box_of_points([a, b], tol)
                    mids = []
                    for (_, name) in self.closure_index.query(bx):
                        c = self.closures[name].pos
                        d, t = seg_distance(c, a, b)
                        if d <= tol and 1e-6 < t < 1 - 1e-6:
                            mids.append((t, c, name))
                    for t, c, name in sorted(mids):
                        enriched.append((c, name))
            idx = [i for i, (_, nm) in enumerate(enriched) if nm]
            plist = [p for p, _ in enriched]
            if not idx:
                self.pieces.append(CablePiece(e.dxf.handle, e.dxf.layer, plist, None, None))
                continue
            if idx[0] != 0:
                self.pieces.append(CablePiece(e.dxf.handle, e.dxf.layer, plist[:idx[0] + 1], None, enriched[idx[0]][1]))
            for i0, i1 in zip(idx, idx[1:]):
                a_n, b_n = enriched[i0][1], enriched[i1][1]
                if a_n == b_n:
                    continue
                piece = CablePiece(e.dxf.handle, e.dxf.layer, plist[i0:i1 + 1], a_n, b_n)
                piece.pts[0] = self.closures[a_n].pos
                piece.pts[-1] = self.closures[b_n].pos
                self.pieces.append(piece)
            if idx[-1] != len(plist) - 1:
                self.pieces.append(CablePiece(e.dxf.handle, e.dxf.layer, plist[idx[-1]:], enriched[idx[-1]][1], None))
        for i, pc in enumerate(self.pieces):
            if pc.a and pc.b:
                self.adj[pc.a].append(i)
                self.adj[pc.b].append(i)
        n_cc = sum(1 for p in self.pieces if p.a and p.b)
        self.rep.log(f"Cable topology: {len(self.pieces)} pieces, {n_cc} closure-to-closure links")

    def other_end(self, piece_idx: int, node: str) -> str:
        pc = self.pieces[piece_idx]
        return pc.b if pc.a == node else pc.a  # type: ignore[return-value]

    def piece_from(self, piece_idx: int, node: str) -> List[Vec2]:
        pc = self.pieces[piece_idx]
        return list(pc.pts) if pc.a == node else list(reversed(pc.pts))

    def find_piece_toward(self, start: str, target: str, max_nodes: int = 5000) -> Optional[int]:
        """Index of the piece leaving `start` whose side of the network contains `target`."""
        for pi in self.adj.get(start, []):
            first = self.other_end(pi, start)
            if first == target:
                return pi
        for pi in self.adj.get(start, []):
            seen = {start}
            stack = [self.other_end(pi, start)]
            while stack and len(seen) < max_nodes:
                n = stack.pop()
                if n == target:
                    return pi
                if n in seen:
                    continue
                seen.add(n)
                for pj in self.adj.get(n, []):
                    m = self.other_end(pj, n)
                    if m not in seen:
                        stack.append(m)
        return None

    def handhole_for(self, closure: Closure) -> Optional[Handhole]:
        r = self.spec.handhole_search_radius
        p = closure.pos
        best, bd = None, r
        for (_, h) in self.hh_index.query((p.x - r, p.y - r, p.x + r, p.y + r)):
            d = (h.pos - p).magnitude
            if d <= bd:
                best, bd = h, d
        return best


# =============================================================================
# Callout planning
# =============================================================================
@dataclass
class CalloutPlan:
    uid: str
    kind: str                     # TERMINAL | CABLE
    olm: str
    text: str                     # MText content (uses \P)
    layer: str
    tip: Vec2
    last: Vec2
    dog: int
    anchor: str
    closure: str
    numbers: List[int]
    calculation: str
    sources: str
    status: str = "PLANNED"
    notes: str = ""
    handle: str = ""
    existing_match: str = ""
    existing_text: str = ""

    @property
    def location(self) -> Vec2:
        return Vec2(0, 0)


@dataclass
class SegmentPlanInfo:
    closure: str
    cable: str
    seg: int
    count: Optional[int]
    end: Optional[str]
    live: List[int]
    unresolved: List[str]
    report_file: str
    olm: str
    text: str = ""
    geometry: str = ""


class CalloutPlanner:
    def __init__(self, model: DrawingModel, book: AddressBook, reports: List[SpliceReport],
                 spec: CalloutSpec, issues: IssueLog, rep: Reporter) -> None:
        self.m = model
        self.book = book
        self.reports = reports
        self.spec = spec
        self.issues = issues
        self.rep = rep
        self.placed = GridIndex(50.0)
        self.leaders: List[Tuple[Vec2, Vec2]] = []
        self.leader_index = GridIndex(50.0)
        self.plans: List[CalloutPlan] = []
        self.segments: List[SegmentPlanInfo] = []
        self.terminal_rows: List[Dict[str, Any]] = []

    # ---------------------------------------------------------------- text box
    def text_box(self, last: Vec2, dog: int, lines: List[str]) -> Box:
        h = self.spec.text_height
        width = max(len(x) for x in lines) * h * 0.9
        height = h + (len(lines) - 1) * h * 1.667 + 0.2 * h
        loc = self.text_location(last, dog)
        if dog > 0:
            return (loc.x, loc.y - height, loc.x + width, loc.y)
        return (loc.x - width, loc.y - height, loc.x, loc.y)

    def text_location(self, last: Vec2, dog: int) -> Vec2:
        s = self.spec
        return Vec2(last.x + dog * (s.dogleg_length + s.landing_gap), last.y + s.text_y_factor * s.text_height)

    def _score(self, tip: Vec2, last: Vec2, dog: int, lines: List[str], away_from: Optional[Vec2]) -> float:
        """Lower is better: penalises real overlaps only (grid queries return cell neighbours)."""
        box = self.text_box(last, dog, lines)
        margin = 1.0
        clear = (box[0] - margin, box[1] - margin, box[2] + margin, box[3] + margin)
        score = 0.0
        for (b, _item) in self.placed.query(clear):
            ov = box_overlap(clear, b)
            if ov > 0:
                score += 400.0 + 60.0 * ov
        for (b, _item) in self.m.obstacles.query(clear):
            ov = box_overlap(clear, b)
            if ov > 0:
                score += 150.0 + 20.0 * ov
        n_pts = sum(1 for (b, _p) in self.m.cable_points.query(box)
                    if box[0] <= b[0] <= box[2] and box[1] <= b[1] <= box[3])
        score += 25.0 * n_pts
        lb = box_of_points([tip, last], 0.5)
        for (_b, seg) in self.leader_index.query(lb):
            if segments_intersect(tip, last, seg[0], seg[1]):
                score += 300.0
        for (b, _item) in self.placed.query(lb):
            if segment_hits_box(tip, last, b):
                score += 150.0
        if away_from is not None:
            d = (last - tip)
            v = (tip - away_from)
            if d.magnitude > 1e-9 and v.magnitude > 1e-9:
                score += 20.0 * (1.0 - d.normalize().dot(v.normalize()))
        return score

    def placement_qa(self) -> Dict[str, int]:
        """Flag text/text overlaps, text/symbol overlaps and leader crossings on every placed plan."""
        live = [p for p in self.plans if p.status != "FAILED"]
        boxes = {p.uid: self.text_box(p.last, p.dog, p.text.split("\\P")) for p in live}
        tgrid = GridIndex(50.0)
        lgrid = GridIndex(50.0)
        for p in live:
            tgrid.insert(boxes[p.uid], p.uid)
            lgrid.insert(box_of_points([p.tip, p.last], 0.1), (p.uid, p.tip, p.last))
        counts = Counter()
        for p in live:
            b = boxes[p.uid]
            flags = []
            if any(u != p.uid and box_overlap(b, ob) > 0 for (ob, u) in tgrid.query(b)):
                flags.append("text overlaps another callout")
                counts["text_overlap"] += 1
            if any(box_overlap(b, ob) > 0 for (ob, _i) in self.m.obstacles.query(b)):
                flags.append("text overlaps a symbol")
                counts["symbol_overlap"] += 1
            if any(u != p.uid and segments_intersect(p.tip, p.last, a, c)
                   for (_bb, (u, a, c)) in lgrid.query(box_of_points([p.tip, p.last], 0.1))):
                flags.append("leader crosses another leader")
                counts["leader_cross"] += 1
            if flags:
                p.notes = "; ".join(x for x in [p.notes, "PLACEMENT REVIEW: " + ", ".join(flags)] if x)
        counts["clean"] = sum(1 for p in live if "PLACEMENT REVIEW" not in p.notes)
        return dict(counts)

    def _register(self, plan: CalloutPlan, lines: List[str]) -> None:
        box = self.text_box(plan.last, plan.dog, lines)
        self.placed.insert(box, plan.uid)
        self.leader_index.insert(box_of_points([plan.tip, plan.last], 0.5), (plan.tip, plan.last))
        self.plans.append(plan)

    # ---------------------------------------------------------------- terminals
    def plan_terminals(self) -> None:
        spec = self.spec
        by_closure: Dict[str, List[Terminal]] = defaultdict(list)
        for _key, t in sorted(self.book.terminals.items(), key=lambda kv: (kv[0][0], min(kv[1].numbers))):
            by_closure[t.splice].append(t)
        uid = 0
        for closure_name, terms in by_closure.items():
            self.rep.check_cancel()
            closure = self.m.closures.get(closure_name)
            for t in terms:
                runs = number_runs(t.numbers) if spec.split_terminal_runs else [(min(t.numbers), max(t.numbers))]
                text_runs = [fmt_run(r, spec.zero_pad) for r in runs] if spec.split_terminal_runs else \
                    [fmt_numbers(t.numbers, "exact", spec.zero_pad)]
                calc = "; ".join(f"{r.address}: PON {r.pon} port {r.port} -> ({r.pon}-1)*{self.book.ports_per_pon}+{r.port} = {r.number}"
                                 for r in sorted(t.records, key=lambda x: x.number))
                srcs = ", ".join(sorted({r.source for r in t.records}))
                row = dict(OLM=t.olm, Terminal=t.tap, Closure=t.splice, Numbers=fmt_numbers(t.numbers, "exact", 0),
                           Addresses=len(t.records), Callouts=", ".join(text_runs), Status="")
                if closure is None:
                    self.issues.add("ERROR", "Missing Closure", t.splice or "(blank)",
                                    f"Drop closure of terminal {t.tap} ({fmt_numbers(t.numbers, 'exact', 0)}) not found in drawing")
                    row["Status"] = "NOT PLACED - closure missing"
                    self.terminal_rows.append(row)
                    for tr in text_runs:
                        uid += 1
                        self.plans.append(CalloutPlan(uid=f"T{uid:05d}", kind="TERMINAL", olm=t.olm, text=tr,
                                                      layer=spec.terminal_layer, tip=Vec2(), last=Vec2(), dog=1,
                                                      anchor="-", closure=t.splice, numbers=sorted(t.numbers),
                                                      calculation=calc, sources=srcs, status="FAILED",
                                                      notes="Drop closure not found in drawing"))
                    continue
                hh = self.m.handhole_for(closure)
                for tr, run in zip(text_runs, runs):
                    uid += 1
                    nums = [n for n in t.numbers if run[0] <= n <= run[1]] if spec.split_terminal_runs else sorted(t.numbers)
                    plan = self._place_terminal(f"T{uid:05d}", t, tr, closure, hh, nums, calc, srcs)
                    self._register(plan, [tr])
                row["Status"] = "PLACED" + ("" if hh else " (no handhole - anchored on closure symbol)")
                self.terminal_rows.append(row)

    def _place_terminal(self, uid: str, t: Terminal, text: str, closure: Closure, hh: Optional[Handhole],
                        nums: List[int], calc: str, srcs: str) -> CalloutPlan:
        spec = self.spec
        if hh is not None:
            corners = hh.corners
            center = Vec2(sum(c.x for c in corners) / 4, sum(c.y for c in corners) / 4)
            anchor = f"{hh.block} {hh.handle}"
            note = ""
        else:
            xmin, ymin, xmax, ymax = closure.box
            corners = [Vec2(xmin, ymin), Vec2(xmax, ymin), Vec2(xmax, ymax), Vec2(xmin, ymax)]
            center = closure.pos
            anchor = f"closure {closure.name}"
            note = "No handhole within search radius; arrow placed on closure symbol"
            self.issues.add("WARNING", "Placement", t.tap,
                            f"No handhole within {spec.handhole_search_radius:g} of closure {closure.name}; arrow on closure symbol")
        candidates: List[Tuple[Vec2, Vec2]] = []
        for i in range(4):
            c = corners[i]
            d = c - center
            if d.magnitude > 1e-9:
                candidates.append((c, d.normalize()))
            mid = (corners[i] + corners[(i + 1) % 4]) * 0.5
            dm = mid - center
            if dm.magnitude > 1e-9:
                candidates.append((mid, dm.normalize()))
        best = None
        for tip, direction in candidates:
            for len_factor in (1.0, 1.6):
                last = tip + direction * (spec.terminal_leader_length * len_factor)
                dogs = [1 if direction.x > 1e-6 else -1] if abs(direction.x) > 1e-6 else [1, -1]
                for dog in dogs:
                    sc = self._score(tip, last, dog, [text], closure.pos) + (len_factor - 1.0) * 20.0
                    if best is None or sc < best[0]:
                        best = (sc, tip, last, dog)
        _, tip, last, dog = best  # type: ignore[misc]
        return CalloutPlan(uid=uid, kind="TERMINAL", olm=t.olm, text=text, layer=spec.terminal_layer,
                           tip=tip, last=last, dog=dog, anchor=anchor, closure=closure.name, numbers=nums,
                           calculation=f"Terminal {t.tap}: " + calc, sources=srcs, notes=note)

    # ---------------------------------------------------------------- cables
    def _report_olm(self, r: SpliceReport) -> Optional[str]:
        votes: Counter = Counter()
        for c in r.conns:
            if c.address:
                rec = self.book.resolve(c.tap, c.address, None, c.port_idx)
                if rec:
                    votes[rec.olm] += 1
        if votes:
            return votes.most_common(1)[0][0]
        closure_olms = Counter(rec.olm for rec in self.book.records if rec.splice == r.name)
        if closure_olms:
            return closure_olms.most_common(1)[0][0]
        return r.olm

    def plan_cables(self) -> None:
        spec = self.spec
        merged: Dict[str, List[SpliceReport]] = defaultdict(list)
        for r in self.reports:
            merged[r.name].append(r)
        uid = 0
        all_segment_ends: Dict[Tuple[str, int], Tuple[Optional[str], Optional[str]]] = {}
        for r in self.reports:
            for s in r.segments:
                key = (s.cable, s.seg)
                st, en = all_segment_ends.get(key, (None, None))
                all_segment_ends[key] = (st or s.start, en or s.end)
        for n_rep, (name, reps) in enumerate(sorted(merged.items())):
            self.rep.check_cancel()
            if n_rep % 25 == 0:
                self.rep.progress(55 + 25 * n_rep / max(1, len(merged)), f"Planning cable callouts ({n_rep}/{len(merged)})")
            closure = self.m.closures.get(name)
            if closure is None:
                self.issues.add("ERROR", "Missing Closure", name, "Splice report closure not found in drawing")
            for r in reps:
                olm = self._report_olm(r) or "?"
                r.olm = olm
                for sg in r.segments:
                    if sg.direction != "out":
                        continue
                    dest = sg.cable.split(" TO ", 1)[1] if " TO " in sg.cable else ""
                    dest = re.sub(r"\s+FC\s*\d+\s*$", "", dest).strip()
                    if sg.kind.lower() == "pigtail" or dest in self.book.tap_names:
                        continue
                    key = (sg.cable, sg.seg)
                    end = sg.end or all_segment_ends.get(key, (None, None))[1]
                    piece = self._segment_piece(name, end, dest)
                    live: List[int] = []
                    unresolved: List[str] = []
                    ambiguous: Dict[Tuple[int, ...], List[str]] = defaultdict(list)
                    for c in r.conns:
                        if key not in (c.a_seg, c.b_seg):
                            continue
                        state = norm_text(c.state)
                        if any(s_ in state for s_ in DEAD_STATES):
                            continue
                        if not c.address or NON_ADDRESS_RE.search(c.usage or c.address):
                            continue
                        cands = self.book.resolve_candidates(c.tap, c.address, olm, c.port_idx)
                        if len(cands) > 1 and piece is not None:
                            down = [x for x in cands if x.splice in self.m.closures and
                                    self.m.find_piece_toward(name, x.splice) == piece]
                            cands = down or cands
                        if len(cands) == 1:
                            live.append(cands[0].number)
                        elif cands:
                            ambiguous[tuple(sorted(x.number for x in cands))].append(c.usage or c.address)
                        else:
                            unresolved.append(c.usage or c.address)
                    for nums, usages in ambiguous.items():
                        # the same address on as many live fibers as it has address-list rows -> all of them
                        if len(usages) == len(nums):
                            live.extend(nums)
                        else:
                            unresolved.extend(u + " (ambiguous)" for u in usages)
                    info = SegmentPlanInfo(closure=name, cable=sg.cable, seg=sg.seg, count=sg.count, end=end,
                                           live=sorted(set(live)), unresolved=unresolved, report_file=r.file, olm=olm)
                    self.segments.append(info)
                    real_unres = list(unresolved)
                    if real_unres:
                        self.issues.add("WARNING", "Splice Mismatch", f"{name} {sg.cable}:Seg{sg.seg}",
                                        f"{len(real_unres)} live fiber usage(s) not resolvable against the address list, "
                                        f"e.g. {real_unres[0]!r}")
                    if not info.live:
                        info.geometry = "skipped - no live fibers"
                        continue
                    if sg.count is None:
                        self.issues.add("WARNING", "Splice Report", f"{name} {sg.cable}",
                                        "Segment fiber count unknown (no 'Count' and no FCxxx suffix)")
                    count_txt = str(sg.count) if sg.count is not None else "?"
                    text = f"{spec.cable_prefix} {count_txt}\\P{fmt_numbers(info.live, spec.cable_range_mode, spec.zero_pad)}"
                    info.text = text
                    uid += 1
                    plan = self._place_cable(f"C{uid:05d}", info, closure, piece, text)
                    info.geometry = plan.anchor if plan.status != "FAILED" else plan.notes
                    if plan.status == "FAILED":
                        self.plans.append(plan)
                    else:
                        self._register(plan, text.split("\\P"))

    def _segment_piece(self, start: str, end: Optional[str], dest: str) -> Optional[int]:
        """Cable piece leaving `start` that carries the segment (toward its end closure / cable destination)."""
        if start not in self.m.closures:
            return None
        pi = None
        if end and end in self.m.closures:
            pi = self.m.find_piece_toward(start, end)
        if pi is None and dest and dest in self.m.closures and dest != start:
            pi = self.m.find_piece_toward(start, dest)
        if pi is None:
            cands = self.m.adj.get(start, [])
            if len(cands) == 1:
                pi = cands[0]
        return pi

    def _place_cable(self, uid: str, info: SegmentPlanInfo, closure: Optional[Closure], pi: Optional[int],
                     text: str) -> CalloutPlan:
        spec = self.spec
        calc = (f"Segment {info.seg} of {info.cable} leaving {info.closure}: {len(info.live)} live fibers "
                f"{fmt_numbers(info.live, 'exact', 0)}; count={info.count}; "
                f"text = '{spec.cable_prefix} <count>' + {'min-max' if spec.cable_range_mode == 'span' else 'runs'}")
        base = dict(uid=uid, kind="CABLE", olm=info.olm, text=text, layer=spec.cable_layer, closure=info.closure,
                    numbers=info.live, calculation=calc, sources=info.report_file)
        if closure is None:
            return CalloutPlan(tip=Vec2(), last=Vec2(), dog=1, anchor="-", status="FAILED",
                               notes="Closure not in drawing", **base)
        note = ""
        if pi is not None:
            pts = self.m.piece_from(pi, info.closure)
            piece_desc = f"cable {self.m.pieces[pi].handle} toward {self.m.other_end(pi, info.closure)}"
        elif info.end and info.end in self.m.closures:
            pts = [closure.pos, self.m.closures[info.end].pos]
            piece_desc = f"straight line toward {info.end}"
            note = "No cable polyline between closures; arrow placed on straight line"
            self.issues.add("WARNING", "Placement", f"{info.closure} {info.cable}:Seg{info.seg}", note)
        else:
            self.issues.add("WARNING", "Placement", f"{info.closure} {info.cable}:Seg{info.seg}",
                            "Cannot identify the cable geometry of this segment; callout not placed")
            return CalloutPlan(tip=Vec2(), last=Vec2(), dog=1, anchor="-", status="FAILED",
                               notes="Segment geometry not found", **base)
        L = polyline_length(pts)
        s_along = min(spec.cable_offset_along, 0.45 * L) if L > 0 else 0.0
        s_along = max(s_along, min(5.0, 0.5 * L))
        lines = text.split("\\P")
        best = None
        for frac in (1.0, 0.7, 1.4, 2.0, 2.8):
            s = min(max(s_along * frac, 0.05 * L), 0.9 * L)
            tip, tan = point_at_length(pts, s)
            nrm = Vec2(-tan.y, tan.x)
            for side in (1, -1):
                for skew in (0.0, 35.0, -35.0):
                    d = (nrm * side).rotate_deg(skew)
                    for len_factor in (1.0, 1.5):
                        last = tip + d * (spec.cable_leader_length * len_factor)
                        dog = 1 if d.x >= 0 else -1
                        sc = (self._score(tip, last, dog, lines, closure.pos) + abs(frac - 1.0) * 12.0
                              + (len_factor - 1.0) * 20.0)
                        if best is None or sc < best[0]:
                            best = (sc, tip, last, dog)
        _, tip, last, dog = best  # type: ignore[misc]
        return CalloutPlan(tip=tip, last=last, dog=dog, anchor=piece_desc, notes=note, **base)


# =============================================================================
# Existing-callout comparison and CAD writing
# =============================================================================
def compare_existing(plans: List[CalloutPlan], existing: List[ExistingCallout], radius: float = 150.0) -> List[Dict[str, Any]]:
    idx = GridIndex(100.0)
    for i, ex in enumerate(existing):
        idx.insert((ex.tip.x, ex.tip.y, ex.tip.x, ex.tip.y), i)
    used: Set[int] = set()
    rows: List[Dict[str, Any]] = []
    for p in plans:
        if p.status == "FAILED":
            continue
        want = canonical_callout_text(p.text)
        cands = []
        for (_, i) in idx.query((p.tip.x - radius, p.tip.y - radius, p.tip.x + radius, p.tip.y + radius)):
            if i in used:
                continue
            d = (existing[i].tip - p.tip).magnitude
            if d <= radius:
                cands.append((d, i))
        exact = sorted((d, i) for d, i in cands if canonical_callout_text(existing[i].text) == want)
        if exact:
            d, i = exact[0]
            used.add(i)
            strict = existing[i].text.replace("\n", "\\P") == p.text
            p.existing_match = "MATCH" if strict else "MATCH (formatting differs)"
            p.existing_text = existing[i].text
            rows.append(dict(Callout=p.uid, Kind=p.kind, Computed=p.text, Existing=existing[i].text,
                             ExistingHandle=existing[i].handle, Distance=round(d, 2), Result=p.existing_match))
        else:
            p.existing_match = "NEW"
            near = sorted(cands)[:1]
            rows.append(dict(Callout=p.uid, Kind=p.kind, Computed=p.text,
                             Existing=existing[near[0][1]].text if near else "",
                             ExistingHandle=existing[near[0][1]].handle if near else "",
                             Distance=round(near[0][0], 2) if near else None,
                             Result="NOT IN DRAWING" if not near else "TEXT DIFFERS / NOT FOUND"))
    for i, ex in enumerate(existing):
        if i not in used:
            rows.append(dict(Callout="", Kind="", Computed="", Existing=ex.text, ExistingHandle=ex.handle,
                             Distance=None, Result="EXISTING ONLY (no computed counterpart)"))
    return rows


class CalloutWriter:
    def __init__(self, doc: Any, spec: CalloutSpec, issues: IssueLog, rep: Reporter) -> None:
        self.doc = doc
        self.msp = doc.modelspace()
        self.spec = spec
        self.issues = issues
        self.rep = rep
        self._ensure_resources()

    def _ensure_resources(self) -> None:
        s = self.spec
        for layer in {s.cable_layer, s.terminal_layer}:
            if layer not in self.doc.layers:
                self.doc.layers.add(layer, color=s.layer_color, linetype=s.layer_linetype)
                self.rep.log(f"Created layer '{layer}' (ACI {s.layer_color}, {s.layer_linetype})")
        if s.text_style not in self.doc.styles:
            self.doc.styles.add(s.text_style, font=s.text_style_font)
            self.rep.log(f"Created text style '{s.text_style}' ({s.text_style_font})")
        else:
            st = self.doc.styles.get(s.text_style)
            font = (st.dxf.get("font", "") or "").lower()
            if font and not font.startswith(s.text_style_font.lower()):
                self.issues.add("WARNING", "Style", s.text_style,
                                f"Existing text style font is '{font}' (reference uses '{s.text_style_font}')")
        if self.doc.mleader_styles.get(s.mleader_style) is None:
            self.doc.mleader_styles.new(s.mleader_style)
            self.rep.log(f"Created MLEADER style '{s.mleader_style}'")

    def add(self, plan: CalloutPlan) -> str:
        s = self.spec
        dog = 1 if plan.dog >= 0 else -1
        text_loc = Vec2(plan.last.x + dog * (s.dogleg_length + s.landing_gap),
                        plan.last.y + s.text_y_factor * s.text_height)
        builder = self.msp.add_multileader_mtext(s.mleader_style, dxfattribs={"layer": plan.layer})
        builder.set_content(plan.text, char_height=s.text_height,
                            alignment=mlr.TextAlignment.left if dog > 0 else mlr.TextAlignment.right,
                            style=s.text_style)
        builder.set_connection_properties(landing_gap=s.landing_gap, dogleg_length=s.dogleg_length)
        builder.set_connection_types(left=mlr.HorizontalConnection(s.side_attachment),
                                     right=mlr.HorizontalConnection(s.side_attachment),
                                     top=mlr.VerticalConnection(s.vertical_attachment),
                                     bottom=mlr.VerticalConnection(s.vertical_attachment))
        builder.set_arrow_properties(name="", size=s.arrow_size)
        builder.set_leader_properties(color=ezcolors.BYLAYER, linetype="BYBLOCK", lineweight=-2,
                                      leader_type=mlr.LeaderType.straight_lines)
        side = mlr.ConnectionSide.left if dog > 0 else mlr.ConnectionSide.right
        builder.add_leader_line(side, [Vec2(plan.tip.x, plan.tip.y)])
        builder.build(insert=text_loc)
        ml = builder.multileader
        # ---- exact reference geometry & properties --------------------------
        ml.dxf.color = ezcolors.BYLAYER
        ml.dxf.linetype = "BYLAYER"
        ml.dxf.lineweight = -1
        ml.dxf.leader_type = 1
        ml.dxf.leader_line_color = ezcolors.BY_LAYER_RAW_VALUE
        ml.dxf.leader_lineweight = -2
        ml.dxf.has_landing = 1
        ml.dxf.has_dogleg = 1
        ml.dxf.dogleg_length = s.dogleg_length
        ml.dxf.arrow_head_size = s.arrow_size
        ml.dxf.content_type = 2
        ml.dxf.text_left_attachment_type = s.side_attachment
        ml.dxf.text_right_attachment_type = s.side_attachment
        ml.dxf.text_top_attachment_type = s.vertical_attachment
        ml.dxf.text_bottom_attachment_type = s.vertical_attachment
        ml.dxf.text_attachment_direction = 0
        ml.dxf.text_angle_type = 1
        ml.dxf.text_alignment_type = 2
        ml.dxf.text_attachment_point = 1 if dog > 0 else 3
        ml.dxf.text_color = ezcolors.BY_LAYER_RAW_VALUE
        ml.dxf.has_text_frame = 0
        ml.dxf.is_annotative = 0
        ml.dxf.scale = 1.0
        ctx = ml.context
        ctx.scale = 1.0
        ctx.char_height = s.text_height
        ctx.arrow_head_size = s.arrow_size
        ctx.landing_gap_size = s.landing_gap
        ctx.left_attachment = s.side_attachment
        ctx.right_attachment = s.side_attachment
        ctx.top_attachment = s.vertical_attachment
        ctx.bottom_attachment = s.vertical_attachment
        ctx.text_align_type = 0 if dog > 0 else 2
        ctx.attachment_type = 0
        lines = plan.text.split("\\P")
        est_w = max(len(x) for x in lines) * s.text_height * 0.9
        ctx.base_point = Vec3(text_loc.x if dog > 0 else text_loc.x - est_w, plan.last.y, 0)
        mt = ctx.mtext
        mt.default_content = plan.text
        mt.insert = Vec3(text_loc.x, text_loc.y, 0)
        mt.text_direction = Vec3(1, 0, 0)
        mt.rotation = 0.0
        mt.width = 0.0
        mt.defined_height = 0.0
        mt.line_spacing_factor = 1.0
        mt.line_spacing_style = 1
        mt.color = ezcolors.BY_BLOCK_RAW_VALUE
        mt.alignment = 1 if dog > 0 else 3
        mt.flow_direction = 1
        mt.has_bg_fill = 0
        mt.use_window_bg_color = 0
        mt.bg_scale_factor = 1.5
        mt.column_type = 0
        if ctx.leaders:
            ld = ctx.leaders[0]
            ld.has_last_leader_line = 1
            ld.has_dogleg_vector = 1
            ld.last_leader_point = Vec3(plan.last.x, plan.last.y, 0)
            ld.dogleg_vector = Vec3(-1 if dog < 0 else 1, 0, 0)
            ld.dogleg_length = s.dogleg_length
            ld.attachment_direction = 0
            if ld.lines:
                ld.lines[0].vertices = [Vec3(plan.tip.x, plan.tip.y, 0)]
                ld.lines[0].color = ezcolors.BY_LAYER_RAW_VALUE
        try:
            ml.update_proxy_graphic()
        except Exception:  # noqa: BLE001 - proxy graphic is optional display data
            pass
        return ml.dxf.handle


# =============================================================================
# QA / QC workbook
# =============================================================================
def write_qa_workbook(path: Path, cfg: RunConfig, plans: List[CalloutPlan], segments: List[SegmentPlanInfo],
                      terminals: List[Dict[str, Any]], issues: IssueLog, existing_rows: List[Dict[str, Any]],
                      summary: Dict[str, Any]) -> None:
    callouts = [dict(ID=p.uid, Kind="Fiber number (terminal)" if p.kind == "TERMINAL" else "BFO (cable segment)",
                     OLM=p.olm, Text=p.text.replace("\\P", " / "), Layer=p.layer, Closure=p.closure, Anchor=p.anchor,
                     TipX=round(p.tip.x, 4), TipY=round(p.tip.y, 4), LandingX=round(p.last.x, 4),
                     LandingY=round(p.last.y, 4), Dogleg="+X" if p.dog > 0 else "-X",
                     Numbers=fmt_numbers(p.numbers, "exact", 0), Calculation=p.calculation, Sources=p.sources,
                     Status=p.status, Handle=p.handle, ExistingCheck=p.existing_match,
                     ExistingText=p.existing_text.replace("\n", " / "), Notes=p.notes) for p in plans]
    segs = [dict(Closure=s.closure, Cable=s.cable, Segment=s.seg, FiberCount=s.count, EndClosure=s.end, OLM=s.olm,
                 LiveFibers=len(s.live), Numbers=fmt_numbers(s.live, "exact", 0), CalloutText=s.text.replace("\\P", " / "),
                 Geometry=s.geometry, UnresolvedUsages=len(s.unresolved),
                 UnresolvedExample=s.unresolved[0] if s.unresolved else "", Report=s.report_file) for s in segments]
    iss = [dict(Severity=i.severity, Category=i.category, Object=i.obj, Message=i.message) for i in issues.items]
    summ = [dict(Item=k, Value=v) for k, v in summary.items()]
    settings = [dict(Section="run", Key=k, Value=json.dumps(v) if isinstance(v, (list, dict)) else v)
                for k, v in cfg.to_dict().items() if k not in ("columns", "cad", "spec")]
    for sec in ("columns", "cad", "spec"):
        for k, v in getattr(cfg, sec).to_dict().items():
            settings.append(dict(Section=sec, Key=k, Value=json.dumps(v) if isinstance(v, (list, dict)) else v))
    sheets = [("Summary", summ), ("Callouts", callouts), ("Terminals", terminals), ("Cable Segments", segs),
              ("Validation Issues", iss), ("Existing vs Computed", existing_rows), ("Settings", settings)]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, rows in sheets:
            df = pd.DataFrame(rows if rows else [{"Info": "No rows"}])
            df.to_excel(xw, sheet_name=name, index=False)
    wb = openpyxl.load_workbook(path)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    sev_fill = {"ERROR": PatternFill("solid", fgColor="F8CBAD"), "WARNING": PatternFill("solid", fgColor="FFE699")}
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            width = max((len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 400) + 1)), default=8)
            ws.column_dimensions[letter].width = min(max(10, width + 2), 80)
        if ws.title == "Validation Issues":
            for r in range(2, ws.max_row + 1):
                f = sev_fill.get(str(ws.cell(r, 1).value))
                if f:
                    for c in ws[r]:
                        c.fill = f
    wb.save(path)


# =============================================================================
# Engine (GUI-independent)
# =============================================================================
@dataclass
class ScanResult:
    headers: List[str]
    columns: ColumnMap
    cad: Optional[CadMap]
    layers: List[str]
    insert_pairs: List[Tuple[str, str, int]]
    insert_tags: Dict[str, List[str]]
    poly_layers: List[Tuple[str, int]]
    issues: IssueLog


class CalloutEngine:
    def __init__(self, cfg: RunConfig, rep: Reporter) -> None:
        self.cfg = cfg
        self.rep = rep
        self.collector = InputCollector()

    # ------------------------------------------------------------------ scan
    def scan(self) -> ScanResult:
        issues = IssueLog()
        try:
            addr_files = self.collector.expand(self.cfg.address_paths, (".xlsx", ".xlsm", ".xls", ".csv", ".txt"), issues)
            headers = AddressBook.scan_headers(addr_files, issues)
            cols = auto_map_columns(headers)
            self.rep.log(f"Address list headers: {', '.join(headers) if headers else '(none)'}")
            for k in ("pon", "port", "address", "tap", "splice", "fiber", "cable", "olm"):
                self.rep.log(f"  map {k:<8} -> {getattr(cols, k)!r}", "DEBUG")
            cad = None
            layers: List[str] = []
            pairs: List[Tuple[str, str, int]] = []
            tags: Dict[str, List[str]] = {}
            polys: List[Tuple[str, int]] = []
            if self.cfg.drawing_path:
                self.rep.progress(20, "Reading drawing")
                doc = open_drawing(self.cfg.drawing_path, self.cfg.oda_path, self.rep)
                inv = drawing_inventory(doc)
                layers = inv["layers"]
                pairs = [(l, b, n) for (l, b), n in inv["inserts"].most_common()]
                tags = {f"{l}|{b}": sorted(t) for (l, b), t in inv["tags"].items()}
                polys = inv["poly_layers"].most_common()
                closure_names: Set[str] = set()
                if addr_files and cols.splice:
                    tmp_book = AddressBook()
                    tmp_issues = IssueLog()
                    tmp_book.load(addr_files, cols, self.cfg.ports_per_pon, tmp_issues, Reporter(lambda m, l: None))
                    closure_names = {r.splice for r in tmp_book.records if r.splice}
                spl_files = self.collector.expand(self.cfg.splice_paths, (".xlsx", ".xlsm", ".xls", ".csv"), issues)
                closure_names |= {Path(f).stem.upper() for f in spl_files}
                cad = self.auto_detect_cad(doc, inv, closure_names)
                self.rep.log(f"Auto-detected CAD mapping: closures '{cad.closure_layer}'/'{cad.closure_block}' tag "
                             f"'{cad.closure_name_tag}', handholes '{cad.handhole_layer}' {cad.handhole_blocks}, "
                             f"cables {cad.cable_layers}")
            self.rep.progress(100, "Scan complete")
            return ScanResult(headers, cols, cad, layers, pairs, tags, polys, issues)
        finally:
            self.collector.cleanup()

    def auto_detect_cad(self, doc: Any, inv: Dict[str, Any], closure_names: Set[str]) -> CadMap:
        """Detect closure / handhole / terminal / cable content from geometry and the input data.

        closures  : INSERT (layer, block, tag) whose attribute values best match closure names
                    from the address lists ('Fiber Splice') and splice report file names
        terminals : blocks joined to closures by a pigtail (linework from closure to block)
        handholes : non-terminal blocks found 8-45 units from a large share of closures
        cables    : layers whose linework endpoints coincide with closures (pigtail/drop excluded)
        """
        msp = doc.modelspace()
        cad = CadMap()
        values: Dict[Tuple[str, str, str], Set[str]] = defaultdict(set)
        positions: Dict[Tuple[str, str], List[Vec2]] = defaultdict(list)
        for e in msp.query("INSERT"):
            key = (e.dxf.layer, e.dxf.name)
            positions[key].append(v2(e.dxf.insert))
            for a in e.attribs:
                values[(e.dxf.layer, e.dxf.name, a.dxf.tag.upper())].add(norm_text(a.dxf.text))
        best = (0, None, None, None)
        for (layer, block, tag), vals in values.items():
            score = len(vals & closure_names)
            if score > best[0]:
                best = (score, layer, block, tag)
        if best[1]:
            cad.closure_layer, cad.closure_block, cad.closure_name_tag = best[1], best[2], best[3]
            self.rep.log(f"Closure detection: {best[0]} closure names matched on '{best[1]}'/'{best[2]}' tag '{best[3]}'")
        closure_pts = positions.get((cad.closure_layer, cad.closure_block), [])
        if not closure_pts:
            self.rep.log("Closure blocks could not be auto-detected - set them in the Layer Mapper", "WARNING")
            return cad
        cgrid = GridIndex(25.0)
        for p in closure_pts:
            cgrid.insert((p.x, p.y, p.x, p.y), p)

        def near_closure(p: Vec2, tol: float) -> bool:
            return any((q - p).magnitude <= tol for (_, q) in cgrid.query((p.x - tol, p.y - tol, p.x + tol, p.y + tol)))

        # linework endpoints, used for pigtail connectivity and cable-layer detection
        end_grid = GridIndex(25.0)
        layer_hits: Counter = Counter()
        layer_total: Counter = Counter()
        for e in msp:
            if e.dxftype() not in ("LWPOLYLINE", "POLYLINE", "LINE"):
                continue
            pts = _entity_points(e)
            if len(pts) < 2:
                continue
            a, b = pts[0], pts[-1]
            end_grid.insert((a.x, a.y, a.x, a.y), (a, b))
            end_grid.insert((b.x, b.y, b.x, b.y), (b, a))
            layer_total[e.dxf.layer] += 1
            if near_closure(a, 2.0) or near_closure(b, 2.0):
                layer_hits[e.dxf.layer] += 1
        sample = closure_pts[:: max(1, len(closure_pts) // 300)]
        stats: List[Tuple[str, str, int, float, float, float]] = []
        for (layer, block), pts in positions.items():
            if (layer, block) == (cad.closure_layer, cad.closure_block) or len(pts) < 3:
                continue
            pgrid = GridIndex(50.0)
            for p in pts:
                pgrid.insert((p.x, p.y, p.x, p.y), p)
            dists = []
            for c in sample:
                ds = [(q - c).magnitude for (_, q) in pgrid.query((c.x - 50, c.y - 50, c.x + 50, c.y + 50))]
                dists.append(min(ds) if ds else 1e9)
            window = sum(1 for d in dists if 8.0 <= d <= 45.0) / max(1, len(dists))
            close = sum(1 for d in dists if d <= 20.0) / max(1, len(dists))
            conn = 0
            inst = pts[:: max(1, len(pts) // 200)]
            for p in inst:
                for (_, (here, other)) in end_grid.query((p.x - 1, p.y - 1, p.x + 1, p.y + 1)):
                    if (here - p).magnitude <= 1.0 and near_closure(other, 2.0):
                        conn += 1
                        break
            stats.append((layer, block, len(pts), window, close, conn / max(1, len(inst))))
        terminals = sorted([s_ for s_ in stats if s_[5] >= 0.5], key=lambda s_: -s_[2])
        if terminals:
            cad.terminal_layer = terminals[0][0]
        hh = [s_ for s_ in stats if s_[5] < 0.2 and s_[3] >= 0.05]
        if hh:
            by_layer: Counter = Counter()
            for layer, block, n, window, close, conn in hh:
                by_layer[layer] += window
            top = by_layer.most_common(1)[0][0]
            cad.handhole_layer = top
            cad.handhole_blocks = [b for (l, b, n, w, c, k) in sorted(hh, key=lambda s_: -s_[3]) if l == top and w >= 0.05]
        cables = [l for l, n in layer_hits.items() if n >= 3 and n / max(1, layer_total[l]) >= 0.3
                  and "PIGTAIL" not in l.upper() and "DROP" not in l.upper()]
        if cables:
            cad.cable_layers = sorted(cables, key=lambda l: -layer_hits[l])
        return cad

    # ------------------------------------------------------------------ core
    def _load_all(self, issues: IssueLog) -> Tuple[Any, AddressBook, List[SpliceReport], DrawingModel]:
        cfg = self.cfg
        self.rep.progress(2, "Collecting input files")
        addr_files = self.collector.expand(cfg.address_paths, (".xlsx", ".xlsm", ".xls", ".csv", ".txt"), issues)
        spl_files = self.collector.expand(cfg.splice_paths, (".xlsx", ".xlsm", ".xls", ".csv"), issues)
        if not addr_files:
            issues.add("ERROR", "Input", "Address List", "No address list files found")
        if not spl_files:
            issues.add("WARNING", "Input", "Splice Report", "No splice report files found - no BFO callouts will be produced")
        self.rep.log(f"Inputs: {len(addr_files)} address list file(s), {len(spl_files)} splice report file(s)")
        self.rep.progress(5, "Reading address lists")
        book = AddressBook()
        cols = cfg.columns
        if not all(getattr(cols, k) for k, (req, _, _) in ADDRESS_FIELDS.items() if req):
            headers = AddressBook.scan_headers(addr_files, issues)
            auto = auto_map_columns(headers)
            for k in ADDRESS_FIELDS:
                if not getattr(cols, k):
                    setattr(cols, k, getattr(auto, k))
        book.load(addr_files, cols, cfg.ports_per_pon, issues, self.rep)
        reports = load_splice_reports(spl_files, issues, self.rep)
        self.rep.progress(28, "Reading drawing")
        doc = open_drawing(cfg.drawing_path, cfg.oda_path, self.rep)
        self.rep.progress(40, "Building drawing model")
        model = DrawingModel(doc, cfg.cad, cfg.spec, issues, self.rep)
        return doc, book, reports, model

    def _cross_validate(self, book: AddressBook, reports: List[SpliceReport], model: DrawingModel,
                        issues: IssueLog) -> None:
        for c in sorted({r.splice for r in book.records if r.splice}):
            if c not in model.closures:
                n = sum(1 for r in book.records if r.splice == c)
                issues.add("ERROR", "Missing Closure", c, f"Referenced by {n} address row(s) but not found in drawing")
        for r in reports:
            if r.name not in model.closures:
                issues.add("ERROR", "Missing Closure", r.name, f"Splice report {r.file} has no closure in drawing")
        used: Set[Tuple[str, int]] = set()
        report_closures = {r.name for r in reports}
        mism = 0
        for r in reports:
            for c in r.conns:
                if not c.address or any(s in norm_text(c.state) for s in DEAD_STATES):
                    continue
                rec = book.resolve(c.tap, c.address, None, c.port_idx)
                if rec:
                    used.add((rec.olm, rec.number))
                    for sg in r.segments:
                        if sg.kind.lower() == "pigtail" and sg.start == r.name and sg.direction == "out" \
                                and (c.a_seg == (sg.cable, sg.seg) or c.b_seg == (sg.cable, sg.seg)) \
                                and rec.splice and rec.splice != r.name:
                            mism += 1
                            issues.add("WARNING", "Splice Mismatch", rec.address,
                                       f"Address list drop splice {rec.splice} but splice report {r.file} drops it at {r.name}")
        for rec in book.records:
            if rec.splice in report_closures and (rec.olm, rec.number) not in used:
                sev = "INFO" if "POWER" in rec.address else "WARNING"
                issues.add(sev, "Orphan Address", f"{rec.olm} #{rec.number}",
                           f"{rec.address} ({rec.tap}) not carried by any live fiber in the splice reports")
        for name, cl in model.closures.items():
            if name in {r.splice for r in book.records} and model.handhole_for(cl) is None:
                issues.add("INFO", "Placement", name, "No handhole near this drop closure; closure symbol will be used")
        if mism:
            self.rep.log(f"{mism} splice/address-list drop mismatches found", "WARNING")

    def validate(self) -> IssueLog:
        issues = IssueLog()
        try:
            _doc, book, reports, model = self._load_all(issues)
            self.rep.progress(60, "Cross-validating")
            self._cross_validate(book, reports, model, issues)
            planner = CalloutPlanner(model, book, reports, self.cfg.spec, issues, self.rep)
            planner.plan_terminals()
            planner.plan_cables()
            pq = planner.placement_qa()
            self.rep.log(f"Placement QA: {pq.get('clean', 0)} clean, {pq.get('text_overlap', 0)} text/text overlaps, "
                         f"{pq.get('symbol_overlap', 0)} text/symbol overlaps, {pq.get('leader_cross', 0)} leader crossings")
            n_t = sum(1 for p in planner.plans if p.kind == "TERMINAL" and p.status != "FAILED")
            n_c = sum(1 for p in planner.plans if p.kind == "CABLE" and p.status != "FAILED")
            self.rep.log(f"Dry run: {n_t} fiber-number callouts and {n_c} BFO callouts can be placed")
            self.rep.progress(100, "Validation complete")
            self._log_issues(issues)
            return issues
        finally:
            self.collector.cleanup()

    def _log_issues(self, issues: IssueLog) -> None:
        by = Counter((i.severity, i.category) for i in issues.items)
        for (sev, cat), n in sorted(by.items()):
            self.rep.log(f"{sev:<7} {cat}: {n}", sev)
        shown = 0
        for i in issues.items:
            if i.severity in ("ERROR", "WARNING") and shown < 40:
                self.rep.log(f"  {i.severity}: [{i.category}] {i.obj} - {i.message}", i.severity)
                shown += 1
        if len([i for i in issues.items if i.severity != "INFO"]) > shown:
            self.rep.log("  ... see the 'Validation Issues' sheet of the QA workbook for the full list")
        self.rep.log("Validation summary: " + issues.summary())

    def generate(self) -> Dict[str, Any]:
        cfg = self.cfg
        issues = IssueLog()
        try:
            doc, book, reports, model = self._load_all(issues)
            if not book.records:
                self._log_issues(issues)
                raise RuntimeError("Cannot continue: no usable address-list rows (see the issues above)")
            self.rep.progress(48, "Cross-validating")
            self._cross_validate(book, reports, model, issues)
            planner = CalloutPlanner(model, book, reports, cfg.spec, issues, self.rep)
            self.rep.progress(50, "Planning fiber-number callouts")
            planner.plan_terminals()
            planner.plan_cables()
            pq = planner.placement_qa()
            self.rep.log(f"Placement QA: {pq.get('clean', 0)} clean, {pq.get('text_overlap', 0)} text/text overlaps, "
                         f"{pq.get('symbol_overlap', 0)} text/symbol overlaps, {pq.get('leader_cross', 0)} leader crossings "
                         f"(flagged 'PLACEMENT REVIEW' in the QA workbook)")
            plans = planner.plans
            self.rep.progress(82, "Comparing with existing callouts")
            existing_rows = compare_existing(plans, model.existing) if model.existing else []
            if model.existing:
                ok = sum(1 for p in plans if p.existing_match.startswith("MATCH"))
                self.rep.log(f"Existing callouts: {len(model.existing)}; computed callouts matching an existing one: {ok}")
            msp = doc.modelspace()
            if cfg.existing_mode == "replace" and model.existing:
                for ex in model.existing:
                    ent = doc.entitydb.get(ex.handle)
                    if ent is not None and ent.is_alive:
                        msp.delete_entity(ent)
                self.rep.log(f"Removed {len(model.existing)} existing callout(s) from output layer(s) (mode: replace)")
            self.rep.progress(85, "Writing callouts")
            writer = CalloutWriter(doc, cfg.spec, issues, self.rep)
            placed = 0
            for i, p in enumerate(plans):
                self.rep.check_cancel()
                if p.status == "FAILED":
                    continue
                if cfg.existing_mode == "skip" and p.existing_match.startswith("MATCH"):
                    p.status = "SKIPPED (already in drawing)"
                    continue
                try:
                    p.handle = writer.add(p)
                    p.status = "PLACED"
                    placed += 1
                except Exception as exc:  # noqa: BLE001 - one bad callout must not stop the run
                    p.status = "FAILED"
                    p.notes = f"Write error: {exc}"
                    issues.add("ERROR", "CAD Write", p.uid, str(exc))
                if i % 200 == 0:
                    self.rep.progress(85 + 8 * i / max(1, len(plans)), f"Writing callouts ({i}/{len(plans)})")
            out_dir = Path(cfg.output_dir or Path(cfg.drawing_path).parent)
            out_dir.mkdir(parents=True, exist_ok=True)
            stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
            stem = Path(cfg.drawing_path).stem
            dxf_path = out_dir / f"{stem}_FIBER_CALLOUTS_{stamp}.dxf"
            self.rep.progress(94, "Saving drawing")
            doc.saveas(str(dxf_path))
            outputs = {"dxf": str(dxf_path)}
            self.rep.log(f"Saved drawing: {dxf_path}")
            if cfg.export_dwg:
                try:
                    from ezdxf.addons import odafc
                    if cfg.oda_path:
                        ezdxf.options.set("odafc-addon", "win_exec_path" if os.name == "nt" else "unix_exec_path", cfg.oda_path)
                    if odafc.is_installed():
                        dwg_path = out_dir / f"{stem}_FIBER_CALLOUTS_{stamp}.dwg"
                        odafc.export_dwg(doc, str(dwg_path), version=cfg.dwg_version, replace=True)
                        outputs["dwg"] = str(dwg_path)
                        self.rep.log(f"Saved DWG ({cfg.dwg_version}): {dwg_path}")
                    else:
                        issues.add("WARNING", "Export", "DWG", "ODA File Converter not installed - DWG export skipped (DXF written)")
                except Exception as exc:  # noqa: BLE001
                    issues.add("WARNING", "Export", "DWG", f"DWG export failed: {exc}")
            n_t = sum(1 for p in plans if p.kind == "TERMINAL" and p.status == "PLACED")
            n_c = sum(1 for p in plans if p.kind == "CABLE" and p.status == "PLACED")
            summary = {
                "Tool": f"{APP_NAME} {APP_VERSION}",
                "Run at": _dt.datetime.now().isoformat(timespec="seconds"),
                "Drawing (input)": cfg.drawing_path,
                "Drawing (output)": str(dxf_path),
                "Address rows (PON assigned)": len(book.records),
                "Terminals": len(book.terminals),
                "OLMs": ", ".join(book.olms),
                "Ports per PON": book.ports_per_pon,
                "Splice reports parsed": len(reports),
                "Closures in drawing": len(model.closures),
                "Handholes in drawing": len(model.handholes),
                "Cable segments evaluated": len(planner.segments),
                "Fiber-number callouts placed": n_t,
                "BFO callouts placed": n_c,
                "Callouts failed": sum(1 for p in plans if p.status == "FAILED"),
                "Callouts skipped (existing)": sum(1 for p in plans if p.status.startswith("SKIPPED")),
                "Existing callouts found": len(model.existing),
                "Existing callouts matched": sum(1 for p in plans if p.existing_match.startswith("MATCH")),
                "Existing-callout mode": cfg.existing_mode,
                "Placement - clean callouts": pq.get("clean", 0),
                "Placement - text/text overlaps": pq.get("text_overlap", 0),
                "Placement - text/symbol overlaps": pq.get("symbol_overlap", 0),
                "Placement - leader crossings": pq.get("leader_cross", 0),
                "Errors": issues.count("ERROR"),
                "Warnings": issues.count("WARNING"),
                "Formula": f"N = (PON - 1) x {book.ports_per_pon} + PON PORT (per OLM)",
            }
            qa_path = out_dir / f"{stem}_FIBER_CALLOUTS_QA_{stamp}.xlsx"
            self.rep.progress(97, "Writing QA/QC workbook")
            write_qa_workbook(qa_path, cfg, plans, planner.segments, planner.terminal_rows, issues, existing_rows, summary)
            outputs["qa"] = str(qa_path)
            self.rep.log(f"Saved QA/QC workbook: {qa_path}")
            self._log_issues(issues)
            self.rep.log(f"DONE - {n_t} fiber-number + {n_c} BFO callouts placed", "SUCCESS")
            self.rep.progress(100, "Done")
            return dict(outputs=outputs, summary=summary, issues=issues)
        finally:
            self.collector.cleanup()


# =============================================================================
# GUI (PyQt6)
# =============================================================================
def run_gui(on_ready: Optional[Callable[[Any, Any], None]] = None) -> int:
    """Start the PyQt6 GUI. `on_ready(window, app)` is invoked once the event loop runs (automation hook)."""
    try:
        from PyQt6.QtCore import QObject, Qt, QThread, QUrl, pyqtSignal
        from PyQt6.QtGui import QColor, QDesktopServices, QFont, QTextCharFormat, QTextCursor
        from PyQt6.QtWidgets import (QAbstractItemView, QApplication, QCheckBox, QComboBox, QDoubleSpinBox,
                                     QFileDialog, QFormLayout, QGridLayout, QGroupBox, QHBoxLayout, QLabel,
                                     QLineEdit, QListWidget, QListWidgetItem, QMainWindow, QMessageBox,
                                     QProgressBar, QPushButton, QSpinBox, QSplitter, QTabWidget,
                                     QTextEdit, QVBoxLayout, QWidget)
    except ImportError:
        sys.stderr.write("PyQt6 is required for the GUI: pip install PyQt6   (or use --cli)\n")
        return 2

    LEVEL_COLORS = {"ERROR": "#C00000", "WARNING": "#B07000", "SUCCESS": "#00703C", "DEBUG": "#808080", "INFO": "#202020"}

    class Worker(QObject):
        log = pyqtSignal(str, str)
        progress = pyqtSignal(int, str)
        finished = pyqtSignal(str, object)
        failed = pyqtSignal(str, str)

        def __init__(self, task: str, cfg: RunConfig, cancel: threading.Event) -> None:
            super().__init__()
            self.task, self.cfg, self.cancel = task, cfg, cancel

        def run(self) -> None:
            rep = Reporter(lambda m, l: self.log.emit(m, l), lambda p, m: self.progress.emit(p, m), self.cancel)
            try:
                eng = CalloutEngine(self.cfg, rep)
                if self.task == "scan":
                    res: Any = eng.scan()
                elif self.task == "validate":
                    res = eng.validate()
                else:
                    res = eng.generate()
                self.finished.emit(self.task, res)
            except Cancelled:
                self.failed.emit(self.task, "Cancelled by user")
            except Exception as exc:  # noqa: BLE001 - surface every failure to the GUI
                self.log.emit(traceback.format_exc(), "DEBUG")
                self.failed.emit(self.task, f"{type(exc).__name__}: {exc}")

    class FileList(QWidget):
        def __init__(self, title: str, filt: str) -> None:
            super().__init__()
            self.filt = filt
            lay = QVBoxLayout(self)
            lay.setContentsMargins(0, 0, 0, 0)
            self.list = QListWidget()
            self.list.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
            self.list.setMinimumHeight(70)
            row = QHBoxLayout()
            for label, fn in (("Add files...", self.add_files), ("Add folder...", self.add_folder),
                              ("Remove", self.remove), ("Clear", self.list.clear)):
                b = QPushButton(label)
                b.clicked.connect(fn)
                row.addWidget(b)
            row.addStretch(1)
            lay.addWidget(QLabel(title))
            lay.addWidget(self.list)
            lay.addLayout(row)

        def add_files(self) -> None:
            files, _ = QFileDialog.getOpenFileNames(self, "Select files", "", self.filt)
            for f in files:
                self.add(f)

        def add_folder(self) -> None:
            d = QFileDialog.getExistingDirectory(self, "Select folder")
            if d:
                self.add(d)

        def add(self, path: str) -> None:
            if path and path not in self.paths():
                self.list.addItem(QListWidgetItem(path))

        def remove(self) -> None:
            for it in self.list.selectedItems():
                self.list.takeItem(self.list.row(it))

        def paths(self) -> List[str]:
            return [self.list.item(i).text() for i in range(self.list.count())]

    class MainWindow(QMainWindow):
        def __init__(self) -> None:
            super().__init__()
            self.setWindowTitle(f"{APP_NAME} {APP_VERSION}")
            self.resize(1180, 860)
            self.cfg = RunConfig()
            self.thread: Optional[QThread] = None
            self.worker: Optional[Worker] = None
            self.cancel = threading.Event()
            self.last_outputs: Dict[str, str] = {}
            self._build_ui()

        # ------------------------------------------------------------- layout
        def _build_ui(self) -> None:
            central = QWidget()
            root = QVBoxLayout(central)
            self.tabs = QTabWidget()
            self.tabs.addTab(self._tab_inputs(), "1  Inputs")
            self.tabs.addTab(self._tab_columns(), "2  Column Mapper")
            self.tabs.addTab(self._tab_layers(), "3  Layer Mapper")
            self.tabs.addTab(self._tab_style(), "4  Callout Style")
            splitter = QSplitter(Qt.Orientation.Vertical)
            splitter.addWidget(self.tabs)
            bottom = QWidget()
            bl = QVBoxLayout(bottom)
            bl.setContentsMargins(0, 0, 0, 0)
            btns = QHBoxLayout()
            self.btn_scan = QPushButton("Scan Inputs / Auto-Detect")
            self.btn_validate = QPushButton("Validate Data")
            self.btn_generate = QPushButton("Generate Callouts")
            self.btn_cancel = QPushButton("Cancel")
            self.btn_open = QPushButton("Open Output Folder")
            self.btn_generate.setStyleSheet("font-weight: bold;")
            self.btn_cancel.setEnabled(False)
            self.btn_open.setEnabled(False)
            self.btn_scan.clicked.connect(lambda: self.start("scan"))
            self.btn_validate.clicked.connect(lambda: self.start("validate"))
            self.btn_generate.clicked.connect(lambda: self.start("generate"))
            self.btn_cancel.clicked.connect(self.cancel_job)
            self.btn_open.clicked.connect(self.open_output)
            for b in (self.btn_scan, self.btn_validate, self.btn_generate, self.btn_cancel, self.btn_open):
                btns.addWidget(b)
            bl.addLayout(btns)
            self.progress = QProgressBar()
            self.progress.setRange(0, 100)
            self.status = QLabel("Ready")
            pr = QHBoxLayout()
            pr.addWidget(self.progress, 3)
            pr.addWidget(self.status, 2)
            bl.addLayout(pr)
            self.console = QTextEdit()
            self.console.setReadOnly(True)
            self.console.setFont(QFont("Consolas" if os.name == "nt" else "Monospace", 9))
            bl.addWidget(self.console)
            splitter.addWidget(bottom)
            splitter.setSizes([480, 380])
            root.addWidget(splitter)
            self.setCentralWidget(central)

        def _path_row(self, edit: QLineEdit, cb: Callable[[], None]) -> QWidget:
            w = QWidget()
            h = QHBoxLayout(w)
            h.setContentsMargins(0, 0, 0, 0)
            h.addWidget(edit)
            b = QPushButton("Browse...")
            b.clicked.connect(cb)
            h.addWidget(b)
            return w

        def _tab_inputs(self) -> QWidget:
            w = QWidget()
            lay = QVBoxLayout(w)
            form = QFormLayout()
            self.ed_drawing = QLineEdit()
            self.ed_output = QLineEdit()
            self.ed_oda = QLineEdit()
            self.ed_oda.setPlaceholderText("Optional: ODAFileConverter executable (needed for .dwg)")
            form.addRow("Target drawing (.dxf / .dwg):", self._path_row(self.ed_drawing, self.pick_drawing))
            form.addRow("Output folder:", self._path_row(self.ed_output, self.pick_output))
            form.addRow("ODA File Converter:", self._path_row(self.ed_oda, self.pick_oda))
            lay.addLayout(form)
            self.fl_splice = FileList("Splice reports (files, folders or .zip):",
                                      "Splice reports (*.xlsx *.xlsm *.xls *.csv *.zip);;All files (*)")
            self.fl_addr = FileList("Address lists (files, folders or .zip):",
                                    "Address lists (*.xlsx *.xlsm *.xls *.csv *.zip);;All files (*)")
            lay.addWidget(self.fl_splice)
            lay.addWidget(self.fl_addr)
            opts = QGroupBox("Run options")
            g = QGridLayout(opts)
            self.cb_existing = QComboBox()
            self.cb_existing.addItems(["replace", "skip", "append"])
            self.cb_existing.setToolTip("replace: delete existing callouts on the output layer(s) first\n"
                                        "skip: keep existing, do not duplicate matching callouts\n"
                                        "append: keep existing and add all computed callouts")
            self.chk_dwg = QCheckBox("Also export DWG (needs ODA File Converter)")
            self.cb_dwgver = QComboBox()
            self.cb_dwgver.addItems(["R2018", "R2013", "R2010", "R2007", "R2004", "R2000"])
            self.sp_ports = QSpinBox()
            self.sp_ports.setRange(0, 1024)
            self.sp_ports.setSpecialValueText("auto")
            self.sp_ports.setToolTip("Ports per PON used in N = (PON-1) x ports + port. 'auto' = next power of two >= max port")
            g.addWidget(QLabel("Existing callouts:"), 0, 0)
            g.addWidget(self.cb_existing, 0, 1)
            g.addWidget(QLabel("Ports per PON:"), 0, 2)
            g.addWidget(self.sp_ports, 0, 3)
            g.addWidget(self.chk_dwg, 1, 0, 1, 2)
            g.addWidget(QLabel("DWG version:"), 1, 2)
            g.addWidget(self.cb_dwgver, 1, 3)
            lay.addWidget(opts)
            return w

        def _tab_columns(self) -> QWidget:
            w = QWidget()
            lay = QVBoxLayout(w)
            lay.addWidget(QLabel("Map address-list columns. 'Scan Inputs' auto-detects; override any field here. "
                                 "Required fields are marked *."))
            form = QFormLayout()
            self.col_boxes: Dict[str, QComboBox] = {}
            for k, (req, _syn, desc) in ADDRESS_FIELDS.items():
                cb = QComboBox()
                cb.setEditable(True)
                cb.addItem("")
                self.col_boxes[k] = cb
                form.addRow(f"{k.upper()}{' *' if req else ''}  ({desc}):", cb)
            lay.addLayout(form)
            lay.addStretch(1)
            return w

        def _tab_layers(self) -> QWidget:
            w = QWidget()
            lay = QVBoxLayout(w)
            lay.addWidget(QLabel("Map drawing content. 'Scan Inputs' auto-detects from the drawing + splice/address data."))
            form = QFormLayout()
            self.cb_cl_layer = QComboBox(); self.cb_cl_layer.setEditable(True)
            self.cb_cl_block = QComboBox(); self.cb_cl_block.setEditable(True)
            self.cb_cl_tag = QComboBox(); self.cb_cl_tag.setEditable(True)
            self.cb_hh_layer = QComboBox(); self.cb_hh_layer.setEditable(True)
            self.cb_term_layer = QComboBox(); self.cb_term_layer.setEditable(True)
            self.cb_out_cable = QComboBox(); self.cb_out_cable.setEditable(True)
            self.cb_out_term = QComboBox(); self.cb_out_term.setEditable(True)
            d = CadMap()
            s = CalloutSpec()
            for cb, val in ((self.cb_cl_layer, d.closure_layer), (self.cb_cl_block, d.closure_block),
                            (self.cb_cl_tag, d.closure_name_tag), (self.cb_hh_layer, d.handhole_layer),
                            (self.cb_term_layer, d.terminal_layer), (self.cb_out_cable, s.cable_layer),
                            (self.cb_out_term, s.terminal_layer)):
                cb.addItem(val)
                cb.setCurrentText(val)
            self.cb_cl_layer.currentTextChanged.connect(self._refresh_blocks)
            self.cb_hh_layer.currentTextChanged.connect(self._refresh_hh_blocks)
            self.lst_hh = QListWidget()
            self.lst_hh.setMaximumHeight(110)
            self.lst_cables = QListWidget()
            self.lst_cables.setMaximumHeight(130)
            for b in d.handhole_blocks:
                it = QListWidgetItem(b); it.setFlags(it.flags() | Qt.ItemFlag.ItemIsUserCheckable); it.setCheckState(Qt.CheckState.Checked)
                self.lst_hh.addItem(it)
            for l in d.cable_layers:
                it = QListWidgetItem(l); it.setFlags(it.flags() | Qt.ItemFlag.ItemIsUserCheckable); it.setCheckState(Qt.CheckState.Checked)
                self.lst_cables.addItem(it)
            self.sp_snap = QDoubleSpinBox(); self.sp_snap.setRange(0.01, 50); self.sp_snap.setValue(d.snap_tolerance)
            form.addRow("Closure layer:", self.cb_cl_layer)
            form.addRow("Closure block (blank = any):", self.cb_cl_block)
            form.addRow("Closure name attribute tag:", self.cb_cl_tag)
            form.addRow("Handhole layer:", self.cb_hh_layer)
            form.addRow("Handhole blocks (checked):", self.lst_hh)
            form.addRow("Terminal layer (obstacle only):", self.cb_term_layer)
            form.addRow("Cable layers (checked):", self.lst_cables)
            form.addRow("Cable->closure snap tolerance:", self.sp_snap)
            form.addRow("Output layer - BFO (cable) callouts:", self.cb_out_cable)
            form.addRow("Output layer - fiber number (terminal) callouts:", self.cb_out_term)
            lay.addLayout(form)
            self._scan: Optional[ScanResult] = None
            return w

        def _tab_style(self) -> QWidget:
            w = QWidget()
            lay = QVBoxLayout(w)
            lay.addWidget(QLabel("Defaults reproduce the reverse-engineered DES_FIBER_SPL standard exactly."))
            form = QFormLayout()
            s = CalloutSpec()
            self.style_widgets: Dict[str, Any] = {}

            def dspin(name: str, val: float, lo: float, hi: float, dec: int = 3) -> None:
                sp = QDoubleSpinBox(); sp.setRange(lo, hi); sp.setDecimals(dec); sp.setValue(val)
                self.style_widgets[name] = sp
                form.addRow(name.replace("_", " ").capitalize() + ":", sp)

            def ispin(name: str, val: int, lo: int, hi: int) -> None:
                sp = QSpinBox(); sp.setRange(lo, hi); sp.setValue(val)
                self.style_widgets[name] = sp
                form.addRow(name.replace("_", " ").capitalize() + ":", sp)

            def text(name: str, val: str) -> None:
                ed = QLineEdit(val)
                self.style_widgets[name] = ed
                form.addRow(name.replace("_", " ").capitalize() + ":", ed)

            dspin("text_height", s.text_height, 0.01, 1000)
            dspin("arrow_size", s.arrow_size, 0.0, 1000)
            dspin("landing_gap", s.landing_gap, 0.0, 100)
            dspin("dogleg_length", s.dogleg_length, 0.0, 100)
            dspin("terminal_leader_length", s.terminal_leader_length, 0.5, 1000, 2)
            dspin("cable_leader_length", s.cable_leader_length, 0.5, 1000, 2)
            dspin("cable_offset_along", s.cable_offset_along, 0.0, 1000, 2)
            dspin("handhole_search_radius", s.handhole_search_radius, 1.0, 5000, 1)
            ispin("layer_color", s.layer_color, 1, 255)
            ispin("zero_pad", s.zero_pad, 0, 6)
            text("cable_prefix", s.cable_prefix)
            text("text_style", s.text_style)
            text("text_style_font", s.text_style_font)
            text("mleader_style", s.mleader_style)
            text("layer_linetype", s.layer_linetype)
            cb = QComboBox(); cb.addItems(["span", "exact"]); cb.setCurrentText(s.cable_range_mode)
            self.style_widgets["cable_range_mode"] = cb
            form.addRow("BFO range format (span = min-max as reference):", cb)
            ch = QCheckBox("Split non-contiguous terminal numbers into separate callouts (reference behaviour)")
            ch.setChecked(s.split_terminal_runs)
            self.style_widgets["split_terminal_runs"] = ch
            form.addRow("", ch)
            lay.addLayout(form)
            lay.addStretch(1)
            return w

        # ------------------------------------------------------------- pickers
        def pick_drawing(self) -> None:
            f, _ = QFileDialog.getOpenFileName(self, "Target drawing", "", "CAD (*.dxf *.dwg);;All files (*)")
            if f:
                self.ed_drawing.setText(f)
                if not self.ed_output.text():
                    self.ed_output.setText(str(Path(f).parent / "callout_output"))

        def pick_output(self) -> None:
            d = QFileDialog.getExistingDirectory(self, "Output folder")
            if d:
                self.ed_output.setText(d)

        def pick_oda(self) -> None:
            f, _ = QFileDialog.getOpenFileName(self, "ODA File Converter executable")
            if f:
                self.ed_oda.setText(f)

        def _refresh_blocks(self) -> None:
            if not self._scan:
                return
            layer = self.cb_cl_layer.currentText()
            cur = self.cb_cl_block.currentText()
            self.cb_cl_block.blockSignals(True)
            self.cb_cl_block.clear()
            self.cb_cl_block.addItem("")
            for l, b, n in self._scan.insert_pairs:
                if l == layer:
                    self.cb_cl_block.addItem(b)
            self.cb_cl_block.setCurrentText(cur)
            self.cb_cl_block.blockSignals(False)
            tags = set()
            for key, t in self._scan.insert_tags.items():
                if key.split("|")[0] == layer:
                    tags.update(t)
            curt = self.cb_cl_tag.currentText()
            self.cb_cl_tag.clear()
            self.cb_cl_tag.addItems(sorted(tags) or [curt])
            self.cb_cl_tag.setCurrentText(curt)

        def _refresh_hh_blocks(self, checked: Optional[List[str]] = None) -> None:
            if not self._scan:
                return
            layer = self.cb_hh_layer.currentText()
            keep = set(checked) if checked is not None else {self.lst_hh.item(i).text() for i in range(self.lst_hh.count())
                                                              if self.lst_hh.item(i).checkState() == Qt.CheckState.Checked}
            self.lst_hh.clear()
            for l, b, n in self._scan.insert_pairs:
                if l == layer:
                    it = QListWidgetItem(b)
                    it.setFlags(it.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                    it.setCheckState(Qt.CheckState.Checked if b in keep else Qt.CheckState.Unchecked)
                    it.setToolTip(f"{n} instance(s)")
                    self.lst_hh.addItem(it)

        # ------------------------------------------------------------- config
        def collect_config(self) -> RunConfig:
            cfg = RunConfig()
            cfg.drawing_path = self.ed_drawing.text().strip()
            cfg.output_dir = self.ed_output.text().strip()
            cfg.oda_path = self.ed_oda.text().strip()
            cfg.splice_paths = self.fl_splice.paths()
            cfg.address_paths = self.fl_addr.paths()
            cfg.existing_mode = self.cb_existing.currentText()
            cfg.export_dwg = self.chk_dwg.isChecked()
            cfg.dwg_version = self.cb_dwgver.currentText()
            cfg.ports_per_pon = self.sp_ports.value()
            for k, cb in self.col_boxes.items():
                setattr(cfg.columns, k, cb.currentText().strip() or None)
            cfg.cad = CadMap(
                closure_layer=self.cb_cl_layer.currentText().strip(),
                closure_block=self.cb_cl_block.currentText().strip(),
                closure_name_tag=self.cb_cl_tag.currentText().strip() or "NAME",
                handhole_layer=self.cb_hh_layer.currentText().strip(),
                handhole_blocks=[self.lst_hh.item(i).text() for i in range(self.lst_hh.count())
                                 if self.lst_hh.item(i).checkState() == Qt.CheckState.Checked],
                terminal_layer=self.cb_term_layer.currentText().strip(),
                cable_layers=[self.lst_cables.item(i).text() for i in range(self.lst_cables.count())
                              if self.lst_cables.item(i).checkState() == Qt.CheckState.Checked],
                snap_tolerance=self.sp_snap.value())
            sp = CalloutSpec()
            for name, wdg in self.style_widgets.items():
                if isinstance(wdg, (QDoubleSpinBox, QSpinBox)):
                    setattr(sp, name, wdg.value())
                elif isinstance(wdg, QLineEdit):
                    setattr(sp, name, wdg.text().strip())
                elif isinstance(wdg, QComboBox):
                    setattr(sp, name, wdg.currentText())
                elif isinstance(wdg, QCheckBox):
                    setattr(sp, name, wdg.isChecked())
            sp.cable_layer = self.cb_out_cable.currentText().strip() or sp.cable_layer
            sp.terminal_layer = self.cb_out_term.currentText().strip() or sp.terminal_layer
            cfg.spec = sp
            return cfg

        def apply_scan(self, res: ScanResult) -> None:
            self._scan = res
            for k, cb in self.col_boxes.items():
                cur = getattr(res.columns, k) or ""
                cb.clear()
                cb.addItem("")
                cb.addItems(res.headers)
                cb.setCurrentText(cur)
            if res.cad:
                layers = res.layers
                for cb in (self.cb_cl_layer, self.cb_hh_layer, self.cb_term_layer, self.cb_out_cable, self.cb_out_term):
                    cur = cb.currentText()
                    cb.blockSignals(True)
                    cb.clear()
                    cb.addItems(layers)
                    cb.setCurrentText(cur)
                    cb.blockSignals(False)
                self.cb_cl_layer.setCurrentText(res.cad.closure_layer)
                self._refresh_blocks()
                self.cb_cl_block.setCurrentText(res.cad.closure_block)
                self.cb_cl_tag.setCurrentText(res.cad.closure_name_tag)
                self.cb_hh_layer.setCurrentText(res.cad.handhole_layer)
                self._refresh_hh_blocks(res.cad.handhole_blocks)
                self.cb_term_layer.setCurrentText(res.cad.terminal_layer)
                self.lst_cables.clear()
                for l, n in res.poly_layers:
                    it = QListWidgetItem(l)
                    it.setFlags(it.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                    it.setCheckState(Qt.CheckState.Checked if l in res.cad.cable_layers else Qt.CheckState.Unchecked)
                    it.setToolTip(f"{n} linear entities")
                    self.lst_cables.addItem(it)

        # ------------------------------------------------------------- jobs
        def start(self, task: str) -> None:
            if self.thread is not None:
                return
            cfg = self.collect_config()
            if task != "scan" and not cfg.drawing_path:
                QMessageBox.warning(self, APP_NAME, "Select a target drawing first.")
                return
            if task == "scan" and not (cfg.drawing_path or cfg.address_paths):
                QMessageBox.warning(self, APP_NAME, "Select a drawing and/or address list(s) to scan.")
                return
            if task == "generate" and not cfg.output_dir:
                cfg.output_dir = str(Path(cfg.drawing_path).parent / "callout_output")
                self.ed_output.setText(cfg.output_dir)
            self.cancel.clear()
            self.progress.setValue(0)
            self._log(f"===== {task.upper()} started =====", "INFO")
            self._busy(True)
            self.thread = QThread(self)
            self.worker = Worker(task, cfg, self.cancel)
            self.worker.moveToThread(self.thread)
            self.thread.started.connect(self.worker.run)
            self.worker.log.connect(self._log)
            self.worker.progress.connect(self._progress)
            self.worker.finished.connect(self._done)
            self.worker.failed.connect(self._failed)
            self.worker.finished.connect(self.thread.quit)
            self.worker.failed.connect(self.thread.quit)
            self.thread.finished.connect(self._cleanup_thread)
            self.thread.start()

        def _cleanup_thread(self) -> None:
            if self.worker is not None:
                self.worker.deleteLater()
            if self.thread is not None:
                self.thread.deleteLater()
            self.worker = None
            self.thread = None
            self._busy(False)

        def _busy(self, busy: bool) -> None:
            for b in (self.btn_scan, self.btn_validate, self.btn_generate):
                b.setEnabled(not busy)
            self.btn_cancel.setEnabled(busy)
            self.tabs.setEnabled(not busy)

        def cancel_job(self) -> None:
            self.cancel.set()
            self._log("Cancellation requested ...", "WARNING")

        def _log(self, msg: str, level: str) -> None:
            fmt = QTextCharFormat()
            fmt.setForeground(QColor(LEVEL_COLORS.get(level, "#202020")))
            cur = self.console.textCursor()
            cur.movePosition(QTextCursor.MoveOperation.End)
            stamp = _dt.datetime.now().strftime("%H:%M:%S")
            cur.insertText(f"[{stamp}] {level:<7} {msg}\n", fmt)
            self.console.setTextCursor(cur)
            self.console.ensureCursorVisible()

        def _progress(self, pct: int, msg: str) -> None:
            self.progress.setValue(pct)
            if msg:
                self.status.setText(msg)

        def _done(self, task: str, res: Any) -> None:
            if task == "scan" and isinstance(res, ScanResult):
                self.apply_scan(res)
                self.status.setText("Scan complete - review the mapper tabs")
            elif task == "validate" and isinstance(res, IssueLog):
                self.status.setText("Validation: " + res.summary())
                if res.count("ERROR"):
                    QMessageBox.warning(self, APP_NAME, "Validation found errors:\n" + res.summary() +
                                        "\n\nSee the console. You can still generate; affected callouts are skipped and listed in the QA report.")
            elif task == "generate" and isinstance(res, dict):
                self.last_outputs = res.get("outputs", {})
                self.btn_open.setEnabled(True)
                s = res.get("summary", {})
                self.status.setText(f"Done: {s.get('Fiber-number callouts placed', 0)} fiber-number + "
                                    f"{s.get('BFO callouts placed', 0)} BFO callouts")
                QMessageBox.information(self, APP_NAME, "Finished.\n\n" + "\n".join(f"{k}: {v}" for k, v in self.last_outputs.items()))

        def _failed(self, task: str, msg: str) -> None:
            self._log(f"{task} failed: {msg}", "ERROR")
            self.status.setText(f"{task} failed")
            if msg != "Cancelled by user":
                QMessageBox.critical(self, APP_NAME, f"{task} failed:\n{msg}")

        def open_output(self) -> None:
            folder = Path(self.last_outputs.get("dxf", self.ed_output.text() or ".")).parent
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(folder)))

        def closeEvent(self, event: Any) -> None:  # noqa: N802 - Qt API
            if self.thread is not None:
                self.cancel.set()
                self.thread.quit()
                self.thread.wait(5000)
            event.accept()

    app = QApplication.instance() or QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    win = MainWindow()
    win.show()
    if on_ready is not None:
        from PyQt6.QtCore import QTimer
        QTimer.singleShot(0, lambda: on_ready(win, app))
    return app.exec()


# =============================================================================
# CLI
# =============================================================================
def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=f"{APP_NAME} {APP_VERSION}")
    p.add_argument("--cli", action="store_true", help="run without GUI")
    p.add_argument("--drawing", help="target .dxf/.dwg")
    p.add_argument("--splice", nargs="*", default=[], help="splice report files/folders/zips")
    p.add_argument("--address", nargs="*", default=[], help="address list files/folders/zips")
    p.add_argument("--out", help="output folder")
    p.add_argument("--validate-only", action="store_true")
    p.add_argument("--scan-only", action="store_true")
    p.add_argument("--existing", choices=["replace", "skip", "append"], default="replace")
    p.add_argument("--ports-per-pon", type=int, default=0)
    p.add_argument("--export-dwg", action="store_true")
    p.add_argument("--oda", default="", help="ODA File Converter executable")
    p.add_argument("--config", help="JSON file with 'columns', 'cad' and 'spec' overrides")
    return p


def config_from_args(a: argparse.Namespace) -> RunConfig:
    cfg = RunConfig(drawing_path=a.drawing or "", splice_paths=a.splice, address_paths=a.address,
                    output_dir=a.out or "", ports_per_pon=a.ports_per_pon, existing_mode=a.existing,
                    export_dwg=a.export_dwg, oda_path=a.oda)
    if a.config:
        data = json.loads(Path(a.config).read_text(encoding="utf-8"))
        for sec, cls in (("columns", ColumnMap), ("cad", CadMap), ("spec", CalloutSpec)):
            if sec in data:
                base = dataclasses.asdict(getattr(cfg, sec))
                base.update(data[sec])
                setattr(cfg, sec, cls(**base))
    return cfg


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_arg_parser().parse_args(argv)
    if not args.cli:
        return run_gui()
    cfg = config_from_args(args)
    rep = Reporter()
    eng = CalloutEngine(cfg, rep)
    try:
        if args.scan_only:
            res = eng.scan()
            print(json.dumps(dict(columns=res.columns.to_dict(), cad=res.cad.to_dict() if res.cad else None), indent=2))
            return 0
        if args.validate_only:
            issues = eng.validate()
            return 1 if issues.count("ERROR") else 0
        if not cfg.drawing_path:
            rep.log("--drawing is required", "ERROR")
            return 2
        res = eng.generate()
        print(json.dumps(res["outputs"], indent=2))
        return 0
    except Cancelled:
        return 130
    except Exception as exc:  # noqa: BLE001 - CLI reports any failure with a non-zero exit
        rep.log(f"{type(exc).__name__}: {exc}", "ERROR")
        rep.log(traceback.format_exc(), "DEBUG")
        return 1


if __name__ == "__main__":
    sys.exit(main())
