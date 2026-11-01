# -*- coding: utf-8 -*-
"""
Built-in self-test: generates a synthetic FTTx region that shares nothing with Credit River (different
prefixes, layers, blocks, attribute tags, header names, 1x32 splitters, rotated rectangular handholes,
both splice-report layouts, a closure named only by a text label, a closure located only from splice-report
coordinates, commercial "C-" usages, a non-contiguous terminal and a closure typo) and checks every callout
against expectations computed independently from the network's construction.
"""
from __future__ import annotations

import math
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import ezdxf
import openpyxl
from ezdxf.math import Vec2

PREFIX = "WBK"
OLM = "WBK-R07"
PORTS = 32
COLORS = ["BL", "OR", "GR", "BR", "SL", "WH", "RD", "BK", "YL", "VI", "RO", "AQ"]
LONG = {"BL": "blue", "OR": "orange", "GR": "green", "BR": "brown", "SL": "slate", "WH": "white", "RD": "red",
        "BK": "black", "YL": "yellow", "VI": "violet", "RO": "rose", "AQ": "aqua"}


def _c(n: int) -> str:
    return f"{PREFIX}-{n:05d}"


class Region:
    """Synthetic network and its expected callouts."""

    def __init__(self) -> None:
        x0, y0 = 2450000.0, 650000.0
        self.hub = _c(7001) + "G"
        self.pos: Dict[str, Tuple[float, float]] = {self.hub: (x0, y0)}
        F1 = [self.hub] + [_c(n) for n in range(7002, 7010)] + [_c(7010) + "G"]
        for i, n in enumerate(F1[1:], start=1):
            self.pos[n] = (x0 + 400.0 * i, y0 + (25.0 if i % 2 else -15.0))
        B1 = [_c(7003)] + [_c(n) for n in (7011, 7012, 7013, 7014)] + [_c(7015) + "G"]
        bx, by = self.pos[B1[0]]
        for i, n in enumerate(B1[1:], start=1):
            self.pos[n] = (bx + (12.0 if i % 2 else -8.0), by + 350.0 * i)
        B2 = [_c(7006)] + [_c(n) for n in (7016, 7017, 7018)] + [_c(7019) + "G"]
        bx, by = self.pos[B2[0]]
        for i, n in enumerate(B2[1:], start=1):
            self.pos[n] = (bx - 10.0 * i, by - 300.0 * i)
        F2 = [self.hub] + [_c(n) for n in (7021, 7022, 7023)] + [_c(7024) + "G"]
        for i, n in enumerate(F2[1:], start=1):
            self.pos[n] = (x0 - 450.0 * i, y0 + 30.0 * i)
        self.cables = [
            (f"{self.hub} TO {F1[-1]} FC144", 144, F1, "MAPLE AVE"),
            (f"{B1[0]} TO {B1[-1]} FC048", 48, B1, "OAK ST"),
            (f"{B2[0]} TO {B2[-1]} FC024", 24, B2, "PINE RD"),
            (f"{self.hub} TO {F2[-1]} FC096", 96, F2, "CEDAR LN"),
        ]
        # tree: child closure -> (parent, cable, seg)
        self.parent: Dict[str, Tuple[str, str, int]] = {}
        self.children: Dict[str, List[Tuple[str, str, int]]] = {n: [] for n in self.pos}
        for name, _fc, path, _st in self.cables:
            for k in range(1, len(path)):
                self.parent[path[k]] = (path[k - 1], name, k)
                self.children[path[k - 1]].append((path[k], name, k))
        self.fc = {name: fc for name, fc, _p, _s in self.cables}
        self.street = {}
        for name, _fc, path, st in self.cables:
            for n in path[1:]:
                self.street[n] = st
        self.street[self.hub] = "MAIN ST"
        # terminals in DFS order -> sequential fiber numbers
        self.terminals: List[Dict[str, Any]] = []
        order: List[str] = []

        def dfs(n: str) -> None:
            order.append(n)
            for ch, _cab, _k in sorted(self.children[n], key=lambda c: (c[1] != self.cables[0][0], c[0])):
                dfs(ch)
        dfs(self.hub)
        num = 0
        house = 100
        for n in order:
            n_terms = 2 if n in (_c(7004), _c(7012), _c(7017)) else 1
            for t in range(n_terms):
                n_addr = 1 + (house // 7) % 3
                addrs = []
                for _a in range(n_addr):
                    house += 4
                    num += 1
                    addrs.append((num, f"{house} {self.street[n]}"))
                tap = f"{PREFIX}-T-{addrs[0][1]}"
                self.terminals.append(dict(tap=tap, closure=n, fibers=addrs))
        self.total = num
        # non-contiguous terminal: swap one fiber between two neighbouring terminals
        ta = next(t for t in self.terminals if t["closure"] == _c(7008))
        tb = next(t for t in self.terminals if t["closure"] == _c(7009))
        if len(ta["fibers"]) > 1 and tb["fibers"]:
            (na, aa), (nb, ab) = ta["fibers"][-1], tb["fibers"][0]
            ta["fibers"][-1], tb["fibers"][0] = (nb, aa), (na, ab)
        for t in self.terminals:
            t["fibers"].sort()
        self.number_closure = {num: t["closure"] for t in self.terminals for num, _a in t["fibers"]}
        self.typo_terminal = next(t for t in self.terminals if t["closure"] == _c(7022))
        self.commercial_tap = next(t for t in self.terminals if t["closure"] == _c(7005))
        self.no_attr_label = _c(7012)          # closure with a TEXT label instead of an attribute
        self.no_attr_coords = _c(7014)         # closure found only via splice-report coordinates
        self.guide_closure = _c(7006)          # splice report in "Fiber Optic Splice Guide" layout
        self.report_closures = [self.hub, _c(7002), _c(7003), _c(7005), _c(7006), _c(7011), _c(7013), _c(7014),
                                _c(7021), _c(7023)]

    def downstream(self, n: str) -> List[int]:
        out = [num for num, c in self.number_closure.items() if c == n]
        for ch, _cab, _k in self.children[n]:
            out += self.downstream(ch)
        return sorted(out)

    # ------------------------------------------------------------------ expectations
    def expected(self) -> List[Dict[str, Any]]:
        exp = []
        for t in self.terminals:
            if t is self.typo_terminal:
                continue
            nums = [n for n, _a in t["fibers"]]
            runs: List[List[int]] = []
            for n in sorted(nums):
                if runs and n == runs[-1][-1] + 1:
                    runs[-1].append(n)
                else:
                    runs.append([n])
            for r in runs:
                exp.append(dict(kind="TERMINAL", text=str(r[0]) if len(r) == 1 else f"{r[0]}-{r[-1]}", closure=t["closure"]))
        for s in self.report_closures:
            for ch, cab, _k in self.children[s]:
                d = self.downstream(ch)
                if d:
                    rng = str(d[0]) if d[0] == d[-1] else f"{d[0]}-{d[-1]}"
                    exp.append(dict(kind="CABLE", text=f"BFO {self.fc[cab]}\\P{rng}", closure=s, next=ch))
        return exp


# ---------------------------------------------------------------------------------------------- files
def _write_drawing(reg: Region, path: Path) -> Dict[str, Any]:
    doc = ezdxf.new("R2018", setup=True)
    doc.header["$INSUNITS"] = 2
    for name, color in (("FO_SPLICE_CLOSURE", 3), ("UG_STRUCTURES", 8), ("FO_TERMINAL", 5), ("FO_CABLE_UG", 84),
                        ("FO_CABLE_AER", 21), ("FO_PIGTAIL", 7), ("ROADS", 4), ("ANNO_TEXT", 7), ("ADDRESS_PTS", 104)):
        doc.layers.add(name, color=color)
    blk = doc.blocks.new("SC_BOWTIE")
    blk.add_lwpolyline([(-10, -5), (10, 5), (10, -5), (-10, 5), (-10, -5)])
    blk.add_attdef("SPLICE_ID", (-10, -9), dxfattribs={"height": 3.0})
    hh = doc.blocks.new("HH_2436")
    hh.add_lwpolyline([(-5, -7), (5, -7), (5, 7), (-5, 7)], close=True)
    hh.add_line((-5, -7), (5, 7))
    hh.add_attdef("HH_ID", (0, 0), dxfattribs={"height": 1.0})
    term = doc.blocks.new("TERM_MST")
    term.add_circle((0, 0), 2.5)
    addr = doc.blocks.new("ADDR_PT")
    addr.add_circle((0, 0), 1.0)
    addr.add_attdef("HOUSE_NO", (2, 0), dxfattribs={"height": 2.0})
    msp = doc.modelspace()
    hh_geom: Dict[str, List[Vec2]] = {}
    for i, (n, (x, y)) in enumerate(sorted(reg.pos.items())):
        ins = msp.add_blockref("SC_BOWTIE", (x, y), dxfattribs={"layer": "FO_SPLICE_CLOSURE", "rotation": (i * 7) % 20})
        if n == reg.no_attr_label:
            msp.add_text(n, height=3.0, dxfattribs={"layer": "ANNO_TEXT", "insert": (x - 12, y + 14)})
        elif n != reg.no_attr_coords:
            ins.add_attrib("SPLICE_ID", n, insert=(x - 10, y - 9), dxfattribs={"height": 3.0, "layer": "FO_SPLICE_CLOSURE"})
        rot = (i * 13) % 35
        hpos = (x + 18.0, y + 14.0)
        h = msp.add_blockref("HH_2436", hpos, dxfattribs={"layer": "UG_STRUCTURES", "rotation": rot})
        h.add_attrib("HH_ID", f"HH{i}", insert=hpos, dxfattribs={"height": 1.0}).dxf.invisible = 1
        a = math.radians(rot)
        hh_geom[n] = [Vec2(hpos[0] + px * math.cos(a) - py * math.sin(a), hpos[1] + px * math.sin(a) + py * math.cos(a))
                      for px, py in ((-5, -7), (5, -7), (5, 7), (-5, 7))]
        # noise: a far-away structure of the same block that must not be chosen
        msp.add_blockref("HH_2436", (x + 160.0, y - 120.0), dxfattribs={"layer": "UG_STRUCTURES"})
    for t in reg.terminals:
        x, y = reg.pos[t["closure"]]
        k = sum(1 for u in reg.terminals if u["closure"] == t["closure"] and reg.terminals.index(u) < reg.terminals.index(t))
        tx, ty = x - 4.0 + 9.0 * k, y + 10.0
        msp.add_blockref("TERM_MST", (tx, ty), dxfattribs={"layer": "FO_TERMINAL"})
        msp.add_lwpolyline([(x, y), (tx, ty)], dxfattribs={"layer": "FO_PIGTAIL"})
        for j, (_n, address) in enumerate(t["fibers"]):
            ap = msp.add_blockref("ADDR_PT", (tx + 40 + 15 * j, ty + 60), dxfattribs={"layer": "ADDRESS_PTS"})
            ap.add_attrib("HOUSE_NO", address.split(" ")[0], insert=(tx + 42 + 15 * j, ty + 60), dxfattribs={"height": 2.0})
    seg_geom: Dict[Tuple[str, str], List[Vec2]] = {}
    for ci, (name, _fc, route, street) in enumerate(reg.cables):
        pts: List[Tuple[float, float]] = []
        for k in range(len(route) - 1):
            a, b = reg.pos[route[k]], reg.pos[route[k + 1]]
            mid = ((a[0] + b[0]) / 2 + (6 if k % 2 else -6), (a[1] + b[1]) / 2 + (5 if k % 2 else -4))
            piece = [a, mid, b]
            seg_geom[(route[k], route[k + 1])] = [Vec2(p) for p in piece]
            if ci == 1:                    # branch cable drawn as one polyline per segment
                msp.add_lwpolyline(piece, dxfattribs={"layer": "FO_CABLE_UG"})
            else:
                pts += piece if not pts else piece[1:]
            msp.add_text(f"{int(math.dist(a, b))}'", height=4.0,
                         dxfattribs={"layer": "ANNO_TEXT", "insert": (mid[0] + 8, mid[1] + 8)})
        if ci != 1:
            msp.add_lwpolyline(pts, dxfattribs={"layer": "FO_CABLE_AER" if ci == 3 else "FO_CABLE_UG"})
        a, b = reg.pos[route[0]], reg.pos[route[-1]]
        msp.add_lwpolyline([(a[0], a[1] - 40), (b[0], b[1] - 40)], dxfattribs={"layer": "ROADS"})
        msp.add_text(street, height=8.0, dxfattribs={"layer": "ROADS", "insert": ((a[0] + b[0]) / 2, (a[1] + b[1]) / 2 - 55)})
    doc.saveas(path)
    return dict(hh=hh_geom, seg=seg_geom)


def _write_address_list(reg: Region, path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = OLM
    ws.append(["OLT PON", "Splitter Port", "Strand", "Premise Address", "Tube", "Fiber Color", "Cable",
               "Terminal ID", "Splice Enclosure"])
    for k in range(1, 5):
        ws.append([None, None, k, f"{OLM} TRANSPORT", "BL", COLORS[k - 1], reg.cables[0][0], None, reg.hub])
    for t in reg.terminals:
        for n, address in t["fibers"]:
            pon, port = (n - 1) // PORTS + 1, (n - 1) % PORTS + 1
            closure = "WKB-07022" if t is reg.typo_terminal else t["closure"]
            ws.append([f"PON {pon}", port, n + 12, address, COLORS[(n // 12) % 12], COLORS[n % 12],
                       reg.cables[0][0], t["tap"], closure])
    ws.append([None, None, 99, "Spare (dark)", "AQ", "AQ", reg.cables[0][0], None, None])
    wb.save(path)


def _usage(reg: Region, n: int) -> str:
    t = next(t for t in reg.terminals if any(f[0] == n for f in t["fibers"]))
    k = [f[0] for f in t["fibers"]].index(n) + 1
    address = dict(t["fibers"])[n]
    if t is reg.commercial_tap:
        address = "C-" + address
    return f"{t['tap']}:{k} {address}"


def _seg_label(cable: str, k: int) -> str:
    return f"{cable}: Seg{k}"


def _write_connections_report(reg: Region, s: str, path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Splice Connections Report"
    x, y = reg.pos[s]
    for row in (["Splice Connections Report"], ["Prepared By:", "Self-test"], ["On: ", "2026-01-01"],
                ["Splice name:", s], ["Location:", "SYNTHETIC"], ["Address:", reg.street[s]],
                ["Coordinates:", x, y], ["Latitude/Longitude:", 0.0, 0.0], ["Enclosure Type:", "FOSC 450D"]):
        ws.append(row)
    ws.append(["Segment", "Count", "Start Splice", "End Splice", "Owner", "Cable Design Length", "Cable Asbuilt Length"])
    segs: List[Tuple[str, int, int, str, str]] = []
    if s == reg.hub:
        segs.append((f"{OLM} TO {reg.hub} FC024", 1, 24, f"OLM Housing: {OLM}", f"Splice Case: {reg.hub}"))
    else:
        par, cab, k = reg.parent[s]
        segs.append((cab, k, reg.fc[cab], f"Splice Case: {par}", f"Splice Case: {s}"))
    for ch, cab, k in reg.children[s]:
        segs.append((cab, k, reg.fc[cab], f"Splice Case: {s}", f"Splice Case: {ch}"))
    terms = [t for t in reg.terminals if t["closure"] == s]
    for t in terms:
        segs.append((f"{s} TO {t['tap']} FC004", 1, 4, f"Splice Case: {s}", f"MAX10 Enclosure(M Bulk): {t['tap']}"))
    for cab, k, cnt, st, en in segs:
        kind = "Pigtail" if cnt == 4 else "Underground"
        ws.append([f"Segment {k} of {cab} - {kind}", cnt, st, en, "Owner", 100, 0])
    if s == reg.hub:
        for p in range(1, (reg.total - 1) // PORTS + 2):
            ws.append([f"PLC COUPLER {PORTS} {p} FE0-F1PLC{PORTS}", PORTS + 1, None, None, "Owner"])
    ws.append([])
    ws.append(["A Cable", None, None, None, None, None, "B Cable", None, None, None, None, "Usage", None, None, None, "Circuit"])
    ws.append(["Segment", "Group 3", "Group 2", "Fiber", "#", "Dir", "#", "Fiber", "Group 2", "Group 3", "Segment",
               "Usage", "State", "Owner", "Carrier type", "Start", "End", "Master Circuit", "Circuit ID"])

    def pos(i: int) -> str:
        return f"{(i - 1) // 12 + 1:02d}.{(i - 1) % 12 + 1:02d}({i})"
    if s == reg.hub:
        incoming = None
    else:
        par, cab, k = reg.parent[s]
        incoming = _seg_label(cab, k)
    strand = 1
    for n in reg.downstream(s):
        c = reg.number_closure[n]
        if c == s:
            t = next(t for t in terms if any(f[0] == n for f in t["fibers"]))
            out = _seg_label(f"{s} TO {t['tap']} FC004", 1)
        else:
            node = c
            while reg.parent[node][0] != s:
                node = reg.parent[node][0]
            _p, cab, k = reg.parent[node]
            out = _seg_label(cab, k)
        a_seg = incoming or f"PLCCOUPLER{PORTS}{(n - 1) // PORTS + 1}"
        ws.append([a_seg, None, "blue", "blue", pos(strand), "<-", pos(strand), "blue", "blue", None, out,
                   _usage(reg, n), "Return", "Node", None, "Enclosure", f"Splice Case: {reg.hub}"])
        strand += 1
    if incoming:
        for j in range(3):
            ws.append([incoming, None, "green", "red", pos(100 + j), "->", pos(100 + j), "red", "green", None, incoming,
                       "No Desc", "Broken", "No Owner", None, f"Splice Case: {reg.hub}", f"Splice Case: {s}"])
        ws.append([incoming, None, "green", "red", pos(110), "->", pos(110), "red", "green", None, incoming,
                   "No Desc", "Spare (dark)", "No Owner", None, f"Splice Case: {reg.hub}", f"Splice Case: {s}"])
    wb.save(path)


def _write_guide_report(reg: Region, s: str, path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Splice Guide"
    par, cab_in, k_in = reg.parent[s]
    ins = [f"{cab_in}: Seg{k_in}-{reg.fc[cab_in]}ct"]
    outs = [f"{cab}: Seg{k}-{reg.fc[cab]}ct" for _ch, cab, k in reg.children[s]]
    terms = [t for t in reg.terminals if t["closure"] == s]
    outs += [f"{s} TO {t['tap']} FC004 : Seg1-4ct" for t in terms]
    rows = [[None, "Fiber Optic Splice Guide"], [None, f"Westbrook {OLM}", None, None, None, None, None, None, None, None, "In", None, None, None, "Out"],
            [None, "Splice Name:", None, None, s], [None, "Date Created:", None, None, "1/1/2026"],
            [None, "Enclosure Make/Model:", None, None, "SPLICE CASE"], [None, "Location:", None, None, reg.street[s]]]
    for r in rows:
        ws.append(r)
    for i in range(max(len(ins), len(outs))):
        ws.cell(3 + i, 11, ins[i] if i < len(ins) else None)
        ws.cell(3 + i, 15, outs[i] if i < len(outs) else None)
    base = 3 + max(len(ins), len(outs)) + 2
    ws.cell(base, 2, "Cable"); ws.cell(base, 6, "Buffer"); ws.cell(base, 7, "Fiber"); ws.cell(base, 8, "Cable")
    ws.cell(base, 12, "Buffer"); ws.cell(base, 13, "Fiber")
    ws.cell(base + 1, 2, "ID"); ws.cell(base + 1, 6, "Color"); ws.cell(base + 1, 7, "Color"); ws.cell(base + 1, 8, "ID")
    ws.cell(base + 1, 12, "Color"); ws.cell(base + 1, 13, "Color"); ws.cell(base + 1, 14, "Hot")
    r = base + 2
    for i, n in enumerate(reg.downstream(s), start=1):
        c = reg.number_closure[n]
        if c == s:
            t = next(t for t in terms if any(f[0] == n for f in t["fibers"]))
            out = f"{s} TO {t['tap']} FC004 SEG-1"
        else:
            node = c
            while reg.parent[node][0] != s:
                node = reg.parent[node][0]
            _p, cab, k = reg.parent[node]
            out = f"{cab} SEG-{k}"
        address = dict(next(t for t in reg.terminals if any(f[0] == n for f in t["fibers"]))["fibers"])[n]
        for col, v in ((1, i), (2, f"{cab_in}  SEG-{k_in}"), (6, "BL"), (7, COLORS[i % 12]), (8, out), (12, "BL"),
                       (13, COLORS[i % 12]), (14, "H"), (15, address)):
            ws.cell(r, col, v)
        r += 1
    wb.save(path)


def build_region(root: Path) -> Tuple[Region, Dict[str, Any], Path, Path, Path]:
    reg = Region()
    root.mkdir(parents=True, exist_ok=True)
    dxf = root / f"{OLM}_design.dxf"
    geom = _write_drawing(reg, dxf)
    addr = root / f"Westbrook_{OLM}_Address List.xlsx"
    _write_address_list(reg, addr)
    rep_dir = root / "reports"
    rep_dir.mkdir(exist_ok=True)
    for s in reg.report_closures:
        if s == reg.guide_closure:
            _write_guide_report(reg, s, rep_dir / f"{s}.xlsx")
        else:
            _write_connections_report(reg, s, rep_dir / f"{s}.xlsx")
    zpath = root / "SpliceReports.zip"
    with zipfile.ZipFile(zpath, "w") as zf:
        for f in sorted(rep_dir.iterdir()):
            zf.write(f, f"Splice Reports/{f.name}")
    return reg, geom, dxf, addr, zpath


# ---------------------------------------------------------------------------------------------- verification
def _dist_to_polyline(p: Vec2, pts: List[Vec2]) -> float:
    best = 1e18
    for a, b in zip(pts, pts[1:]):
        d = b - a
        L2 = d.x * d.x + d.y * d.y
        t = 0.0 if L2 == 0 else max(0.0, min(1.0, (p - a).dot(d) / L2))
        best = min(best, (p - (a + d * t)).magnitude)
    return best


def run_selftest(verbose: bool = True, workdir: Optional[str] = None,
                 log: Optional[Any] = None) -> Tuple[bool, str]:
    from .engine import CalloutEngine, Reporter, RunConfig
    lines: List[str] = []

    def out(msg: str) -> None:
        lines.append(msg)
        if log:
            log(msg)
        elif verbose:
            print(msg, flush=True)
    tmp = Path(workdir) if workdir else Path(tempfile.mkdtemp(prefix="fcs_selftest_"))
    checks: List[Tuple[str, bool, str]] = []
    try:
        reg, geom, dxf, addr, zpath = build_region(tmp / "input")
        exp = reg.expected()
        out(f"Synthetic region {OLM}: {len(reg.pos)} closures, {len(reg.terminals)} terminals, {reg.total} fibers, "
            f"{len(reg.report_closures)} splice reports, 1x{PORTS} splitters; {len(exp)} expected callouts")
        cfg = RunConfig(drawing_path=str(dxf), splice_paths=[str(zpath)], address_paths=[str(addr)],
                        output_dir=str(tmp / "output"))
        res = CalloutEngine(cfg, Reporter(lambda m, l: None)).generate()
        doc = ezdxf.readfile(res.outputs["Callouts DXF"])
        gen = []
        for ml in doc.modelspace().query("MULTILEADER"):
            c = ml.context
            gen.append(dict(ml=ml, text=c.mtext.default_content, tip=Vec2(c.leaders[0].lines[0].vertices[0]),
                            last=Vec2(c.leaders[0].last_leader_point), dog=c.leaders[0].dogleg_vector.x,
                            loc=Vec2(c.mtext.insert), h=c.char_height))
        used = set()
        missing, wrong_anchor = [], []
        for e in exp:
            cands = [i for i, g in enumerate(gen) if i not in used and g["text"] == e["text"]]
            ok_i = None
            for i in cands:
                tip = gen[i]["tip"]
                if e["kind"] == "TERMINAL":
                    corners = geom["hh"][e["closure"]]
                    on = _dist_to_polyline(tip, corners + [corners[0]]) < 1e-6
                else:
                    on = _dist_to_polyline(tip, geom["seg"][(e["closure"], e["next"])]) < 1e-6
                if on:
                    ok_i = i
                    break
            if ok_i is not None:
                used.add(ok_i)
            elif cands:
                wrong_anchor.append(e["text"])
            else:
                missing.append(e["text"])
        extra = [gen[i]["text"] for i in range(len(gen)) if i not in used]
        checks.append(("every expected callout generated with the exact text", not missing, f"missing {missing}"))
        checks.append(("every callout anchored on the right handhole / cable segment", not wrong_anchor, f"{wrong_anchor}"))
        checks.append(("no unexpected callouts", not extra, f"extra {extra}"))
        n_t = sum(1 for e in exp if e["kind"] == "TERMINAL")
        n_c = len(exp) - n_t
        checks.append((f"{n_t} fiber-number + {n_c} BFO callouts",
                       res.summary["Fiber-number callouts"] == n_t and res.summary["BFO callouts"] == n_c,
                       f"got {res.summary['Fiber-number callouts']} + {res.summary['BFO callouts']}"))
        props_ok = True
        detail = ""
        for g in gen:
            ml = g["ml"]
            c = ml.context
            want_loc = Vec2(g["last"].x + g["dog"] * 0.45, g["last"].y + 4.8)
            conds = [ml.dxf.layer == "DES_FIBER_SPL", ml.dxf.color == 256, c.char_height == 4.0,
                     ml.dxf.arrow_head_size == 4.0, abs(c.landing_gap_size - 0.09) < 1e-9,
                     abs(c.leaders[0].dogleg_length - 0.36) < 1e-9, ml.dxf.text_left_attachment_type == 6,
                     ml.dxf.text_right_attachment_type == 6, ml.dxf.text_top_attachment_type == 9,
                     c.mtext.alignment == (1 if g["dog"] > 0 else 3), (g["loc"] - want_loc).magnitude < 1e-6,
                     doc.entitydb[ml.dxf.style_handle].dxf.name == "Standard"]
            if not all(conds):
                props_ok = False
                detail = f"{g['text']}: {conds}"
                break
        checks.append(("MULTILEADER properties identical to the DES_FIBER_SPL standard", props_ok, detail))
        layer = doc.layers.get("DES_FIBER_SPL")
        checks.append(("layer DES_FIBER_SPL = ACI 6 / Continuous", layer.dxf.color == 6 and layer.dxf.linetype.upper() == "CONTINUOUS", ""))
        checks.append(("output DXF passes audit", len(doc.audit().errors) == 0, ""))
        errs = [i for i in res.issues.items if i.severity == "ERROR"]
        typo_flagged = any("WKB-07022" in i.obj and "did you mean" in i.message for i in errs)
        checks.append(("closure typo detected with suggestion", typo_flagged, str([i.message for i in errs])))
        checks.append(("no other errors", all("WKB-07022" in i.obj for i in errs), str([(i.obj, i.message) for i in errs])))
        checks.append(("splitter size 1x32 detected from splice reports", "32 (splitter size" in res.summary["Ports per PON"],
                       res.summary["Ports per PON"]))
        checks.append(("QA workbook written", Path(res.outputs.get("QA workbook", "")).exists(), ""))
        clean = res.placement.get("clean", 0)
        checks.append(("placement QA clean for all callouts", clean == len(gen), f"{clean}/{len(gen)}"))
        # second pass: CSV address list with the same data must give the same callouts
        import csv
        csv_path = tmp / "input" / f"Westbrook_{OLM}_Address List.csv"
        wb = openpyxl.load_workbook(addr)
        with open(csv_path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            for row in wb.active.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])
        cfg2 = RunConfig(drawing_path=str(dxf), splice_paths=[str(tmp / "input" / "reports")],
                         address_paths=[str(csv_path)], output_dir=str(tmp / "output_csv"))
        res2 = CalloutEngine(cfg2, Reporter(lambda m, l: None)).validate()
        same = sorted(p.text for p in res2.plans if p.status != "FAILED") == sorted(p.text for p in res.plans if p.status not in ("FAILED",))
        checks.append(("CSV address list + folder of reports give identical callouts", same, ""))
    except Exception as exc:  # noqa: BLE001 - a crash is a failed self-test
        import traceback
        checks.append(("self-test ran without exceptions", False, f"{exc}\n{traceback.format_exc()}"))
    finally:
        if not workdir:
            shutil.rmtree(tmp, ignore_errors=True)
    ok = all(c[1] for c in checks)
    for name, passed, detail in checks:
        out(f"  [{'PASS' if passed else 'FAIL'}] {name}" + ("" if passed or not detail else f"  -> {detail}"))
    out("SELF-TEST " + ("PASSED" if ok else "FAILED"))
    return ok, "\n".join(lines)
