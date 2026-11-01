# -*- coding: utf-8 -*-
"""
Callout engine: reads Address Lists, Splice Reports and the drawing (DXF, DWG via ODA, or DWG via the built-in
reader), reconstructs the FTTx network and plans/writes the DES_FIBER_SPL callouts.

Reverse-engineered drafting rules (docs/PHASE1_REVERSE_ENGINEERING.md):

  Fiber number   N = (PON - 1) * PORTS_PER_PON + PON PORT          (per OLM; PORTS_PER_PON = splitter size)

  Terminal callout  one MULTILEADER per terminal + drop closure (+ per contiguous run): "N" or "Nmin-Nmax";
                    arrowhead on the boundary of the handhole block hosting the drop closure.

  Cable ("BFO") callout  one MULTILEADER per live cable segment leaving a closure that has a splice report:
                    "BFO <segment fiber count>\\P<Nmin>-<Nmax>" over the live fibers of that segment;
                    arrowhead on the cable polyline near the closure.
"""
from __future__ import annotations

import dataclasses
import datetime as _dt
import difflib
import glob
import json
import math
import os
import re
import shutil
import sys
import tempfile
import threading
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Set, Tuple

import ezdxf
import openpyxl
from ezdxf import bbox as ezbbox
from ezdxf import colors as ezcolors
from ezdxf import path as ezpath
from ezdxf.math import Vec2, Vec3
from ezdxf.render import mleader as mlr
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import APP_NAME, APP_VERSION
from .dwg_native import DwgUnsupported, NativeDwg

# =============================================================================
# Configuration
# =============================================================================
LENGTH_FIELDS = ("text_height", "arrow_size", "landing_gap", "dogleg_length", "terminal_leader_length",
                 "cable_leader_length", "cable_offset_along", "handhole_search_radius")


@dataclass
class CalloutSpec:
    """Drafting standard. Lengths are in FEET and converted automatically for metric drawings."""

    cable_layer: str = "DES_FIBER_SPL"
    terminal_layer: str = "DES_FIBER_SPL"
    layer_color: int = 6
    layer_linetype: str = "Continuous"
    mleader_style: str = "Standard"
    text_style: str = "Standard"
    text_style_font: str = "txt"
    text_height: float = 4.0
    arrow_size: float = 4.0
    landing_gap: float = 0.09
    dogleg_length: float = 0.36
    side_attachment: int = 6
    vertical_attachment: int = 9
    text_y_factor: float = 1.2
    cable_prefix: str = "BFO"
    terminal_leader_length: float = 17.0
    cable_leader_length: float = 30.0
    cable_offset_along: float = 30.0
    handhole_search_radius: float = 75.0
    zero_pad: int = 0
    cable_range_mode: str = "span"
    split_terminal_runs: bool = True

    def scaled(self, factor: float) -> "CalloutSpec":
        s = dataclasses.replace(self)
        for k in LENGTH_FIELDS:
            setattr(s, k, getattr(self, k) * factor)
        return s


@dataclass
class ColumnMap:
    pon: Optional[str] = None
    port: Optional[str] = None
    address: Optional[str] = None
    tap: Optional[str] = None
    splice: Optional[str] = None
    fiber: Optional[str] = None
    cable: Optional[str] = None
    olm: Optional[str] = None


@dataclass
class CadMap:
    """Empty values mean 'auto-detect from the drawing and the input data'."""

    closure_layer: str = ""
    closure_block: str = ""
    closure_name_tag: str = ""
    handhole_layer: str = ""
    handhole_blocks: List[str] = field(default_factory=list)
    terminal_layer: str = ""
    cable_layers: List[str] = field(default_factory=list)
    snap_tolerance: float = 0.0


@dataclass
class RunConfig:
    drawing_path: str = ""
    splice_paths: List[str] = field(default_factory=list)
    address_paths: List[str] = field(default_factory=list)
    output_dir: str = ""
    columns: ColumnMap = field(default_factory=ColumnMap)
    cad: CadMap = field(default_factory=CadMap)
    spec: CalloutSpec = field(default_factory=CalloutSpec)
    ports_per_pon: int = 0             # 0 = splitter size from the splice reports
    units: str = "auto"                # auto ($INSUNITS of a DXF, feet for DWG) | feet | meters | inches

    def to_dict(self) -> Dict[str, Any]:
        return dataclasses.asdict(self)

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> "RunConfig":
        cfg = cls()
        for k, v in data.items():
            if k == "columns" and isinstance(v, dict):
                cfg.columns = ColumnMap(**{a: b for a, b in v.items() if a in ColumnMap.__dataclass_fields__})
            elif k == "cad" and isinstance(v, dict):
                cfg.cad = CadMap(**{a: b for a, b in v.items() if a in CadMap.__dataclass_fields__})
            elif k == "spec" and isinstance(v, dict):
                cfg.spec = CalloutSpec(**{a: b for a, b in v.items() if a in CalloutSpec.__dataclass_fields__})
            elif k in cls.__dataclass_fields__:
                setattr(cfg, k, v)
        return cfg


ADDRESS_EXTS = (".xlsx", ".xlsm", ".xls", ".csv", ".txt")
REPORT_EXTS = (".xlsx", ".xlsm", ".xls", ".csv")

ADDRESS_FIELDS: Dict[str, Tuple[bool, List[str], str]] = {
    "pon": (True, ["PON NAME", "PON", "PON #", "PON NO", "PON NUMBER", "PON ID", "SPLITTER", "SPLITTER #",
                   "SPLITTER NAME", "SPLITTER ID", "COUPLER", "PLC", "OLT PON"], "PON / splitter index"),
    "port": (True, ["PON PORT", "PORT", "PORT #", "SPLITTER PORT", "OUTPUT PORT", "PON PORT #", "COUPLER PORT",
                    "SPLITTER OUTPUT", "LEG"], "Port on that PON / splitter"),
    "address": (True, ["ADDRESS", "SERVICE ADDRESS", "SERVICE LOCATION", "LOCATION", "FULL ADDRESS", "PREMISE",
                       "PREMISE ADDRESS", "SITE ADDRESS", "CUSTOMER ADDRESS"], "Served address"),
    "tap": (True, ["FIBER TAP NAME", "TAP NAME", "TAP", "TERMINAL", "TERMINAL NAME", "TERMINAL ID", "MST", "NAP",
                   "DROP TERMINAL", "FIBER TAP", "FAT", "FDP"], "Terminal / tap name"),
    "splice": (True, ["FIBER SPLICE", "SPLICE", "SPLICE NAME", "CLOSURE", "SPLICE CLOSURE", "SPLICE CASE",
                      "DROP SPLICE", "ENCLOSURE", "SPLICE ENCLOSURE", "SPLICE POINT"], "Drop splice closure"),
    "fiber": (False, ["FIBER #", "FIBER", "FIBER NO", "FIBER NUMBER", "STRAND", "FIBER NUM"],
              "Fiber position in first cable"),
    "cable": (False, ["CABLE NAME", "CABLE", "CABLE ID"], "Feeder / distribution cable"),
    "olm": (False, ["OLM", "OLT", "HUB", "CABINET", "FDH", "OLM NAME", "SERVICE AREA"], "OLM / hub id"),
}


# =============================================================================
# Infrastructure
# =============================================================================
class Cancelled(Exception):
    """Raised when the user cancels a running job."""


class Reporter:
    def __init__(self, log_fn: Optional[Callable[[str, str], None]] = None,
                 progress_fn: Optional[Callable[[int, str], None]] = None,
                 cancel_event: Optional[threading.Event] = None) -> None:
        self._log_fn = log_fn
        self._progress_fn = progress_fn
        self._cancel = cancel_event or threading.Event()

    def log(self, msg: str, level: str = "INFO") -> None:
        if self._log_fn:
            self._log_fn(msg, level)
        else:
            stamp = _dt.datetime.now().strftime("%H:%M:%S")
            print(f"[{stamp}] {level:<7} {msg}", flush=True)

    def progress(self, pct: float, msg: str = "") -> None:
        if self._progress_fn:
            self._progress_fn(int(max(0, min(100, pct))), msg)

    def check_cancel(self) -> None:
        if self._cancel.is_set():
            raise Cancelled("Cancelled by user")


@dataclass
class Issue:
    severity: str
    category: str
    obj: str
    message: str


class IssueLog:
    def __init__(self) -> None:
        self.items: List[Issue] = []

    def add(self, severity: str, category: str, obj: Any, message: str) -> None:
        self.items.append(Issue(severity, category, str(obj), message))

    def count(self, severity: str) -> int:
        return sum(1 for i in self.items if i.severity == severity)

    def summary(self) -> str:
        return f"{self.count('ERROR')} error(s), {self.count('WARNING')} warning(s), {self.count('INFO')} info"


# =============================================================================
# Text helpers
# =============================================================================
def norm_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    s = str(value).replace("\u00a0", " ").strip().upper()
    s = re.sub(r"\s+", " ", s)
    return s.rstrip(".")


def norm_header(value: Any) -> str:
    return re.sub(r"[^A-Z0-9#]", "", norm_text(value))


def norm_cable(value: Any) -> str:
    return re.sub(r"\s*:\s*", ":", norm_text(value))


def strip_kind(value: Any) -> str:
    """'Splice Case: NMK-02001G' -> 'NMK-02001G'."""
    s = norm_text(value)
    if ":" in s and not re.search(r":\d+\s", s):
        s = s.split(":", 1)[1].strip()
    return s


def to_int(value: Any) -> Optional[int]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isnan(value) or abs(value - round(value)) > 1e-6:
            return None
        return int(round(value))
    text = str(value).strip().replace(",", "")
    if re.fullmatch(r"\d+", text):
        return int(text)
    try:
        f = float(text)
    except ValueError:
        return None
    if math.isnan(f) or abs(f - round(f)) > 1e-6:
        return None
    return int(round(f))


def to_index(value: Any) -> Optional[int]:
    """Integer index from 1, 1.0, '1', 'PON 3', 'P-03', 'Port 12'."""
    v = to_int(value)
    if v is not None:
        return v
    groups = re.findall(r"\d+", str(value or ""))
    return int(groups[-1]) if groups else None


def number_runs(nums: Iterable[int]) -> List[Tuple[int, int]]:
    runs: List[Tuple[int, int]] = []
    for n in sorted(set(nums)):
        if runs and n == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], n)
        else:
            runs.append((n, n))
    return runs


def fmt_run(run: Tuple[int, int], pad: int = 0) -> str:
    a, b = run
    fa = str(a).zfill(pad) if pad else str(a)
    fb = str(b).zfill(pad) if pad else str(b)
    return fa if a == b else f"{fa}-{fb}"


def fmt_numbers(nums: Iterable[int], mode: str = "exact", pad: int = 0) -> str:
    runs = number_runs(nums)
    if not runs:
        return ""
    if mode == "span":
        return fmt_run((runs[0][0], runs[-1][1]), pad)
    return ",".join(fmt_run(r, pad) for r in runs)


def canonical_callout_text(text: str) -> str:
    t = (text or "").replace("\\P", "\n").replace("\r", "")
    t = re.sub(r"\\[A-Za-z][^;\\]*;", "", t)
    t = re.sub(r"[{}]", "", t)
    t = re.sub(r"\d+", lambda m: str(int(m.group(0))), t)
    return "|".join(norm_text(x) for x in t.split("\n") if x.strip())


def is_fiber_callout_text(text: str, prefix: str = "BFO") -> bool:
    c = canonical_callout_text(text)
    rng = r"\d+(?:-\d+)?(?:,\d+(?:-\d+)?)*"
    return bool(re.fullmatch(rf"(?:{re.escape(norm_text(prefix))} [\d?]+\|)?{rng}", c))


# =============================================================================
# Geometry
# =============================================================================
Box = Tuple[float, float, float, float]


def v2(p: Any) -> Vec2:
    return Vec2(float(p[0]), float(p[1]))


def seg_distance(p: Vec2, a: Vec2, b: Vec2) -> Tuple[float, float]:
    d = b - a
    L2 = d.x * d.x + d.y * d.y
    t = 0.0 if L2 == 0 else max(0.0, min(1.0, ((p.x - a.x) * d.x + (p.y - a.y) * d.y) / L2))
    return (p - (a + d * t)).magnitude, t


def polyline_length(pts: Sequence[Vec2]) -> float:
    return sum((pts[i + 1] - pts[i]).magnitude for i in range(len(pts) - 1))


def point_at_length(pts: Sequence[Vec2], s: float) -> Tuple[Vec2, Vec2]:
    acc = 0.0
    for i in range(len(pts) - 1):
        seg = pts[i + 1] - pts[i]
        L = seg.magnitude
        if L <= 1e-9:
            continue
        if acc + L >= s:
            return pts[i] + seg * ((s - acc) / L), seg / L
        acc += L
    last = pts[-1] - pts[-2] if len(pts) > 1 else Vec2(1, 0)
    return pts[-1], (last.normalize() if last.magnitude > 1e-9 else Vec2(1, 0))


def box_overlap(a: Box, b: Box) -> float:
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0.0


def box_of_points(pts: Iterable[Vec2], pad: float = 0.0) -> Box:
    pts = list(pts)
    return (min(p.x for p in pts) - pad, min(p.y for p in pts) - pad,
            max(p.x for p in pts) + pad, max(p.y for p in pts) + pad)


def segments_intersect(p1: Vec2, p2: Vec2, q1: Vec2, q2: Vec2) -> bool:
    def orient(a: Vec2, b: Vec2, c: Vec2) -> float:
        return (b.x - a.x) * (c.y - a.y) - (b.y - a.y) * (c.x - a.x)
    o1, o2 = orient(p1, p2, q1), orient(p1, p2, q2)
    o3, o4 = orient(q1, q2, p1), orient(q1, q2, p2)
    return (o1 * o2 < 0) and (o3 * o4 < 0)


def segment_hits_box(a: Vec2, b: Vec2, box: Box) -> bool:
    def inside(p: Vec2) -> bool:
        return box[0] <= p.x <= box[2] and box[1] <= p.y <= box[3]
    if inside(a) or inside(b):
        return True
    c = [Vec2(box[0], box[1]), Vec2(box[2], box[1]), Vec2(box[2], box[3]), Vec2(box[0], box[3])]
    return any(segments_intersect(a, b, c[i], c[(i + 1) % 4]) for i in range(4))


def text_extent_box(pos: Vec2, height: float, text: str, rotation: float = 0.0, attachment: int = 7) -> Box:
    """Approximate world box of a TEXT (attachment 7 = baseline-left) or MTEXT (attachment 1..9)."""
    lines = [x for x in re.sub(r"\\[A-Za-z][^;\\]*;|[{}]", "", (text or "").replace("\\P", "\n")).split("\n")] or [""]
    w = max(1, max(len(x) for x in lines)) * height * 0.85
    h = height + (len(lines) - 1) * height * 1.667
    col = (attachment - 1) % 3
    row = (attachment - 1) // 3
    x0 = -w * (0, 0.5, 1)[col]
    y0 = (-h, -h / 2, 0)[row]
    ang = math.radians(rotation)
    ca, sa = math.cos(ang), math.sin(ang)
    corners = []
    for dx, dy in ((x0, y0), (x0 + w, y0), (x0 + w, y0 + h), (x0, y0 + h)):
        corners.append(Vec2(pos.x + dx * ca - dy * sa, pos.y + dx * sa + dy * ca))
    return box_of_points(corners)


class GridIndex:
    def __init__(self, cell: float = 50.0) -> None:
        self.cell = max(cell, 1e-6)
        self.cells: Dict[Tuple[int, int], List[Tuple[Box, Any]]] = defaultdict(list)

    def _keys(self, b: Box) -> Iterable[Tuple[int, int]]:
        c = self.cell
        x0, x1 = int(math.floor(b[0] / c)), int(math.floor(b[2] / c))
        y0, y1 = int(math.floor(b[1] / c)), int(math.floor(b[3] / c))
        if (x1 - x0 + 1) * (y1 - y0 + 1) > 10000:
            x1 = x0 + 99
            y1 = y0 + 99
        for ix in range(x0, x1 + 1):
            for iy in range(y0, y1 + 1):
                yield ix, iy

    def insert(self, b: Box, item: Any) -> None:
        entry = (b, item)
        for k in self._keys(b):
            self.cells[k].append(entry)

    def query(self, b: Box) -> List[Tuple[Box, Any]]:
        seen: Set[int] = set()
        out: List[Tuple[Box, Any]] = []
        for k in self._keys(b):
            for entry in self.cells.get(k, ()):
                if id(entry) not in seen:
                    seen.add(id(entry))
                    out.append(entry)
        return out

    def nearest(self, p: Vec2, radius: float, key: Callable[[Any], Vec2]) -> Tuple[Optional[Any], float]:
        best, bd = None, radius
        for (_b, item) in self.query((p.x - radius, p.y - radius, p.x + radius, p.y + radius)):
            d = (key(item) - p).magnitude
            if d <= bd:
                best, bd = item, d
        return best, bd


# =============================================================================
# Inputs
# =============================================================================
class InputCollector:
    def __init__(self) -> None:
        self._tmp: List[str] = []

    def expand(self, paths: Sequence[str], exts: Sequence[str], issues: IssueLog) -> List[Path]:
        out: List[Path] = []
        seen: Set[str] = set()
        exts_l = tuple(e.lower() for e in exts)

        hashes: Dict[str, str] = {}

        def add(p: Path) -> None:
            key = str(p.resolve()).lower()
            if key in seen or p.name.startswith("~$") or p.name.startswith("."):
                return
            if p.suffix.lower() in exts_l:
                seen.add(key)
                try:
                    import hashlib
                    digest = hashlib.sha1(p.read_bytes()).hexdigest()
                except OSError:
                    digest = key
                if digest in hashes:
                    issues.add("INFO", "Input", p.name, f"Identical to {hashes[digest]} - duplicate skipped")
                    return
                hashes[digest] = p.name
                out.append(p)

        def walk(root: Path) -> None:
            for f in sorted(root.rglob("*")):
                if f.is_file():
                    if f.suffix.lower() == ".zip":
                        for z in self._extract(f, issues):
                            walk(Path(z))
                    else:
                        add(f)

        for raw in paths:
            if not raw:
                continue
            p = Path(raw)
            if not p.exists():
                issues.add("ERROR", "Input", raw, "Path does not exist")
            elif p.is_dir():
                walk(p)
            elif p.suffix.lower() == ".zip":
                for z in self._extract(p, issues):
                    walk(Path(z))
            else:
                add(p)
        return out

    def _extract(self, zpath: Path, issues: IssueLog) -> List[str]:
        try:
            tmp = tempfile.mkdtemp(prefix="fcs_zip_")
            self._tmp.append(tmp)
            root = Path(tmp).resolve()
            with zipfile.ZipFile(zpath) as zf:
                for member in zf.infolist():
                    target = (root / member.filename).resolve()
                    if root != target and root not in target.parents:
                        issues.add("WARNING", "Input", zpath.name, f"Skipped unsafe ZIP member {member.filename}")
                        continue
                    zf.extract(member, root)
            return [tmp]
        except (zipfile.BadZipFile, OSError) as exc:
            issues.add("ERROR", "Input", zpath.name, f"Cannot open ZIP: {exc}")
            return []

    def cleanup(self) -> None:
        for t in self._tmp:
            shutil.rmtree(t, ignore_errors=True)
        self._tmp.clear()


def read_tables(path: Path) -> List[Tuple[str, List[List[Any]]]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        finally:
            wb.close()
    if suffix == ".xls":
        try:
            import xlrd  # type: ignore
        except ImportError as exc:
            raise ValueError("legacy .xls needs the 'xlrd' package - save the file as .xlsx") from exc
        book = xlrd.open_workbook(str(path))
        return [(sh.name, [sh.row_values(r) for r in range(sh.nrows)]) for sh in book.sheets()]
    if suffix in (".csv", ".txt"):
        import csv
        raw = path.read_bytes()
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("cannot decode text file")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [[(c if c != "" else None) for c in r] for r in csv.reader(text.splitlines(), dialect)]
        return [(path.stem, rows)]
    raise ValueError(f"unsupported file type {suffix}")


# =============================================================================
# Address lists
# =============================================================================
@dataclass
class AddressRecord:
    olm: str
    pon: int
    port: int
    number: int
    address: str
    tap: str
    splice: str
    fiber: Optional[int]
    cable: str
    source: str
    row: int


@dataclass
class Terminal:
    olm: str
    tap: str
    splice: str
    numbers: List[int] = field(default_factory=list)
    records: List[AddressRecord] = field(default_factory=list)


def _match_header(headers: List[str], synonyms: List[str]) -> Optional[str]:
    normed = {norm_header(h): h for h in headers if h}
    syn_n = [norm_header(s) for s in synonyms]
    for s in syn_n:
        if s in normed:
            return normed[s]
    for s in syn_n:
        for hn, h in normed.items():
            if len(s) >= 3 and (hn.startswith(s) or s in hn):
                return h
    best, best_r = None, 0.0
    for s in syn_n:
        for hn, h in normed.items():
            r = difflib.SequenceMatcher(None, s, hn).ratio()
            if r > best_r:
                best, best_r = h, r
    return best if best_r >= 0.82 else None


def detect_header_row(rows: List[List[Any]]) -> Tuple[int, List[str], int]:
    syn_n = {norm_header(s) for (_, syns, _) in ADDRESS_FIELDS.values() for s in syns}
    best_i, best_score = 0, -1
    for i, r in enumerate(rows[:30]):
        score = sum(1 for c in r if c is not None and norm_header(c) in syn_n)
        if score > best_score:
            best_i, best_score = i, score
    headers = [str(c).strip() if c is not None else "" for c in rows[best_i]] if rows else []
    return best_i, headers, max(0, best_score)


def auto_map_columns(headers: List[str]) -> ColumnMap:
    cm = ColumnMap()
    used: Set[str] = set()
    splice_syn = {norm_header(s) for s in ADDRESS_FIELDS["splice"][1]}
    for key in ("pon", "port", "address", "tap", "splice", "fiber", "cable", "olm"):
        h = _match_header([x for x in headers if x not in used], ADDRESS_FIELDS[key][1])
        if key == "olm" and h is not None and norm_header(h) in splice_syn:
            h = None
        if h:
            used.add(h)
        setattr(cm, key, h)
    return cm


OLM_TOKEN = re.compile(r"(?<![A-Z0-9])([A-Z]{2,}-[A-Z]{0,3}\d{1,4}[A-Z]?)(?![A-Z0-9])")
GENERIC_SHEETS = {"SHEET", "SHEET1", "SHEET2", "SHEET3", "DATA", "ADDRESS LIST", "ADDRESSES", "EXPORT", "REPORT"}


def looks_like_splice_report(rows: List[List[Any]]) -> bool:
    head = " ".join(str(c) for r in rows[:14] for c in r if c is not None).lower()
    return "splice name" in head or "splice connections report" in head or "fiber optic splice guide" in head


def derive_olm(sheet: str, filename: str) -> str:
    s = norm_text(sheet)
    stem = norm_text(Path(filename).stem)
    if s and s not in GENERIC_SHEETS and s != stem:
        return s
    toks = OLM_TOKEN.findall(stem)
    return toks[-1] if toks else stem


class AddressBook:
    def __init__(self) -> None:
        self.records: List[AddressRecord] = []
        self.terminals: Dict[Tuple[str, str, str], Terminal] = {}
        self.by_tap_addr: Dict[Tuple[str, str], List[AddressRecord]] = defaultdict(list)
        self.by_addr: Dict[str, List[AddressRecord]] = defaultdict(list)
        self.by_tap: Dict[str, List[AddressRecord]] = defaultdict(list)
        self.by_number: Dict[Tuple[str, int], AddressRecord] = {}
        self.tap_names: Set[str] = set()
        self.ports_per_pon = 64
        self.ports_source = "default"
        self.olms: List[str] = []
        self.max_port = 0

    @staticmethod
    def scan_headers(files: List[Path], issues: IssueLog) -> List[str]:
        headers: List[str] = []
        for f in files:
            try:
                for _sheet, rows in read_tables(f):
                    _, h, score = detect_header_row(rows)
                    if score < 2:
                        continue
                    for x in h:
                        if x and x not in headers:
                            headers.append(x)
            except Exception as exc:  # noqa: BLE001 - every unreadable file is reported
                issues.add("ERROR", "Address List", f.name, f"Unreadable: {exc}")
        return headers

    def load(self, files: List[Path], cmap: ColumnMap, ports_per_pon: int, report_splitter: int,
             pigtail_splices: Dict[str, str], issues: IssueLog, rep: Reporter, quiet_cross: bool = False) -> None:
        raw: List[Tuple[str, int, int, Dict[str, Any], str, int]] = []
        for f in files:
            rep.check_cancel()
            try:
                tables = read_tables(f)
            except Exception as exc:  # noqa: BLE001
                issues.add("ERROR", "Address List", f.name, f"Unreadable: {exc}")
                continue
            for sheet, rows in tables:
                if not rows:
                    continue
                if looks_like_splice_report(rows):
                    if not quiet_cross:
                        issues.add("INFO", "Address List", f"{f.name}:{sheet}", "This is a splice report - not used as address list")
                    continue
                hi, headers, score = detect_header_row(rows)
                if score < 2:
                    issues.add("INFO", "Address List", f"{f.name}:{sheet}", "Sheet has no address-list headers - skipped")
                    continue
                idx = {h: i for i, h in enumerate(headers) if h}
                auto = auto_map_columns(headers)
                colmap: Dict[str, Optional[str]] = {}
                for k in ADDRESS_FIELDS:
                    user = getattr(cmap, k)
                    colmap[k] = user if user in idx else getattr(auto, k)
                missing = [k for k, (req, _, _) in ADDRESS_FIELDS.items()
                           if req and k != "splice" and not colmap[k]]
                if missing:
                    issues.add("ERROR", "Address List", f"{f.name}:{sheet}",
                               "Required column(s) not found: " + ", ".join(missing) + " - map them in the Column Mapper")
                    continue
                default_olm = derive_olm(sheet, f.name)
                for ri, r in enumerate(rows[hi + 1:], start=hi + 2):
                    def get(key: str) -> Any:
                        col = colmap.get(key)
                        if not col or idx[col] >= len(r):
                            return None
                        return r[idx[col]]
                    pon, port = to_index(get("pon")), to_index(get("port"))
                    if pon is None and port is None:
                        continue
                    if pon is None or port is None or pon < 1 or port < 1:
                        issues.add("WARNING", "Address List", f"{f.name}:{sheet} row {ri}",
                                   f"Invalid PON/port ({get('pon')!r}, {get('port')!r}) - row skipped")
                        continue
                    olm = norm_text(get("olm")) or default_olm
                    raw.append((olm, pon, port, {k: get(k) for k in ADDRESS_FIELDS}, f"{f.name}:{sheet}", ri))
        if not raw:
            issues.add("ERROR", "Address List", "-", "No PON-assigned rows found in the address list(s)")
            return
        self.max_port = max(p for _, _, p, _, _, _ in raw)
        if ports_per_pon and ports_per_pon > 0:
            self.ports_per_pon, self.ports_source = ports_per_pon, "user setting"
        elif report_splitter and report_splitter >= self.max_port:
            self.ports_per_pon, self.ports_source = report_splitter, "splitter size in splice reports"
        else:
            self.ports_per_pon = next((s for s in (8, 16, 32, 64, 128, 256) if s >= self.max_port), self.max_port)
            self.ports_source = "address list (next power of two >= highest port)"
        if self.max_port > self.ports_per_pon:
            issues.add("ERROR", "Address List", "-",
                       f"PON PORT {self.max_port} exceeds ports-per-PON {self.ports_per_pon}")
        no_splice: Counter = Counter()
        for olm, pon, port, vals, src, ri in raw:
            number = (pon - 1) * self.ports_per_pon + port
            tap = norm_text(vals["tap"])
            rec = AddressRecord(olm=olm, pon=pon, port=port, number=number, address=norm_text(vals["address"]),
                                tap=tap, splice=strip_kind(vals["splice"]), fiber=to_int(vals["fiber"]),
                                cable=norm_text(vals["cable"]), source=src, row=ri)
            key = (olm, number)
            if key in self.by_number:
                issues.add("ERROR", "Address List", f"{src} row {ri}",
                           f"Duplicate fiber number {number} in {olm} (PON {pon} port {port})")
                continue
            if not rec.tap:
                rec.tap = rec.address
                issues.add("WARNING", "Address List", f"{src} row {ri}",
                           f"No terminal name for {rec.address}; the address is used as terminal")
            if not rec.splice and rec.tap in pigtail_splices:
                rec.splice = pigtail_splices[rec.tap]
                issues.add("INFO", "Address List", f"{src} row {ri}",
                           f"Drop closure of {rec.tap} taken from splice report pigtail ({rec.splice})")
            if not rec.splice:
                no_splice[src] += 1
            self.by_number[key] = rec
            self.records.append(rec)
            other = [k for k in self.terminals if k[0] == olm and k[1] == rec.tap and k[2] != rec.splice]
            if other:
                issues.add("WARNING", "Splice Mismatch", rec.tap,
                           f"Terminal rows reference different drop closures ({other[0][2]} vs {rec.splice}); "
                           f"one callout is placed at each closure")
            t = self.terminals.setdefault((olm, rec.tap, rec.splice), Terminal(olm=olm, tap=rec.tap, splice=rec.splice))
            t.numbers.append(number)
            t.records.append(rec)
            self.by_tap_addr[(rec.tap, rec.address)].append(rec)
            self.by_addr[rec.address].append(rec)
            self.by_tap[rec.tap].append(rec)
            self.tap_names.add(rec.tap)
        for src, n in no_splice.items():
            issues.add("WARNING", "Address List", src,
                       f"{n} row(s) have no drop splice closure (map the 'Fiber Splice' column, or add splice reports "
                       f"with pigtail segments) - their fiber-number callouts cannot be placed")
        self.olms = sorted({r.olm for r in self.records})
        for lst in self.by_tap.values():
            lst.sort(key=lambda r: (r.olm, r.number))
        rep.log(f"Address lists: {len(self.records)} PON-assigned rows, {len(self.terminals)} terminal groups, "
                f"OLMs {', '.join(self.olms)}; ports per PON = {self.ports_per_pon} ({self.ports_source})")

    @staticmethod
    def address_variants(address: str) -> List[str]:
        a = norm_text(address)
        out = [a]
        b = re.sub(r"^(?:[A-Z]{1,2}-)(?=\d)", "", a)
        c = re.sub(r"\s+(?:#\s*\d+[A-Z]?|BLDG\s*\d+[A-Z]?|UNIT\s*\w+|APT\s*\w+|STE\s*\w+|LOT\s*\w+)\b", "", b)
        for v in (b, c):
            v = re.sub(r"\s+", " ", v).strip()
            if v and v not in out:
                out.append(v)
        return out

    def resolve_candidates(self, tap: str, address: str, olm_hint: Optional[str],
                           port_idx: Optional[int] = None) -> List[AddressRecord]:
        tap_n, addr_n = norm_text(tap), norm_text(address)

        def scope(c: List[AddressRecord]) -> List[AddressRecord]:
            if olm_hint:
                in_olm = [x for x in c if x.olm == olm_hint]
                if in_olm:
                    return in_olm
            return c
        if tap_n:
            c = scope(self.by_tap_addr.get((tap_n, addr_n), []))
            if len(c) == 1:
                return c
            recs = scope(self.by_tap.get(tap_n, []))
            if recs and port_idx and 1 <= port_idx <= len(recs) and len({r.olm for r in recs}) == 1:
                return [recs[port_idx - 1]]
            if len(c) > 1:
                return c
        for variant in self.address_variants(addr_n):
            c = scope(self.by_addr.get(variant, []))
            if c:
                return c
        return []

    def resolve(self, tap: str, address: str, olm_hint: Optional[str], port_idx: Optional[int] = None) -> Optional[AddressRecord]:
        c = self.resolve_candidates(tap, address, olm_hint, port_idx)
        return c[0] if len(c) == 1 else None


# =============================================================================
# Splice reports
# =============================================================================
@dataclass
class ReportSegment:
    cable: str
    seg: int
    count: Optional[int]
    kind: str
    start: Optional[str]
    end: Optional[str]
    direction: str


@dataclass
class ReportConn:
    a_seg: Optional[Tuple[str, int]]
    b_seg: Optional[Tuple[str, int]]
    usage: str
    state: str
    tap: str
    address: str
    port_idx: Optional[int] = None


@dataclass
class SpliceReport:
    name: str
    file: str
    fmt: str
    xy: Optional[Tuple[float, float]] = None
    segments: List[ReportSegment] = field(default_factory=list)
    conns: List[ReportConn] = field(default_factory=list)
    olm: Optional[str] = None
    splitters: List[int] = field(default_factory=list)


SEG_ROW_RE = re.compile(r"^\s*Segment\s+(\d+)\s+of\s+(.+?)\s+-\s+([A-Za-z][\w /]*)\s*$", re.I)
CONN_SEG_RE = re.compile(r"^(.+?)\s*:\s*Seg\s*(\d+)\s*$", re.I)
GUIDE_SEG_RE = re.compile(r"^(.+?)\s*:\s*Seg\s*(\d+)(?:\s*-\s*(\d+)\s*ct)?\s*$", re.I)
GUIDE_CABLE_RE = re.compile(r"^(.+?)\s+SEG\s*-?\s*(\d+)\s*$", re.I)
FC_RE = re.compile(r"\bFC\s*0*(\d+)\s*$", re.I)
USAGE_RE = re.compile(r"^(.*?):(\d+)\s+(.*)$")
SPLITTER_RE = re.compile(r"PLC\s+COUPLER\s+(\d+)\s+\d+|\bF\d?PLC(\d+)\b|SPLITTER\s*1\s*[xX:]\s*(\d+)|"
                         r"\b1\s*[xX]\s*(\d+)\s*(?:PLC|SPLITTER)", re.I)
DEAD_STATES = ("BROKEN", "SPARE", "DARK", "CUT", "DEAD")
NON_ADDRESS_RE = re.compile(r"TRANSPORT|SPARES?\b|POWER PORT|^NO DESC|COUPLER|\bOLM\b", re.I)


def _fc_of(cable: str) -> Optional[int]:
    m = FC_RE.search(cable)
    return int(m.group(1)) if m else None


def _label_value(row: List[Any], label_idx: int) -> Any:
    for v in row[label_idx + 1:]:
        if v is not None and str(v).strip() != "":
            return v
    return None


def _find_col(low: List[str], *names: str) -> Optional[int]:
    for n in names:
        if n in low:
            return low.index(n)
    for n in names:
        for j, c in enumerate(low):
            if c and n in c:
                return j
    return None


def parse_connections_report(rows: List[List[Any]], fname: str) -> Optional[SpliceReport]:
    name = None
    xy = None
    seg_hdr = conn_hdr = None
    splitters: List[int] = []
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in r]
        low = [c.lower() for c in cells]
        if not any(low):
            continue
        if low[0].startswith("splice name") or low[0].startswith("closure name") or low[0].startswith("enclosure name"):
            name = strip_kind(_label_value(r, 0))
        elif low[0].startswith("coordinates"):
            nums = [c for c in r[1:] if isinstance(c, (int, float))]
            if len(nums) < 2:
                nums = [float(x) for x in re.findall(r"-?\d+(?:\.\d+)?", " ".join(cells[1:]))]
            if len(nums) >= 2:
                xy = (float(nums[0]), float(nums[1]))
        for c in cells:
            m = SPLITTER_RE.search(c)
            if m:
                splitters.append(int(next(g for g in m.groups() if g)))
        if seg_hdr is None and "count" in low and _find_col(low, "start splice", "start enclosure", "from") is not None \
                and _find_col(low, "end splice", "end enclosure", "to") is not None and "segment" in low:
            seg_hdr = (i, dict(segment=low.index("segment"), count=low.index("count"),
                               start=_find_col(low, "start splice", "start enclosure", "from"),
                               end=_find_col(low, "end splice", "end enclosure", "to")))
        elif conn_hdr is None and "usage" in low and low.count("segment") >= 2:
            segs = [j for j, c in enumerate(low) if c == "segment"]
            conn_hdr = (i, dict(a_seg=segs[0], b_seg=segs[1], usage=low.index("usage"),
                                state=low.index("state") if "state" in low else None))
    if not name:
        return None
    rep = SpliceReport(name=name, file=fname, fmt="Splice Connections Report", xy=xy, splitters=splitters)
    if seg_hdr:
        i0, c = seg_hdr
        for r in rows[i0 + 1:]:
            if not r or c["segment"] >= len(r) or r[c["segment"]] is None or str(r[c["segment"]]).strip() == "":
                break
            m = SEG_ROW_RE.match(str(r[c["segment"]]))
            if not m:
                continue
            cable = norm_cable(m.group(2))
            start = strip_kind(r[c["start"]]) if c["start"] < len(r) else None
            end = strip_kind(r[c["end"]]) if c["end"] < len(r) else None
            cnt = to_int(r[c["count"]]) if c["count"] < len(r) else None
            direction = "out" if start == name else ("in" if end == name else "thru")
            rep.segments.append(ReportSegment(cable=cable, seg=int(m.group(1)), count=cnt or _fc_of(cable),
                                              kind=m.group(3).strip().title(), start=start or None, end=end or None,
                                              direction=direction))
    if conn_hdr:
        i0, c = conn_hdr

        def segkey(v: Any) -> Optional[Tuple[str, int]]:
            m = CONN_SEG_RE.match(str(v or "").strip())
            return (norm_cable(m.group(1)), int(m.group(2))) if m else None
        for r in rows[i0 + 1:]:
            if not any(v is not None and str(v).strip() for v in r):
                break
            usage = str(r[c["usage"]] or "").strip() if c["usage"] < len(r) else ""
            state = str(r[c["state"]] or "").strip() if c["state"] is not None and c["state"] < len(r) else ""
            m = USAGE_RE.match(usage)
            tap, addr = (norm_text(m.group(1)), norm_text(m.group(3))) if m else ("", norm_text(usage))
            rep.conns.append(ReportConn(a_seg=segkey(r[c["a_seg"]]) if c["a_seg"] < len(r) else None,
                                        b_seg=segkey(r[c["b_seg"]]) if c["b_seg"] < len(r) else None,
                                        usage=usage, state=state, tap=tap, address=addr,
                                        port_idx=int(m.group(2)) if m else None))
    return rep


def parse_guide_report(rows: List[List[Any]], fname: str) -> Optional[SpliceReport]:
    name = None
    title_olm = None
    in_col = out_col = io_row = hdr_row = None
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in r]
        low = [c.lower() for c in cells]
        for j, c in enumerate(low):
            if c.startswith("splice name") and name is None:
                name = strip_kind(_label_value(r, j))
            if c == "in" and in_col is None:
                in_col, io_row = j, i
            if c == "out" and out_col is None:
                out_col = j
        if i < 4 and title_olm is None:
            toks = OLM_TOKEN.findall(norm_text(" ".join(cells)))
            title_olm = toks[-1] if toks else None
        if hdr_row is None and low.count("cable") >= 2 and "buffer" in low:
            hdr_row = i
    if not name:
        return None
    rep = SpliceReport(name=name, file=fname, fmt="Fiber Optic Splice Guide", olm=title_olm)
    if io_row is not None:
        stop = hdr_row if hdr_row is not None else len(rows)
        for col, direction in ((in_col, "in"), (out_col, "out")):
            if col is None:
                continue
            for r in rows[io_row + 1:stop]:
                if col < len(r) and r[col]:
                    m = GUIDE_SEG_RE.match(str(r[col]).strip())
                    if m:
                        cable = norm_cable(m.group(1))
                        cnt = int(m.group(3)) if m.group(3) else _fc_of(cable)
                        rep.segments.append(ReportSegment(
                            cable=cable, seg=int(m.group(2)), count=cnt,
                            kind="Pigtail" if (cnt is not None and cnt <= 4) else "Cable",
                            start=name if direction == "out" else None, end=name if direction == "in" else None,
                            direction=direction))
    if hdr_row is not None:
        hdr = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_row]]
        sub = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_row + 1]] if hdr_row + 1 < len(rows) else []
        cable_cols = [j for j, c in enumerate(hdr) if c == "cable"]
        hot_col = sub.index("hot") if "hot" in sub else None
        a_col, b_col = cable_cols[0], cable_cols[1]

        def ck(v: Any) -> Optional[Tuple[str, int]]:
            m = GUIDE_CABLE_RE.match(str(v or "").strip())
            return (norm_cable(m.group(1)), int(m.group(2))) if m else None
        for r in rows[hdr_row + 2:]:
            if not any(v is not None and str(v).strip() for v in r):
                continue
            hot = str(r[hot_col] or "").strip().upper() if hot_col is not None and hot_col < len(r) else ""
            addr = ""
            if hot_col is not None:
                for v in r[hot_col + 1:]:
                    if v is not None and str(v).strip():
                        addr = norm_text(v)
                        break
            rep.conns.append(ReportConn(a_seg=ck(r[a_col]) if a_col < len(r) else None,
                                        b_seg=ck(r[b_col]) if b_col < len(r) else None, usage=addr,
                                        state="Live" if hot in ("H", "HOT", "Y", "YES", "1", "LIVE") else "Dark",
                                        tap="", address=addr))
    return rep


def load_splice_reports(files: List[Path], issues: IssueLog, rep: Reporter, quiet_cross: bool = False) -> List[SpliceReport]:
    reports: List[SpliceReport] = []
    for n, f in enumerate(files):
        rep.check_cancel()
        if n % 20 == 0:
            rep.progress(4 + 14 * n / max(1, len(files)), f"Reading splice reports ({n + 1}/{len(files)})")
        try:
            tables = read_tables(f)
        except Exception as exc:  # noqa: BLE001
            issues.add("ERROR", "Splice Report", f.name, f"Unreadable: {exc}")
            continue
        parsed = False
        for sheet, rows in tables:
            flat = " ".join(str(c) for r in rows[:6] for c in r if c is not None).lower()
            try:
                r = parse_guide_report(rows, f.name) if "splice guide" in flat else parse_connections_report(rows, f.name)
                if r is None and "splice guide" not in flat:
                    r = parse_guide_report(rows, f.name)
            except Exception as exc:  # noqa: BLE001
                issues.add("ERROR", "Splice Report", f"{f.name}:{sheet}", f"Parse failure: {exc}")
                continue
            if r is not None:
                reports.append(r)
                parsed = True
        if not parsed:
            if any(detect_header_row(rows)[2] >= 3 for _sheet, rows in tables if rows):
                if not quiet_cross:
                    issues.add("INFO", "Splice Report", f.name, "This is an address list - not used as splice report")
            else:
                issues.add("ERROR", "Splice Report", f.name, "Unrecognised splice report layout")
    names = Counter(r.name for r in reports)
    for nm, c in names.items():
        if c > 1:
            issues.add("WARNING", "Splice Report", nm, f"{c} reports share this closure name; merged")
    fmts = Counter(r.fmt for r in reports)
    if reports:
        rep.log(f"Splice reports: {len(reports)} parsed ({', '.join(f'{k}: {v}' for k, v in fmts.items())})")
    return reports


def splitter_size(reports: List[SpliceReport]) -> int:
    c = Counter(s for r in reports for s in r.splitters if 2 <= s <= 512)
    return c.most_common(1)[0][0] if c else 0


def pigtail_splices(reports: List[SpliceReport], tap_names: Optional[Set[str]] = None) -> Dict[str, str]:
    """terminal name -> closure, from pigtail segments 'CLOSURE TO TERMINAL FCnnn'."""
    out: Dict[str, str] = {}
    for r in reports:
        for s in r.segments:
            if s.direction == "out" and (s.kind.lower() == "pigtail") and " TO " in s.cable:
                dest = re.sub(r"\s+FC\s*\d+\s*$", "", s.cable.split(" TO ", 1)[1]).strip()
                out.setdefault(dest, r.name)
    return out


# =============================================================================
# Drawing sources (DXF / DWG-ODA via ezdxf, DWG via built-in reader)
# =============================================================================
@dataclass
class GeoInsert:
    handle: str
    layer: str
    block: str
    pos: Vec2
    attribs: Dict[str, str]
    corners: List[Vec2]
    box: Box


@dataclass
class GeoLine:
    handle: str
    layer: str
    pts: List[Vec2]


@dataclass
class GeoText:
    handle: str
    layer: str
    dxftype: str
    text: str
    pos: Vec2
    tip: Vec2
    box: Box
    height: float


@dataclass
class NetworkSource:
    kind: str                       # dxf | dwg-oda | dwg-native
    path: str
    version: str
    doc: Any
    layers: List[str]
    inserts: List[GeoInsert]
    lines: List[GeoLine]
    texts: List[GeoText]
    insunits: int = 0


def find_oda_converter(user_path: str = "") -> str:
    if user_path and Path(user_path).is_file():
        return user_path
    cands: List[str] = []
    if os.name == "nt":
        for root in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                     os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            cands += sorted(glob.glob(os.path.join(root, "ODA", "ODAFileConverter*", "ODAFileConverter.exe")), reverse=True)
    else:
        cands += [p for p in ("/usr/bin/ODAFileConverter", "/usr/local/bin/ODAFileConverter",
                              "/Applications/ODAFileConverter.app/Contents/MacOS/ODAFileConverter") if Path(p).exists()]
    which = shutil.which("ODAFileConverter")
    if which:
        cands.append(which)
    return cands[0] if cands else ""


def _configure_oda(path: str) -> bool:
    if not path:
        return False
    from ezdxf.addons import odafc
    ezdxf.options.set("odafc-addon", "win_exec_path" if os.name == "nt" else "unix_exec_path", path)
    try:
        return bool(odafc.is_installed())
    except Exception:  # noqa: BLE001
        return False


def _block_local_box(doc: Any, name: str, cache: Dict[str, Any]) -> Optional[Tuple[Vec2, Vec2]]:
    if name in cache:
        return cache[name]
    res = None
    try:
        blk = doc.blocks.get(name)
        if blk is not None:
            ents = [e for e in blk if e.dxftype() in ("LINE", "LWPOLYLINE", "POLYLINE", "CIRCLE", "ARC", "SOLID",
                                                         "ELLIPSE", "SPLINE", "HATCH", "TRACE", "3DFACE")]
            if ents:
                ext = ezbbox.extents(ents, fast=True)
                if ext.has_data:
                    base = blk.block.dxf.base_point
                    res = (Vec2(ext.extmin.x - base.x, ext.extmin.y - base.y), Vec2(ext.extmax.x - base.x, ext.extmax.y - base.y))
    except Exception:  # noqa: BLE001 - malformed block -> treated as point symbol
        res = None
    cache[name] = res
    return res


def _entity_points(e: Any) -> List[Vec2]:
    t = e.dxftype()
    if t == "LINE":
        return [v2(e.dxf.start), v2(e.dxf.end)]
    try:
        pts = [v2(v) for v in ezpath.make_path(e).flattening(distance=0.25)]
    except Exception:  # noqa: BLE001
        if t == "LWPOLYLINE":
            pts = [Vec2(x, y) for x, y in e.get_points("xy")]
        elif t == "POLYLINE":
            pts = [v2(v.dxf.location) for v in e.vertices]
        else:
            pts = []
    out: List[Vec2] = []
    for p in pts:
        if not out or (p - out[-1]).magnitude > 1e-6:
            out.append(p)
    return out


def block_key(name: str) -> str:
    """Comparison key of a block name; unresolved anonymous (dynamic-block) variants form one family."""
    n = (name or "").upper()
    return "*U*" if n.startswith("*U") else n


def match_closure_name(value: Any, names: Set[str]) -> str:
    """Closure name from an attribute / label value ('DMI-10007', 'DMI-10007 (FOSC)', 'SP: DMI-10007')."""
    v = norm_text(value)
    if not v or v in names:
        return v
    for tok in re.split(r"[\s,;:()\[\]/]+", v):
        if tok in names:
            return tok
    return v


def source_from_ezdxf(doc: Any, path: str, kind: str) -> NetworkSource:
    msp = doc.modelspace()
    cache: Dict[str, Any] = {}
    inserts: List[GeoInsert] = []
    lines: List[GeoLine] = []
    texts: List[GeoText] = []
    for e in msp:
        t = e.dxftype()
        try:
            if t == "INSERT":
                loc = _block_local_box(doc, e.dxf.name, cache)
                lo, hi = loc if loc else (Vec2(-1, -1), Vec2(1, 1))
                m = e.matrix44()
                blk = doc.blocks.get(e.dxf.name)
                base = blk.block.dxf.base_point if blk is not None else Vec3()
                corners = [v2(m.transform(Vec3(x + base.x, y + base.y, 0)))
                           for x, y in ((lo.x, lo.y), (hi.x, lo.y), (hi.x, hi.y), (lo.x, hi.y))]
                att = {a.dxf.tag.upper(): a.dxf.text for a in e.attribs}
                name = e.dxf.name
                if name.startswith("*"):
                    try:
                        from ezdxf.dynblkhelper import get_dynamic_block_definition
                        dyn = get_dynamic_block_definition(e, doc)
                        if dyn is not None:
                            name = dyn.name
                    except Exception:  # noqa: BLE001 - keep the anonymous name
                        pass
                inserts.append(GeoInsert(e.dxf.handle, e.dxf.layer, name, v2(e.dxf.insert), att, corners,
                                         box_of_points(corners)))
                for a in e.attribs:
                    if a.dxf.text and not a.is_invisible:
                        texts.append(GeoText(a.dxf.handle, a.dxf.layer, "ATTRIB", a.dxf.text, v2(a.dxf.insert),
                                             v2(a.dxf.insert), text_extent_box(v2(a.dxf.insert), a.dxf.height, a.dxf.text,
                                                                               a.dxf.rotation), a.dxf.height))
            elif t in ("LWPOLYLINE", "POLYLINE", "LINE", "SPLINE", "ARC"):
                pts = _entity_points(e)
                if len(pts) >= 2:
                    lines.append(GeoLine(e.dxf.handle, e.dxf.layer, pts))
            elif t == "TEXT":
                p = v2(e.dxf.insert)
                texts.append(GeoText(e.dxf.handle, e.dxf.layer, t, e.dxf.text, p, p,
                                     text_extent_box(p, e.dxf.height, e.dxf.text, e.dxf.rotation), e.dxf.height))
            elif t == "MTEXT":
                p = v2(e.dxf.insert)
                texts.append(GeoText(e.dxf.handle, e.dxf.layer, t, e.text, p, p,
                                     text_extent_box(p, e.dxf.char_height, e.text, e.dxf.get("rotation", 0.0),
                                                     e.dxf.attachment_point), e.dxf.char_height))
            elif t == "MULTILEADER":
                ctx = e.context
                if ctx.mtext is None:
                    continue
                txt = ctx.mtext.default_content
                loc = v2(ctx.mtext.insert)
                tip = None
                for ld in ctx.leaders:
                    for ln in ld.lines:
                        if ln.vertices:
                            tip = v2(ln.vertices[0])
                            break
                    if tip is not None:
                        break
                h = ctx.char_height or 1.0
                att = 3 if ctx.mtext.alignment == 3 else 1
                texts.append(GeoText(e.dxf.handle, e.dxf.layer, t, txt, loc, tip or loc,
                                     text_extent_box(loc, h, txt, 0.0, att), h))
        except Exception:  # noqa: BLE001 - one malformed entity must not stop the analysis
            continue
    return NetworkSource(kind=kind, path=path, version=doc.dxfversion, doc=doc,
                         layers=sorted(l.dxf.name for l in doc.layers), inserts=inserts, lines=lines, texts=texts,
                         insunits=int(doc.header.get("$INSUNITS", 0) or 0))


def source_from_native(path: str, rep: Reporter) -> NetworkSource:
    dwg = NativeDwg(path, progress=lambda p, m: rep.progress(20 + p * 0.15, m))
    ents = dwg.entities()
    if dwg.errors:
        rep.log("Built-in DWG reader skipped undecodable objects: " +
                ", ".join(f"{k} {v}" for k, v in dwg.errors.items()), "WARNING")
    block_pts: Dict[str, List[Tuple[float, float]]] = defaultdict(list)
    for e in ents:
        if e["space"] != "block" or e.get("owner") not in dwg.blocks:
            continue
        name = dwg.blocks[e["owner"]]
        if e["type"] in ("LINE", "LWPOLYLINE", "SOLID"):
            block_pts[name] += list(e.get("points", []))
        elif e["type"] in ("CIRCLE", "ARC"):
            cx, cy = e["center"]
            r = e["radius"]
            block_pts[name] += [(cx - r, cy - r), (cx + r, cy + r)]
    local: Dict[str, Tuple[Vec2, Vec2]] = {}
    for name, pts in block_pts.items():
        bx, by = dwg.block_base.get(name, (0.0, 0.0))
        local[name] = (Vec2(min(p[0] for p in pts) - bx, min(p[1] for p in pts) - by),
                       Vec2(max(p[0] for p in pts) - bx, max(p[1] for p in pts) - by))
    inserts: List[GeoInsert] = []
    lines: List[GeoLine] = []
    texts: List[GeoText] = []
    for e in ents:
        if e["space"] != "model":
            continue
        t = e["type"]
        h = format(e["handle"], "X")
        if t == "INSERT":
            lo, hi = local.get(e.get("block_def") or e["block"], (Vec2(-1, -1), Vec2(1, 1)))
            sx, sy = e["scale"]
            if e.get("mirrored"):
                sx = -sx
            ang = math.radians(e["rotation"])
            ca, sa = math.cos(ang), math.sin(ang)
            px, py = e["insert"]
            corners = []
            for x, y in ((lo.x, lo.y), (hi.x, lo.y), (hi.x, hi.y), (lo.x, hi.y)):
                x, y = x * sx, y * sy
                corners.append(Vec2(px + x * ca - y * sa, py + x * sa + y * ca))
            inserts.append(GeoInsert(h, e["layer"], e["block"], Vec2(px, py), e.get("attribs", {}), corners,
                                     box_of_points(corners)))
        elif t == "LWPOLYLINE" or t == "LINE":
            pts = [Vec2(x, y) for x, y in e.get("points", [])]
            if len(pts) >= 2:
                if e.get("closed"):
                    pts.append(pts[0])
                lines.append(GeoLine(h, e["layer"], pts))
        elif t in ("TEXT", "MTEXT"):
            p = Vec2(*e["insert"])
            att = e.get("attachment", 7) if t == "MTEXT" else 7
            texts.append(GeoText(h, e["layer"], t, e.get("text", ""), p, p,
                                 text_extent_box(p, e.get("height", 1.0), e.get("text", ""), e.get("rotation", 0.0), att),
                                 e.get("height", 1.0)))
        elif t == "MULTILEADER":
            loc = Vec2(*e["insert"])
            tip = Vec2(*e["tip"]) if e.get("tip") else loc
            hgt = e.get("height") or 1.0
            texts.append(GeoText(h, e["layer"], t, e.get("text", ""), loc, tip,
                                 text_extent_box(loc, hgt, e.get("text", ""), 0.0, 1 if e.get("dog", 1) > 0 else 3), hgt))
    rep.log(f"Built-in DWG reader: {dwg.release} ({dwg.version}), {len(dwg.layers)} layers, {len(inserts):,} block "
            f"references, {len(lines):,} linework entities, {len(texts):,} text/multileaders")
    return NetworkSource(kind="dwg-native", path=path, version=dwg.version, doc=None, layers=dwg.layer_names,
                         inserts=inserts, lines=lines, texts=texts, insunits=0)


def open_source(cfg: RunConfig, rep: Reporter) -> NetworkSource:
    p = Path(cfg.drawing_path)
    if not p.exists():
        raise FileNotFoundError(f"Drawing not found: {cfg.drawing_path}")
    if p.suffix.lower() == ".dwg":
        try:
            return source_from_native(str(p), rep)
        except DwgUnsupported as exc:
            oda = find_oda_converter()
            if oda and _configure_oda(oda):
                from ezdxf.addons import odafc
                rep.log(f"{exc} - converting with the ODA File Converter found at {oda}", "WARNING")
                return source_from_ezdxf(odafc.readfile(str(p)), str(p), "dwg-oda")
            raise RuntimeError(f"{exc}.") from exc
    if p.suffix.lower() == ".dxf":
        from ezdxf import recover
        doc, auditor = recover.readfile(str(p))
        if auditor.has_errors:
            rep.log(f"DXF recovered with {len(auditor.errors)} structural error(s) fixed", "WARNING")
        return source_from_ezdxf(doc, str(p), "dxf")
    raise RuntimeError(f"Unsupported drawing type '{p.suffix}' - select a .dwg or .dxf file")


def units_factor(source: NetworkSource, setting: str, rep: Reporter) -> float:
    table = {"feet": 1.0, "meters": 0.3048, "inches": 12.0, "millimeters": 304.8, "centimeters": 30.48}
    if setting in table:
        return table[setting]
    code = source.insunits
    by_code = {1: ("inches", 12.0), 2: ("feet", 1.0), 4: ("millimeters", 304.8), 5: ("centimeters", 30.48),
               6: ("meters", 0.3048), 21: ("US survey feet", 1.0)}
    if code in by_code:
        name, f = by_code[code]
        rep.log(f"Drawing units: {name} ($INSUNITS={code})")
        return f
    rep.log("Drawing units not stored in the file - feet assumed (change in Callout Style if metric)")
    return 1.0


# =============================================================================
# Auto-detection of the CAD mapping
# =============================================================================
def auto_detect_cad(src: NetworkSource, closure_names: Set[str], report_xy: List[Tuple[float, float]],
                    rep: Reporter, scale: float = 1.0) -> CadMap:
    cad = CadMap()
    values: Dict[Tuple[str, str, str], Set[str]] = defaultdict(set)
    positions: Dict[Tuple[str, str], List[Vec2]] = defaultdict(list)
    for ins in src.inserts:
        key = (ins.layer, block_key(ins.block))
        positions[key].append(ins.pos)
        for tag, val in ins.attribs.items():
            values[(ins.layer, block_key(ins.block), tag)].add(match_closure_name(val, closure_names))
    best = (0, None, None, None)
    for (layer, block, tag), vals in values.items():
        score = len(vals & closure_names)
        if score > best[0]:
            best = (score, layer, block, tag)
    if best[1]:
        cad.closure_layer, cad.closure_block, cad.closure_name_tag = best[1], best[2], best[3]
        rep.log(f"Closures: {best[0]} names matched on layer '{best[1]}', block '{best[2]}', attribute '{best[3]}'")
    elif report_xy:
        votes: Counter = Counter()
        grid = GridIndex(50.0)
        for ins in src.inserts:
            grid.insert((ins.pos.x, ins.pos.y, ins.pos.x, ins.pos.y), ins)
        for x, y in report_xy:
            it, d = grid.nearest(Vec2(x, y), 3.0 * scale, lambda i: i.pos)
            if it is not None:
                votes[(it.layer, block_key(it.block))] += 1
        if votes:
            (cad.closure_layer, cad.closure_block), n = votes.most_common(1)[0]
            rep.log(f"Closures: no name attribute matched; {n} splice-report coordinates matched block "
                    f"'{cad.closure_block}' on layer '{cad.closure_layer}'")
    else:
        texts = [t for t in src.texts if match_closure_name(t.text, closure_names) in closure_names]
        votes = Counter()
        grid = GridIndex(50.0)
        for ins in src.inserts:
            grid.insert((ins.pos.x, ins.pos.y, ins.pos.x, ins.pos.y), ins)
        for t in texts:
            it, _ = grid.nearest(t.pos, 60.0 * scale, lambda i: i.pos)
            if it is not None:
                votes[(it.layer, block_key(it.block))] += 1
        if votes:
            (cad.closure_layer, cad.closure_block), n = votes.most_common(1)[0]
            rep.log(f"Closures: identified from {n} name labels next to block '{cad.closure_block}'")
    closure_pts = positions.get((cad.closure_layer, cad.closure_block), [])
    if not closure_pts:
        rep.log("Closure blocks could not be auto-detected - set them in the CAD Mapping page", "WARNING")
        return cad
    cgrid = GridIndex(25.0)
    for p in closure_pts:
        cgrid.insert((p.x, p.y, p.x, p.y), p)

    def nearest_closure(p: Vec2, tol: float) -> float:
        _it, d = cgrid.nearest(p, tol, lambda q: q)
        return d if _it is not None else 1e18
    end_grid = GridIndex(25.0)
    vtx_d: Dict[str, List[float]] = defaultdict(list)        # distance of every vertex near a closure
    touch: Counter = Counter()                                # lines with >= 1 vertex on a closure
    total: Counter = Counter()
    for ln in src.lines:
        a, b = ln.pts[0], ln.pts[-1]
        end_grid.insert((a.x, a.y, a.x, a.y), (a, b))
        end_grid.insert((b.x, b.y, b.x, b.y), (b, a))
        total[ln.layer] += 1
        hit = False
        for p in ln.pts:
            d = nearest_closure(p, 3.0 * scale)
            if d < 1e17:
                vtx_d[ln.layer].append(d)
                hit = True
        if hit:
            touch[ln.layer] += 1
    sample = closure_pts[:: max(1, len(closure_pts) // 300)]
    stats = []
    for (layer, block), pts in positions.items():
        if (layer, block) == (cad.closure_layer, cad.closure_block) or len(pts) < 3:
            continue
        pgrid = GridIndex(50.0)
        for p in pts:
            pgrid.insert((p.x, p.y, p.x, p.y), p)
        dists = []
        for c in sample:
            _it, d = pgrid.nearest(c, 50.0 * scale, lambda q: q)
            dists.append(d if _it is not None else 1e9)
        window = sum(1 for d in dists if 8.0 * scale <= d <= 45.0 * scale) / max(1, len(dists))
        conn = 0
        inst = pts[:: max(1, len(pts) // 200)]
        for p in inst:
            for (_b, (here, other)) in end_grid.query((p.x - 1, p.y - 1, p.x + 1, p.y + 1)):
                if (here - p).magnitude <= 1.0 * scale and nearest_closure(other, 2.0 * scale) < 1e17:
                    conn += 1
                    break
        stats.append((layer, block, len(pts), window, conn / max(1, len(inst))))
    terminals = sorted([s for s in stats if s[4] >= 0.5], key=lambda s: -s[2])
    if terminals:
        cad.terminal_layer = terminals[0][0]
    hh = [s for s in stats if s[4] < 0.2 and s[3] >= 0.05]
    if hh:
        by_layer: Counter = Counter()
        for layer, _b, _n, window, _c in hh:
            by_layer[layer] += window
        top = by_layer.most_common(1)[0][0]
        cad.handhole_layer = top
        cad.handhole_blocks = [b for (l, b, n, w, c) in sorted(hh, key=lambda s: -s[3]) if l == top and w >= 0.05]
    # pigtail layers: linework running from a closure to a terminal block
    tgrid = GridIndex(25.0)
    for ins in src.inserts:
        if cad.terminal_layer and ins.layer == cad.terminal_layer:
            tgrid.insert((ins.pos.x, ins.pos.y, ins.pos.x, ins.pos.y), ins.pos)
    pig: Counter = Counter()
    for ln in src.lines:
        a, b = ln.pts[0], ln.pts[-1]
        for here, other in ((a, b), (b, a)):
            if nearest_closure(here, 2.0 * scale) < 1e17 and tgrid.nearest(other, 2.0 * scale, lambda q: q)[0] is not None:
                pig[ln.layer] += 1
                break
    cables = []
    for layer, ds in vtx_d.items():
        hits = len(ds)
        if hits >= 2 and touch[layer] / max(1, total[layer]) >= 0.3 and pig[layer] < 0.5 * touch[layer]:
            if not (cad.terminal_layer and layer == cad.terminal_layer):
                cables.append((hits, layer))
    if cables:
        cad.cable_layers = [l for _h, l in sorted(cables, reverse=True)]
        close = sorted(d for l in cad.cable_layers for d in vtx_d[l])
        if close:
            p95 = close[max(0, int(len(close) * 0.95) - 1)]
            cad.snap_tolerance = round(min(3.0 * scale, max(0.5 * scale, p95 * 1.5 + 0.01)), 3)
    if not cad.snap_tolerance:
        cad.snap_tolerance = 2.0 * scale
    return cad


def merge_cad(user: CadMap, auto: CadMap) -> CadMap:
    out = dataclasses.replace(auto)
    for f in CadMap.__dataclass_fields__:
        val = getattr(user, f)
        if val:
            setattr(out, f, val)
    return out


# =============================================================================
# Network model
# =============================================================================
@dataclass
class Closure:
    name: str
    pos: Vec2
    handle: str
    box: Box
    origin: str = "attribute"


@dataclass
class Handhole:
    pos: Vec2
    handle: str
    block: str
    corners: List[Vec2]
    box: Box


@dataclass
class CablePiece:
    handle: str
    layer: str
    pts: List[Vec2]
    a: Optional[str]
    b: Optional[str]


class DrawingModel:
    def __init__(self, src: NetworkSource, cad: CadMap, spec: CalloutSpec, issues: IssueLog, rep: Reporter,
                 known_names: Set[str], report_xy: Dict[str, Tuple[float, float]], scale: float) -> None:
        self.src = src
        self.cad = cad
        self.spec = spec
        self.issues = issues
        self.rep = rep
        self.scale = scale
        self.closures: Dict[str, Closure] = {}
        self.handholes: List[Handhole] = []
        self.pieces: List[CablePiece] = []
        self.adj: Dict[str, List[int]] = defaultdict(list)
        self.existing: List[GeoText] = []
        self.obstacles = GridIndex(50.0)
        self.text_obstacles = GridIndex(50.0)
        self.cable_points = GridIndex(50.0)
        self.hh_index = GridIndex(100.0)
        self.closure_index = GridIndex(50.0)
        self._build(known_names, report_xy)

    def _add_closure(self, c: Closure) -> None:
        self.closures[c.name] = c
        self.obstacles.insert(c.box, ("closure", c.name))
        self.closure_index.insert((c.pos.x, c.pos.y, c.pos.x, c.pos.y), c.name)

    def _build(self, known_names: Set[str], report_xy: Dict[str, Tuple[float, float]]) -> None:
        cad, spec = self.cad, self.spec
        cl_layer = cad.closure_layer.upper()
        cl_block = block_key(cad.closure_block)
        tag = cad.closure_name_tag.upper()
        hh_layer = cad.handhole_layer.upper()
        hh_blocks = {block_key(b) for b in cad.handhole_blocks}
        term_layer = cad.terminal_layer.upper()
        out_layers = {spec.cable_layer.upper(), spec.terminal_layer.upper()}
        unnamed: List[GeoInsert] = []
        for ins in self.src.inserts:
            layer = ins.layer.upper()
            if cl_layer and layer == cl_layer and (not cl_block or block_key(ins.block) == cl_block):
                name = match_closure_name(ins.attribs.get(tag, ""), known_names) if tag else ""
                if not name:
                    unnamed.append(ins)
                    continue
                if name in self.closures:
                    self.issues.add("WARNING", "Drawing", name, "Duplicate closure name in drawing; first kept")
                    continue
                self._add_closure(Closure(name, ins.pos, ins.handle, ins.box))
            elif hh_layer and layer == hh_layer and block_key(ins.block) in hh_blocks:
                h = Handhole(ins.pos, ins.handle, ins.block, ins.corners, ins.box)
                self.handholes.append(h)
                self.hh_index.insert(ins.box, h)
                self.obstacles.insert(ins.box, ("handhole", ins.handle))
            else:
                self.obstacles.insert(ins.box, ("symbol" if layer != term_layer else "terminal", ins.handle))
        # fallback 1: closure name labels (TEXT/MTEXT) next to unnamed closure blocks
        if unnamed:
            ugrid = GridIndex(50.0)
            for ins in unnamed:
                ugrid.insert((ins.pos.x, ins.pos.y, ins.pos.x, ins.pos.y), ins)
            named = 0
            taken: Set[str] = set()
            for t in self.src.texts:
                nm = match_closure_name(t.text, known_names)
                if nm in known_names and nm not in self.closures:
                    ins, _d = ugrid.nearest(t.pos, 60.0 * self.scale, lambda i: i.pos)
                    if ins is not None and ins.handle not in taken:
                        taken.add(ins.handle)
                        self._add_closure(Closure(nm, ins.pos, ins.handle, ins.box, origin="label"))
                        named += 1
            if named:
                self.rep.log(f"{named} closure(s) named from text labels (blocks without a name attribute)")
            rest = len(unnamed) - named
            if rest:
                self.issues.add("INFO", "Drawing", cad.closure_layer, f"{rest} closure block(s) without a name were ignored")
        # fallback 2: splice report coordinates
        located = virtual = 0
        cand_grid = GridIndex(50.0)
        for ins in self.src.inserts:
            if cl_layer and ins.layer.upper() == cl_layer:
                cand_grid.insert((ins.pos.x, ins.pos.y, ins.pos.x, ins.pos.y), ins)
        used_handles = {c.handle for c in self.closures.values()}
        for name, (x, y) in report_xy.items():
            if name in self.closures:
                d = (self.closures[name].pos - Vec2(x, y)).magnitude
                if d > 50.0 * self.scale:
                    self.issues.add("WARNING", "Drawing", name,
                                    f"Closure symbol is {d:,.0f} units from the splice report coordinates")
                continue
            ins, _d = cand_grid.nearest(Vec2(x, y), 5.0 * self.scale, lambda i: i.pos)
            if ins is not None and ins.handle not in used_handles:
                used_handles.add(ins.handle)
                self._add_closure(Closure(name, ins.pos, ins.handle, ins.box, origin="report coordinates"))
                located += 1
            else:
                s = 6.0 * self.scale
                self._add_closure(Closure(name, Vec2(x, y), "", (x - s, y - s, x + s, y + s), origin="report (no symbol)"))
                virtual += 1
                self.issues.add("WARNING", "Drawing", name,
                                "Closure symbol not found in drawing - located from splice report coordinates")
        if located:
            self.rep.log(f"{located} closure(s) located from splice report coordinates")
        for t in self.src.texts:
            layer = t.layer.upper()
            if layer in out_layers and t.dxftype in ("MULTILEADER", "MTEXT", "TEXT") and \
                    is_fiber_callout_text(t.text, spec.cable_prefix):
                self.existing.append(t)
            else:
                self.text_obstacles.insert(t.box, ("text", t.handle))
        self.rep.log(f"Drawing: {len(self.closures)} closures, {len(self.handholes)} handholes, "
                     f"{len(self.existing)} existing fiber callouts on the output layer(s)")
        if not self.closures:
            self.issues.add("ERROR", "Drawing", cad.closure_layer or "(closure layer)",
                            "No named closures found - check the CAD Mapping page")
        if not self.handholes:
            self.issues.add("WARNING", "Drawing", cad.handhole_layer or "(handhole layer)",
                            "No handholes mapped - terminal callouts will point at closure symbols")
        cable_layers = {l.upper() for l in cad.cable_layers}
        self._split_cables([ln for ln in self.src.lines if ln.layer.upper() in cable_layers])

    def nearest_closure(self, p: Vec2, tol: float) -> Optional[str]:
        name, _d = self.closure_index.nearest(p, tol, lambda n: self.closures[n].pos)
        return name

    def _split_cables(self, lines: List[GeoLine]) -> None:
        tol = self.cad.snap_tolerance or 2.0 * self.scale
        step = 5.0 * self.scale
        for ln in lines:
            pts = ln.pts
            for i in range(len(pts) - 1):
                a, b = pts[i], pts[i + 1]
                self.cable_points.insert((a.x, a.y, a.x, a.y), a)
                L = (b - a).magnitude
                for k in range(1, int(L // step) + 1):
                    q = a + (b - a) * (k * step / L)
                    self.cable_points.insert((q.x, q.y, q.x, q.y), q)
            enriched: List[Tuple[Vec2, Optional[str]]] = []
            for i in range(len(pts)):
                enriched.append((pts[i], self.nearest_closure(pts[i], tol)))
                if i + 1 < len(pts):
                    a, b = pts[i], pts[i + 1]
                    mids = []
                    for (_bx, name) in self.closure_index.query(box_of_points([a, b], tol)):
                        c = self.closures[name].pos
                        d, t = seg_distance(c, a, b)
                        if d <= tol and 1e-6 < t < 1 - 1e-6:
                            mids.append((t, c, name))
                    for _t, c, name in sorted(mids, key=lambda m: m[0]):
                        enriched.append((c, name))
            idx = [i for i, (_, nm) in enumerate(enriched) if nm]
            plist = [p for p, _ in enriched]
            if not idx:
                self.pieces.append(CablePiece(ln.handle, ln.layer, plist, None, None))
                continue
            if idx[0] != 0:
                self.pieces.append(CablePiece(ln.handle, ln.layer, plist[:idx[0] + 1], None, enriched[idx[0]][1]))
            for i0, i1 in zip(idx, idx[1:]):
                a_n, b_n = enriched[i0][1], enriched[i1][1]
                if a_n == b_n:
                    continue
                piece = CablePiece(ln.handle, ln.layer, plist[i0:i1 + 1], a_n, b_n)
                piece.pts[0] = self.closures[a_n].pos
                piece.pts[-1] = self.closures[b_n].pos
                self.pieces.append(piece)
            if idx[-1] != len(plist) - 1:
                self.pieces.append(CablePiece(ln.handle, ln.layer, plist[idx[-1]:], enriched[idx[-1]][1], None))
        for i, pc in enumerate(self.pieces):
            if pc.a and pc.b:
                self.adj[pc.a].append(i)
                self.adj[pc.b].append(i)
        self.rep.log(f"Cable topology: {len(lines)} cable entities -> {sum(1 for p in self.pieces if p.a and p.b)} "
                     f"closure-to-closure links (snap tolerance {tol:g})")

    def other_end(self, pi: int, node: str) -> str:
        pc = self.pieces[pi]
        return pc.b if pc.a == node else pc.a  # type: ignore[return-value]

    def piece_from(self, pi: int, node: str) -> List[Vec2]:
        pc = self.pieces[pi]
        return list(pc.pts) if pc.a == node else list(reversed(pc.pts))

    def find_piece_toward(self, start: str, target: str, max_nodes: int = 20000) -> Optional[int]:
        for pi in self.adj.get(start, []):
            if self.other_end(pi, start) == target:
                return pi
        for pi in self.adj.get(start, []):
            seen = {start}
            stack = [self.other_end(pi, start)]
            while stack and len(seen) < max_nodes:
                n = stack.pop()
                if n == target:
                    return pi
                if n in seen:
                    continue
                seen.add(n)
                for pj in self.adj.get(n, []):
                    m = self.other_end(pj, n)
                    if m not in seen:
                        stack.append(m)
        return None

    def handhole_for(self, closure: Closure) -> Optional[Handhole]:
        h, _d = self.hh_index.nearest(closure.pos, self.spec.handhole_search_radius, lambda x: x.pos)
        return h


# =============================================================================
# Planning
# =============================================================================
@dataclass
class CalloutPlan:
    uid: str
    kind: str
    olm: str
    text: str
    layer: str
    tip: Vec2
    last: Vec2
    dog: int
    anchor: str
    closure: str
    numbers: List[int]
    calculation: str
    sources: str
    status: str = "PLANNED"
    notes: str = ""
    handle: str = ""
    existing_match: str = ""
    existing_text: str = ""


@dataclass
class SegmentPlanInfo:
    closure: str
    cable: str
    seg: int
    count: Optional[int]
    end: Optional[str]
    live: List[int]
    unresolved: List[str]
    report_file: str
    olm: str
    text: str = ""
    geometry: str = ""


class CalloutPlanner:
    def __init__(self, model: DrawingModel, book: AddressBook, reports: List[SpliceReport], spec: CalloutSpec,
                 issues: IssueLog, rep: Reporter) -> None:
        self.m = model
        self.book = book
        self.reports = reports
        self.spec = spec
        self.issues = issues
        self.rep = rep
        self.placed = GridIndex(50.0)
        self.leader_index = GridIndex(50.0)
        self.plans: List[CalloutPlan] = []
        self.segments: List[SegmentPlanInfo] = []
        self.terminal_rows: List[Dict[str, Any]] = []

    # ------------------------------------------------------------ geometry of a callout
    def text_location(self, last: Vec2, dog: int) -> Vec2:
        s = self.spec
        return Vec2(last.x + dog * (s.dogleg_length + s.landing_gap), last.y + s.text_y_factor * s.text_height)

    def text_box(self, last: Vec2, dog: int, lines: List[str]) -> Box:
        h = self.spec.text_height
        width = max(len(x) for x in lines) * h * 0.9
        height = h + (len(lines) - 1) * h * 1.667 + 0.2 * h
        loc = self.text_location(last, dog)
        if dog > 0:
            return (loc.x, loc.y - height, loc.x + width, loc.y)
        return (loc.x - width, loc.y - height, loc.x, loc.y)

    def _score(self, tip: Vec2, last: Vec2, dog: int, lines: List[str], away_from: Optional[Vec2]) -> float:
        box = self.text_box(last, dog, lines)
        m = 0.25 * self.spec.text_height
        clear = (box[0] - m, box[1] - m, box[2] + m, box[3] + m)
        score = 0.0
        for (b, _i) in self.placed.query(clear):
            ov = box_overlap(clear, b)
            if ov > 0:
                score += 400.0 + 60.0 * ov / self.spec.text_height
        for (b, _i) in self.m.obstacles.query(clear):
            ov = box_overlap(clear, b)
            if ov > 0:
                score += 150.0 + 20.0 * ov / self.spec.text_height
        for (b, _i) in self.m.text_obstacles.query(box):
            ov = box_overlap(box, b)
            if ov > 0:
                score += 60.0 + 8.0 * ov / self.spec.text_height
        score += 25.0 * sum(1 for (b, _p) in self.m.cable_points.query(box)
                            if box[0] <= b[0] <= box[2] and box[1] <= b[1] <= box[3])
        lb = box_of_points([tip, last], 0.1)
        for (_b, seg) in self.leader_index.query(lb):
            if segments_intersect(tip, last, seg[0], seg[1]):
                score += 300.0
        for (b, _i) in self.placed.query(lb):
            if segment_hits_box(tip, last, b):
                score += 150.0
        if away_from is not None:
            d, v = last - tip, tip - away_from
            if d.magnitude > 1e-9 and v.magnitude > 1e-9:
                score += 20.0 * (1.0 - d.normalize().dot(v.normalize()))
        return score

    def _register(self, plan: CalloutPlan) -> None:
        self.placed.insert(self.text_box(plan.last, plan.dog, plan.text.split("\\P")), plan.uid)
        self.leader_index.insert(box_of_points([plan.tip, plan.last], 0.1), (plan.tip, plan.last))
        self.plans.append(plan)

    # ------------------------------------------------------------ terminals
    def plan_terminals(self) -> None:
        spec = self.spec
        by_closure: Dict[str, List[Terminal]] = defaultdict(list)
        for _k, t in sorted(self.book.terminals.items(), key=lambda kv: (kv[0][0], min(kv[1].numbers))):
            by_closure[t.splice].append(t)
        uid = 0
        for closure_name, terms in by_closure.items():
            self.rep.check_cancel()
            closure = self.m.closures.get(closure_name)
            hh = self.m.handhole_for(closure) if closure else None
            for t in terms:
                runs = number_runs(t.numbers) if spec.split_terminal_runs else [(min(t.numbers), max(t.numbers))]
                texts = [fmt_run(r, spec.zero_pad) for r in runs] if spec.split_terminal_runs else \
                    [fmt_numbers(t.numbers, "exact", spec.zero_pad)]
                calc = "; ".join(f"{r.address}: PON {r.pon} port {r.port} -> ({r.pon}-1)x{self.book.ports_per_pon}+{r.port} = {r.number}"
                                 for r in sorted(t.records, key=lambda x: x.number))
                srcs = ", ".join(sorted({r.source for r in t.records}))
                row = dict(OLM=t.olm, Terminal=t.tap, Closure=t.splice or "(none)",
                           Numbers=fmt_numbers(t.numbers, "exact", 0), Addresses=len(t.records),
                           Callouts=", ".join(texts), Status="")
                for txt, run in zip(texts, runs):
                    uid += 1
                    nums = sorted(n for n in t.numbers if run[0] <= n <= run[1]) if spec.split_terminal_runs else sorted(t.numbers)
                    if closure is None:
                        self.plans.append(CalloutPlan(f"T{uid:05d}", "TERMINAL", t.olm, txt, spec.terminal_layer,
                                                      Vec2(), Vec2(), 1, "-", t.splice, nums, calc, srcs, status="FAILED",
                                                      notes="Drop closure not found in drawing" if t.splice else
                                                      "Drop splice closure unknown"))
                        continue
                    plan = self._place_terminal(f"T{uid:05d}", t, txt, closure, hh, nums, calc, srcs)
                    self._register(plan)
                if closure is None and not t.splice:
                    self.issues.add("ERROR", "Missing Data", t.tap,
                                    f"Drop splice closure of terminal {t.tap} ({fmt_numbers(t.numbers, 'exact', 0)}) is unknown")
                    row["Status"] = "NOT PLACED - drop closure unknown"
                elif closure is None:
                    self.issues.add("ERROR", "Missing Closure", t.splice,
                                    f"Drop closure of terminal {t.tap} ({fmt_numbers(t.numbers, 'exact', 0)}) not found in drawing")
                    row["Status"] = "NOT PLACED - closure missing"
                else:
                    row["Status"] = "PLACED" if hh else "PLACED on closure symbol (no handhole found)"
                self.terminal_rows.append(row)

    def _place_terminal(self, uid: str, t: Terminal, text: str, closure: Closure, hh: Optional[Handhole],
                        nums: List[int], calc: str, srcs: str) -> CalloutPlan:
        spec = self.spec
        if hh is not None:
            corners = hh.corners
            anchor = f"{hh.block} {hh.handle}"
            note = ""
        else:
            x0, y0, x1, y1 = closure.box
            corners = [Vec2(x0, y0), Vec2(x1, y0), Vec2(x1, y1), Vec2(x0, y1)]
            anchor = f"closure {closure.name}"
            note = "No handhole within search radius; arrow on closure symbol"
            self.issues.add("WARNING", "Placement", t.tap,
                            f"No handhole within {spec.handhole_search_radius:g} of closure {closure.name}; arrow on closure symbol")
        center = Vec2(sum(c.x for c in corners) / 4, sum(c.y for c in corners) / 4)
        cands: List[Tuple[Vec2, Vec2]] = []
        for i in range(4):
            for p in (corners[i], (corners[i] + corners[(i + 1) % 4]) * 0.5):
                d = p - center
                if d.magnitude > 1e-9:
                    cands.append((p, d.normalize()))
        if not cands:
            cands = [(center, Vec2(1, 1).normalize())]
        best = None
        for tip, direction in cands:
            for lf in (1.0, 1.6):
                last = tip + direction * (spec.terminal_leader_length * lf)
                dogs = [1 if direction.x > 0 else -1] if abs(direction.x) > 1e-6 else [1, -1]
                for dog in dogs:
                    sc = self._score(tip, last, dog, [text], closure.pos) + (lf - 1.0) * 20.0
                    if best is None or sc < best[0]:
                        best = (sc, tip, last, dog)
        _, tip, last, dog = best  # type: ignore[misc]
        return CalloutPlan(uid, "TERMINAL", t.olm, text, spec.terminal_layer, tip, last, dog, anchor, closure.name,
                           nums, f"Terminal {t.tap}: " + calc, srcs, notes=note)

    # ------------------------------------------------------------ cables (BFO)
    def _report_olm(self, r: SpliceReport) -> str:
        votes: Counter = Counter()
        for c in r.conns:
            if c.address and not NON_ADDRESS_RE.search(c.usage or c.address):
                rec = self.book.resolve(c.tap, c.address, None, c.port_idx)
                if rec:
                    votes[rec.olm] += 1
        if votes:
            return votes.most_common(1)[0][0]
        c2 = Counter(rec.olm for rec in self.book.records if rec.splice == r.name)
        if c2:
            return c2.most_common(1)[0][0]
        if r.olm and r.olm in self.book.olms:
            return r.olm
        return self.book.olms[0] if len(self.book.olms) == 1 else (r.olm or "?")

    def _segment_piece(self, start: str, end: Optional[str], dest: str) -> Optional[int]:
        if start not in self.m.closures:
            return None
        pi = None
        if end and end in self.m.closures:
            pi = self.m.find_piece_toward(start, end)
        if pi is None and dest and dest in self.m.closures and dest != start:
            pi = self.m.find_piece_toward(start, dest)
        if pi is None and len(self.m.adj.get(start, [])) == 1:
            pi = self.m.adj[start][0]
        return pi

    def plan_cables(self) -> None:
        spec = self.spec
        merged: Dict[str, List[SpliceReport]] = defaultdict(list)
        for r in self.reports:
            merged[r.name].append(r)
        ends: Dict[Tuple[str, int], Tuple[Optional[str], Optional[str]]] = {}
        for r in self.reports:
            for s in r.segments:
                st, en = ends.get((s.cable, s.seg), (None, None))
                ends[(s.cable, s.seg)] = (st or s.start, en or s.end)
        done: Set[Tuple[str, str, int]] = set()
        uid = 0
        for n_rep, (name, reps) in enumerate(sorted(merged.items())):
            self.rep.check_cancel()
            if n_rep % 25 == 0:
                self.rep.progress(60 + 20 * n_rep / max(1, len(merged)), f"Planning cable callouts ({n_rep}/{len(merged)})")
            closure = self.m.closures.get(name)
            if closure is None:
                self.issues.add("ERROR", "Missing Closure", name, "Splice report closure not found in drawing")
            for r in reps:
                olm = self._report_olm(r)
                r.olm = olm
                for sg in r.segments:
                    if sg.direction != "out":
                        continue
                    key = (sg.cable, sg.seg)
                    if (name, sg.cable, sg.seg) in done:
                        continue
                    done.add((name, sg.cable, sg.seg))
                    dest = re.sub(r"\s+FC\s*\d+\s*$", "", sg.cable.split(" TO ", 1)[1]).strip() if " TO " in sg.cable else ""
                    if sg.kind.lower() == "pigtail" or dest in self.book.tap_names:
                        continue
                    end = sg.end or ends.get(key, (None, None))[1]
                    piece = self._segment_piece(name, end, dest)
                    live: List[int] = []
                    unresolved: List[str] = []
                    ambiguous: Dict[Tuple[int, ...], List[str]] = defaultdict(list)
                    for c in r.conns:
                        if key not in (c.a_seg, c.b_seg):
                            continue
                        if any(s_ in norm_text(c.state) for s_ in DEAD_STATES):
                            continue
                        if not c.address or NON_ADDRESS_RE.search(c.usage or c.address):
                            continue
                        cands = self.book.resolve_candidates(c.tap, c.address, olm, c.port_idx)
                        if len(cands) > 1 and piece is not None:
                            down = [x for x in cands if x.splice in self.m.closures
                                    and self.m.find_piece_toward(name, x.splice) == piece]
                            cands = down or cands
                        if len(cands) == 1:
                            live.append(cands[0].number)
                        elif cands:
                            ambiguous[tuple(sorted(x.number for x in cands))].append(c.usage or c.address)
                        else:
                            unresolved.append(c.usage or c.address)
                    for nums, usages in ambiguous.items():
                        if len(usages) == len(nums):
                            live.extend(nums)
                        else:
                            unresolved.extend(u + " (ambiguous)" for u in usages)
                    info = SegmentPlanInfo(name, sg.cable, sg.seg, sg.count, end, sorted(set(live)), unresolved, r.file, olm)
                    self.segments.append(info)
                    if unresolved:
                        self.issues.add("WARNING", "Splice Mismatch", f"{name} {sg.cable}:Seg{sg.seg}",
                                        f"{len(unresolved)} live fiber usage(s) not found in the address list, e.g. {unresolved[0]!r}")
                    if not info.live:
                        info.geometry = "no callout - no live fibers"
                        continue
                    if sg.count is None:
                        self.issues.add("WARNING", "Splice Report", f"{name} {sg.cable}",
                                        "Segment fiber count unknown (no Count column and no FCnnn suffix)")
                    count_txt = str(sg.count) if sg.count is not None else "?"
                    text = f"{spec.cable_prefix} {count_txt}\\P{fmt_numbers(info.live, spec.cable_range_mode, spec.zero_pad)}"
                    info.text = text
                    uid += 1
                    plan = self._place_cable(f"C{uid:05d}", info, closure, piece, text)
                    info.geometry = plan.anchor if plan.status != "FAILED" else plan.notes
                    if plan.status == "FAILED":
                        self.plans.append(plan)
                    else:
                        self._register(plan)

    def _place_cable(self, uid: str, info: SegmentPlanInfo, closure: Optional[Closure], pi: Optional[int],
                     text: str) -> CalloutPlan:
        spec = self.spec
        calc = (f"Segment {info.seg} of {info.cable} leaving {info.closure}: {len(info.live)} live fiber(s) "
                f"{fmt_numbers(info.live, 'exact', 0)}; fiber count {info.count}; text = '{spec.cable_prefix} <count>' + "
                f"{'min-max' if spec.cable_range_mode == 'span' else 'runs'}")
        base = dict(uid=uid, kind="CABLE", olm=info.olm, text=text, layer=spec.cable_layer, closure=info.closure,
                    numbers=info.live, calculation=calc, sources=info.report_file)
        if closure is None:
            return CalloutPlan(tip=Vec2(), last=Vec2(), dog=1, anchor="-", status="FAILED",
                               notes="Closure not in drawing", **base)
        note = ""
        if pi is not None:
            pts = self.m.piece_from(pi, info.closure)
            desc = f"cable {self.m.pieces[pi].handle} toward {self.m.other_end(pi, info.closure)}"
        elif info.end and info.end in self.m.closures:
            pts = [closure.pos, self.m.closures[info.end].pos]
            desc = f"straight line toward {info.end}"
            note = "No cable polyline between the closures; arrow placed on the straight line"
            self.issues.add("WARNING", "Placement", f"{info.closure} {info.cable}:Seg{info.seg}", note)
        else:
            self.issues.add("WARNING", "Placement", f"{info.closure} {info.cable}:Seg{info.seg}",
                            "Cable geometry of this segment not found in the drawing; callout not placed")
            return CalloutPlan(tip=Vec2(), last=Vec2(), dog=1, anchor="-", status="FAILED",
                               notes="Segment geometry not found", **base)
        L = polyline_length(pts)
        if L <= 1e-6:
            return CalloutPlan(tip=Vec2(), last=Vec2(), dog=1, anchor="-", status="FAILED",
                               notes="Zero-length segment geometry", **base)
        s_along = max(min(spec.cable_offset_along, 0.45 * L), min(5.0 * spec.text_height / 4.0, 0.5 * L))
        lines = text.split("\\P")
        best = None
        for frac in (1.0, 0.7, 1.4, 2.0, 2.8):
            s = min(max(s_along * frac, 0.05 * L), 0.9 * L)
            tip, tan = point_at_length(pts, s)
            nrm = Vec2(-tan.y, tan.x)
            for side in (1, -1):
                for skew in (0.0, 35.0, -35.0):
                    d = (nrm * side).rotate_deg(skew)
                    for lf in (1.0, 1.5):
                        last = tip + d * (spec.cable_leader_length * lf)
                        dog = 1 if d.x >= 0 else -1
                        sc = (self._score(tip, last, dog, lines, closure.pos) + abs(frac - 1.0) * 12.0 + (lf - 1.0) * 20.0)
                        if best is None or sc < best[0]:
                            best = (sc, tip, last, dog)
        _, tip, last, dog = best  # type: ignore[misc]
        return CalloutPlan(tip=tip, last=last, dog=dog, anchor=desc, notes=note, **base)

    # ------------------------------------------------------------ QA
    def placement_qa(self) -> Dict[str, int]:
        live = [p for p in self.plans if p.status != "FAILED"]
        boxes = {p.uid: self.text_box(p.last, p.dog, p.text.split("\\P")) for p in live}
        tgrid = GridIndex(50.0)
        lgrid = GridIndex(50.0)
        for p in live:
            tgrid.insert(boxes[p.uid], p.uid)
            lgrid.insert(box_of_points([p.tip, p.last], 0.1), (p.uid, p.tip, p.last))
        counts: Counter = Counter()
        for p in live:
            b = boxes[p.uid]
            flags = []
            if any(u != p.uid and box_overlap(b, ob) > 0 for (ob, u) in tgrid.query(b)):
                flags.append("text overlaps another callout")
                counts["text_overlap"] += 1
            if any(box_overlap(b, ob) > 0 for (ob, _i) in self.m.obstacles.query(b)):
                flags.append("text overlaps a symbol")
                counts["symbol_overlap"] += 1
            if any(u != p.uid and segments_intersect(p.tip, p.last, a, c)
                   for (_bb, (u, a, c)) in lgrid.query(box_of_points([p.tip, p.last], 0.1))):
                flags.append("leader crosses another leader")
                counts["leader_cross"] += 1
            if flags:
                p.notes = "; ".join(x for x in [p.notes, "PLACEMENT REVIEW: " + ", ".join(flags)] if x)
        counts["clean"] = sum(1 for p in live if "PLACEMENT REVIEW" not in p.notes)
        return dict(counts)


def compare_existing(plans: List[CalloutPlan], existing: List[GeoText], radius: float) -> List[Dict[str, Any]]:
    idx = GridIndex(100.0)
    for i, ex in enumerate(existing):
        idx.insert((ex.tip.x, ex.tip.y, ex.tip.x, ex.tip.y), i)
    used: Set[int] = set()
    rows: List[Dict[str, Any]] = []
    for p in plans:
        if p.status == "FAILED":
            continue
        want = canonical_callout_text(p.text)
        cands = []
        for (_b, i) in idx.query((p.tip.x - radius, p.tip.y - radius, p.tip.x + radius, p.tip.y + radius)):
            if i not in used:
                d = (existing[i].tip - p.tip).magnitude
                if d <= radius:
                    cands.append((d, i))
        exact = sorted((d, i) for d, i in cands if canonical_callout_text(existing[i].text) == want)
        if exact:
            d, i = exact[0]
            used.add(i)
            strict = existing[i].text.replace("\n", "\\P") == p.text
            p.existing_match = "MATCH" if strict else "MATCH (formatting differs)"
            p.existing_text = existing[i].text
            rows.append(dict(Callout=p.uid, Kind=p.kind, Computed=p.text.replace("\\P", " / "),
                             Existing=existing[i].text.replace("\\P", " / ").replace("\n", " / "),
                             ExistingHandle=existing[i].handle, Distance=round(d, 2), Result=p.existing_match))
        else:
            p.existing_match = "NEW"
            near = sorted(cands)[:1]
            rows.append(dict(Callout=p.uid, Kind=p.kind, Computed=p.text.replace("\\P", " / "),
                             Existing=existing[near[0][1]].text.replace("\\P", " / ").replace("\n", " / ") if near else "",
                             ExistingHandle=existing[near[0][1]].handle if near else "",
                             Distance=round(near[0][0], 2) if near else None,
                             Result="TEXT DIFFERS / NOT FOUND" if near else "NOT IN DRAWING"))
    for i, ex in enumerate(existing):
        if i not in used:
            rows.append(dict(Callout="", Kind="", Computed="", Existing=ex.text.replace("\\P", " / ").replace("\n", " / "),
                             ExistingHandle=ex.handle, Distance=None, Result="EXISTING ONLY (no computed counterpart)"))
    return rows


# =============================================================================
# Writing
# =============================================================================
class CalloutWriter:
    """Writes MULTILEADER callouts with the exact reference properties. All override flags are set, so the
    callouts look identical whatever the target drawing's MLEADER/Standard style contains."""

    def __init__(self, doc: Any, spec: CalloutSpec, issues: IssueLog, rep: Reporter) -> None:
        self.doc = doc
        self.msp = doc.modelspace()
        self.spec = spec
        self.issues = issues
        self.rep = rep
        s = spec
        for layer in {s.cable_layer, s.terminal_layer}:
            if layer not in doc.layers:
                lt = s.layer_linetype if s.layer_linetype in doc.linetypes else "Continuous"
                doc.layers.add(layer, color=s.layer_color, linetype=lt)
                rep.log(f"Created layer '{layer}' (ACI {s.layer_color}, {lt})")
        if s.text_style not in doc.styles:
            doc.styles.add(s.text_style, font=s.text_style_font)
            rep.log(f"Created text style '{s.text_style}' ({s.text_style_font})")
        else:
            font = (doc.styles.get(s.text_style).dxf.get("font", "") or "").lower()
            if font and not font.startswith(s.text_style_font.lower()):
                issues.add("WARNING", "Style", s.text_style,
                           f"Text style '{s.text_style}' uses font '{font}' in this drawing (reference: {s.text_style_font})")
        if doc.mleader_styles.get(s.mleader_style) is None:
            doc.mleader_styles.new(s.mleader_style)
            rep.log(f"Created MLEADER style '{s.mleader_style}'")

    def add(self, plan: CalloutPlan) -> str:
        s = self.spec
        dog = 1 if plan.dog >= 0 else -1
        loc = Vec2(plan.last.x + dog * (s.dogleg_length + s.landing_gap), plan.last.y + s.text_y_factor * s.text_height)
        b = self.msp.add_multileader_mtext(s.mleader_style, dxfattribs={"layer": plan.layer})
        b.set_content(plan.text, char_height=s.text_height,
                      alignment=mlr.TextAlignment.left if dog > 0 else mlr.TextAlignment.right, style=s.text_style)
        b.set_connection_properties(landing_gap=s.landing_gap, dogleg_length=s.dogleg_length)
        b.set_connection_types(left=mlr.HorizontalConnection(s.side_attachment),
                               right=mlr.HorizontalConnection(s.side_attachment),
                               top=mlr.VerticalConnection(s.vertical_attachment),
                               bottom=mlr.VerticalConnection(s.vertical_attachment))
        b.set_arrow_properties(name="", size=s.arrow_size)
        b.set_leader_properties(color=ezcolors.BYLAYER, linetype="BYBLOCK", lineweight=-2,
                                leader_type=mlr.LeaderType.straight_lines)
        b.add_leader_line(mlr.ConnectionSide.left if dog > 0 else mlr.ConnectionSide.right, [Vec2(plan.tip.x, plan.tip.y)])
        b.build(insert=loc)
        ml = b.multileader
        d = ml.dxf
        d.color = ezcolors.BYLAYER
        d.linetype = "BYLAYER"
        d.lineweight = -1
        d.leader_type = 1
        d.leader_line_color = ezcolors.BY_LAYER_RAW_VALUE
        d.leader_lineweight = -2
        d.has_landing = 1
        d.has_dogleg = 1
        d.dogleg_length = s.dogleg_length
        d.arrow_head_size = s.arrow_size
        d.content_type = 2
        d.text_left_attachment_type = s.side_attachment
        d.text_right_attachment_type = s.side_attachment
        d.text_top_attachment_type = s.vertical_attachment
        d.text_bottom_attachment_type = s.vertical_attachment
        d.text_attachment_direction = 0
        d.text_angle_type = 1
        d.text_alignment_type = 2
        d.text_attachment_point = 1 if dog > 0 else 3
        d.text_color = ezcolors.BY_LAYER_RAW_VALUE
        d.has_text_frame = 0
        d.is_annotative = 0
        d.scale = 1.0
        ctx = ml.context
        ctx.scale = 1.0
        ctx.char_height = s.text_height
        ctx.arrow_head_size = s.arrow_size
        ctx.landing_gap_size = s.landing_gap
        ctx.left_attachment = s.side_attachment
        ctx.right_attachment = s.side_attachment
        ctx.top_attachment = s.vertical_attachment
        ctx.bottom_attachment = s.vertical_attachment
        ctx.text_align_type = 0 if dog > 0 else 2
        ctx.attachment_type = 0
        est_w = max(len(x) for x in plan.text.split("\\P")) * s.text_height * 0.9
        ctx.base_point = Vec3(loc.x if dog > 0 else loc.x - est_w, plan.last.y, 0)
        mt = ctx.mtext
        mt.default_content = plan.text
        mt.insert = Vec3(loc.x, loc.y, 0)
        mt.text_direction = Vec3(1, 0, 0)
        mt.rotation = 0.0
        mt.width = 0.0
        mt.defined_height = 0.0
        mt.line_spacing_factor = 1.0
        mt.line_spacing_style = 1
        mt.color = ezcolors.BY_BLOCK_RAW_VALUE
        mt.alignment = 1 if dog > 0 else 3
        mt.flow_direction = 1
        mt.has_bg_fill = 0
        mt.use_window_bg_color = 0
        mt.bg_scale_factor = 1.5
        mt.column_type = 0
        if ctx.leaders:
            ld = ctx.leaders[0]
            ld.has_last_leader_line = 1
            ld.has_dogleg_vector = 1
            ld.last_leader_point = Vec3(plan.last.x, plan.last.y, 0)
            ld.dogleg_vector = Vec3(dog, 0, 0)
            ld.dogleg_length = s.dogleg_length
            ld.attachment_direction = 0
            if ld.lines:
                ld.lines[0].vertices = [Vec3(plan.tip.x, plan.tip.y, 0)]
                ld.lines[0].color = ezcolors.BY_LAYER_RAW_VALUE
        try:
            ml.update_proxy_graphic()
        except Exception:  # noqa: BLE001 - proxy graphic is optional display data
            pass
        return ml.dxf.handle


# =============================================================================
# QA workbook
# =============================================================================
def _write_sheet(ws: Any, rows: List[Dict[str, Any]]) -> None:
    if not rows:
        rows = [{"Info": "No rows"}]
    headers: List[str] = []
    for r in rows:
        for k in r:
            if k not in headers:
                headers.append(k)
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    fill = PatternFill("solid", fgColor="0B3954")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        width = max((len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 300) + 1)), default=8)
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, width + 2), 80)


def write_qa_workbook(path: Path, cfg: RunConfig, result: "RunResult") -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    callouts = [dict(ID=p.uid, Kind="Fiber number (terminal)" if p.kind == "TERMINAL" else "BFO (cable segment)",
                     OLM=p.olm, Text=p.text.replace("\\P", " / "), Layer=p.layer, Closure=p.closure, Anchor=p.anchor,
                     TipX=round(p.tip.x, 4), TipY=round(p.tip.y, 4), LandingX=round(p.last.x, 4),
                     LandingY=round(p.last.y, 4), Dogleg="+X" if p.dog > 0 else "-X",
                     Numbers=fmt_numbers(p.numbers, "exact", 0), Calculation=p.calculation, Sources=p.sources,
                     Status=p.status, Handle=p.handle, ExistingCheck=p.existing_match,
                     ExistingText=(p.existing_text or "").replace("\n", " / "), Notes=p.notes) for p in result.plans]
    segs = [dict(Closure=s.closure, Cable=s.cable, Segment=s.seg, FiberCount=s.count, EndClosure=s.end, OLM=s.olm,
                 LiveFibers=len(s.live), Numbers=fmt_numbers(s.live, "exact", 0), CalloutText=s.text.replace("\\P", " / "),
                 Geometry=s.geometry, Unresolved=len(s.unresolved),
                 UnresolvedExample=s.unresolved[0] if s.unresolved else "", Report=s.report_file) for s in result.segments]
    issues = [dict(Severity=i.severity, Category=i.category, Object=i.obj, Message=i.message) for i in result.issues.items]
    settings = []
    for k, v in cfg.to_dict().items():
        if isinstance(v, dict):
            for k2, v2 in v.items():
                settings.append(dict(Section=k, Key=k2, Value=json.dumps(v2) if isinstance(v2, (list, dict)) else v2))
        else:
            settings.append(dict(Section="run", Key=k, Value=json.dumps(v) if isinstance(v, (list, dict)) else v))
    for name, rows in (("Summary", [dict(Item=k, Value=v) for k, v in result.summary.items()]),
                       ("Network", result.network), ("Callouts", callouts), ("Terminals", result.terminals),
                       ("Cable Segments", segs), ("Validation Issues", issues),
                       ("Existing vs Computed", result.existing_rows), ("Settings", settings)):
        _write_sheet(wb.create_sheet(name), rows)
    ws = wb["Validation Issues"]
    fills = {"ERROR": PatternFill("solid", fgColor="F8CBAD"), "WARNING": PatternFill("solid", fgColor="FFE699")}
    for r in range(2, ws.max_row + 1):
        f = fills.get(str(ws.cell(r, 1).value))
        if f:
            for c in ws[r]:
                c.fill = f
    wb.save(path)


# =============================================================================
# Engine
# =============================================================================
@dataclass
class ScanResult:
    headers: List[str]
    columns: ColumnMap
    cad: Optional[CadMap]
    layers: List[str]
    drawing_info: str
    ports_hint: str
    issues: IssueLog


@dataclass
class RunResult:
    plans: List[CalloutPlan]
    segments: List[SegmentPlanInfo]
    terminals: List[Dict[str, Any]]
    issues: IssueLog
    summary: Dict[str, Any]
    network: List[Dict[str, Any]]
    existing_rows: List[Dict[str, Any]]
    outputs: Dict[str, str]
    placement: Dict[str, int]
    spec: CalloutSpec = field(default_factory=CalloutSpec)
    signature: str = ""


def input_signature(cfg: RunConfig) -> str:
    """Identifies the inputs of a run (paths + sizes + modification times) so an analysis can be reused."""
    parts = []

    def stamp(f: Path) -> str:
        try:
            st = f.stat()
            return f"{f.resolve()}|{st.st_size}|{st.st_mtime_ns}"
        except OSError:
            return f"{f}|missing"

    for p in [cfg.drawing_path] + list(cfg.splice_paths) + list(cfg.address_paths):
        path = Path(p)
        parts.append(stamp(path))
        if path.is_dir():       # a file edited in place does not change the folder's own time stamp
            parts.extend(stamp(f) for f in sorted(path.rglob("*")) if f.is_file())
    parts.append(json.dumps(dataclasses.asdict(cfg.columns), sort_keys=True))
    parts.append(json.dumps(dataclasses.asdict(cfg.cad), sort_keys=True))
    parts.append(json.dumps(dataclasses.asdict(cfg.spec), sort_keys=True))
    parts.append(f"{cfg.ports_per_pon}|{cfg.units}")
    return "\n".join(parts)


class CalloutEngine:
    def __init__(self, cfg: RunConfig, rep: Reporter) -> None:
        self.cfg = cfg
        self.rep = rep
        self.collector = InputCollector()

    # ------------------------------------------------------------------ scan (CLI diagnostics)
    def scan(self) -> ScanResult:
        issues = IssueLog()
        try:
            addr_files = self.collector.expand(self.cfg.address_paths, ADDRESS_EXTS, issues)
            spl_files = self.collector.expand(self.cfg.splice_paths, REPORT_EXTS, issues)
            headers = AddressBook.scan_headers(addr_files, issues)
            cols = auto_map_columns(headers)
            reports = load_splice_reports(spl_files, issues, self.rep) if spl_files else []
            splitter = splitter_size(reports)
            ports_hint = f"{splitter} (splitter size found in splice reports)" if splitter else "auto (from address list)"
            cad = None
            layers: List[str] = []
            info = ""
            if self.cfg.drawing_path:
                src = open_source(self.cfg, self.rep)
                scale = units_factor(src, self.cfg.units, self.rep)
                book = AddressBook()
                if addr_files:
                    book.load(addr_files, cols, self.cfg.ports_per_pon, splitter, pigtail_splices(reports), IssueLog(),
                              Reporter(lambda m, l: None))
                names = {r.splice for r in book.records if r.splice} | {r.name for r in reports}
                cad = auto_detect_cad(src, names, [r.xy for r in reports if r.xy], self.rep, scale)
                layers = src.layers
                info = f"{Path(self.cfg.drawing_path).name}: {src.kind}, {src.version}, {len(src.layers)} layers"
            return ScanResult(headers, cols, cad, layers, info, ports_hint, issues)
        finally:
            self.collector.cleanup()

    # ------------------------------------------------------------------ analysis
    def _prepare(self, issues: IssueLog) -> Tuple[NetworkSource, AddressBook, List[SpliceReport], DrawingModel,
                                                  CalloutPlanner, CalloutSpec, float]:
        cfg = self.cfg
        self.rep.progress(2, "Collecting input files")
        combined = bool(cfg.splice_paths) and sorted(cfg.splice_paths) == sorted(cfg.address_paths)
        addr_files = self.collector.expand(cfg.address_paths, ADDRESS_EXTS, issues)
        spl_files = self.collector.expand(cfg.splice_paths, REPORT_EXTS, IssueLog() if combined else issues)
        reports = load_splice_reports(spl_files, issues, self.rep, quiet_cross=combined)
        if not addr_files:
            issues.add("ERROR", "Input", "Address List", "No address list files found")
        if not reports:
            issues.add("WARNING", "Input", "Splice Report", "No splice reports found - BFO (cable) callouts cannot be produced")
        self.rep.progress(18, "Reading address lists")
        book = AddressBook()
        book.load(addr_files, cfg.columns, cfg.ports_per_pon, splitter_size(reports), pigtail_splices(reports), issues,
                  self.rep, quiet_cross=combined)
        self.rep.log(f"Inputs: {len(reports)} splice report(s), {len({r.source.split(':')[0] for r in book.records})} "
                     f"address list file(s)")
        if not book.records:
            self._log_issues(issues)
            raise RuntimeError("No usable address-list rows were found. Check that the address lists contain PON, "
                               "PORT, address, terminal and splice columns (details in the log).")
        if book.ports_source.startswith("address list") and reports:
            issues.add("WARNING", "Numbering", "Ports per PON",
                       f"Splitter size not found in the splice reports - {book.ports_per_pon} assumed from the highest "
                       f"PON PORT ({book.max_port}). Check the fiber numbers in the QA workbook.")
        self.rep.progress(22, "Reading drawing")
        src = open_source(cfg, self.rep)
        scale = units_factor(src, cfg.units, self.rep)
        spec = cfg.spec.scaled(scale)
        names = {r.splice for r in book.records if r.splice} | {r.name for r in reports}
        report_xy = {r.name: r.xy for r in reports if r.xy}
        self.rep.progress(40, "Analysing the network")
        auto = auto_detect_cad(src, names, list(report_xy.values()), self.rep, scale)
        cad = merge_cad(cfg.cad, auto)
        model = DrawingModel(src, cad, spec, issues, self.rep, names, report_xy, scale)
        self._cross_validate(book, reports, model, issues)
        planner = CalloutPlanner(model, book, reports, spec, issues, self.rep)
        self.rep.progress(52, "Planning fiber-number callouts")
        planner.plan_terminals()
        planner.plan_cables()
        return src, book, reports, model, planner, spec, scale

    def _cross_validate(self, book: AddressBook, reports: List[SpliceReport], model: DrawingModel, issues: IssueLog) -> None:
        for c in sorted({r.splice for r in book.records if r.splice}):
            if c not in model.closures:
                n = sum(1 for r in book.records if r.splice == c)
                close = difflib.get_close_matches(c, list(model.closures), n=1, cutoff=0.85)
                hint = f" (did you mean {close[0]}?)" if close else ""
                issues.add("ERROR", "Missing Closure", c, f"Referenced by {n} address row(s) but not found in the drawing{hint}")
        report_closures = {r.name for r in reports}
        used: Set[Tuple[str, int]] = set()
        for r in reports:
            for c in r.conns:
                if not c.address or any(s in norm_text(c.state) for s in DEAD_STATES) or NON_ADDRESS_RE.search(c.usage or c.address):
                    continue
                for rec in book.resolve_candidates(c.tap, c.address, None, c.port_idx):
                    used.add((rec.olm, rec.number))
        for rec in book.records:
            if rec.splice in report_closures and (rec.olm, rec.number) not in used:
                sev = "INFO" if "POWER" in rec.address else "WARNING"
                issues.add(sev, "Orphan Address", f"{rec.olm} #{rec.number}",
                           f"{rec.address} ({rec.tap}) is not on any live fiber in the splice reports")

    def _network_rows(self, book: AddressBook, reports: List[SpliceReport], planner: CalloutPlanner) -> List[Dict[str, Any]]:
        rows = []
        for olm in book.olms:
            recs = [r for r in book.records if r.olm == olm]
            plans = [p for p in planner.plans if p.olm == olm]
            rows.append({
                "OLM": olm, "PONs used": len({r.pon for r in recs}), "Fiber numbers": len(recs),
                "Highest number": max(r.number for r in recs),
                "Terminals": len({(r.tap, r.splice) for r in recs}), "Drop closures": len({r.splice for r in recs}),
                "Splice reports": sum(1 for r in reports if r.olm == olm),
                "Cable segments evaluated": sum(1 for s in planner.segments if s.olm == olm),
                "Fiber-number callouts": sum(1 for p in plans if p.kind == "TERMINAL" and p.status != "FAILED"),
                "BFO callouts": sum(1 for p in plans if p.kind == "CABLE" and p.status != "FAILED"),
                "Failed": sum(1 for p in plans if p.status == "FAILED")})
        return rows

    def _log_issues(self, issues: IssueLog) -> None:
        by = Counter((i.severity, i.category) for i in issues.items)
        for (sev, cat), n in sorted(by.items()):
            if sev != "INFO":
                self.rep.log(f"{cat}: {n} {sev.lower()}(s)", sev)
        shown = 0
        for i in issues.items:
            if i.severity in ("ERROR", "WARNING") and shown < 25:
                self.rep.log(f"  [{i.category}] {i.obj} - {i.message}", i.severity)
                shown += 1
        if sum(1 for i in issues.items if i.severity != "INFO") > shown:
            self.rep.log("  ... the complete list is in the Results page and the QA workbook")
        self.rep.log("Validation summary: " + issues.summary())

    def analyze(self) -> RunResult:
        """Read and cross-check all inputs and plan every callout (nothing is written)."""
        issues = IssueLog()
        try:
            src, book, reports, model, planner, spec, _scale = self._prepare(issues)
            placement = planner.placement_qa()
            self.rep.progress(82, "Comparing with existing callouts")
            existing_rows = compare_existing(planner.plans, model.existing, 40.0 * spec.text_height) if model.existing else []
            if model.existing:
                ok = sum(1 for p in planner.plans if p.existing_match.startswith("MATCH"))
                self.rep.log(f"Existing fiber callouts in the drawing: {len(model.existing)}; "
                             f"{ok} identical to a computed callout")
            plans = planner.plans
            n_t = sum(1 for p in plans if p.kind == "TERMINAL" and p.status != "FAILED")
            n_c = sum(1 for p in plans if p.kind == "CABLE" and p.status != "FAILED")
            summary = {
                "Tool": f"{APP_NAME} {APP_VERSION}",
                "Run at": _dt.datetime.now().isoformat(timespec="seconds"),
                "Drawing": self.cfg.drawing_path,
                "Drawing reader": {"dxf": "DXF", "dwg-oda": "DWG (ODA File Converter)",
                                   "dwg-native": "DWG (built-in reader)"}[src.kind],
                "Drawing version": src.version,
                "OLMs": ", ".join(book.olms),
                "Ports per PON": f"{book.ports_per_pon} ({book.ports_source})",
                "Formula": f"N = (PON - 1) x {book.ports_per_pon} + PON PORT (per OLM)",
                "Address rows (PON assigned)": len(book.records),
                "Terminal groups": len(book.terminals),
                "Splice reports": len(reports),
                "Closures in drawing": len(model.closures),
                "Handholes in drawing": len(model.handholes),
                "Cable segments evaluated": len(planner.segments),
                "Fiber-number callouts": n_t,
                "BFO callouts": n_c,
                "Callouts failed": sum(1 for p in plans if p.status == "FAILED"),
                "Existing callouts found": len(model.existing),
                "Existing callouts matched": sum(1 for p in plans if p.existing_match.startswith("MATCH")),
                "Placement - clean": placement.get("clean", 0),
                "Placement - needs review": n_t + n_c - placement.get("clean", 0),
                "Errors": issues.count("ERROR"),
                "Warnings": issues.count("WARNING"),
            }
            self.rep.log(f"Analysis: {n_t} fiber-number + {n_c} BFO callouts; {placement.get('clean', 0)} with clean placement")
            self._log_issues(issues)
            self.rep.progress(100, "Analysis complete")
            return RunResult(plans, planner.segments, planner.terminal_rows, issues, summary,
                             self._network_rows(book, reports, planner), existing_rows, {}, placement, spec,
                             input_signature(self.cfg))
        finally:
            self.collector.cleanup()

    validate = analyze

    def export(self, result: RunResult) -> RunResult:
        """Write the callouts (DXF at the drawing's coordinates) and the QA/QC workbook for an analysis."""
        cfg = self.cfg
        out_dir = Path(cfg.output_dir or Path(cfg.drawing_path).parent / "callout_output")
        try:
            out_dir.mkdir(parents=True, exist_ok=True)
            probe = out_dir / f".fcs_write_test_{os.getpid()}"
            probe.write_bytes(b"")
            probe.unlink()
        except OSError as exc:
            raise RuntimeError(f"Cannot write to the output folder {out_dir} ({exc.strerror or exc}). "
                               "Choose another output folder.") from exc
        stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        stem = Path(cfg.drawing_path).stem
        self.rep.progress(86, "Writing callouts")
        doc = ezdxf.new("R2018", setup=True)
        doc.header["$INSUNITS"] = 0
        doc.header["$MEASUREMENT"] = 0
        writer = CalloutWriter(doc, result.spec, result.issues, self.rep)
        plans = result.plans
        for i, p in enumerate(plans):
            self.rep.check_cancel()
            if p.status == "FAILED":
                continue
            try:
                p.handle = writer.add(p)
                p.status = "PLACED"
            except Exception as exc:  # noqa: BLE001 - one bad callout must not stop the run
                p.status = "FAILED"
                p.notes = f"Write error: {exc}"
                result.issues.add("ERROR", "CAD Write", p.uid, str(exc))
            if i % 200 == 0:
                self.rep.progress(86 + 8 * i / max(1, len(plans)), f"Writing callouts ({i}/{len(plans)})")
        self.rep.progress(95, "Saving outputs")
        dxf_path = out_dir / f"{stem}_FIBER_CALLOUTS_{stamp}.dxf"
        doc.saveas(str(dxf_path))
        qa_path = out_dir / f"{stem}_FIBER_CALLOUTS_QA_{stamp}.xlsx"
        result.outputs = {"Callouts DXF": str(dxf_path), "QA workbook": str(qa_path)}
        result.summary["Callouts written"] = sum(1 for p in plans if p.status == "PLACED")
        result.summary["Errors"] = result.issues.count("ERROR")
        result.summary["Output callouts DXF"] = str(dxf_path)
        result.summary["Output QA workbook"] = str(qa_path)
        write_qa_workbook(qa_path, cfg, result)
        self.rep.log(f"Saved callouts: {dxf_path}")
        self.rep.log(f"Saved QA/QC workbook: {qa_path}")
        self.rep.log(f"DONE - {result.summary['Fiber-number callouts']} fiber-number + {result.summary['BFO callouts']} "
                     f"BFO callouts. In AutoCAD: open the callouts DXF, Ctrl+A, Ctrl+C, switch to the design, PASTEORIG.",
                     "SUCCESS")
        self.rep.progress(100, "Done")
        return result

    def generate(self, previous: Optional[RunResult] = None) -> RunResult:
        """Analyse (or reuse an unchanged analysis) and export."""
        if previous is not None and previous.signature and previous.signature == input_signature(self.cfg):
            self.rep.log("Inputs unchanged since the last analysis - exporting it directly")
            result = previous
        else:
            result = self.analyze()
        return self.export(result)


def resource_path(name: str) -> str:
    """Path of a bundled asset (works from source and from a PyInstaller build)."""
    base = getattr(sys, "_MEIPASS", None) or str(Path(__file__).resolve().parent.parent)
    return str(Path(base) / name)
