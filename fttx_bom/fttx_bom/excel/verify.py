"""Post-generation verification of the BOM workbook.

1. Reopen the saved file and compare structure with the template (sheets,
   part numbers, descriptions, G/F-column formulas).
2. Confirm every written quantity is in the expected cell.
3. Recalculate with LibreOffice (headless) when available and check for Excel
   errors and that Total Quantity (G) equals the sum of the area columns.
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
import warnings
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .writer import WriteResult

ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


@dataclass
class Verification:
    checks: list[tuple[str, str, str]] = field(default_factory=list)  # (name, PASS/FAIL/WARN, detail)
    recalculated: bool = False
    recalc_values: dict[str, float] = field(default_factory=dict)

    def add(self, name, ok, detail, warn=False):
        self.checks.append((name, "PASS" if ok else ("WARN" if warn else "FAIL"), detail))

    @property
    def failed(self):
        return [c for c in self.checks if c[1] == "FAIL"]


def _load(p, data_only=False):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(p, data_only=data_only, read_only=True)


def grid(ws, max_col=20) -> dict[str, object]:
    """Cell values of the first `max_col` columns keyed by coordinate (fast in read-only mode)."""
    out = {}
    with warnings.catch_warnings():  # read-only sheets parse lazily, so warnings surface here
        warnings.simplefilter("ignore")
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max_col):
            for c in row:
                if getattr(c, "coordinate", None) and c.value is not None:
                    out[c.coordinate] = c.value
    return out


class _Cell:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class GridSheet:
    """Minimal worksheet facade over a grid() dict (cell(r, c).value, ws['A1'].value, max_row)."""

    def __init__(self, cells: dict[str, object]):
        from openpyxl.utils.cell import coordinate_from_string

        self.cells = cells
        self.max_row = max((coordinate_from_string(k)[1] for k in cells), default=0)

    def cell(self, row, column):
        from openpyxl.utils import get_column_letter

        return _Cell(self.cells.get(f"{get_column_letter(column)}{row}"))

    def __getitem__(self, coord):
        return _Cell(self.cells.get(coord))


def recalc_with_libreoffice(src: Path, workdir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    workdir.mkdir(parents=True, exist_ok=True)
    profile = workdir / "lo_profile"
    cmd = [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--calc", "--convert-to", "xlsx",
           "--outdir", str(workdir), str(src)]
    try:
        subprocess.run(cmd, capture_output=True, text=True, timeout=900)
    except subprocess.TimeoutExpired:
        return None
    out = workdir / src.name
    return out if out.exists() else None


def verify(res: WriteResult, expected: dict[tuple[str, str], float | None], workdir: Path,
           recalc: bool = True) -> Verification:
    v = Verification()
    out = res.path
    v.add("output file exists", out.exists() and out.stat().st_size > 0, str(out))
    got = _load(out)
    v.add("worksheets preserved", res.template_sheets == got.sheetnames, f"template {res.template_sheets} vs output {got.sheetnames}")
    gs = grid(got[res.sheet])
    struct_diff = []
    for coord, a in res.template_snapshot.items():
        if gs.get(coord) != a:
            struct_diff.append(f"{coord}: {a!r} -> {gs.get(coord)!r}")
    for coord, b in gs.items():
        if coord[0] in "BCDEFG" and coord not in res.template_snapshot and int(coord[1:]) > 5:
            struct_diff.append(f"{coord}: None -> {b!r}")
    v.add("part numbers / descriptions / F-G formulas unchanged", not struct_diff, "; ".join(struct_diff[:10]) or "identical")
    misplaced = []
    for (key, L), val in res.written.items():
        r = res.rows[key]
        cell = gs.get(f"{L}{r}")
        if cell != val:
            misplaced.append(f"{key} {L}{r}: expected {val!r} found {cell!r}")
    v.add("quantities in expected cells", not misplaced, "; ".join(misplaced[:10]) or f"{len(res.written)} cells checked")
    dup = {}
    for key, r in res.rows.items():
        dup.setdefault(r, []).append(key)
    dups = {r: k for r, k in dup.items() if len(k) > 1}
    v.add("no row mapped twice", not dups, str(dups) if dups else "ok")

    if not recalc:
        v.add("formula recalculation", False, "skipped by option", warn=True)
        return v
    rc = recalc_with_libreoffice(out, workdir / "recalc")
    if rc is None:
        v.add("formula recalculation", False, "LibreOffice not available - formulas NOT independently recalculated", warn=True)
        return v
    v.recalculated = True
    vals = _load(rc, data_only=True)
    errs = []
    for ws in vals.worksheets:
        if ws.title not in ("OLM SF-8M BOM", "Jumper Calculator", "Transport BOM"):
            continue
        for coord, val in grid(ws).items():
            if isinstance(val, str) and val.strip() in ERRORS:
                errs.append(f"{ws.title}!{coord}={val}")
    v.add("no Excel error values after recalculation", not errs, "; ".join(errs[:10]) or "none")
    ws = grid(vals[res.sheet])
    bad = []
    for key, r in res.rows.items():
        g = ws.get(f"G{r}")
        parts = [ws.get(f"{L}{r}") for L in res.columns.values()]
        s = sum(p for p in parts if isinstance(p, (int, float)))
        formula = str(gs.get(f"G{r}") or "")
        if formula.upper().startswith("=SUM(H") and isinstance(g, (int, float)) and abs(g - s) > 1e-6:
            bad.append(f"G{r} {g} != sum {s}")
        v.recalc_values[key] = g
    v.add("Total Quantity column = sum of area columns", not bad, "; ".join(bad[:10]) or "ok")
    exp_bad = []
    for k, tot in expected.items():
        g = v.recalc_values.get(k[0])
        if tot is not None and isinstance(g, (int, float)) and abs(g - tot) > 0.51:
            exp_bad.append(f"{k[0]}: workbook {g} vs engine {tot}")
    v.add("recalculated totals = engine totals", not exp_bad, "; ".join(exp_bad[:10]) or "ok")
    return v
