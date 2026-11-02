"""BOM_AUDIT.xlsx and BOM_DISCREPANCIES.xlsx generation + reference BOM reader."""
from __future__ import annotations

import json
import re
import warnings
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engine.network import GRANT, NONGRANT
from ..engine.rules import Rule
from ..engine.takeoff import Takeoff
from .writer import WriteResult, find_rows, norm

HDR = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="305496")
FILLS = {
    "MATCH": PatternFill("solid", fgColor="C6EFCE"),
    "TOTAL_MATCH": PatternFill("solid", fgColor="E2EFDA"),
    "DIFF": PatternFill("solid", fgColor="FFC7CE"),
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "WARN": PatternFill("solid", fgColor="FFEB9C"),
}
CONFIG_DIR = Path(__file__).resolve().parents[2] / "config"


@dataclass
class Reference:
    path: str
    columns: dict[tuple[str, str], str] = field(default_factory=dict)
    values: dict[str, dict[tuple[str, str], float | None]] = field(default_factory=dict)  # rule key -> col -> value
    formulas: dict[str, dict[tuple[str, str], str]] = field(default_factory=dict)


def read_reference(path: Path, sheet: str, rules: list[Rule]) -> Reference:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wv = openpyxl.load_workbook(path, data_only=True, read_only=True)
        wf = openpyxl.load_workbook(path, data_only=False, read_only=True)
    from .verify import GridSheet, grid

    from ..inputs.detect import bom_sheet_name

    name = bom_sheet_name(wv.sheetnames, sheet)  # manual BOMs rename the sheet per OLM ("OLM DMI-R15 BOM")
    if name is None:
        raise ValueError(f"approved BOM {path.name} has no '{sheet}' or 'OLM ... BOM' sheet")
    ws = GridSheet(grid(wv[name], max_col=24))
    fgrid = grid(wf[name], max_col=24)
    ref = Reference(str(path))
    header_row = next(r for r in range(1, 30) if str(ws.cell(r, 2).value or "").strip().lower() == "part number")
    tq = next((c for c in range(1, 24) if str(ws.cell(header_row, c).value or "").strip().lower() == "total quantity"), 7)
    for c in range(tq + 1, tq + 17):
        h = str(ws.cell(header_row, c).value or "").strip()
        m = re.match(r"^(\S+)\s+(Grant|Non\s*-?\s*Grant)$", h, re.I)
        if m:
            area = GRANT if m.group(2).lower() == "grant" else NONGRANT
            ref.columns[(m.group(1), area)] = get_column_letter(c)
    rows, _ = find_rows(ws, rules, header_row)
    for rule in rules:
        r = rows.get(rule.key)
        if r is None:
            continue
        ref.values[rule.key] = {}
        ref.formulas[rule.key] = {}
        for col, L in ref.columns.items():
            v = ws[f"{L}{r}"].value
            ref.values[rule.key][col] = v if isinstance(v, (int, float)) else None
            f = fgrid.get(f"{L}{r}")
            if isinstance(f, str) and f.startswith("="):
                ref.formulas[rule.key][col] = f
    return ref


def _sheet(wb, title, headers, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, _h in enumerate(headers, 1):
        c = ws.cell(1, i)
        c.font, c.fill = HDR, HFILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 14)
    ws.freeze_panes = "A2"
    return ws


def _fmt(v):
    if v is None:
        return None
    return round(v, 4) if isinstance(v, float) and not float(v).is_integer() else (int(v) if isinstance(v, float) else v)


def load_investigations() -> dict:
    p = CONFIG_DIR / "investigations.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def compare(rules: list[Rule], res: WriteResult, ref: Reference, tk: Takeoff) -> list[dict]:
    inv = load_investigations()
    out = []
    for rule in rules:
        if rule.status == "TEMPLATE_FORMULA" or rule.key not in ref.values:
            continue
        auto_vals = res.rule_values.get(rule.key, {})
        tol = 0.01 if rule.unit == "MI" else 0.5  # same tolerance as policy calibration (0.5 mi would hide 2,640 ft)
        for net in tk.networks:
            ref_tot = auto_tot = 0.0
            ref_any = auto_any = False
            for area in (GRANT, NONGRANT):
                col = (net, area)
                rv = ref.values[rule.key].get(col)
                av = auto_vals.get(col, (None, ""))[0] if rule.write else auto_vals.get(col, (None, ""))[0]
                if rv is not None:
                    ref_tot += rv
                    ref_any = True
                if av is not None:
                    auto_tot += av
                    auto_any = True
                if rv in (None, 0) and av in (None, 0):
                    status = "MATCH"
                elif rv is not None and av is not None and abs(rv - av) < tol + 1e-9 * abs(rv):
                    status = "MATCH"
                else:
                    status = "DIFF"
                out.append(dict(rule=rule, net=net, area=area, ref=rv, auto=av, status=status,
                                formula=ref.formulas.get(rule.key, {}).get(col, ""),
                                note=inv.get(f"{rule.key}|{net}|{area}") or inv.get(f"{rule.key}|{net}") or inv.get(rule.key, "")))
            out.append(dict(rule=rule, net=net, area="TOTAL", ref=ref_tot if ref_any else None,
                            auto=auto_tot if auto_any else None,
                            status="MATCH" if (not ref_any and not auto_any) or abs(ref_tot - auto_tot) < tol + 1e-9 * abs(ref_tot) else "DIFF",
                            formula="", note=inv.get(f"{rule.key}|{net}") or inv.get(rule.key, "")))
    return out


def coverage(template_rows: list[tuple[int, str, str]], rules: list[Rule], res: WriteResult, comp: list[dict] | None):
    mapped = {res.rows[r.key]: r for r in rules if r.key in res.rows}
    c = Counter()
    for row, _p, _d in template_rows:
        rule = mapped.get(row)
        c[rule.status if rule else "UNMAPPED (no drawing symbol / not used in design)"] += 1
    disc = sum(1 for x in comp or [] if x["status"] == "DIFF" and x["area"] != "TOTAL")
    return c, disc


def template_material_rows(ws) -> list[tuple[int, str, str]]:
    out = []
    hr = next(r for r in range(1, 30) if str(ws.cell(r, 2).value or "").strip().lower() == "part number")
    for r in range(hr + 1, ws.max_row + 1):
        b, d = ws.cell(r, 2).value, ws.cell(r, 4).value
        if d or (b and str(b).strip().upper() in ("TOTAL ROUTE MILEAGE",)):
            out.append((r, str(b or "").strip(), str(d or "").strip()))
        elif b and isinstance(b, str) and b.strip().lower().endswith(("count", "footage", "mileage")):
            out.append((r, b.strip(), ""))
    return out


def write_audit(path: Path, tk: Takeoff, rules: list[Rule], res: WriteResult, verification, meta: dict,
                cls, reports, ref: Reference | None, comp: list[dict] | None, template_rows, issues=None) -> dict:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cov, ndisc = coverage(template_rows, rules, res, comp)
    fails = [c for c in tk.checks if c.status == "FAIL"] + [c for c in verification.checks if c[1] == "FAIL"]

    ws = _sheet(wb, "Summary", ["Item", "Value"], {1: 48, 2: 110})
    for k, v in meta.items():
        if not k.startswith("_"):
            ws.append([k, str(v)[:32000]])
    ws.append([])
    ws.append(["COVERAGE (template material lines)", ""])
    ws.append(["Total material lines in template", sum(cov.values())])
    for k, v in sorted(cov.items()):
        ws.append([f"  status {k}", v])
    ws.append(["Area-column discrepancies vs reference", ndisc if comp is not None else "no reference supplied"])
    ws.append(["Failed engine checks", sum(1 for c in tk.checks if c.status == "FAIL")])
    ws.append(["Failed workbook verification checks", sum(1 for c in verification.checks if c[1] == "FAIL")])
    ws.append(["Formulas recalculated (LibreOffice)", "yes" if verification.recalculated else "NO - not verified"])
    sev = Counter(i.severity for i in (issues or []))
    ws.append(["Deep check: ERROR / WARN / INFO", f"{sev.get('ERROR', 0)} / {sev.get('WARN', 0)} / {sev.get('INFO', 0)} (see 'Deep Check' sheet)"])
    ws.append(["Policies in use", ", ".join(f"{k}={v}" for k, v in tk.policy_used.items())])
    nlow = sum(1 for f in res.flags.values() if f.level == "LOW")
    ws.append(["Cells to check (highlighted in the BOM)", f"{nlow} low confidence (red), {len(res.flags) - nlow} review (amber) - see 'Review' sheet"])
    excl = tk.extra.get("scope", {}).get("excluded", []) + [("RF plant", h, k, w, p) for k, h, p, w in tk.extra.get("rf_structures", [])]
    ws.append(["Objects excluded from this design", f"{len(excl)} (other fiber network, existing / HFC plant) - see 'Excluded' sheet"])
    ws.append([])
    ws.append(["Status meanings", "VERIFIED = reproduces the NMK reference from independent sources; VERIFIED_TOTAL = network totals verified, "
               "Grant/Non-Grant split inferred; INFERRED = defensible rule, not proven against an approved BOM; UNRESOLVED / NOT_DETECTED = "
               "left blank in the BOM, engineer must supply; SUGGESTED = audit only. A passed software test does not prove an engineering quantity."])

    issues = issues or []
    ws = _sheet(wb, "Deep Check", ["#", "Severity", "Category", "Code", "Network", "Problem", "BOM lines affected", "How to fix", "X", "Y"],
                {3: 14, 4: 24, 6: 90, 7: 30, 8: 70})
    sev_fill = {"ERROR": FILLS["FAIL"], "WARN": FILLS["WARN"], "INFO": PatternFill("solid", fgColor="DDEBF7")}
    for k, i in enumerate(issues, 1):
        ws.append([k, i.severity, i.category, i.code, i.network, i.message, i.affects, i.fix,
                   round(i.x, 2) if i.x is not None else None, round(i.y, 2) if i.y is not None else None])
        ws.cell(ws.max_row, 2).fill = sev_fill.get(i.severity, FILLS["WARN"])
    ws.auto_filter.ref = f"A1:J{max(2, ws.max_row)}"

    ws = _sheet(wb, "Review", ["BOM line", "Network", "Area", "Value", "Confidence", "Why"], {1: 26, 6: 120})
    lvl_fill = {"LOW": PatternFill("solid", fgColor="FFC7CE"), "REVIEW": PatternFill("solid", fgColor="FFEB9C")}
    for (key, net, area), f in sorted(res.flags.items(), key=lambda kv: (kv[1].level != "LOW", kv[0])):
        ws.append([key, net, "Grant" if area == "G" else "Non-Grant", f.value, f.level, "; ".join(f.reasons)])
        ws.cell(ws.max_row, 5).fill = lvl_fill[f.level]
    ws.auto_filter.ref = f"A1:F{max(2, ws.max_row)}"

    ws = _sheet(wb, "Excluded", ["Category", "Object", "Quantity", "Reason", "X", "Y"], {1: 26, 2: 40, 4: 90})
    for cat, what, qty, why, pt in excl:
        ws.append([cat, what, qty, why, round(pt[0], 2) if pt else None, round(pt[1], 2) if pt else None])
    ws.auto_filter.ref = f"A1:F{max(2, ws.max_row)}"

    cols = tk.columns
    ws = _sheet(wb, "BOM Lines", ["Row", "Part Number", "Description", "Status", "Written to BOM", "Rule / basis", "Source"]
                + [f"{n} {a}" for n, a in cols] + ["Total", "Notes"], {2: 26, 3: 50, 6: 60})
    by_row = {res.rows[r.key]: r for r in rules if r.key in res.rows}
    for row, part, desc in template_rows:
        rule = by_row.get(row)
        if rule is None:
            ws.append([row, part, desc, "UNMAPPED", "no", "no rule - left as template (blank)", ""])
            continue
        vals = res.rule_values.get(rule.key, {})
        line = [row, part or rule.label, desc, rule.status, "yes" if rule.write and rule.status != "TEMPLATE_FORMULA" else "no", rule.basis, rule.source]
        nums = [vals.get(c, (None, ""))[0] for c in cols]
        notes = sorted({vals.get(c, (None, ""))[1] for c in cols} - {""})
        tot = sum(x for x in nums if x is not None) if any(x is not None for x in nums) else None
        ws.append(line + [_fmt(x) for x in nums] + [_fmt(tot), "; ".join(notes)])
        ws.cell(ws.max_row, 4).fill = FILLS["PASS"] if rule.status.startswith("VERIFIED") else FILLS["WARN"]

    from ..engine.policies import POLICIES

    ws = _sheet(wb, "Policies", ["Policy", "Decides", "Option", "In use", "Description", "Total (all columns)"]
                + [f"{n} {a}" for n, a in cols] + ["Approved-BOM cells matched"], {1: 16, 2: 30, 3: 16, 5: 70})
    calib = {(r["policy"], r["option"]): r for r in (tk.extra.get("calibration") or [])}
    for p in POLICIES:
        keys = sorted({k.split("@")[0] for k, _n, _a in tk.q if "@" in k and p.owns(k.split("@")[0])})
        for opt, desc in p.options.items():
            for key in keys:
                vals = [tk.get(f"{key}@{opt}", n, a) for n, a in cols]
                tot = sum(v for v in vals if v is not None) if any(v is not None for v in vals) else None
                c = calib.get((p.name, opt))
                ws.append([p.name, f"{p.title} ({key})", opt, "YES" if tk.policy_used.get(p.name) == opt else "",
                           desc, _fmt(tot)] + [_fmt(v) for v in vals] + [f"{c['matched']}/{c['cells']}" if c else ""])
                if tk.policy_used.get(p.name) == opt:
                    ws.cell(ws.max_row, 4).fill = FILLS["PASS"]

    ws = _sheet(wb, "Takeoff", ["Quantity key", "Network", "Area", "Value", "Unit", "Method", "Evidence count", "Evidence (handles / names)"], {6: 60, 8: 80})
    for (k, n, a), q in sorted(tk.q.items()):
        if "@" in k:
            continue  # policy alternatives are on the Policies sheet
        ws.append([k, n, a, _fmt(q.value), q.unit, q.method, len(q.evidence), ", ".join(map(str, q.evidence[:60]))])

    ws = _sheet(wb, "Cables", ["Label handle", "Segment", "Network", "Fibers", "Placement", "Route FTG", "Slack SLK", "TOTAL",
                               "Routed path ft", "Path/FTG", "Non-Grant fraction", "Grant ft", "Non-Grant ft", "Spans (len:area)", "Method"],
                {2: 44, 14: 50, 15: 40})
    for c in tk.cables:
        ng = c.total_ft * (c.ng_fraction or 0)
        ws.append([c.handle, c.name, c.network, c.fibers, c.placement, c.route_ft, c.slack_ft, c.total_ft,
                   round(c.path_ft, 1) if c.path_ft else None,
                   round(c.path_ft / c.route_ft, 3) if c.path_ft and c.route_ft else None,
                   round(c.ng_fraction, 4) if c.ng_fraction is not None else None, round(c.total_ft - ng, 1), round(ng, 1),
                   " ".join(f"{L:.0f}:{a}" for L, a in c.spans)[:300], c.split_method])

    ws = _sheet(wb, "Splice Enclosures", ["Splice", "Network", "Enclosure Type", "Fusion splices", "Pass-through", "Splitters", "File"], {1: 16, 3: 18, 7: 70})
    for r in reports:
        ws.append([r.name, r.network, r.enclosure, r.fusion_splices, r.pass_through, ", ".join(f"1x{x}" for x, _ in r.couplers), r.path])

    ws = _sheet(wb, "Checks", ["Source", "Check", "Status", "Detail"], {2: 40, 4: 120})
    for c in tk.checks:
        ws.append(["engine", c.name, c.status, c.detail])
        ws.cell(ws.max_row, 3).fill = FILLS.get(c.status, FILLS["WARN"])
    for name, st, detail in verification.checks:
        ws.append(["workbook", name, st, detail])
        ws.cell(ws.max_row, 3).fill = FILLS.get(st, FILLS["WARN"])

    ws = _sheet(wb, "Network Assignment", ["Kind", "Object", "Network", "Distance ft / position"], {2: 60})
    for w, n, d in cls.nearest_assigned:
        ws.append(["outside boundary -> nearest network", w, n, d])
    for u in tk.unassigned:
        ws.append(["UNASSIGNED (not counted)", u, "", ""])
    for note in cls.notes:
        ws.append(["classifier note", note, "", ""])

    ws = _sheet(wb, "CAD Census", ["Entity / block", "Count", "Handled"], {1: 40})
    handled = {"INSERT", "ATTRIB", "LWPOLYLINE", "TEXT", "MTEXT", "MULTILEADER", "LEADER", "ACAD_PROXY_ENTITY", "LINE", "CIRCLE"}
    for k, v in sorted(meta.get("_census", {}).items()):
        ws.append([k, v, "yes" if k in handled else "NO - ignored (reported)"])
    for name, cnt in sorted(meta.get("_block_census", {}).items()):
        if cnt.get("INSERT"):
            ws.append([f"nested INSERTs inside block '{name}'", cnt["INSERT"], "not expanded (symbol geometry only) - avoids double counting"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"coverage": dict(cov), "discrepancies": ndisc, "failed_checks": len(fails),
            "issues": dict(Counter(i.severity for i in (issues or [])))}


def write_discrepancies(path: Path, comp: list[dict], ref: Reference, tk: Takeoff):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = _sheet(wb, "By Area Column", ["Part / Line", "Status (rule)", "Network", "Area", "Reference qty", "Automated qty", "Abs diff", "% diff",
                                        "Result", "Reference cell formula", "Investigation / suspected cause"], {1: 28, 2: 16, 10: 18, 11: 110})
    for x in comp:
        if x["area"] == "TOTAL":
            continue
        rv, av = x["ref"], x["auto"]
        diff = (av or 0) - (rv or 0) if (rv is not None or av is not None) else None
        pct = round(100 * diff / rv, 1) if rv not in (None, 0) and diff is not None else None
        ws.append([x["rule"].key, x["rule"].status, x["net"], x["area"], _fmt(rv), _fmt(av), _fmt(abs(diff)) if diff is not None else None,
                   pct, x["status"], x["formula"], x["note"]])
        ws.cell(ws.max_row, 9).fill = FILLS[x["status"]]
    ws = _sheet(wb, "By Network Total", ["Part / Line", "Status (rule)", "Network", "Reference total", "Automated total", "Abs diff", "% diff",
                                          "Result", "Investigation / suspected cause"], {1: 28, 9: 110})
    for x in comp:
        if x["area"] != "TOTAL":
            continue
        rv, av = x["ref"], x["auto"]
        diff = (av or 0) - (rv or 0)
        ws.append([x["rule"].key, x["rule"].status, x["net"], _fmt(rv), _fmt(av), _fmt(abs(diff)),
                   round(100 * diff / rv, 1) if rv else None, x["status"], x["note"]])
        ws.cell(ws.max_row, 8).fill = FILLS[x["status"]]
    calib = tk.extra.get("calibration") or []
    if calib:
        ws = _sheet(wb, "Policy Calibration", ["Policy", "Option", "Cells matched", "Cells compared", "Abs error", "Best", "In use", "Description"],
                    {1: 16, 2: 16, 8: 80})
        for r in calib:
            ws.append([r["policy"], r["option"], r["matched"], r["cells"], r["abs_error"], "BEST" if r["best"] else "", "YES" if r["selected"] else "", r["description"]])
            if r["best"]:
                ws.cell(ws.max_row, 6).fill = FILLS["PASS"]
    ws = _sheet(wb, "Summary", ["Measure", "Value"], {1: 60, 2: 20})
    cells = [x for x in comp if x["area"] != "TOTAL"]
    tots = [x for x in comp if x["area"] == "TOTAL"]
    ws.append(["Reference workbook", ref.path])
    ws.append(["Area cells compared", len(cells)])
    ws.append(["Area cells matching", sum(1 for x in cells if x["status"] == "MATCH")])
    ws.append(["Area cells differing", sum(1 for x in cells if x["status"] == "DIFF")])
    ws.append(["Network totals compared", len(tots)])
    ws.append(["Network totals matching", sum(1 for x in tots if x["status"] == "MATCH")])
    ws.append(["Network totals differing", sum(1 for x in tots if x["status"] == "DIFF")])
    wb.move_sheet("Summary", offset=-2)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
