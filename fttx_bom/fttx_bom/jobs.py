"""Build jobs from auto-detected inputs and run many jobs (batch)."""
from __future__ import annotations

import traceback
from collections import Counter
from pathlib import Path
from typing import Callable

import openpyxl
from openpyxl.styles import Font, PatternFill

from .inputs.detect import Detected, detect, discover_jobs
from .pipeline import JobParams, JobResult, run_job


def params_from_detected(det: Detected, out: Path, template: Path | None = None, reference: Path | None = None,
                         **kw) -> tuple[JobParams, list[str]]:
    """Choose the job inputs from a detection result. Returns (params, notes). Raises ValueError when unusable."""
    notes = []
    if not det.drawings:
        raise ValueError("no drawing (.dwg / .dxf) found")
    if len(det.drawings) > 1:
        raise ValueError("more than one drawing found (" + ", ".join(p.name for p in det.drawings)
                         + ") - put each network's files in its own folder and use batch mode")
    if not det.address_lists:
        notes.append("no address list found - passings come from drawing house counts only")
    if not det.splice_reports:
        notes.append("no splice reports found - enclosures, trays and designed splitters cannot be counted")
    tpl = template or (det.bom_templates[0] if det.bom_templates else None) or (det.bom_completed[0] if det.bom_completed else None)
    if tpl is None:
        raise ValueError("no BOM template found - add the blank BOM template (or a completed BOM to use its layout)")
    if template is None and not det.bom_templates and det.bom_completed:
        notes.append(f"no blank template: using the layout of completed BOM {tpl.name}")
    ref = reference or (det.bom_completed[0] if det.bom_completed else None)
    if ref is not None:
        notes.append(f"approved BOM {ref.name} used as reference (discrepancies + policy calibration)")
    for p, why in det.unknown:
        notes.append(f"ignored {p.name}: {why}")
    params = JobParams(dwg=det.drawings[0], address_lists=det.address_lists, splice_reports=det.splice_reports,
                       template=tpl, out=out, reference=ref, **kw)
    return params, notes


def run_auto(paths: list[Path], out: Path, template: Path | None = None, reference: Path | None = None,
             log: Callable[[str], None] = print, **kw) -> tuple[JobResult, list[str]]:
    det = detect(paths, out / ".extracted", exclude=[out])
    log("detected: " + ", ".join(f"{k} {v}" for k, v in det.summary().items()))
    params, notes = params_from_detected(det, out, template, reference, **kw)
    for n in notes:
        log("note: " + n)
    return run_job(params, log), notes


def run_batch(root: Path, out: Path, template: Path | None = None, log: Callable[[str], None] = print, **kw) -> Path:
    """Every sub-folder of `root` with a drawing is a job. Writes BATCH_SUMMARY.xlsx."""
    jobs = discover_jobs(root)
    if not jobs:
        raise ValueError(f"no job folders with a drawing under {root}")
    rows = []
    for job in jobs:
        log(f"=== job {job.name}")
        jout = out / job.name
        try:
            res, notes = run_auto([job], jout, template=template, log=log, **kw)
            for g in res.groups:
                sev = Counter(i.severity for i in g.issues)  # per BOM, not job-wide
                rows.append([job.name, ", ".join(g.networks), "OK" if not g.workbook_failures else "WORKBOOK FAIL",
                             sev.get("ERROR", 0), sev.get("WARN", 0), sev.get("INFO", 0),
                             f"{g.ref_cells[0]}/{g.ref_cells[1]}" if g.ref_cells else "",
                             sum(res.totals.get("passings", {}).get(n, 0) for n in g.networks),
                             g.bom.name, g.audit.name, "; ".join(notes)])
        except Exception as exc:  # one bad job must not stop the batch
            log(f"job {job.name} FAILED: {exc}")
            rows.append([job.name, "", f"FAILED: {exc}", "", "", "", "", "", "", "", traceback.format_exc(limit=3)[-500:]])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Batch"
    ws.append(["Job", "Networks", "Status", "Errors", "Warnings", "Info", "Reference cells matched", "Passings", "BOM", "Audit", "Notes"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for r in rows:
        ws.append(r)
        st = ws.cell(ws.max_row, 3)
        st.fill = PatternFill("solid", fgColor="C6EFCE" if r[2] == "OK" else "FFC7CE")
    for col, w in zip("ABCDEFGHIJK", (24, 30, 18, 8, 9, 7, 12, 9, 40, 40, 80)):
        ws.column_dimensions[col].width = w
    out.mkdir(parents=True, exist_ok=True)
    path = out / "BATCH_SUMMARY.xlsx"
    wb.save(path)
    log(f"batch summary: {path}")
    return path
