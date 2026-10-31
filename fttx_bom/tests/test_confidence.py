"""Confidence rating of BOM cells and their highlighting in the workbook."""
import openpyxl

from fttx_bom.engine import policies as pol
from fttx_bom.engine.doctor import Issue
from fttx_bom.engine.network import GRANT, NONGRANT
from fttx_bom.engine.rules import Rule, load_rules
from fttx_bom.engine.takeoff import Takeoff
from fttx_bom.excel.confidence import LOW, REVIEW, assess

NET = "N-R01"


def _rule(key, value, status="VERIFIED"):
    return Rule(key, None, value, "EA", status, True, "", "")


def test_levels():
    tk = Takeoff([NET])
    for a in (GRANT, NONGRANT):
        tk.add("ped_max5", NET, a, 3, "EA", "")
        tk.add("pass_res", NET, a, 10, "EA", "")
        tk.add("route_ug_ft", NET, a, 100, "FT", "")
    rules = [_rule("CBP-MAX5DSM24", "ped_max5"), _rule("Residential Count", "pass_res"),
             _rule("UG Footage", "route_ug_ft", "INFERRED"), _rule("FE0-FOSC450D66NT0T3V", "fosc450d")]
    flags = assess(rules, tk, [], {}, [])
    assert ("CBP-MAX5DSM24", NET, GRANT) not in flags  # verified rule, complete inputs
    assert flags[("Residential Count", NET, GRANT)].level == LOW  # no address list
    assert flags[("UG Footage", NET, NONGRANT)].level == REVIEW
    f = flags[("FE0-FOSC450D66NT0T3V", NET, GRANT)]
    assert f.level == LOW and "splice reports" in f.reasons[0]  # plain-language reason


def test_issue_and_edge_effects():
    tk = Takeoff([NET])
    for a in (GRANT, NONGRANT):
        tk.add("ped_max5", NET, a, 3, "EA", "")
        tk.add("cable_024", NET, a, 500, "FT", "")
    rules = [_rule("CBP-MAX5DSM24", "ped_max5"), _rule("FFA-F902402G5B", "cable_024")]
    issues = [Issue("INFO", "Boundaries", "NEAR_GRANT_EDGE", NET, "CX1-MAX5 ABC is 3.0 ft from a Non-Grant edge (counted NG)", "", ""),
              Issue("ERROR", "Cable", "CABLE_TOTAL_MISMATCH", NET, "X TO Y FC024: FTG 10 + SLK 5 != TOTAL 20", "", "")]
    flags = assess(rules, tk, issues, {}, [])
    assert flags[("CBP-MAX5DSM24", NET, GRANT)].level == REVIEW
    assert flags[("FFA-F902402G5B", NET, NONGRANT)].level == LOW


def test_unsettled_convention_flagged():
    tk = Takeoff([NET])
    for a, v in ((GRANT, 10194), (NONGRANT, 25923)):
        tk.add("drop_ft", NET, a, v, "FT", "")
        tk.add("drop_ft_bom", NET, a, v if a == NONGRANT else 0, "FT", "")
    pol.compute_variants(tk)
    pol.apply(tk, {})
    flags = assess([_rule("Total Drop Footage", "drop_ft_bom")], tk, [], {}, [])
    f = flags[("Total Drop Footage", NET, GRANT)]
    assert f.level == REVIEW and "not settled" in f.reasons[0]


def test_workbook_highlight_and_review_sheet(tmp_path):
    from test_review_regressions import _mini_template

    from fttx_bom.excel.writer import write_bom

    cfg, _rules = load_rules()
    rule = Rule("CBP-MAX5DSM24", None, "ped_max5", "EA", "INFERRED", True, "", "")
    tk = Takeoff([NET])
    tk.add("ped_max5", NET, GRANT, 5, "EA", "")
    tk.add("ped_max5", NET, NONGRANT, 7, "EA", "")
    flags = assess([rule], tk, [], {}, [])
    res = write_bom(_mini_template(tmp_path / "t.xlsx", "G"), tmp_path / "o.xlsx", tk, [rule], cfg, {}, flags)
    wb = openpyxl.load_workbook(tmp_path / "o.xlsx")
    assert wb.sheetnames[-1] == "BOM Review" and res.added_sheets == ["BOM Review"]
    cell = wb["OLM SF-8M BOM"]["H6"]
    assert cell.value == 5 and cell.fill.fgColor.rgb.endswith("FFEB9C") and "REVIEW" in cell.comment.text
    review = list(wb["BOM Review"].iter_rows(min_row=5, values_only=True))
    assert {r[0] for r in review} == {"H6", "I6"} and all(r[6] == REVIEW for r in review)


def _cable(name, fibers, method="route spans", areas=()):
    from fttx_bom.engine.takeoff import CableSegment
    c = CableSegment("h", name, "A", "B", fibers, "UG", 100, 0, 100, NET, split_method=method)
    c.fallback_areas = areas
    return c


def test_cable_issue_flags_only_its_fiber_count():
    tk = Takeoff([NET])
    for a in (GRANT, NONGRANT):
        tk.add("cable_024", NET, a, 500, "FT", "")
        tk.add("cable_048", NET, a, 500, "FT", "")
    tk.cables = [_cable("N-01001 TO N-01002 FC048", 48)]
    rules = [_rule("FFA-F902402G5B", "cable_024"), _rule("FFA-F904802G5B", "cable_048")]
    # the issue names the cable but carries no network (splice names do not contain the OLM name)
    issues = [Issue("WARN", "Cable", "CABLE_ROUTE_LENGTH", "", "N-01001 TO N-01002 FC048: drawn route 900 ft vs label FTG 300 ft", "", ""),
              Issue("WARN", "Cable", "CABLE_LABEL_DUPLICATE", "", "N-01001 TO N-01002 FC024 labelled more than once (1A); identical values - counted once", "", "")]
    flags = assess(rules, tk, issues, {}, [])
    assert ("FFA-F902402G5B", NET, GRANT) not in flags  # duplicate with identical values: quantity not affected
    assert flags[("FFA-F904802G5B", NET, NONGRANT)].level == REVIEW


def test_untraced_cable_flagged_only_when_its_ends_disagree():
    tk = Takeoff([NET])
    for a in (GRANT, NONGRANT):
        tk.add("cable_024", NET, a, 500, "FT", "")
    rules = [_rule("FFA-F902402G5B", "cable_024")]
    tk.cables = [_cable("A TO B FC024", 24, "label point (route not found)", (GRANT, GRANT, GRANT))]
    assert not assess(rules, tk, [], {}, [])
    tk.cables = [_cable("A TO B FC024", 24, "label point (route not found)", (GRANT, GRANT, NONGRANT))]
    assert assess(rules, tk, [], {}, [])[("FFA-F902402G5B", NET, GRANT)].level == REVIEW
    tk.cables = [_cable("A TO B FC024", 24, "label point (route not found)", (GRANT,))]  # an end not found
    assert ("FFA-F902402G5B", NET, GRANT) in assess(rules, tk, [], {}, [])


def test_address_layer_conflict_flags_residential_and_only_changed_splitters():
    class AL:
        served = ["x"]
    tk = Takeoff([NET])
    tk.extra["ratio"][NET] = 32
    for a, p in ((GRANT, 64), (NONGRANT, 100)):
        tk.add("passings", NET, a, p, "EA", "")
        tk.add("pass_res", NET, a, p, "EA", "")
        tk.add("splitter_1x32", NET, a, 2 if a == GRANT else 2, "EA", "")
        tk.add("pon_ports", NET, a, 1 if a == GRANT else 2, "EA", "")
    rules = [_rule("Residential Count", "pass_res"), _rule("FE0-SPLIT32", "splitter_1x32"), _rule("SFP", "pon_ports")]
    conflict = Issue("WARN", "Addresses", "ADDRESS_LAYER_VS_BOUNDARY", NET,
                     "Served address 1 MAIN is on layer Address but outside every Non-Grant polygon", "", "")
    flags = assess(rules, tk, [conflict], {NET: AL()}, [])
    assert flags[("Residential Count", NET, GRANT)].level == REVIEW
    # Grant 64 -> 65 passings: ROUNDUP(65/32) = 3 splitters instead of 2, SFP ROUNDUP(3/2) = 2 instead of 1
    assert ("FE0-SPLIT32", NET, GRANT) in flags and ("SFP", NET, GRANT) in flags
    tk.q[("passings", NET, GRANT)].value = 40  # 41 still needs 2 splitters and 1 SFP
    flags = assess(rules, tk, [conflict], {NET: AL()}, [])
    assert ("FE0-SPLIT32", NET, GRANT) not in flags and ("SFP", NET, GRANT) not in flags


def test_zero_rows_ignore_unmapped_aerial_blocks_and_reasons_aggregate():
    tk = Takeoff([NET])
    for a in (GRANT, NONGRANT):
        tk.add("ped_max5", NET, a, 3, "EA", "")
    rules = [_rule("CBP-GFP103160002874", "0"), _rule("CBP-MAX5DSM24", "ped_max5")]
    aerial = Issue("WARN", "Drawing", "UNMAPPED_BLOCK", "", "4 x 'CX1-AELBOX' on RTE_AEPOLES is not mapped to any BOM line", "", "")
    assert not assess(rules, tk, [aerial], {}, [])
    ug = Issue("WARN", "Drawing", "UNMAPPED_BLOCK", "", "2 x 'CX1-GOAT' on RTE_UGPOLES is not mapped to any BOM line", "", "")
    edge = [Issue("INFO", "Boundaries", "NEAR_GRANT_EDGE", NET, f"CX1-MAX5 P{i} is 3.0 ft from a Non-Grant edge (counted NG)", "", "")
            for i in range(3)]
    flags = assess(rules, tk, [aerial, ug] + edge, {}, [])
    assert flags[("CBP-GFP103160002874", NET, GRANT)].level == REVIEW
    reasons = flags[("CBP-MAX5DSM24", NET, GRANT)].reasons
    assert len(reasons) == 1 and reasons[0].endswith("(+2 more like this)")
