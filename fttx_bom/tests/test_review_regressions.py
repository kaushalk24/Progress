"""Regression tests for the independent code review findings (see CHANGELOG 0.5.0).

Synthetic inputs only, so these run without the project files."""
import json
import zipfile
from pathlib import Path

import openpyxl
import pytest

from fttx_bom.engine.network import GRANT, NONGRANT, build_classifier
from fttx_bom.engine.rules import Rule, evaluate, load_rules, safe_eval
from fttx_bom.engine.takeoff import Takeoff, count_blocks, enclosure_letter
from fttx_bom.extract.model import Drawing, Insert, Polyline, Text
from fttx_bom.inputs.address_list import read_address_lists
from fttx_bom.inputs.detect import GENERATED_MARK, classify_xlsx, detect, safe_extract
from fttx_bom.inputs.splice_reports import _number, read_splice_reports

CONV = json.loads((Path(__file__).resolve().parents[1] / "config" / "conventions.json").read_text())
HDR = ["PON NAME", "PON PORT", "FIBER #", "ADDRESS", "COLOR TUBE", "COLOR FIBER", "CABLE NAME", "Fiber Tap Name", "Fiber Splice"]


def _address_book(path, header, extra_sheet=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "X-R01"
    ws.append(header)
    ws.append([1, 1, 1, "10 MAIN ST", "BL", "BL", "X-01001 TO X-01002 FC024", "T1", "S1"])
    ws.append([1, 2, 2, "12 MAIN ST", "BL", "OR", "X-01001 TO X-01002 FC024", "T1", "S1"])
    if extra_sheet:
        wb.create_sheet("Notes")["A1"] = "exported by tool X"
    wb.save(path)
    return path


def test_address_list_header_variant_and_extra_sheet(tmp_path):
    # MAJOR 4: a header differing only in case gave 0 served passings; any extra sheet became a network
    hdr = list(HDR)
    hdr[7] = "Fiber Tap name"
    skipped = []
    als = read_address_lists([_address_book(tmp_path / "a.xlsx", hdr, extra_sheet=True)], skipped)
    assert sorted(als) == ["X-R01"] and als["X-R01"].passings == 2
    assert any("Notes" in s for s in skipped)
    assert als["X-R01"].issues and als["X-R01"].issues[0].startswith("header differs")


def test_address_list_columns_matched_by_name(tmp_path):
    hdr = ["PON PORT", "PON NAME", "ADDRESS"] + HDR[2:3] + HDR[4:]  # reordered columns
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "X-R01"
    ws.append(hdr)
    ws.append([3, 1, "10 MAIN ST", 1])
    wb.save(tmp_path / "b.xlsx")
    al = read_address_lists([tmp_path / "b.xlsx"])["X-R01"]
    assert (al.served[0].pon, al.served[0].port, al.served[0].address) == (1, 3, "10 MAIN ST")


def test_enclosure_spelling_variants():
    # MAJOR 6: 'FOSC-450B' / '450D' were dropped silently
    for text, want in (("FOSC 450D", "D"), ("FOSC-450B", "B"), ("450D", "D"), ("FOSC450-D", "D"), ("Fosc 450-b", "B"),
                       ("FOSC450-B6-6-NT-0-B3V", "B"), ("SPLICE CASE", None), ("FOSC 4500", None), ("", None)):
        assert enclosure_letter(text) == want, text


def test_generated_bom_never_used_as_reference(tmp_path):
    # MAJOR 8: a rerun took the tool's own FINAL_BOM as the approved BOM (and --learn learned from it)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OLM SF-8M BOM"
    ws.append([None, "Residential Count", None, None, None, None, None, 17, 210])
    wb.properties.keywords = f"{GENERATED_MARK} 0.5.0"
    wb.save(tmp_path / "renamed.xlsx")
    assert classify_xlsx(tmp_path / "renamed.xlsx").startswith("unknown")
    out = tmp_path / "out"
    out.mkdir()
    wb.properties.keywords = ""
    wb.save(out / "X_FINAL_BOM.xlsx")  # inside the output folder
    wb.save(tmp_path / "Y_FINAL_BOM.xlsx")  # generated file name
    det = detect([tmp_path], tmp_path / "x", exclude=[out])
    assert not det.bom_completed


def test_zip_same_stem_not_merged_and_bomb_rejected(tmp_path):
    # MINOR 14
    for sub in ("a", "b"):
        (tmp_path / sub).mkdir()
        with zipfile.ZipFile(tmp_path / sub / "Splice Reports.zip", "w") as z:
            z.writestr("r.txt", sub)
    det = detect([tmp_path / "a", tmp_path / "b"], tmp_path / "x")
    assert sorted(p.parent.name for p, _why in det.unknown) == ["Splice Reports", "Splice Reports_2"]
    bomb = tmp_path / "bomb.zip"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("zeros.bin", b"\0" * (30 * 1024 * 1024))
    with pytest.raises(ValueError, match="zip bomb"):
        safe_extract(bomb, tmp_path / "bomb")


def test_rule_expressions_cannot_run_code():
    # MINOR 13: eval with empty builtins was escapable
    assert safe_eval("3*olm", {"olm": 2}, set())[0] == 6
    assert safe_eval("ceil(x/2) + max(a, 1)", {"x": 3, "a": 0}, set())[0] == 3
    with pytest.raises(ValueError):
        safe_eval("().__class__.__base__.__subclasses__()", {}, set())
    with pytest.raises(ValueError):
        safe_eval("[c for c in (1, 2)]", {}, set())
    tk = Takeoff(["N"])
    tk.add("olm", "N", GRANT, 1, "EA", "")
    v, note = evaluate(Rule("P", None, "__import__('os').getpid()", "EA", "VERIFIED", True, "", ""), tk, "N", GRANT)
    assert v is None and "not allowed" in note


def test_negative_quantity_left_blank():
    tk = Takeoff(["N"])
    tk.add("pass_res", "N", NONGRANT, -2, "EA", "")
    v, note = evaluate(Rule("P", None, "pass_res", "EA", "VERIFIED", True, "", ""), tk, "N", NONGRANT)
    assert v is None and "< 0" in note


def test_renamed_block_is_blank_not_zero():
    # MAJOR 11(b): a renamed vault block wrote 0 / 0 with no signal
    d = Drawing("t", "test")
    d.block_census = {"CX1-HANDHOLE": {"LINE": 4}}
    tk = Takeoff(["N"])

    class C:
        def network_of(self, p, what="", strict=False):
            return "N"

        def area_of(self, p):
            return GRANT

        def foreign_of(self, p, strict=False):
            return None

    count_blocks(d, CONV, C(), tk, "vault")
    assert ("vault", "N", GRANT) not in tk.q
    assert any(c.name.startswith("block not defined") and c.status == "FAIL" for c in tk.checks)
    d.block_census["CX1-VAULT"] = {"LINE": 4}  # defined but not inserted: a real zero
    count_blocks(d, CONV, C(), tk, "vault")
    assert tk.get("vault", "N", GRANT) == 0


def test_text_numbers_in_splice_report(tmp_path):
    # MAJOR 16: one text number aborted the whole job
    assert _number("1,234") == 1234 and _number("48 ct") == 48 and _number(72) == 72 and _number("n/a") is None
    (tmp_path / "bad.xlsx").write_bytes(b"not a workbook")
    errors = []
    assert read_splice_reports([tmp_path], errors) == [] and "bad.xlsx" in errors[0]
    with pytest.raises(Exception):
        read_splice_reports([tmp_path])


def test_doctor_does_not_log_into_classifier():
    # MINOR 17: the deep check added ~28 rows to the audit's Network Assignment sheet
    from fttx_bom.engine.doctor import run_doctor

    sq = [(0, 0), (1000, 0), (1000, 1000), (0, 1000)]
    d = Drawing("t", "test")
    d.polylines.append(Polyline("b", CONV["network_boundary_layers"][0], sq, True))
    d.texts.append(Text("t", CONV["network_label_layers"][0], "X-R01", (500, 500)))
    d.inserts.append(Insert("v", "CX1-VAULT", "x", (1050, 500)))  # 50 ft outside: nearest-assigned
    cls = build_classifier(d, CONV, ["X-R01"])
    tk = Takeoff(["X-R01"])
    before = (len(cls.unassigned), len(cls.nearest_assigned), len(cls.foreign_assigned))
    run_doctor(d, CONV, cls, tk, {}, [])
    assert (len(cls.unassigned), len(cls.nearest_assigned), len(cls.foreign_assigned)) == before


def _mini_template(path, total_col):
    """BOM sheet whose Total Quantity sits in `total_col` (H = one column more than the standard G)."""
    from openpyxl.utils import column_index_from_string, get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OLM SF-8M BOM"
    t = column_index_from_string(total_col)
    hdr = [None, "Part Number", "Manufacturer", "Description", "Unit", "Qty"] + ["Extra"] * (t - 7) + ["Total Quantity",
                                                                                                      "A Grant", "A Non Grant"]
    for r in range(1, 5):
        ws.append([None, f"header {r}"])
    ws.append(hdr)
    a, b = get_column_letter(t + 1), get_column_letter(t + 2)
    ws.append([None, "CBP-MAX5DSM24", "RaDD", "ped", "EA", None] + [None] * (t - 7) + [f"=SUM({a}6:{b}6)"])
    wb.save(path)
    return path


@pytest.mark.parametrize("total_col", ["G", "H"])
def test_writer_finds_area_columns_after_total_quantity(tmp_path, total_col):
    # MAJOR 10: area columns were assumed to start at H; a shifted template lost its Total Quantity formulas
    from openpyxl.utils import column_index_from_string, get_column_letter

    from fttx_bom.excel.writer import write_bom

    cfg, _rules = load_rules()
    rule = Rule("CBP-MAX5DSM24", None, "ped_max5", "EA", "VERIFIED", True, "", "")
    tk = Takeoff(["N-R01"])
    tk.add("ped_max5", "N-R01", GRANT, 5, "EA", "")
    tk.add("ped_max5", "N-R01", NONGRANT, 7, "EA", "")
    res = write_bom(_mini_template(tmp_path / "t.xlsx", total_col), tmp_path / "o.xlsx", tk, [rule], cfg, {})
    t = column_index_from_string(total_col)
    ws = openpyxl.load_workbook(tmp_path / "o.xlsx")["OLM SF-8M BOM"]
    assert res.total_col == total_col
    assert str(ws.cell(6, t).value).startswith("=SUM(")  # Total Quantity formula kept
    assert (ws.cell(6, t + 1).value, ws.cell(6, t + 2).value) == (5, 7)
    assert ws.cell(5, t).value == "Total Quantity"
    assert GENERATED_MARK in (openpyxl.load_workbook(tmp_path / "o.xlsx").properties.keywords or "")
    assert res.columns == {("N-R01", GRANT): get_column_letter(t + 1), ("N-R01", NONGRANT): get_column_letter(t + 2)}


def test_discrepancy_mileage_tolerance():
    # regression: route-mileage cells 0.2 mi (1,056 ft) apart were reported as MATCH (0.5 tolerance for every unit)
    from fttx_bom.excel.reports import Reference, compare
    from fttx_bom.excel.writer import WriteResult

    rule = Rule(None, "Total Route Mileage", "route_miles", "MI", "INFERRED", True, "", "")
    tk = Takeoff(["N"])
    res = WriteResult(Path("x.xlsx"), "S")
    res.rule_values[rule.key] = {("N", GRANT): (7.9991, ""), ("N", NONGRANT): (6.9663, "")}
    ref = Reference("r", columns={("N", GRANT): "H", ("N", NONGRANT): "I"},
                    values={rule.key: {("N", GRANT): 7.7992, ("N", NONGRANT): 7.166}})
    rows = {r["area"]: r["status"] for r in compare([rule], res, ref, tk)}
    assert rows == {GRANT: "DIFF", NONGRANT: "DIFF", "TOTAL": "MATCH"}


def test_zero_vaults_flagged_when_block_absent():
    # code review MAJOR 11(b), second half: inserts renamed but the old definition left in the drawing
    from fttx_bom.engine.doctor import run_doctor

    d = Drawing("t", "test")
    d.block_census = {"CX1-VAULT": {"LINE": 4}, "CX1-HANDHOLE": {"LINE": 4}}
    d.inserts.append(Insert("h", "CX1-HANDHOLE", "RTE_UGPOLES", (10, 10)))
    tk = Takeoff(["N"])
    issues = run_doctor(d, CONV, build_classifier(d, CONV, ["N"]), tk, {}, [])
    assert any(i.code == "BLOCK_MISSING" and "CX1-VAULT" in i.message for i in issues)
    assert any(i.code == "UNMAPPED_BLOCK" and "CX1-HANDHOLE" in i.message for i in issues)
