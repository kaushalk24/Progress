"""Tests for auto-detection, policies/calibration, deep check output, jobs and the app."""
import json
import zipfile
from pathlib import Path

import openpyxl
import pytest

from fttx_bom.engine import policies as pol
from fttx_bom.engine.doctor import Issue, write_issue_files
from fttx_bom.engine.network import GRANT, NONGRANT
from fttx_bom.engine.rules import Rule
from fttx_bom.engine.takeoff import CableSegment, Takeoff
from fttx_bom.excel.reports import Reference
from fttx_bom.inputs.detect import Detected, classify_xlsx, detect, safe_extract
from fttx_bom.jobs import params_from_detected

APP = Path(__file__).resolve().parents[1] / "app" / "bom_app.py"


def _xlsx(path, rows, title="Sheet"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def test_classify_inputs(tmp_path):
    al = _xlsx(tmp_path / "al.xlsx", [["PON NAME", "PON PORT", "FIBER #", "ADDRESS"]], "X-R01")
    sr = _xlsx(tmp_path / "sr.xlsx", [["Splice Connections Report"]])
    sg = _xlsx(tmp_path / "sg.xlsx", [[None, "Fiber Optic Splice Guide"]])
    other = _xlsx(tmp_path / "o.xlsx", [["hello"]])
    tpl = _xlsx(tmp_path / "t.xlsx", [[None, "Residental Count", None, None, None, None, None, None]], "OLM SF-8M BOM")
    done = _xlsx(tmp_path / "d.xlsx", [[None, "Residential Count", None, None, None, None, None, 17, 210]], "OLM SF-8M BOM")
    assert classify_xlsx(al) == "address_list"
    assert classify_xlsx(sr) == "splice_report"
    assert classify_xlsx(sg) == "splice_report"  # regression: second report layout
    assert classify_xlsx(other).startswith("unknown")
    assert classify_xlsx(tpl) == "bom_template"
    assert classify_xlsx(done) == "bom_completed"


def test_detect_extracts_nested_zips(tmp_path):
    inner = tmp_path / "inner.zip"
    sr = _xlsx(tmp_path / "rep.xlsx", [["Splice Connections Report"]])
    with zipfile.ZipFile(inner, "w") as z:
        z.write(sr, "Splice Reports/rep.xlsx")
    outer = tmp_path / "job.zip"
    with zipfile.ZipFile(outer, "w") as z:
        z.write(inner, "inner.zip")
        z.writestr("design.dwg", b"AC1032")
    sr.unlink()
    inner.unlink()
    d = detect([outer], tmp_path / "x")
    assert [p.name for p in d.drawings] == ["design.dwg"]
    assert [p.name for p in d.splice_reports] == ["rep.xlsx"]


def test_zip_slip_rejected(tmp_path):
    bad = tmp_path / "bad.zip"
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("../../evil.txt", b"x")
    with pytest.raises(ValueError):
        safe_extract(bad, tmp_path / "out")
    assert not (tmp_path / "evil.txt").exists()


def test_params_from_detected_rules(tmp_path):
    det = Detected(drawings=[tmp_path / "a.dwg"], bom_completed=[tmp_path / "done.xlsx"])
    p, notes = params_from_detected(det, tmp_path / "out")
    assert p.template.name == "done.xlsx" and p.reference.name == "done.xlsx"  # completed BOM gives layout + reference
    assert any("no address list" in n for n in notes)
    with pytest.raises(ValueError):
        params_from_detected(Detected(drawings=[tmp_path / "a.dwg", tmp_path / "b.dwg"]), tmp_path)
    with pytest.raises(ValueError):
        params_from_detected(Detected(drawings=[tmp_path / "a.dwg"]), tmp_path)  # no template


def _tk():
    tk = Takeoff(["N-R01"])
    for a, v in ((GRANT, 1), (NONGRANT, 3)):
        tk.add("pon_ports", "N-R01", a, v, "EA", "base")
        tk.add("pon_ports@jumper_half", "N-R01", a, v, "EA", "")
        tk.add("pon_ports@split_ratio", "N-R01", a, v + 1, "EA", "")
        tk.add("pon_ports@designed_pons", "N-R01", a, v, "EA", "")
    return tk


def test_policy_apply_and_calibrate(tmp_path):
    tk = _tk()
    pol.apply(tk, {"sfp": "split_ratio"})
    assert tk.get("pon_ports", "N-R01", NONGRANT) == 4 and tk.policy_used["sfp"] == "split_ratio"
    rule = Rule("FM0-3FE47548AD", None, "pon_ports", "EA", "VERIFIED", True, "", "")
    ref = Reference("ref", columns={("N-R01", GRANT): "H", ("N-R01", NONGRANT): "I"},
                    values={rule.key: {("N-R01", GRANT): 1, ("N-R01", NONGRANT): 3}})
    rows = [r for r in pol.calibrate(tk, [rule], ref) if r["policy"] == "sfp"]
    best = [r for r in rows if r["best"]]
    assert len(best) == 1 and best[0]["option"] == "jumper_half"  # ties with designed_pons; total_minus_ng not computed here
    path = tmp_path / "policies.json"
    sel = pol.learn(rows, "job1", path)
    assert sel["sfp"] == "jumper_half"
    assert json.loads(path.read_text())["evidence"]["sfp"]["split_ratio"]["job1"]["matched"] == 0


def test_unknown_policy_option_falls_back():
    tk = _tk()
    pol.apply(tk, {"sfp": "nonsense"})
    assert tk.policy_used["sfp"] == pol.BY_NAME["sfp"].default
    assert any(c.name == "policy sfp" for c in tk.checks)


def test_issue_files(tmp_path):
    issues = [Issue("ERROR", "Cable", "X", "N", "msg, with comma", "lines", "fix", 100.0, 200.0),
              Issue("INFO", "Drawing", "Y", "", "no location", "", "")]
    csv_path, scr = write_issue_files(issues, tmp_path / "job")
    assert "msg, with comma" in csv_path.read_text()
    s = scr.read_text()
    assert s.count("CIRCLE") == 1 and "100.000,200.000" in s
    assert '(setvar "OSMODE" 0)' in s and '(setvar "OSMODE" bomchk_os)' in s


def test_report_only_cable_keywords_regression():
    # regression: a new dataclass field shifted positional args and put text into path_ft
    c = CableSegment(handle="report:X", name="A TO B FC024", a="A", b="B", fibers=24, placement="UG", route_ft=100,
                     slack_ft=0, total_ft=100, network="N", ng_fraction=0.0, ng_fraction_length=0.0, split_method="splice report only")
    assert c.path_ft is None and c.split_method == "splice report only"


def test_app_pages_load():
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(APP), default_timeout=60)
    at.run()
    assert not at.exception
    assert at.title[0].value == "New BOM"
    for page in ("Batch", "Design standards", "Settings"):
        at.sidebar.radio[0].set_value(page).run()
        assert not at.exception, page
        assert at.title[0].value == page
    assert len(at.selectbox) == len(pol.POLICIES)  # one per convention on the Settings page
    text = " ".join(str(m.value) for m in at.markdown) + " ".join(str(c.value) for c in at.caption)
    assert "SME" not in text and "meeting" not in text.lower()


def test_batch_summary_counts_per_bom(tmp_path, monkeypatch):
    # regression: batch rows showed job-wide issue totals instead of each BOM's own
    import fttx_bom.jobs as jobs
    from fttx_bom.pipeline import GroupResult, JobParams, JobResult

    (tmp_path / "jobs" / "J1").mkdir(parents=True)
    (tmp_path / "jobs" / "J1" / "d.dwg").write_bytes(b"x")

    def fake_run_auto(paths, out, **kw):
        g1 = GroupResult(["A-R01"], tmp_path / "a.xlsx", tmp_path / "a_audit.xlsx", None, tmp_path / "a.csv", tmp_path / "a.scr",
                         [Issue("ERROR", "c", "X", "A-R01", "m", "", "")], {}, [], True)
        g2 = GroupResult(["A-R02"], tmp_path / "b.xlsx", tmp_path / "b_audit.xlsx", None, tmp_path / "b.csv", tmp_path / "b.scr",
                         [], {}, [], True)
        r = JobResult(params=JobParams(tmp_path / "d.dwg", [], [], tmp_path / "t.xlsx", out), networks=["A-R01", "A-R02"], groups=[g1, g2])
        r.totals = {"passings": {"A-R01": 10, "A-R02": 20}}
        return r, []

    monkeypatch.setattr(jobs, "run_auto", fake_run_auto)
    path = jobs.run_batch(tmp_path / "jobs", tmp_path / "out", log=lambda m: None)
    rows = list(openpyxl.load_workbook(path).active.iter_rows(min_row=2, values_only=True))
    assert [(r[1], r[3], r[7]) for r in rows] == [("A-R01", 1, 10), ("A-R02", 0, 20)]
